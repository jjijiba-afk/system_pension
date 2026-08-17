"""회사별 지급규정 설정 — 임원 정년·직군 제외·위로금.

실제 케이스 6건의 `일반사항` 6번 항목을 반영하며 추가된 기능들이다.
"""

from __future__ import annotations

import datetime as _dt

import pytest

from pension.actuarial import normal_retirement_age
from pension.assumptions import Assumptions, DiscountCurve, RateCurve
from pension.config import CalculationConfig, JobGroupRule
from pension.errors import IssueLog
from pension.models import ActiveMember
from pension.normalize import BenefitPlan, EmployeeType
from pension.validation import validate_active
from pension.valuation import value_member

BASE = _dt.date(2025, 12, 31)


def _rule(**kw) -> JobGroupRule:
    defaults = dict(source_name="정규직", mapped_name="정규직",
                    severance_nra=60, longterm_nra=60, over_nra_add_age=2)
    defaults.update(kw)
    return JobGroupRule(**defaults)


def _member(rule: JobGroupRule, **kw) -> ActiveMember:
    member = ActiveMember(seq=1, row=26)
    member.employee_id = "A1"
    member.job_group_raw = rule.source_name
    member.job_group = rule.mapped_name
    member.job_group_index = 0
    member.birth_date = _dt.date(1980, 1, 1)
    member.hire_date = _dt.date(2005, 1, 1)
    member.settlement_date = member.hire_date
    member.monthly_wage = 5_000_000
    member.plan = BenefitPlan.DB
    member.min_service_years = rule.min_service_years
    member.excluded_group = rule.excluded
    for key, value in kw.items():
        setattr(member, key, value)
    return member


class TestRetirementAgeIsTheSameForEveryone:
    """정년은 **직군 표에 적은 그대로** 쓴다.

    임직원 구분으로 값을 바꾸지 않는다 — 임원 정년이 다르면 직군 표에 임원
    줄을 만들어 그 줄에 적는다. 한 줄만 보면 그 사람의 정년을 알 수 있어야
    하고, 화면에 안 보이는 규칙이 숫자를 바꾸면 안 된다.
    """

    def test_executive_and_staff_get_the_same_age(self) -> None:
        rule = _rule(severance_nra=60, over_nra_add_age=2)
        for age, expected in ((50, 60), (63, 65)):
            assert normal_retirement_age(age, rule, is_executive=True) == expected
            assert normal_retirement_age(age, rule, is_executive=False) == expected

    def test_a_blank_age_means_no_retirement_age(self) -> None:
        """비운 칸은 0 세다 — 이미 넘긴 셈이라 현재 연령에 가산연수를 더한다."""
        rule = _rule(severance_nra=0, over_nra_add_age=2)
        assert normal_retirement_age(45, rule) == 47
        assert normal_retirement_age(63, rule, is_executive=True) == 65

    def test_the_add_on_is_used_as_written(self) -> None:
        assert normal_retirement_age(63, _rule(severance_nra=60, over_nra_add_age=1)) == 64
        assert normal_retirement_age(63, _rule(severance_nra=60, over_nra_add_age=3)) == 66

    def test_wage_peak_still_wins_when_ahead(self) -> None:
        rule = _rule(severance_nra=60, over_nra_add_age=2)
        assert normal_retirement_age(50, rule, wage_peak_age=58) == 58

    def test_validation_applies_the_group_rule(self) -> None:
        rule = _rule(severance_nra=39, over_nra_add_age=2)
        config = CalculationConfig(base_date=BASE, job_group_rules=[rule])
        member = _member(rule, employee_type=EmployeeType.EXECUTIVE)
        validate_active([member], config, IssueLog())
        assert member.severance_nra == member.age + 2


class TestExcludedJobGroup:
    """'계약직은 퇴직금 대상에서 제외' 같은 직군 단위 제외."""

    def test_excluded_group_produces_no_liability(self) -> None:
        rule = _rule(source_name="계약직", mapped_name="계약직", excluded=True)
        config = CalculationConfig(base_date=BASE, job_group_rules=[rule])
        member = _member(rule)
        member.age = 45
        member.severance_nra = 60

        assumptions = Assumptions(
            discount=DiscountCurve(spot=RateCurve({1: 0.045}), flat=0.045)
        )
        result = value_member(member, config, assumptions)

        assert result.dbo == 0.0
        assert result.accrued_benefit == 0.0
        assert "제외" in result.excluded_reason

    def test_included_group_is_unaffected(self) -> None:
        rule = _rule(excluded=False)
        config = CalculationConfig(base_date=BASE, job_group_rules=[rule])
        member = _member(rule)
        member.age = 45
        member.severance_nra = 60

        assumptions = Assumptions(
            discount=DiscountCurve(spot=RateCurve({1: 0.045}), flat=0.045)
        )
        assert value_member(member, config, assumptions).dbo > 0.0

    def test_excluded_flag_reads_several_spellings(self) -> None:
        import openpyxl

        from pension.config import read_config

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Input"
        ws["C3"] = _dt.datetime(2025, 12, 31)
        for offset, token in enumerate(("Y", "제외", "N", "")):
            row = 12 + offset
            ws.cell(row, 2, f"직군{offset}")
            ws.cell(row, 3, f"직군{offset}")
            ws.cell(row, 18, token)

        config = read_config(wb)
        assert [r.excluded for r in config.job_group_rules] == [True, True, False, False]


class TestSeveranceBonus:
    """명부의 추가지급 기본급은 **적어 두기만** 한다.

    그 칸에 담긴 것이 회사마다 다르다(사망 위로금·명퇴 가산금). 모든 사유에
    정액으로 얹으면 채무가 배로 뛰므로, 엔진은 값을 결과에 남길 뿐 급여에
    더하지 않는다. 실제 가산은 [퇴직사유] 표가 사유별로 정한다.
    """

    def _valued(self, extra: float):
        rule = _rule()
        config = CalculationConfig(base_date=BASE, job_group_rules=[rule])
        member = _member(rule, extra_pay_base_wage=extra)
        member.age = 45
        member.severance_nra = 60
        assumptions = Assumptions(
            discount=DiscountCurve(spot=RateCurve({1: 0.045}), flat=0.045)
        )
        return value_member(member, config, assumptions)

    def test_the_amount_is_recorded_but_not_added(self) -> None:
        without = self._valued(0.0)
        with_bonus = self._valued(10_000_000)

        assert with_bonus.extra_payment == 10_000_000
        assert with_bonus.accrued_benefit == pytest.approx(without.accrued_benefit)
        assert with_bonus.dbo == pytest.approx(without.dbo)

    def test_blank_means_not_eligible(self) -> None:
        result = self._valued(0.0)
        assert result.extra_payment == 0.0

    def test_negative_is_ignored(self) -> None:
        assert self._valued(-5_000_000).extra_payment == 0.0


class TestExclusionSummary:
    """DC 전환을 마친 회사는 채무가 0 이다. 이유를 함께 보여 줘야 한다."""

    def test_summarises_reasons(self) -> None:
        from pension.valuation import MemberValuation, ValuationResult

        result = ValuationResult(base_date=BASE, label="당기")
        for index in range(3):
            member = MemberValuation(
                employee_id=f"A{index}", name="", job_group="정규직", gender="남자",
                age=40, past_service=10.0, projection_years=20, monthly_wage=0.0,
            )
            member.excluded_reason = "DC 가입자 (확정기여제도는 확정급여채무 없음)"
            result.members.append(member)

        assert result.headcount == 0
        assert result.dbo == 0.0
        summary = result.exclusion_summary()
        assert summary == {"DC 가입자 (확정기여제도는 확정급여채무 없음)": 3}

    def test_empty_when_everyone_is_valued(self) -> None:
        from pension.valuation import ValuationResult

        assert ValuationResult(base_date=BASE, label="당기").exclusion_summary() == {}
