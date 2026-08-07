"""지급률 방식 세 가지와 기초율 워크북 왕복."""

from __future__ import annotations

import itertools
from pathlib import Path

import openpyxl
import pytest

from pension.assumptions import (
    BENEFIT_RULE_SHEET,
    CUMULATIVE,
    FORMULA,
    PROGRESSIVE,
    BenefitScale,
    RateCurve,
    load_assumptions,
    write_assumptions,
    write_template,
)
from pension.formula import Formula


class TestCumulative:
    """표 값이 그대로 누적 배수. 기존 동작이며 방식을 적지 않으면 이쪽이다."""

    def test_step_lookup(self) -> None:
        scale = BenefitScale(curves={"정규직": RateCurve({1: 1.0, 10: 12.0, 20: 26.0})})
        assert scale.multiple("정규직", 1) == 1.0
        assert scale.multiple("정규직", 9) == 1.0  # 10년 도달 전에는 직전 구간
        assert scale.multiple("정규직", 10) == 12.0
        assert scale.multiple("정규직", 25) == 26.0

    def test_default_mode_is_cumulative(self) -> None:
        scale = BenefitScale(curves={"정규직": RateCurve({10: 12.0})})
        assert scale.mode("정규직") == CUMULATIVE

    def test_missing_rule_falls_back_to_statutory(self) -> None:
        """법정 퇴직금은 계속근로 1년당 30일분이므로 배수 = 근속연수."""
        scale = BenefitScale(statutory_when_missing=True)
        assert scale.multiple("없는규정", 7.5) == 7.5

    def test_missing_rule_is_zero_for_longterm(self) -> None:
        scale = BenefitScale(statutory_when_missing=False)
        assert scale.multiple("없는규정", 7.5) == 0.0


class TestProgressive:
    """표 값이 구간별 연 배수. 누진제 퇴직금 규정을 표만으로 적을 수 있다."""

    @pytest.fixture
    def scale(self) -> BenefitScale:
        return BenefitScale(
            curves={"누진": RateCurve({0: 1.0, 5: 1.5, 10: 2.0})},
            modes={"누진": PROGRESSIVE},
        )

    @pytest.mark.parametrize(
        ("service", "expected"),
        [
            (0, 0.0),
            (3, 3.0),          # 3 × 1.0
            (5, 5.0),          # 5 × 1.0
            (7, 8.0),          # 5×1.0 + 2×1.5
            (10, 12.5),        # 5×1.0 + 5×1.5
            (12, 16.5),        # 5×1.0 + 5×1.5 + 2×2.0
            (20, 32.5),        # 5×1.0 + 5×1.5 + 10×2.0
        ],
    )
    def test_bands_accumulate(self, scale, service, expected) -> None:
        assert scale.multiple("누진", service) == pytest.approx(expected)

    def test_fractional_service(self, scale) -> None:
        """투영 중에는 근속이 소수로 나온다(중간 시점 퇴직 가정)."""
        assert scale.multiple("누진", 5.5) == pytest.approx(5.0 + 0.5 * 1.5)

    def test_is_monotonic(self, scale) -> None:
        values = [scale.multiple("누진", t / 4) for t in range(120)]
        assert all(b >= a for a, b in itertools.pairwise(values))


class TestFormulaMode:
    def test_uses_formula_over_table(self) -> None:
        scale = BenefitScale(
            curves={"임원": RateCurve({1: 99.0})},
            formulas={"임원": Formula("=t * 3")},
            modes={"임원": FORMULA},
        )
        assert scale.multiple("임원", 10) == 30.0

    def test_receives_member_context(self) -> None:
        scale = BenefitScale(
            formulas={"임원": Formula('=IF(제도="DB", t * 2, t)')},
            modes={"임원": FORMULA},
        )
        assert scale.multiple("임원", 10, 제도="DB") == 20.0
        assert scale.multiple("임원", 10, 제도="퇴직금제도") == 10.0

    def test_has_rule_sees_formula_only_rules(self) -> None:
        scale = BenefitScale(formulas={"임원": Formula("=t")}, modes={"임원": FORMULA})
        assert scale.has_rule("임원")


class TestWorkbookRoundTrip:
    def test_template_includes_rule_sheet(self, tmp_path: Path) -> None:
        path = write_template(tmp_path / "기초율.xlsx", job_groups=["정규직", "임원"])
        wb = openpyxl.load_workbook(path)
        assert BENEFIT_RULE_SHEET in wb.sheetnames
        ws = wb[BENEFIT_RULE_SHEET]
        assert [ws.cell(r, 1).value for r in (2, 3)] == ["정규직", "임원"]
        assert ws.cell(2, 2).value == CUMULATIVE

    def _base_sheets(self) -> dict:
        return {
            "할인율": (["연차", "할인율"], [[1, 0.045]]),
            "임금상승률": (["연차", "Base-up 상승률"], [[1, 0.03]]),
            "승급률": (["연령", "정규직", "임원"], [[20, 0.02, 0.01]]),
            "퇴직률": (["연령", "정규직", "임원"], [[20, 0.10, 0.05]]),
            "사망률": (["연령", "남자", "여자"], [[20, 0.0004, 0.0002]]),
            "지급률": (["근속연수", "정규직", "임원"], [[0, 1.0, 1.0], [5, 1.5, 2.0]]),
            "장기급여지급률": (["근속연수", "정규직", "임원"], [[10, 10, 10]]),
        }

    def test_modes_and_formulas_survive_a_round_trip(self, tmp_path: Path) -> None:
        path = write_assumptions(
            tmp_path / "기초율.xlsx",
            self._base_sheets(),
            {"정규직": (PROGRESSIVE, ""), "임원": (FORMULA, "=IF(t<10, t*2, 20+(t-10)*3)")},
        )

        loaded = load_assumptions(path)
        benefit = loaded.severance_benefit

        assert benefit.mode("정규직") == PROGRESSIVE
        assert benefit.multiple("정규직", 7) == pytest.approx(5 * 1.0 + 2 * 1.5)

        assert benefit.mode("임원") == FORMULA
        assert benefit.multiple("임원", 12) == pytest.approx(20 + 2 * 3)

    def test_bad_formula_is_reported_with_its_cell(self, tmp_path: Path) -> None:
        path = write_assumptions(
            tmp_path / "기초율.xlsx",
            self._base_sheets(),
            {"정규직": (FORMULA, "=VLOOKUP(t, A, 2)")},
        )
        with pytest.raises(ValueError, match=r"지급률규정!C2"):
            load_assumptions(path)

    def test_unknown_mode_is_reported(self, tmp_path: Path) -> None:
        path = write_assumptions(
            tmp_path / "기초율.xlsx", self._base_sheets(), {"정규직": ("누적식", "")}
        )
        with pytest.raises(ValueError, match="방식"):
            load_assumptions(path)

    def test_formula_mode_without_formula_is_reported(self, tmp_path: Path) -> None:
        path = write_assumptions(
            tmp_path / "기초율.xlsx", self._base_sheets(), {"정규직": (FORMULA, "")}
        )
        with pytest.raises(ValueError, match="수식이 비어"):
            load_assumptions(path)

    def test_table_mode_with_stray_formula_is_reported(self, tmp_path: Path) -> None:
        """방식은 누진인데 수식이 남아 있으면 어느 쪽이 의도인지 알 수 없다."""
        path = write_assumptions(
            tmp_path / "기초율.xlsx", self._base_sheets(), {"정규직": (PROGRESSIVE, "=t*2")}
        )
        with pytest.raises(ValueError, match="수식을 쓰지 않습니다"):
            load_assumptions(path)

    def test_old_workbook_without_rule_sheet_still_loads(self, tmp_path: Path) -> None:
        """지급률규정 시트가 없던 파일도 그대로 읽혀야 한다."""
        path = write_assumptions(tmp_path / "기초율.xlsx", self._base_sheets(), None)
        loaded = load_assumptions(path)
        assert loaded.severance_benefit.mode("정규직") == CUMULATIVE
