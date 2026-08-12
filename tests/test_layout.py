"""명부 열 배치 인식과 구 서식 읽기."""

from __future__ import annotations

import datetime as _dt
from pathlib import Path

import openpyxl
import pytest

from pension.config import read_config
from pension.errors import IssueLog
from pension.layout import (
    ACTIVE_HEADER_ALIASES,
    REQUIRED_ACTIVE,
    find_data_start,
    find_header_row,
    normalize_header,
    resolve_layout,
)
from pension.readers import ACTIVE_COLUMNS
from pension.workbook import find_sheet, open_workbook


class TestNormalizeHeader:
    @pytest.mark.parametrize(
        ("raw", "expected"),
        [
            ("사번", "사번"),
            ("임직원구분\n(임원,직원)", "임직원구분"),
            ("30일 평균임금", "30일평균임금"),
            ("1日기본급", "1일기본급"),
            ("성별\n(남자/여자)", "성별"),
            ("휴직 \n차감일수", "휴직차감일수"),
            ("퇴직금  지급배수", "퇴직금지급배수"),
            (None, ""),
        ],
    )
    def test_strips_noise(self, raw, expected) -> None:
        assert normalize_header(raw) == expected


def _sheet(headers: list[str], rows: list[list[object]], header_row: int = 3):
    """머리글과 데이터를 가진 임시 시트."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for col, title in enumerate(headers, start=2):
        ws.cell(header_row, col, title)
    ws.cell(header_row + 1, 2, "TYPE")
    ws.cell(header_row + 2, 2, "작성 샘플")
    for offset, values in enumerate(rows, start=header_row + 3):
        ws.cell(offset, 2, offset - header_row - 2)
        for col, value in enumerate(values, start=3):
            ws.cell(offset, col, value)
    return ws


class TestHeaderDiscovery:
    def test_finds_header_row(self) -> None:
        ws = _sheet(["No.", "사번", "직군", "생년월일"], [["A1", "정규직", "1980-01-01"]])
        assert find_header_row(ws, ACTIVE_HEADER_ALIASES) == 3

    def test_returns_zero_when_absent(self) -> None:
        ws = _sheet(["가", "나", "다"], [[1, 2, 3]])
        assert find_header_row(ws, ACTIVE_HEADER_ALIASES) == 0

    def test_data_starts_after_type_and_sample_rows(self) -> None:
        ws = _sheet(["No.", "사번", "직군", "생년월일"], [["A1", "정규직", "1980-01-01"]])
        assert find_data_start(ws, 3) == 6


class TestShiftedColumns:
    """열이 끼어들어도 밀리지 않아야 한다.

    실제 명부 6건 중 하나는 재직자명부가 33열이었다. `연봉제 전환 추계일`,
    `누진적용 근속연수`, `누진적용 율` 세 열이 중간에 끼어 뒤쪽이 세 칸 밀린다.
    열 번호를 고정해 읽으면 오류 없이 엉뚱한 값을 읽는다.
    """

    def test_resolves_by_header_not_position(self) -> None:
        headers = [
            "No.", "사번", "임직원구분", "직군", "성명", "성별", "생년월일", "입사일자",
            "중간정산일", "30일 평균임금", "명예퇴직 산정용 임금", "퇴직급여추계액",
            "1日기본급",
            "연봉제 전환 추계일", "누진적용 근속연수", "누진적용 율",   # ← 끼어든 세 열
            "휴직차감일수", "퇴직급여 제도구분",
        ]
        ws = _sheet(headers, [])
        log = IssueLog()
        defaults = {name: col.index for name, col in ACTIVE_COLUMNS.items()}
        layout = resolve_layout(ws, ACTIVE_HEADER_ALIASES, defaults, REQUIRED_ACTIVE, log)

        # 제도구분은 표준 자리가 아니라 **머리글이 실제로 있는 자리** 로 잡혀야
        # 한다. 앞에 세 열이 끼어들었으므로 목록에서의 자리 그대로다.
        # 머리글은 2열부터 놓이므로 목록에서의 자리 + 2 가 실제 열이다.
        def where(title: str) -> int:
            return headers.index(title) + 2

        assert layout.columns["plan"] == where("퇴직급여 제도구분")
        assert layout.columns["monthly_wage"] == where("30일 평균임금")
        assert not log.has_errors()

    def test_absent_columns_are_not_faked_with_defaults(self) -> None:
        """없는 열을 기본 위치로 메우면 엉뚱한 열을 그 값으로 읽는다.

        구 서식은 28열이 '추가지급 기본급' 인데 신 서식은 '퇴직급여 지급률 규정'
        이다. 폴백하면 기본급이 규정명으로 들어간다.
        """
        headers = ["No.", "사번", "직군", "생년월일", "입사일자", "30일 평균임금",
                   "퇴직급여 제도구분", "추가지급 기본급"]
        ws = _sheet(headers, [])
        log = IssueLog()
        defaults = {name: col.index for name, col in ACTIVE_COLUMNS.items()}
        layout = resolve_layout(ws, ACTIVE_HEADER_ALIASES, defaults, REQUIRED_ACTIVE, log)

        assert "severance_benefit" not in layout.columns
        assert "severance_benefit" in layout.report.missing

    def test_missing_required_column_is_an_error(self) -> None:
        ws = _sheet(["No.", "사번", "생년월일"], [])
        log = IssueLog()
        defaults = {name: col.index for name, col in ACTIVE_COLUMNS.items()}
        resolve_layout(ws, ACTIVE_HEADER_ALIASES, defaults, REQUIRED_ACTIVE, log)
        codes = {i.code for i in log.errors}
        assert "LAYOUT_REQUIRED_COLUMN_MISSING" in codes

    def test_header_with_trailing_description_still_matches(self) -> None:
        """실제 퇴직자명부의 지급사유 머리글에는 1~6 설명이 줄줄이 붙어 있다."""
        from pension.layout import RETIRED_HEADER_ALIASES

        headers = ["No.", "사번", "직군", "생년월일", "입사일", "퇴사일",
                   "지급(퇴직)사유 구분\n1:중도퇴직\n2:사망퇴직\n3:DC전환"]
        ws = _sheet(headers, [])
        log = IssueLog()
        layout = resolve_layout(ws, RETIRED_HEADER_ALIASES, {}, (), log)
        # 머리글은 2열부터 놓았으므로 일곱 번째인 지급사유는 8열이다.
        assert layout.columns["reason"] == 8


class TestSheetAliases:
    def test_finds_numbered_sheet_names(self, tmp_path: Path) -> None:
        wb = openpyxl.Workbook()
        wb.active.title = "2)재직자명부"
        wb.create_sheet("퇴직자명부")
        path = tmp_path / "명부.xlsx"
        wb.save(path)

        book = open_workbook(path)
        try:
            assert find_sheet(book, "재직자명부") is not None
            assert find_sheet(book, "퇴직자명부") is not None
            assert find_sheet(book, "없는시트") is None
        finally:
            book.close()


class TestWorkbookFormats:
    def test_rejects_unknown_extension(self, tmp_path: Path) -> None:
        path = tmp_path / "명부.txt"
        path.write_text("x", encoding="utf-8")
        with pytest.raises(ValueError, match="지원하지 않는"):
            open_workbook(path)

    def test_missing_file(self, tmp_path: Path) -> None:
        with pytest.raises(FileNotFoundError):
            open_workbook(tmp_path / "없는파일.xlsx")


class TestInferredConfig:
    """``Input`` 시트가 없는 자료요청서 원본도 읽을 수 있어야 한다."""

    def test_derives_base_date_and_job_groups(self, tmp_path: Path) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "2)재직자명부"
        ws.cell(20, 2, "작성기준일")
        ws.cell(20, 3, _dt.datetime(2025, 12, 31))
        for col, title in enumerate(
            ["No.", "사번", "임직원구분", "직군", "성명", "성별", "생년월일",
             "입사일자", "중간정산일", "30일 평균임금"], start=2,
        ):
            ws.cell(21, col, title)
        ws.cell(22, 2, "TYPE")
        ws.cell(23, 2, "작성 샘플")
        for offset, group in enumerate(["정규직", "계약직", "정규직"], start=24):
            ws.cell(offset, 2, offset - 23)
            ws.cell(offset, 5, group)
            ws.cell(offset, 8, _dt.datetime(1985, 1, 1))
        path = tmp_path / "원본.xlsx"
        wb.save(path)

        book = open_workbook(path)
        try:
            config = read_config(book)
        finally:
            book.close()

        assert config.base_date == _dt.date(2025, 12, 31)
        assert config.inferred is True
        assert [r.source_name for r in config.job_group_rules] == ["정규직", "계약직"]

    def test_raises_when_no_base_date_anywhere(self, tmp_path: Path) -> None:
        wb = openpyxl.Workbook()
        wb.active.title = "2)재직자명부"
        path = tmp_path / "빈.xlsx"
        wb.save(path)

        book = open_workbook(path)
        try:
            with pytest.raises(ValueError, match="산출기준일"):
                read_config(book)
        finally:
            book.close()


class TestMinimumService:
    """가입자격(최소 근속연수)."""

    def test_benefit_is_zero_below_the_minimum(self) -> None:
        from pension.assumptions import Assumptions, DiscountCurve, RateCurve
        from pension.config import CalculationConfig, JobGroupRule
        from pension.models import ActiveMember
        from pension.normalize import BenefitPlan
        from pension.valuation import value_member

        rule = JobGroupRule("정규직", "정규직", 60, 60, 2, min_service_years=3.0)
        config = CalculationConfig(base_date=_dt.date(2025, 12, 31), job_group_rules=[rule])
        assumptions = Assumptions(discount=DiscountCurve(spot=RateCurve({1: 0.045}), flat=0.045))

        member = ActiveMember(seq=1, row=26)
        member.job_group_index = 0
        member.birth_date = _dt.date(1995, 1, 1)
        member.hire_date = _dt.date(2024, 6, 1)
        member.settlement_date = member.hire_date
        member.monthly_wage = 3_000_000
        member.plan = BenefitPlan.DB
        member.age = 30
        member.severance_nra = 60
        member.min_service_years = 3.0

        result = value_member(member, config, assumptions)
        # 근속 1.6년이라 지금 나가면 못 받는다.
        assert result.accrued_benefit == 0.0
        # 정년까지 남으면 요건을 채우므로 부채는 0 이 아니다.
        assert result.dbo > 0.0

    def test_zero_minimum_means_no_restriction(self) -> None:
        from pension.config import JobGroupRule

        assert JobGroupRule("정규직", "정규직").min_service_years == 0.0
