"""실제 명부에서 드러난 사례들.

275명 재직·247명 퇴직 규모의 실제 명부를 돌려 보고 고친 것들이다. 샘플 명부로는
나오지 않던 형태라, 같은 오탐이 되살아나지 않도록 여기에 고정해 둔다.
"""

from __future__ import annotations

import datetime as _dt

import pytest

from pension.actuarial import attained_age
from pension.config import CalculationConfig, JobGroupRule
from pension.errors import IssueLog
from pension.models import ActiveMember, RetiredMember, Roster
from pension.normalize import BenefitPlan, EmployeeType, RetirementReason
from pension.readers import _multiple
from pension.validation import validate_active, validate_retired

BASE = _dt.date(2025, 12, 31)


@pytest.fixture
def config() -> CalculationConfig:
    return CalculationConfig(
        base_date=BASE,
        job_group_rules=[JobGroupRule("정규직", "정규직", 60, 60, 2)],
    )


def _active(**kw) -> ActiveMember:
    member = ActiveMember(seq=1, row=26)
    member.employee_id = "A1"
    member.name = "홍길동"
    member.job_group_raw = "정규직"
    member.job_group = "정규직"
    member.job_group_index = 0
    member.birth_date = _dt.date(1985, 5, 1)
    member.hire_date = _dt.date(2010, 3, 1)
    member.settlement_date = member.hire_date
    member.monthly_wage = 5_000_000
    member.plan = BenefitPlan.DB
    for key, value in kw.items():
        setattr(member, key, value)
    return member


def _retired(**kw) -> RetiredMember:
    member = RetiredMember(seq=1, row=22)
    member.employee_id = "T1"
    member.name = "홍길동"
    member.job_group_raw = "정규직"
    member.job_group = "정규직"
    member.job_group_index = 0
    member.birth_date = _dt.date(1985, 5, 1)
    member.hire_date = _dt.date(2015, 1, 1)
    member.exit_date = _dt.date(2025, 6, 30)
    member.reason = RetirementReason.VOLUNTARY
    member.plan = BenefitPlan.DB
    member.total_payment = 30_000_000
    for key, value in kw.items():
        setattr(member, key, value)
    return member


class TestDcMemberWage:
    """DC 가입자는 평균임금란을 비워 오는 일이 흔하다.

    확정기여제도는 확정급여채무가 생기지 않으므로 임금이 없어도 산출에 지장이
    없다. 오류로 막으면 명부 전체가 산출되지 않는다.
    """

    def test_blank_wage_is_only_a_warning_for_dc(self, config) -> None:
        log = IssueLog()
        validate_active([_active(plan=BenefitPlan.DC, monthly_wage=0)], config, log)
        assert not log.has_errors()
        assert any(i.code == "JAE_WAGE_MISSING_DC" for i in log.warnings)

    def test_blank_wage_is_still_an_error_for_db(self, config) -> None:
        log = IssueLog()
        validate_active([_active(plan=BenefitPlan.DB, monthly_wage=0)], config, log)
        assert any(i.code == "JAE_WAGE_MISSING" for i in log.errors)


class TestShortServiceRetiree:
    """근로자퇴직급여보장법 제4조 단서 — 계속근로 1년 미만은 지급 대상이 아니다.

    명부 작성요령도 "1년 미만 근무 후 퇴직자도 포함" 이라고 안내하므로, 지급액과
    제도구분이 비어 있는 것이 정상이다.
    """

    def test_missing_amount_is_a_warning_under_one_year(self, config) -> None:
        log = IssueLog()
        validate_retired(
            [_retired(hire_date=_dt.date(2025, 3, 1), exit_date=_dt.date(2025, 9, 1),
                      total_payment=0, plan=None)],
            config, log,
        )
        assert not log.has_errors()
        codes = {i.code for i in log.warnings}
        assert "TOI_TOTAL_MISSING_SHORT" in codes
        assert "TOI_PLAN_MISSING_SHORT" in codes

    def test_missing_amount_is_an_error_over_one_year(self, config) -> None:
        log = IssueLog()
        validate_retired([_retired(total_payment=0)], config, log)
        assert any(i.code == "TOI_TOTAL_MISSING" for i in log.errors)


class TestTransferAndDisposal:
    """전출·사업처분에서는 적립자산이 통째로 승계된다."""

    def test_fund_over_total_is_a_warning_on_disposal(self, config) -> None:
        log = IssueLog()
        validate_retired(
            [_retired(reason=RetirementReason.DISPOSAL,
                      total_payment=383_166_670, fund_payment=430_843_575,
                      transfer_out_payment=383_166_670)],
            config, log,
        )
        assert not log.has_errors()
        assert any(i.code == "TOI_FUND_OVER_TOTAL_TRANSFER" for i in log.warnings)

    def test_fund_over_total_is_an_error_on_ordinary_exit(self, config) -> None:
        log = IssueLog()
        validate_retired(
            [_retired(total_payment=30_000_000, fund_payment=40_000_000)], config, log
        )
        assert any(i.code == "TOI_FUND_OVER_TOTAL" for i in log.errors)

    def test_zero_payment_is_fine_when_paid_as_transfer(self, config) -> None:
        log = IssueLog()
        validate_retired(
            [_retired(reason=RetirementReason.DISPOSAL, total_payment=0,
                      transfer_out_payment=50_000_000)],
            config, log,
        )
        assert not any(i.code == "TOI_TOTAL_MISSING" for i in log.errors)


class TestPayoutMultiple:
    """임원 누진배수는 명부에 '2배' / '현재 3배' 처럼 글자로 적혀 온다."""

    @pytest.mark.parametrize(
        ("raw", "expected"),
        [
            ("2배", 2.0),
            ("현재 3배", 3.0),
            ("1.5배", 1.5),
            (3, 3.0),
            (3.0, 3.0),
            ("", 1.0),
            (None, 1.0),
            ("배수적용", 1.0),
            (0, 1.0),
        ],
    )
    def test_parses_a_number_out_of_the_cell(self, raw, expected) -> None:
        assert _multiple(raw) == expected

    def test_formula_can_use_the_multiple(self) -> None:
        from pension.assumptions import FORMULA, BenefitScale
        from pension.formula import Formula

        scale = BenefitScale(
            formulas={"정규직": Formula("=t * 배수")}, modes={"정규직": FORMULA}
        )
        assert scale.multiple("정규직", 20, 배수=1.0) == 20.0
        assert scale.multiple("정규직", 20, 배수=3.0) == 60.0

    def test_formula_can_branch_on_employee_type(self) -> None:
        from pension.assumptions import FORMULA, BenefitScale
        from pension.formula import Formula

        scale = BenefitScale(
            formulas={"정규직": Formula('=IF(임직원="임원", t*3, t)')},
            modes={"정규직": FORMULA},
        )
        assert scale.multiple("정규직", 10, 임직원="임원") == 30.0
        assert scale.multiple("정규직", 10, 임직원="직원") == 10.0


class TestWagePeakRetirementAge:
    """임금피크 연령이 정년보다 앞서면 그 나이를 정년으로 본다(VBA 동일)."""

    def test_wage_peak_age_wins_when_it_is_ahead(self, config) -> None:
        log = IssueLog()
        member = _active(wage_peak_age=58)
        validate_active([member], config, log)
        assert member.age == attained_age(member.birth_date, BASE)
        assert member.severance_nra == 58

    def test_falls_back_to_group_nra_when_already_past_peak(self, config) -> None:
        log = IssueLog()
        member = _active(birth_date=_dt.date(1960, 5, 1), wage_peak_age=58)
        validate_active([member], config, log)
        # 65세는 정년(60)도 임금피크(58)도 지났으므로 현재 연령 + 가산연수
        assert member.severance_nra == member.age + 2


def test_roster_with_mixed_issues_still_produces_a_roster(config) -> None:
    """오류가 섞여 있어도 명부 객체 자체는 만들어져야 검증 결과를 모두 볼 수 있다."""
    log = IssueLog()
    roster = Roster(
        active=[_active(), _active(plan=BenefitPlan.DC, monthly_wage=0)],
        retired=[_retired(), _retired(total_payment=0)],
    )
    for member in roster.active:
        member.employee_type = EmployeeType.STAFF
    validate_active(roster.active, config, log)
    validate_retired(roster.retired, config, log)
    assert len(roster.active) == 2
    assert len(roster.retired) == 2


class TestNoiseReduction:
    """같은 말을 수백 번 반복하면 정작 봐야 할 항목이 묻힌다.

    실제 케이스 6건에서 경고 1,452건 중 1,211건이 '성명이 비어 있습니다' 였다.
    여섯 파일 모두 성명이 **한 명도** 없었다 — 개인정보를 지우고 사번만 남겨
    보내는 표준 관행이지 누락이 아니다.
    """

    def test_fully_masked_names_collapse_to_one_notice(self, config) -> None:
        from pension.errors import Severity

        log = IssueLog()
        members = [_active(seq=i, employee_id=f"A{i}", name="") for i in range(1, 51)]
        validate_active(members, config, log)

        named = [i for i in log if i.code == "JAE_NAME_MISSING"]
        assert len(named) == 1
        assert named[0].severity is Severity.INFO
        assert "50명" in named[0].message
        assert not log.warnings

    def test_partial_blanks_are_still_reported_per_person(self, config) -> None:
        """일부만 비었으면 진짜 누락일 수 있다."""
        log = IssueLog()
        members = [_active(seq=i, employee_id=f"A{i}") for i in range(1, 6)]
        members[2].name = ""
        validate_active(members, config, log)

        named = [i for i in log if i.code == "JAE_NAME_MISSING"]
        assert len(named) == 1
        assert named[0].seq == 3
        assert named[0] in log.warnings

    def test_notices_are_not_counted_as_warnings(self, config) -> None:
        log = IssueLog()
        validate_active([_active(name="")], config, log)
        assert log.notices and not log.warnings and not log.has_errors()


class TestRetirementReasonColumn:
    """퇴직사유는 경험퇴직률과 지급액 집계에 쓰이므로 비면 오류다.

    다만 케이스 10 은 24명 전부 비어 있었다. 사람별 누락이 아니라 회사가 그
    칸을 안 채운 것이므로 한 줄로 말한다.
    """

    def test_whole_column_blank_is_one_error(self, config) -> None:
        log = IssueLog()
        members = [
            _retired(seq=i, employee_id=f"T{i}", reason=None) for i in range(1, 25)
        ]
        validate_retired(members, config, log)

        found = [i for i in log if i.code == "TOI_REASON_MISSING"]
        assert len(found) == 1
        assert found[0] in log.errors
        assert "24명" in found[0].message

    def test_partial_blanks_are_reported_per_person(self, config) -> None:
        log = IssueLog()
        members = [_retired(seq=i, employee_id=f"T{i}") for i in range(1, 5)]
        members[1].reason = None
        validate_retired(members, config, log)

        found = [i for i in log if i.code == "TOI_REASON_MISSING"]
        assert len(found) == 1
        assert found[0].seq == 2


class TestDcRetiredPayment:
    """DC 는 회사가 부담금 납입으로 의무가 끝나고, 퇴직급여는 운용사가 지급한다.

    회사 명부의 지급액이 0 인 것이 정상이다. 실제 케이스에서 이 오류 39건 중
    31건이 DC 가입자였다.
    """

    def test_zero_payment_is_only_a_warning_for_dc(self, config) -> None:
        log = IssueLog()
        validate_retired([_retired(plan=BenefitPlan.DC, total_payment=0)], config, log)
        assert not log.has_errors()
        assert any(i.code == "TOI_TOTAL_MISSING_DC" for i in log.warnings)

    def test_zero_payment_is_still_an_error_for_db(self, config) -> None:
        log = IssueLog()
        validate_retired([_retired(plan=BenefitPlan.DB, total_payment=0)], config, log)
        assert any(i.code == "TOI_TOTAL_MISSING" for i in log.errors)


class TestUnknownPlanValue:
    """'퇴직금 예외' 처럼 해석할 수 없는 값이 실제로 들어온다.

    무엇이 적혀 있었는지 보여 주지 않으면 담당자가 고칠 수 없다.
    """

    def test_message_shows_what_was_written(self, config) -> None:
        log = IssueLog()
        validate_active([_active(plan=None, plan_raw="퇴직금 예외")], config, log)

        found = [i for i in log.errors if i.code == "JAE_PLAN_MISSING"]
        assert len(found) == 1
        assert "퇴직금 예외" in found[0].message
        assert found[0].value == "퇴직금 예외"

    def test_blank_says_blank(self, config) -> None:
        log = IssueLog()
        validate_active([_active(plan=None, plan_raw="")], config, log)
        assert "비어 있습니다" in log.errors[0].message


class TestPerGroupAssumptionToggles:
    """직군마다 쓰지 않는 가정이 있다.

    임원을 정년까지 근무한다고 보아 퇴직률을 빼거나, 호봉표가 없는 계약직에
    승급률을 주지 않거나, 임금이 계약으로 고정돼 Base-up 을 반영하지 않는 식이다.
    """

    def _assumptions(self):
        from pension.assumptions import (
            Assumptions,
            BenefitScale,
            DiscountCurve,
            RateCurve,
            RateTable,
            SalaryScale,
        )

        return Assumptions(
            discount=DiscountCurve(spot=RateCurve({1: 0.045}), flat=0.045),
            salary=SalaryScale(
                base_up=RateCurve({1: 0.03}),
                promotion=RateTable(curves={"정규직": RateCurve({20: 0.02})}),
            ),
            withdrawal=RateTable(curves={"정규직": RateCurve({20: 0.05})}),
            severance_benefit=BenefitScale(curves={"정규직": RateCurve({0: 1.0})}),
        )

    def _valued(self, config, **flags):
        from pension.valuation import value_member

        member = _active()
        member.age = 40
        member.severance_nra = 60
        for key, value in flags.items():
            setattr(member, key, value)
        return value_member(member, config, self._assumptions())

    def test_turning_off_withdrawal_changes_the_obligation(self, config) -> None:
        """퇴직률을 끄면 급여 지급이 전부 정년 시점으로 밀린다.

        퇴직금은 중도퇴직해도 그때 지급된다. 퇴직률이 있으면 일부가 이른 시점에
        빠져나가고 그만큼 할인을 덜 받아 현재가치가 커진다. 퇴직률을 끄면 전원이
        정년까지 남아 20년치 할인을 받으므로 채무가 **줄어든다**.
        """
        on = self._valued(config)
        off = self._valued(config, apply_withdrawal=False)

        assert off.dbo < on.dbo
        # 끈 쪽은 현금흐름이 정년 시점 한 곳에만 남는다.
        assert list(off.cash_flows) == [20.0]
        assert len(on.cash_flows) > 1

    def test_turning_off_pay_growth_lowers_the_obligation(self, config) -> None:
        on = self._valued(config)
        no_base_up = self._valued(config, apply_base_up=False)
        no_promotion = self._valued(config, apply_promotion=False)

        assert no_base_up.dbo < on.dbo
        assert no_promotion.dbo < on.dbo
        # 둘 다 끄면 임금이 그대로라 가장 작다.
        assert self._valued(config, apply_base_up=False, apply_promotion=False).dbo < min(
            no_base_up.dbo, no_promotion.dbo
        )

    def test_default_is_everything_applied(self, config) -> None:
        """이 칸이 없던 기존 파일과 동작이 같아야 한다."""
        member = _active()
        assert member.apply_base_up
        assert member.apply_promotion
        assert member.apply_withdrawal
        assert member.apply_mortality

    def test_sheet_column_drives_the_flag(self, tmp_path) -> None:
        import openpyxl

        from pension.config import PAYOUT_SHEET, read_payout_rules

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = PAYOUT_SHEET
        ws.append([f"c{i}" for i in range(1, 27)])
        ws.append(["임원", "임원", 60, 60, 2, *[""] * 8, 0, 0, 2, "", "일할", "그대로",
                   0, "반올림", "임원", "반영", "미반영", "미반영", "반영"])
        path = tmp_path / "기초율.xlsx"
        wb.save(path)

        rule = read_payout_rules(openpyxl.load_workbook(path, data_only=True))[0]
        assert rule.apply_base_up is True
        assert rule.apply_promotion is False
        assert rule.apply_withdrawal is False
        assert rule.apply_mortality is True

    def test_blank_column_means_applied(self, tmp_path) -> None:
        """이 열이 아예 없던 기존 파일도 그대로 읽혀야 한다."""
        import openpyxl

        from pension.config import PAYOUT_SHEET, read_payout_rules

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = PAYOUT_SHEET
        ws.append([f"c{i}" for i in range(1, 23)])
        ws.append(["정규직", "정규직", 60, 60, 2, *[""] * 8, 0, 0, 2, "", "일할",
                   "그대로", 0, "반올림", ""])
        path = tmp_path / "기초율.xlsx"
        wb.save(path)

        rule = read_payout_rules(openpyxl.load_workbook(path, data_only=True))[0]
        assert rule.apply_base_up
        assert rule.apply_withdrawal
