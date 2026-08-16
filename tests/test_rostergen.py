"""난수로 만드는 시험용 명부 세 가지.

시험 자료가 시험 대상보다 먼저 틀리면 곤란하다. 세 사례가 **의도한 성격대로**
나오는지를 여기서 못박는다 — 표준은 오류 없이 돌고, 복합제도는 특수 경로를
실제로 지나가고, 특이케이스는 문제지가 기초자료에 그대로 실려야 한다.
"""

from __future__ import annotations

import pytest

from pension.errors import PensionDataError, Severity
from pension.pipeline import RunOptions, run_valuation
from pension.rostertemplate import FIRST_DATA_ROW, HEADER_ROW
from pension.rostergen import CASES, write_case_pack


@pytest.fixture(scope="module")
def pack(tmp_path_factory) -> dict[str, tuple]:
    """세 사례를 한 번만 만들어 나눠 쓴다. 만드는 데 몇 초 걸린다."""
    directory = tmp_path_factory.mktemp("사례")
    write_case_pack(directory)
    return {
        spec.key: (
            directory / f"{spec.title}.xlsx",
            directory / f"{spec.title}_기초율.xlsx",
        )
        for spec in CASES
    }


def _run(pack, key, tmp_path, *, force=False):
    roster, assumptions = pack[key]
    return run_valuation(RunOptions(
        roster_path=roster, assumptions_path=assumptions,
        output_path=tmp_path / "결과.xlsx",
        include_sensitivity=False, include_longterm=True, allow_errors=force,
    ))


class TestStandardCase:
    def test_runs_without_a_single_error(self, pack, tmp_path) -> None:
        """기준선이 되는 사례다. 여기서 오류가 나면 시험 자체가 흔들린다."""
        run = _run(pack, "표준", tmp_path)
        assert run.issues.errors == []
        assert run.issues.warnings == []

    def test_headcount_and_shape(self, pack, tmp_path) -> None:
        run = _run(pack, "표준", tmp_path)
        valuation = run.valuation
        assert valuation.headcount == 290
        assert len(run.roster.retired) == 28
        assert set(valuation.by_job_group()) == {"정규직", "계약직", "임원"}
        # 채무·듀레이션이 상식적인 범위에 있어야 한다.
        assert 5e9 < valuation.dbo < 1e11
        assert 3 < valuation.duration < 15
        assert valuation.service_cost > 0

    def test_longterm_benefit_is_not_empty(self, pack, tmp_path) -> None:
        """장기급여 표가 비면 기타장기종업원급여를 눌러도 볼 것이 없다."""
        run = _run(pack, "표준", tmp_path)
        assert run.longterm is not None
        assert run.longterm.dbo > 0


class TestComplexCase:
    def test_runs_clean_despite_complex_structure(self, pack, tmp_path) -> None:
        """자료는 옳은 사례다 — 구조가 복잡하다고 오류가 나면 안 된다."""
        run = _run(pack, "복합제도", tmp_path)
        assert run.issues.errors == []

    def test_exercises_the_special_paths(self, pack, tmp_path) -> None:
        run = _run(pack, "복합제도", tmp_path)
        active = run.roster.active

        assert any(m.settlement_date for m in active), "중간정산자가 있어야 한다"
        assert any(m.wage_peak_age for m in active), "임금피크 대상이 있어야 한다"
        assert any(m.transfer_in_date for m in active), "전입자가 있어야 한다"
        assert any(m.payout_multiple for m in active), "개별 지급배수가 있어야 한다"
        assert any(m.age > 60 for m in active), "정년 초과 재고용자가 있어야 한다"
        # 임금피크는 정년연령을 실제로 당긴다 — 값이 흘러들어갔는지 본다.
        peak = next(m for m in active if m.wage_peak_age)
        assert peak.severance_nra == peak.wage_peak_age

        excluded = run.valuation.exclusion_summary()
        assert any("DC" in reason for reason in excluded), "DC 가입자가 빠져야 한다"

    def test_raw_job_groups_are_mapped(self, pack, tmp_path) -> None:
        """명부에는 '정규사원·촉탁사원' 이 오지만 산출은 묶음 단위로 나와야 한다."""
        run = _run(pack, "복합제도", tmp_path)
        assert set(run.valuation.by_job_group()) <= {"정규직", "계약직", "임원"}

    def test_retirement_reasons_are_varied(self, pack, tmp_path) -> None:
        run = _run(pack, "복합제도", tmp_path)
        reasons = {m.reason for m in run.roster.retired if m.reason}
        assert len(reasons) >= 4, "정년·사망·전출 등이 섞여 있어야 한다"


class TestSpecialsCase:
    """특이케이스 명부 — 명부는 깨끗하고, 특이사항 표가 문제지다.

    자료불량 명부를 대신한다. 클리닝은 실습하는 사람이 알아서 하므로 명부에
    사례별 장치를 심지 않는다 — 단체가 쓰듯 기본 항목만 적고, 사례 목록을
    [기초자료] 특이사항과 안내문에 싣는다.
    """

    def test_the_roster_itself_is_clean(self, pack, tmp_path) -> None:
        """문제지는 특이사항이지 자료 오류가 아니다 — 강행 없이 돌아야 한다."""
        run = _run(pack, "특이케이스", tmp_path)
        assert run.issues.errors == []
        assert run.valuation.headcount > 0

    def test_every_note_lands_on_the_basics_sheet(self, pack) -> None:
        import openpyxl

        from pension.rostergen import SPECIAL_NOTES

        roster, _ = pack["특이케이스"]
        book = openpyxl.load_workbook(roster)
        ws = book["기초자료"]
        text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row
                          if c.value)
        for title, detail in SPECIAL_NOTES:
            assert title in text, f"{title} 이 기초자료에 없다"
            assert detail[:20] in text, f"{title} 의 내용이 잘렸다"
        book.close()

    def test_the_report_repeats_the_question_sheet(self, pack) -> None:
        from pension.rostergen import SPECIAL_NOTES

        roster, _ = pack["특이케이스"]
        written = (roster.parent / "시험명부3_특이케이스_특이사항.txt").read_text(
            encoding="utf-8")
        for title, _detail in SPECIAL_NOTES:
            assert title in written

    def test_no_case_gadgets_are_planted_in_the_roster(self, pack) -> None:
        """명부는 단체가 적는 기본 항목만. 사례 장치(구간분할·누진 열)가 심겨
        있으면 '깨끗한 명부 + 문제지' 라는 약속이 깨진다."""
        import openpyxl

        from pension.rostertemplate import HEADER_ROW

        roster, _ = pack["특이케이스"]
        ws = openpyxl.load_workbook(roster)["재직자명부"]
        heads = {str(c.value) for c in ws[HEADER_ROW] if c.value}
        for gadget in ("누진적용 근속연수", "누진적용 율", "지급률 기산일",
                       "연봉제 전환 추계일"):
            assert gadget not in heads, f"'{gadget}' 열이 심겨 있다"


class TestGeneralSheet:
    """명부만 올려도 회계기간·신용등급·사외적립자산이 저절로 들어가야 한다.

    담당자가 채워 보내는 항목을 시험 명부에도 넣어 두지 않으면, '다시 적을
    필요가 없다' 는 것을 시험할 방법이 없다.
    """

    @pytest.mark.parametrize("key", [spec.key for spec in CASES])
    def test_the_period_and_grade_reach_the_run(self, pack, tmp_path, key) -> None:
        run = _run(pack, key, tmp_path, force=True)
        info = run.general_info
        assert info is not None
        assert info.period_end == run.config.base_date
        assert info.period_start is not None
        assert info.period_start < info.period_end
        assert info.credit_grade

    @pytest.mark.parametrize("key", [spec.key for spec in CASES])
    def test_the_asset_table_reconciles(self, pack, tmp_path, key) -> None:
        """시험 자료가 스스로 안 맞으면 시스템을 거짓으로 고발하게 된다."""
        run = _run(pack, key, tmp_path, force=True)
        assert round(run.general_info.assets.difference) == 0
        codes = [i.code for i in run.issues.warnings]
        assert "GEN_ASSET_NOT_BALANCED" not in codes

    @pytest.mark.parametrize("key", [spec.key for spec in CASES])
    def test_it_becomes_the_plan_asset_statement(self, pack, tmp_path, key) -> None:
        run = _run(pack, key, tmp_path, force=True)
        assets = run.plan_assets
        assert assets is not None
        assert assets.opening_fair_value > 0
        assert assets.closing_fair_value > 0
        assert assets.contributions > 0
        # 순확정급여부채 = 채무 − 자산.
        assert assets.net_liability == pytest.approx(
            run.valuation.dbo - assets.closing_fair_value, rel=1e-9
        )

    def test_the_obligation_table_matches_the_retiree_roster(
        self, pack, tmp_path
    ) -> None:
        """지급액을 아무 숫자로 넣으면 증감표가 어긋난다 — 명부에서 뽑아야 한다."""
        run = _run(pack, "복합제도", tmp_path, force=True)
        book = run.general_info.obligation
        from pension.normalize import RetirementReason

        converted = sum(
            m.total_payment for m in run.roster.retired
            if m.reason is RetirementReason.DC_CONVERSION
        )
        assert book.dc_converted == pytest.approx(converted, abs=1)
        assert book.benefits_paid > 0

    def test_the_breakdown_is_disclosed(self, pack, tmp_path) -> None:
        """문단 142 는 자산을 분류별로 공시하라고 한다."""
        run = _run(pack, "표준", tmp_path, force=True)
        breakdown = run.general_info.assets.breakdown
        assert breakdown
        assert "합계" not in breakdown
        assert sum(breakdown.values()) == pytest.approx(
            run.general_info.assets.closing, rel=1e-6
        )

    def test_the_report_says_what_was_filled_in(self, tmp_path) -> None:
        from pension.rostergen import write_case_roster

        report = tmp_path / "특이사항.txt"
        write_case_roster(CASES[0], tmp_path / "명부.xlsx", report_path=report)
        text = report.read_text(encoding="utf-8")
        assert "[기본정보]" in text
        assert "사외적립자산 변동내역" in text


def test_same_seed_gives_the_same_file(tmp_path) -> None:
    """산출 결과를 비교하려면 명부가 매번 같아야 한다."""
    first = write_case_pack(tmp_path / "a")[0]
    second = write_case_pack(tmp_path / "b")[0]
    third = write_case_pack(tmp_path / "c", seed=7)[0]

    import openpyxl

    def rows(path):
        ws = openpyxl.load_workbook(path)["재직자명부"]
        return [
            [ws.cell(r, c).value for c in range(2, 11)]
            for r in range(FIRST_DATA_ROW, FIRST_DATA_ROW + 16)
        ]

    assert rows(first) == rows(second)
    assert rows(first) != rows(third)


@pytest.fixture(scope="module")
def practice(tmp_path_factory):
    """practice 플래그로 만든 연습 명부.

    시험명부 3종에서는 빠졌지만(특이케이스 명부는 문제지만 싣는다), 규정을
    읽어야 풀리는 사례를 심는 장치 자체는 플래그로 계속 만들 수 있어야 한다.
    """
    from pension.rostergen import CaseSpec, write_case_assumptions, write_case_roster

    spec = CaseSpec(
        key="연습", title="연습명부", summary="practice 장치 시험용.",
        active=150, retired=20,
        groups=(("정규직", 0.75), ("계약직", 0.13), ("임원", 0.12)),
        flags={"practice": True, "dc_share": 0.06, "settlement_share": 0.05,
               "longterm_share": 0.35, "multiple_share": 0.03, "leave_share": 0.04},
    )
    directory = tmp_path_factory.mktemp("연습")
    report = directory / "연습명부_특이사항.txt"
    roster = write_case_roster(spec, directory / "연습명부.xlsx", report_path=report)
    assumptions = write_case_assumptions(spec, directory / "연습명부_기초율.xlsx")
    return {"연습": (roster, assumptions), "_report": report}


class TestPracticeCases:
    """practice 플래그가 심는 '규정을 읽어야 풀리는' 사례들.

    자료 오류와 성격이 다르다. 검증이 잡아 주는 것이 아니라, 담당자가 비고를
    읽고 산출에 반영해야 하는 것들이라 **실제로 산출 경로를 지나가는지** 를
    본다. 안내문이 짚어 준 사번이 명부에 없으면 시험 자료가 거짓말을 한다.
    """

    def _actives(self, practice):
        import openpyxl

        roster, _ = practice["연습"]
        wb = openpyxl.load_workbook(roster, data_only=True)
        ws = wb["재직자명부"]
        head = {c.value: c.column for c in ws[HEADER_ROW] if c.value}
        rows = [
            {name: ws.cell(r, col).value for name, col in head.items()}
            for r in range(FIRST_DATA_ROW, ws.max_row + 1)
            if ws.cell(r, head["사번"]).value
        ]
        wb.close()
        return rows, head

    def test_every_case_is_named_in_the_report(self, practice) -> None:
        """안내문이 사례마다 사번을 짚어 주고, 그 사번이 명부에 있어야 한다."""
        import re

        from pension.rostergen import PRACTICE_CASES

        roster, _ = practice["연습"]
        report = practice["_report"].read_text(encoding="utf-8")
        rows, _ = self._actives(practice)
        ids = {str(r["사번"]) for r in rows}

        for title, _detail in PRACTICE_CASES:
            assert title in report, f"안내문에 '{title}' 이(가) 없다"
        cited = set(re.findall(r"\(사번 ([AT]\d{4})\)", report))
        assert len(cited) >= 8, f"짚어 준 사번이 너무 적다: {cited}"
        # 재직자 사번은 실제로 명부에 있어야 한다(T… 는 퇴직자명부).
        assert {c for c in cited if c.startswith("A")} <= ids

    def test_extra_columns_are_written_and_read(self, practice, tmp_path) -> None:
        """누진·지급구간 열은 고정 서식 밖이라 머리글로 찾아 읽는 경로를 탄다."""
        rows, head = self._actives(practice)
        for label in ("지급률기산일", "지급률종료일", "누진적용근속연수", "누진적용율"):
            assert label in head, f"'{label}' 열이 만들어지지 않았다"

        run = _run(practice, "연습", tmp_path, force=True)
        progressive = [m for m in run.valuation.members if m.progressive_service]
        assert progressive, "누진 보전이 산출까지 닿지 않았다"
        assert progressive[0].progressive_rate > 1.0

    def test_period_split_is_not_a_duplicate_error(self, practice, tmp_path) -> None:
        """세법한도 동결은 같은 사번 두 줄이지만 오류가 아니라 기간 분할이다."""
        run = _run(practice, "연습", tmp_path, force=True)
        by_id: dict[str, int] = {}
        for m in run.valuation.members:
            by_id[m.employee_id] = by_id.get(m.employee_id, 0) + 1
        split = [emp for emp, n in by_id.items() if n == 2]
        assert split, "지급구간으로 나뉜 사번이 없다"

        flagged = {
            str(issue).split("사번=")[1].split(":")[0]
            for issue in run.issues.errors if issue.code == "JAE_DUP_ID"
        }
        assert not (set(split) & flagged), "기간 분할을 사번 중복 오류로 잡았다"

    def test_the_dc_leaver_appears_on_both_sheets(self, practice, tmp_path) -> None:
        """DC 전환 후 퇴직자는 재직·퇴직 양쪽에 같은 사번으로 있다."""
        run = _run(practice, "연습", tmp_path, force=True)
        active_ids = {m.employee_id for m in run.roster.active}
        retired_ids = {m.employee_id for m in run.roster.retired}
        assert active_ids & retired_ids, "양쪽에 걸친 사번이 없다"

    def test_practice_cases_reach_the_valuation(self, practice, tmp_path) -> None:
        """휴직차감·추가지급·개별배수가 실제 산출 결과에 나타나야 한다."""
        run = _run(practice, "연습", tmp_path, force=True)
        members = run.valuation.members
        assert any(m.extra_payment for m in members), "추가지급 기본급이 안 잡혔다"
        assert any(m.rounding_unit >= 0 for m in members)
        assert run.valuation.dbo > 0

    def test_standard_case_stays_clean(self, pack, tmp_path) -> None:
        """특이사항은 3번 명부에만 심는다 — 1번은 여전히 오류 0 이어야 한다."""
        run = _run(pack, "표준", tmp_path)
        assert not run.issues.errors
