"""웹앱 JSON API 와 가정 state 변환.

브라우저(Pyodide)에서 도는 코드지만 순수 파이썬이므로 PC 에서 그대로 검증한다.
여기서 깨지면 아이패드에서도 깨진다 — 반대로 여기서 통과하면 브라우저 쪽은
표 그리기만 남는다.
"""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from pension import assumption_form as form
from pension import clients
from pension.assumptions import DISCOUNT_SHEET, FORMULA, load_assumptions
from pension.webui import api


@pytest.fixture(autouse=True)
def home(tmp_path, monkeypatch):
    """등록 폴더를 시험용으로 갈아 끼운다. 사용자 폴더를 건드리면 안 된다."""
    monkeypatch.setenv("PENSION_HOME", str(tmp_path / "home"))
    # 마지막 산출 캐시(보고서·사번 조회가 쓴다)가 시험 사이에 새면 안 된다.
    from pension import webui
    webui._LAST_RUN.clear()
    return tmp_path


def call(op: str, **kwargs):
    """api() 를 부르고 성공 응답을 dict 로 돌려받는다."""
    response = json.loads(api(json.dumps({"op": op, **kwargs}, ensure_ascii=False)))
    assert response["ok"], response.get("error")
    return response


def call_error(op: str, **kwargs) -> str:
    response = json.loads(api(json.dumps({"op": op, **kwargs}, ensure_ascii=False)))
    assert not response["ok"]
    return response["error"]


def _kis_book(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "KIS_NET금리"
    ws.append(["No", "기준일자", "구분", "등급", "3월", "1년", "5년", "20년"])
    ws.append([1, None, "공모 무보증회사채", "AA0", 3.01, 3.122, 3.62, 5.26])
    ws.append([2, None, "공모 무보증회사채", "AA-", 3.11, 3.23, 3.72, 5.36])
    wb.save(path)
    return path


# ── 가정 state ───────────────────────────────────────────────────

class TestAssumptionForm:
    def test_round_trip_preserves_everything(self, tmp_path) -> None:
        """state → 워크북 → state 왕복에서 입력이 그대로 남아야 한다."""
        state = form.example_state(["생산직", "관리직", "임원"])
        state["payout"]["임원"]["withdrawal"] = "미반영"
        state["payout"]["임원"]["excluded"] = True
        state["benefit_rules"]["관리직"] = {"mode": "수식", "formula": "=MIN(t, 30)"}
        state["longterm_rules"]["생산직"] = {
            "kind": "현물", "escalation": "3%", "note": "현물 포상",
        }
        # 매핑이 있으면 매핑된 조합의 행만 나간다 — 묶음마다 최소 한 행씩 둔다.
        state["mapping"] = [
            ["사원", "정규사원", "생산직"], ["촉탁", "촉탁사원", "관리직"],
            ["대표", "임원(별정)", "임원"],
        ]

        path = form.write_state(state, tmp_path / "기초율.xlsx")
        back = form.read_state(path)

        assert back["job_groups"] == ["생산직", "관리직", "임원"]
        assert back["grids"][DISCOUNT_SHEET]["rows"] == [["1", "0.045"]]
        assert back["payout"]["임원"]["withdrawal"] == "미반영"
        assert back["payout"]["임원"]["excluded"] is True
        assert back["benefit_rules"]["관리직"]["mode"] == "수식"
        # 워크북에는 엑셀이 수식으로 오해하지 않게 '=' 없이 저장된다.
        assert back["benefit_rules"]["관리직"]["formula"] == "MIN(t, 30)"
        assert back["longterm_rules"]["생산직"]["kind"] == "현물"
        assert back["longterm_rules"]["생산직"]["escalation"] == "3%"
        assert back["mapping"] == state["mapping"]

    def test_every_setting_still_reaches_the_engine(self, tmp_path) -> None:
        """화면을 합쳐도 **넣을 수 있는 값은 하나도 줄지 않아야 한다.**

        탭을 묶고 표를 합치는 것은 보기 좋자는 일이다. 그 과정에서 칸이 하나라도
        사라지면 그 규정을 쓰는 회사는 산출을 못 하게 된다. 한 벌을 통째로 채워
        엔진까지 닿는지 본다.
        """
        from pension.assumptions import load_assumptions

        state = form.example_state(["정규직", "임원"])
        state["payout"]["임원"].update(
            withdrawal="미반영", excluded=False, min_service="3",
            nra="65", executive_nra="70", add_age="3",
            basis="월할", fraction="절사", unit="10원",
        )
        state["mapping"] = [["사원", "정규사원", "정규직"], ["대표", "등기임원", "임원"]]
        state["benefit_rules"]["임원"] = {"mode": "수식", "formula": "=t*배수"}

        # 지급률에 직군 아닌 열을 하나 더 만들고 퇴직사유가 그것을 가리킨다.
        state["grids"]["지급률"]["extra"] = ["사망가산"]
        state["grids"]["지급률"]["rows"] = [
            ["0", "1.0", "1.0", "3.0"], ["10", "10.0", "20.0", "5.0"],
        ]
        state["benefit_rules"]["사망가산"] = {"mode": "누적", "formula": ""}
        state["exit_causes"] = [
            ["정규직", "사망", "", "사망가산", "50000000", "1", "즉시"],
            ["정규직", "정년", "임원", "", "", "", "근속비례"],
        ]

        # 장기급여도 항목 열을 더 만들고 항목마다 다르게 준다.
        state["grids"]["장기급여지급률"]["extra"] = ["금"]
        state["grids"]["장기급여지급률"]["rows"] = [["10", "3", "3", "5000000"]]
        state["longterm_rules"]["정규직"].update(anniversary="10-01", every="5")
        state["longterm_rules"]["금"] = dict(
            form.default_longterm_rule("정규직"),
            kind="현물", escalation="3%", timing="퇴직시",
            accumulate=True, note="현물 포상",
        )

        loaded = load_assumptions(form.write_state(state, tmp_path / "기초율.xlsx"))

        # 지급률 — 직군 열과 따로 만든 열이 모두 살아 있다.
        assert loaded.severance_benefit.multiple("정규직", 12) == 10.0
        assert loaded.severance_benefit.multiple("사망가산", 12) == 5.0
        assert loaded.severance_benefit.mode("임원") == "수식"

        # 퇴직사유 — 두 줄 모두, 모든 칸이.
        death = loaded.exit_causes.get("정규직", "사망")
        assert death.extra_rule == "사망가산"
        assert death.extra_amount == 50_000_000
        assert death.min_service == 1.0
        assert death.attribution == "즉시"
        normal = loaded.exit_causes.get("정규직", "정년")
        assert normal.benefit_rule == "임원"
        assert normal.attribution == "근속비례"

        # 장기급여 — 직군 항목과 추가 항목이 모두.
        items = {i.column("정규직"): i for i in loaded.longterm_items("정규직")}
        assert items["정규직"].anniversary == "10-01"
        assert items["정규직"].every_years == 5
        assert items["금"].kind == "현물"
        assert items["금"].escalation == pytest.approx(0.03)
        assert items["금"].timing == "퇴직시"
        assert items["금"].accumulate is True
        assert "현물 포상" in items["금"].note

        # 지급규정 — 직군 규칙이 명부 조합마다 한 줄씩.
        from pension.pipeline import _read_payout_rules

        rules = {r.source_name: r for r in _read_payout_rules(
            tmp_path / "기초율.xlsx")}
        assert rules["대표"].severance_nra == 65
        assert rules["대표"].executive_nra == 70
        assert rules["대표"].min_service_years == 3
        assert rules["대표"].service_basis == "월할"
        assert rules["대표"].service_fraction == "절사"
        assert rules["대표"].benefit_rounding_unit == 10
        assert rules["대표"].apply_withdrawal is False
        assert rules["대표"].employee_type_filter == "등기임원"

    def test_exit_causes_survive_the_round_trip(self, tmp_path) -> None:
        """퇴직사유 차등이 왕복에서 사라지면 엉뚱한 급여로 산출된다."""
        state = form.example_state(["생산직", "임원"])
        state["exit_causes"] = [
            ["생산직", "사망", "", "", "50000000", "1", ""],
            ["임원", "정년", "생산직", "", "", "", "근속비례"],
        ]

        back = form.read_state(form.write_state(state, tmp_path / "기초율.xlsx"))
        assert back["exit_causes"] == [
            ["생산직", "사망", "", "", "50000000", "1", ""],
            ["임원", "정년", "생산직", "", "", "", "근속비례"],
        ]

    def test_a_row_with_only_a_cause_is_dropped(self, tmp_path) -> None:
        """규정과 사유만 고르고 값을 안 넣은 줄을 남기면 '차등 있음' 으로 읽힌다."""
        state = form.example_state(["생산직"])
        state["exit_causes"] = [["생산직", "사망", "", "", "", "", ""]]
        back = form.read_state(form.write_state(state, tmp_path / "기초율.xlsx"))
        assert back["exit_causes"] == []

    def test_unknown_cause_blocks_saving(self, tmp_path) -> None:
        state = form.example_state(["생산직"])
        state["exit_causes"] = [["생산직", "명예퇴직", "", "", "1000000", "", ""]]
        assert any("명예퇴직" in p for p in form.state_problems(state))

    def test_duplicate_cause_rows_block_saving(self) -> None:
        state = form.example_state(["생산직"])
        state["exit_causes"] = [
            ["생산직", "사망", "", "", "1000000", "", ""],
            ["생산직", "사망", "", "", "2000000", "", ""],
        ]
        assert any("두 번" in p for p in form.state_problems(state))

    def test_longterm_extra_columns_survive_the_round_trip(self, tmp_path) -> None:
        """항목 열이 왕복에서 사라지면 포상금·금이 통째로 빠진다."""
        state = form.example_state(["생산직"])
        state["longterm_rules"]["생산직"]["anniversary"] = "10-01"
        state["grids"]["장기급여지급률"]["extra"] = ["금"]
        state["grids"]["장기급여지급률"]["rows"] = [["10", "10", "5000000"]]
        state["longterm_rules"]["금"] = dict(
            form.default_longterm_rule("생산직"),
            kind="현물", escalation="3%", timing="퇴직시",
            accumulate=True, note="현물 포상",
        )

        back = form.read_state(form.write_state(state, tmp_path / "기초율.xlsx"))
        assert back["grids"]["장기급여지급률"]["extra"] == ["금"]
        assert back["grids"]["장기급여지급률"]["rows"] == [["10", "10", "5000000"]]
        assert back["longterm_rules"]["생산직"]["anniversary"] == "10-01"
        gold = back["longterm_rules"]["금"]
        assert gold["rule"] == "생산직"
        assert gold["kind"] == "현물"
        assert gold["escalation"] == "3%"
        assert gold["timing"] == "퇴직시"
        assert gold["accumulate"] is True
        assert gold["note"] == "현물 포상"

    def test_longterm_extra_columns_reach_the_engine(self, tmp_path) -> None:
        from pension.assumptions import load_assumptions

        state = form.example_state(["생산직"])
        state["grids"]["장기급여지급률"]["extra"] = ["금"]
        state["grids"]["장기급여지급률"]["rows"] = [["10", "10", "5000000"]]
        state["longterm_rules"]["금"] = dict(
            form.default_longterm_rule("생산직"),
            kind="현금", timing="퇴직시", every="5", accumulate=True,
        )
        path = form.write_state(state, tmp_path / "기초율.xlsx")

        items = load_assumptions(path).longterm_items("생산직")
        assert len(items) == 2
        assert items[1].item == "금"
        assert items[1].timing == "퇴직시"
        assert items[1].every_years == 5
        assert items[1].accumulate is True
        assert items[1].column("생산직") == "금"

    def test_an_unknown_longterm_timing_blocks_saving(self) -> None:
        state = form.example_state(["생산직"])
        state["longterm_rules"]["생산직"]["timing"] = "아무때나"
        assert any("아무때나" in p for p in form.state_problems(state))

    def test_an_extra_column_needs_a_job_group(self) -> None:
        """어느 직군의 급여인지 없으면 그 열이 아무에게도 안 걸린다."""
        state = form.example_state(["생산직"])
        state["grids"]["장기급여지급률"]["extra"] = ["금"]
        state["longterm_rules"]["금"] = form.default_longterm_rule("")
        assert any("어느 직군" in p for p in form.state_problems(state))

    def test_benefit_extra_columns_become_their_own_scale(self, tmp_path) -> None:
        """'사망가산' 처럼 직군이 아닌 지급률을 만들 수 있어야 한다.

        [퇴직사유] 의 가산 규정이 이 이름을 가리킨다. 직군 열만 있으면
        '사망 시 기본급 3개월분' 을 적을 자리가 아예 없다.
        """
        from pension.assumptions import load_assumptions

        state = form.example_state(["생산직"])
        state["grids"]["지급률"]["extra"] = ["사망가산"]
        state["grids"]["지급률"]["rows"] = [["0", "1.0", "3.0"], ["10", "10.0", "5.0"]]
        state["benefit_rules"]["사망가산"] = dict(form.default_benefit_rule(), mode="누적")
        state["exit_causes"] = [["생산직", "사망", "", "사망가산", "", "", ""]]

        loaded = load_assumptions(form.write_state(state, tmp_path / "기초율.xlsx"))
        assert loaded.severance_benefit.multiple("사망가산", 12) == 5.0
        assert loaded.exit_causes.get("생산직", "사망").extra_rule == "사망가산"

    def test_causes_reach_the_engine(self, tmp_path) -> None:
        """만든 워크북을 산출 엔진이 그대로 읽어야 한다."""
        from pension.assumptions import load_assumptions

        state = form.example_state(["생산직"])
        state["exit_causes"] = [["생산직", "사망", "", "", "50000000", "1", "즉시"]]
        path = form.write_state(state, tmp_path / "기초율.xlsx")

        entry = load_assumptions(path).exit_causes.get("생산직", "사망")
        assert entry.extra_amount == 50_000_000
        assert entry.min_service == 1.0
        assert entry.attribution == "즉시"

    def test_written_file_feeds_the_engine(self, tmp_path) -> None:
        """만든 워크북이 산출 엔진의 로더·규정 리더로 그대로 읽혀야 한다."""
        from pension.pipeline import _read_payout_rules

        state = form.example_state()
        state["payout"]["임원"]["mortality"] = "미반영"
        path = form.write_state(state, tmp_path / "기초율.xlsx")

        loaded = load_assumptions(path)
        assert loaded.discount.flat == pytest.approx(0.045)
        assert loaded.withdrawal.curve("정규직").rate(20) == pytest.approx(0.15)

        rules = {r.mapped_name: r for r in _read_payout_rules(path)}
        assert rules["임원"].apply_mortality is False
        assert rules["정규직"].apply_mortality is True

    def test_percent_and_comma_parsing(self) -> None:
        state = form.empty_state(["정규직"])
        state["grids"][DISCOUNT_SHEET]["rows"] = [["1", "4.5%"]]
        state["grids"]["지급률"]["rows"] = [["10", "1,000"]]
        sheets, _rules, _lt = form.state_to_sheets(state)
        assert sheets[DISCOUNT_SHEET][1] == [[1.0, 0.045]]
        assert sheets["지급률"][1] == [[10.0, 1000.0]]

    def test_problems_block_saving(self, tmp_path) -> None:
        state = form.empty_state()
        problems = form.state_problems(state)
        assert any("할인율" in p for p in problems)

        state["grids"][DISCOUNT_SHEET]["rows"] = [["1", "4.5%"]]
        state["benefit_rules"]["정규직"] = {"mode": FORMULA, "formula": ""}
        problems = form.state_problems(state)
        assert any("수식이 비어" in p for p in problems)
        with pytest.raises(ValueError):
            form.write_state(state, tmp_path / "기초율.xlsx")

    def test_mapping_rows_copy_target_rules(self, tmp_path) -> None:
        """매핑이 있으면 명부 조합마다 한 행씩, 규정 값은 묶음 것을 복사한다."""
        state = form.example_state(["정규직", "임원"])
        state["mapping"] = [
            ["사원", "정규사원", "정규직"], ["부장", "정규사원", "정규직"],
            ["대표", "임원", "임원"],
        ]
        path = form.write_state(state, tmp_path / "기초율.xlsx")
        ws = openpyxl.load_workbook(path)[form.PAYOUT_SHEET]
        rows = [
            [ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 22).value]
            for r in range(2, ws.max_row + 1)
        ]
        assert rows == [
            ["사원", "정규직", "정규사원"], ["부장", "정규직", "정규사원"],
            ["대표", "임원", "임원"],
        ]

    def test_note_lines_are_skipped_on_read(self, tmp_path) -> None:
        """양식 파일이 표 밑에 적는 설명 줄은 값으로 읽히면 안 된다."""
        state = form.example_state()
        path = form.write_state(state, tmp_path / "기초율.xlsx")
        wb = openpyxl.load_workbook(path)
        ws = wb[DISCOUNT_SHEET]
        ws.cell(ws.max_row + 2, 1, "· 결산일마다 새로 받아 바꾸세요.")
        wb.save(path)

        back = form.read_state(path)
        assert back["grids"][DISCOUNT_SHEET]["rows"] == [["1", "0.045"]]


# ── api() 경계 ───────────────────────────────────────────────────

class TestApi:
    def test_unknown_op(self) -> None:
        assert "모르는 요청" in call_error("없는것")

    def test_meta_lists_form_layout(self) -> None:
        meta = call("meta")
        assert [s["tab"] for s in meta["sheets"]] == [
            "할인율", "Base-up", "승급률", "퇴직률", "사망률", "지급률", "장기급여",
        ]
        assert meta["apply_choices"] == ["반영", "미반영"]
        assert "AA0" in meta["grades"]
        assert meta["exit_causes"] == ["중도", "사망", "정년"]
        assert meta["attributions"] == ["근속비례", "즉시"]
        assert len(meta["exit_cause_headers"]) == 7
        assert meta["longterm_timings"] == ["근속도달", "퇴직시", "정년시"]
        panels = {s["sheet"]: s["column_panel"] for s in meta["sheets"]}
        assert panels["지급률"] == "benefit"
        assert panels["장기급여지급률"] == "longterm"

    def test_editor_groups_place_every_sheet(self) -> None:
        """묶음에 빠진 표가 있으면 그 표가 화면에서 통째로 사라진다.

        규정이 늘 때마다 탭을 하나씩 붙이던 것을 네 묶음으로 접었다. 새 표를
        만들고 묶음에 넣는 것을 잊으면, 값을 넣을 자리가 없는 채로 산출된다.
        """
        meta = call("meta")
        placed = [
            section["sheet"]
            for group in meta["editor_groups"]
            for section in group["sections"]
            if section.get("sheet")
        ]
        assert sorted(placed) == sorted(s["sheet"] for s in meta["sheets"])
        assert len(placed) == len(set(placed)), "같은 표가 두 묶음에 들어갔습니다"

    def test_editor_groups_only_name_panels_the_screen_knows(self) -> None:
        """화면이 모르는 패널 이름을 넣으면 그 자리가 조용히 빈다."""
        known = {"map", "payout", "cause"}
        panels = [
            section["panel"]
            for group in call("meta")["editor_groups"]
            for section in group["sections"]
            if section.get("panel")
        ]
        assert set(panels) <= known
        assert sorted(panels) == sorted(known)

    def test_the_editor_fits_in_four_tabs(self) -> None:
        """아이패드에서 가로로 밀지 않고 다 보여야 한다."""
        groups = call("meta")["editor_groups"]
        assert [g["name"] for g in groups] == ["직군", "기초율", "퇴직급여", "장기급여"]

    def test_state_write_reports_problems_instead_of_raising(self, tmp_path) -> None:
        state = form.empty_state()
        result = call("state_write", state=state, path=str(tmp_path / "a.xlsx"))
        assert result["written"] is False
        assert result["problems"]

    def test_state_write_and_read(self, tmp_path) -> None:
        state = form.example_state()
        path = str(tmp_path / "기초율.xlsx")
        assert call("state_write", state=state, path=path)["written"] is True
        back = call("state_read", path=path)["state"]
        assert back["job_groups"] == ["정규직", "계약직", "임원"]

    def test_standard_state_has_full_table(self) -> None:
        """내장 표준률 — 15~70세 표가 직군 열로 펼쳐져야 한다."""
        state = call("standard_state", groups=["생산직", "임원"])["state"]
        assert state["job_groups"] == ["생산직", "임원"]
        mortality = state["grids"]["사망률"]["rows"]
        assert mortality[0][0] == "15" and mortality[-1][0] == "70"
        withdrawal = state["grids"]["퇴직률"]["rows"]
        assert len(withdrawal[0]) == 3  # 연령 + 직군 2
        assert state["grids"][DISCOUNT_SHEET]["rows"]  # 기본 할인율 한 줄

    def test_standard_state_follows_the_workplace_size(self) -> None:
        """승급률·퇴직률은 300인 미만/이상으로 갈린 원표를 그대로 가져온다."""
        small = call("standard_state", groups=["정규직"], size="300인 미만")
        large = call("standard_state", groups=["정규직"], size="300인 이상")
        assert small["size"] == "300인 미만" and large["size"] == "300인 이상"
        for sheet in ("퇴직률", "승급률"):
            assert (small["state"]["grids"][sheet]["rows"]
                    != large["state"]["grids"][sheet]["rows"])
        # 사망률은 규모로 갈리지 않는다.
        assert (small["state"]["grids"]["사망률"]["rows"]
                == large["state"]["grids"]["사망률"]["rows"])

    def test_unknown_size_falls_back_instead_of_failing(self) -> None:
        """화면이 빈 값을 보내도 산출 출발점은 나와야 한다."""
        blank = call("standard_state", groups=["정규직"], size="")
        assert blank["size"] == "300인 미만"
        assert blank["state"]["grids"]["퇴직률"]["rows"]

    def test_meta_offers_the_sizes(self) -> None:
        meta = call("meta")
        assert meta["standard_sizes"] == ["300인 미만", "300인 이상"]
        assert meta["standard_size_threshold"] == 300
        assert meta["standard_year"]

    def test_formula_ops(self) -> None:
        assert call("formula_check", source="=MIN(t, 30)")["error"] == ""
        assert call("formula_check", source="=WRONG(")["error"]
        preview = call("formula_preview", formulas={"임원": "=t*2"})["preview"]
        assert preview["columns"] == ["근속", "임원"]
        assert preview["rows"][0] == ["1", "2.000"]


class TestLibraryOps:
    def test_register_list_pin_remove(self, tmp_path) -> None:
        source = _kis_book(tmp_path / "금리.xlsx")
        result = call(
            "library_register", kind="금리표", path=str(source), name="KIS_2025",
        )
        assert result["registered"] == "KIS_2025"
        assert result["library"]["금리표"]["default"] == "KIS_2025"

        listing = call("library_list")["library"]
        assert [e["name"] for e in listing["금리표"]["entries"]] == ["KIS_2025"]

        call("library_pin", kind="금리표", name="KIS_2025")
        assert call("library_list")["library"]["금리표"]["pinned"] == "KIS_2025"

        call("library_remove", kind="금리표", name="KIS_2025")
        after = call("library_list")["library"]["금리표"]
        # 내장 금리표는 처음 한 번 심어 두므로 목록에 남는다. 등록해 지운 것만
        # 사라지면 된다.
        user_made = [e for e in after["entries"] if not e["name"].startswith("내장")]
        assert user_made == [] and after["pinned"] == ""

    def test_curve_rows_from_registered(self, tmp_path) -> None:
        source = _kis_book(tmp_path / "금리.xlsx")
        call("library_register", kind="금리표", path=str(source), name="KIS_2025")

        grades = call("curve_grades")["grades"]
        assert any("AA0" in g for g in grades)

        curve = call("curve_rows", grade="AA0")
        assert curve["rows"][0] == ["0.25", "3.0100%"]  # 3월 만기가 살아 있어야 한다
        assert curve["rows"][-1] == ["20", "5.2600%"]

    def test_rates_state_roundtrip(self, tmp_path) -> None:
        """등록한 표준률 워크북을 편집 state 로 불러온다."""
        state = form.example_state(["정규직", "임원"])
        path = form.write_state(state, tmp_path / "표준.xlsx")
        call("library_register", kind="표준률", path=str(path), name="회사표준")

        back = call("rates_state")["state"]
        assert back["job_groups"] == ["정규직", "임원"]

    def test_missing_registration_is_a_clear_error(self) -> None:
        assert "등록" in call_error("curve_rows", grade="AA0")


class TestRosterLibrary:
    """올렸던 명부를 목록에 두고 다음 결산에 그대로 꺼내 쓴다."""

    def test_register_and_run_from_library(self, tmp_path, roster_path) -> None:
        call("library_register", kind="명부", path=str(roster_path), name="1번단체 2412")
        listing = call("library_list")["library"]["명부"]
        assert [e["name"] for e in listing["entries"]] == ["1번단체 2412"]

        stored = call("library_path", kind="명부", name="1번단체 2412")["path"]
        state = form.example_state(["1정규직", "2임원", "3계약직"])
        assumptions = str(tmp_path / "기초율.xlsx")
        call("state_write", state=state, path=assumptions)

        result = call(
            "run", roster=stored, assumptions=assumptions, work=str(tmp_path),
            force=True, sensitivity=False, longterm=False,
        )
        assert result["run"] is True

    def test_old_xls_keeps_its_suffix(self, tmp_path) -> None:
        """.xls 를 .xlsx 로 이름만 바꿔 두면 여는 쪽이 서식을 잘못 짚는다."""
        source = tmp_path / "옛날명부.xls"
        source.write_bytes(b"\xd0\xcf\x11\xe0")  # BIFF 헤더 흉내
        call("library_register", kind="명부", path=str(source), name="옛날명부")
        stored = call("library_path", kind="명부", name="옛날명부")["path"]
        assert stored.endswith(".xls")

    def test_reregistering_replaces_rather_than_duplicates(self, tmp_path, roster_path) -> None:
        call("library_register", kind="명부", path=str(roster_path), name="같은이름")
        old = tmp_path / "같은이름.xls"
        old.write_bytes(b"\xd0\xcf\x11\xe0")
        call("library_register", kind="명부", path=str(old), name="같은이름")
        entries_after = call("library_list")["library"]["명부"]["entries"]
        assert [e["name"] for e in entries_after] == ["같은이름"]
        assert entries_after[0]["suffix"] == ".xls"


class TestDefaults:
    """빈 화면에서 시작해도 말이 되는 산출이 나와야 한다."""

    def test_base_up_is_two_percent_for_all_years(self) -> None:
        state = call("state_new")["state"]
        assert state["grids"]["임금상승률"]["rows"] == [["1", "2%"]]

    def test_benefit_defaults_to_statutory(self, tmp_path) -> None:
        """지급률 표가 비면 법정(배수 = 근속연수, 연속)이어야 한다."""
        state = call("state_new", groups=["정규직"])["state"]
        assert state["benefit_rules"]["정규직"]["mode"] == "법정"
        assert state["grids"]["지급률"]["rows"] == []

        state["grids"]["할인율"]["rows"] = [["1", "4.5%"]]
        path = str(tmp_path / "기초율.xlsx")
        assert call("state_write", state=state, path=path)["written"] is True

        loaded = load_assumptions(path)
        # 12.6년 근속이면 12.6 배 — 정수로 끊기면 안 된다.
        assert loaded.severance_benefit.multiple("정규직", 12.6) == pytest.approx(12.6)

    def test_filled_table_wins_over_statutory(self, tmp_path) -> None:
        """누진·수식을 쓰지 않아도 지급률 탭에 넣은 값이 그대로 적용돼야 한다."""
        state = call("state_new", groups=["정규직"])["state"]
        state["grids"]["할인율"]["rows"] = [["1", "4.5%"]]
        state["grids"]["지급률"]["rows"] = [["0", "1"], ["10", "12"], ["20", "26"]]
        path = str(tmp_path / "기초율.xlsx")
        call("state_write", state=state, path=path)

        scale = load_assumptions(path).severance_benefit
        # 법정이면 5·12.6·25 였을 자리. 표가 이겨서 보간값이 나와야 한다.
        assert scale.multiple("정규직", 5) == pytest.approx(1 + 11 * 5 / 10)
        assert scale.multiple("정규직", 12.6) == pytest.approx(12 + 14 * 2.6 / 10)
        assert scale.multiple("정규직", 25) == pytest.approx(26)

    def test_benefit_split_op_goes_both_ways(self) -> None:
        """화면 둘이 같은 함수를 부르도록 web API 로도 열어 둔다."""
        state = form.example_state(["정규직", "임원"])
        split = call("benefit_split", state=state, rule="정규직")
        assert split["split"] == ["정규직"]
        assert "정규직·정년" in form.grid_columns(split["state"], "지급률")

        merged = call("benefit_split", state=split["state"], rule="정규직", merge=True)
        assert merged["split"] == []
        assert "정규직·정년" not in form.grid_columns(merged["state"], "지급률")

    def test_standard_state_keeps_statutory_default(self) -> None:
        state = call("standard_state", groups=["정규직", "임원"])["state"]
        assert state["benefit_rules"]["정규직"]["mode"] == "법정"
        assert state["grids"]["지급률"]["rows"] == []
        assert state["grids"]["임금상승률"]["rows"] == [["1", "0.02"]]


class TestDates:
    """산출기준일·시작일 입력란."""

    def test_base_date_override(self, tmp_path, roster_path) -> None:
        """명부를 고치지 않고 기준일만 바꿔 산출할 수 있어야 한다."""
        state = form.example_state(["1정규직", "2임원", "3계약직"])
        assumptions = str(tmp_path / "기초율.xlsx")
        call("state_write", state=state, path=assumptions)

        base = call("run", roster=str(roster_path), assumptions=assumptions,
                    work=str(tmp_path), force=True, sensitivity=False, longterm=False)
        moved = call("run", roster=str(roster_path), assumptions=assumptions,
                     work=str(tmp_path / "b"), force=True, sensitivity=False,
                     longterm=False, base_date="2026-12-31")

        assert base["values"]["base_date"] == "2025-12-31"
        assert moved["values"]["base_date"] == "2026-12-31"
        # 한 해 더 근무했으므로 채무가 늘어야 한다.
        assert moved["values"]["dbo"] > base["values"]["dbo"]

    def test_period_start_scales_interest(self, tmp_path, roster_path) -> None:
        """기간이 1년이 아니면 이자원가도 그 기간만큼이어야 한다."""
        state = form.example_state(["1정규직", "2임원", "3계약직"])
        assumptions = str(tmp_path / "기초율.xlsx")
        call("state_write", state=state, path=assumptions)

        common = dict(roster=str(roster_path), assumptions=assumptions, force=True,
                      sensitivity=False, longterm=False, prior_dbo=10_000_000_000,
                      prior_rate=0.045)
        full = call("run", work=str(tmp_path / "1"), **common)
        half = call("run", work=str(tmp_path / "2"),
                    base_date="2025-12-31", period_start="2025-07-01", **common)

        def gain(report):
            return next(v for label, v in report["summary"] if label == "보험수리적손익")

        # 이자원가가 줄면 기대 기말채무가 줄고, 그만큼 손익이 커진다.
        assert gain(half) != gain(full)

    def test_bad_date_is_reported(self, tmp_path, roster_path) -> None:
        assert "날짜를 읽지 못했습니다" in call_error(
            "run", roster=str(roster_path), assumptions=str(roster_path),
            work=str(tmp_path), base_date="어제",
        )


class TestPresets:
    """지급률 등 가정 한 벌을 기본가정에 저장해 두고 다시 쓴다."""

    def test_save_and_load(self, tmp_path) -> None:
        state = call("state_new", groups=["정규직", "임원"])["state"]
        state["grids"]["할인율"]["rows"] = [["1", "4.2%"]]
        state["grids"]["지급률"]["rows"] = [["0", "1", "1"], ["10", "12", "15"]]
        state["payout"]["임원"]["withdrawal"] = "미반영"
        state["mapping"] = [["사원", "정규사원", "정규직"], ["대표", "임원", "임원"]]

        saved = call("preset_save", name="A사 퇴직금규정", state=state)
        assert saved["saved"] is True
        assert [e["name"] for e in saved["library"]["가정세트"]["entries"]] == ["A사 퇴직금규정"]

        back = call("preset_state", name="A사 퇴직금규정")["state"]
        assert back["grids"]["지급률"]["rows"] == [["0", "1", "1"], ["10", "12", "15"]]
        assert back["payout"]["임원"]["withdrawal"] == "미반영"
        assert back["mapping"] == state["mapping"]

    def test_bad_state_is_refused(self) -> None:
        broken = call("state_new", groups=["정규직"])["state"]
        broken["benefit_rules"]["정규직"] = {"mode": "수식", "formula": ""}
        result = call("preset_save", name="틀린것", state=broken)
        assert result["saved"] is False and result["problems"]

    def test_name_is_required(self) -> None:
        state = call("state_new")["state"]
        assert "이름" in call_error("preset_save", name="  ", state=state)


class TestGenerator:
    """시험용 난수 명부를 시스템 안에서 만든다."""

    def test_makes_pack_with_reports(self, tmp_path) -> None:
        result = call("gen_cases", work=str(tmp_path), seed=42)
        assert Path(result["path"]).exists()
        assert len(result["files"]) == 9        # 사례 3종 × (명부·기초율·안내문)
        assert len(result["cases"]) == 3

        dirty = next(c for c in result["cases"] if c["force"])
        assert "일부러 심어 둔 자료 오류" in dirty["report"]
        clean = next(c for c in result["cases"] if not c["force"])
        assert "산출 특이사항" in clean["report"]
        assert Path(clean["roster"]).exists() and Path(clean["assumptions"]).exists()

    def test_generated_case_runs_end_to_end(self, tmp_path) -> None:
        made = call("gen_cases", work=str(tmp_path), seed=42)
        case = next(c for c in made["cases"] if not c["force"])
        report = call("run", roster=case["roster"], assumptions=case["assumptions"],
                      work=str(tmp_path), sensitivity=False, longterm=True)
        assert report["run"] is True
        assert report["values"]["dbo"] > 0

    def test_register_puts_them_in_the_lists(self, tmp_path) -> None:
        call("gen_cases", work=str(tmp_path), seed=42)
        result = call("gen_case_register", work=str(tmp_path))
        assert len(result["registered"]) == 3
        assert len(result["library"]["명부"]["entries"]) == 3
        assert len(result["library"]["가정세트"]["entries"]) == 3

    def test_register_without_generating_says_so(self, tmp_path) -> None:
        assert "먼저" in call_error("gen_case_register", work=str(tmp_path / "빈곳"))


class TestBackup:
    """브라우저 저장소는 지워질 수 있다 — 기기 밖으로 내보내고 되살린다."""

    def test_export_and_restore_everything(self, tmp_path, roster_path, monkeypatch) -> None:
        curve = _kis_book(tmp_path / "금리.xlsx")
        call("library_register", kind="금리표", path=str(curve), name="KIS_2025")
        call("library_register", kind="명부", path=str(roster_path), name="1번단체")
        state = form.example_state(["1정규직", "2임원", "3계약직"])
        assumptions = str(tmp_path / "기초율.xlsx")
        call("state_write", state=state, path=assumptions)
        report = call(
            "run", roster=str(roster_path), assumptions=assumptions,
            work=str(tmp_path), force=True, sensitivity=False, longterm=False,
        )
        call("run_save", name="2412 1번단체", roster=str(roster_path),
             assumptions=assumptions, work=str(tmp_path), report=report)

        exported = call("backup_export", work=str(tmp_path))
        archive = Path(exported["path"])
        assert archive.exists() and exported["filename"].startswith("연금계리보관함_")
        assert call("library_list")["backup"] == exported["stamp"]

        # 저장소를 통째로 잃은 기기(=새 PENSION_HOME)에서 되살린다.
        monkeypatch.setenv("PENSION_HOME", str(tmp_path / "새기기"))
        fresh = call("library_list")["library"]["금리표"]["entries"]
        assert [e["name"] for e in fresh if not e["name"].startswith("내장")] == []

        restored = call("backup_import", path=str(archive))
        names = [e["name"] for e in restored["library"]["금리표"]["entries"]]
        assert "KIS_2025" in names
        assert [e["name"] for e in restored["library"]["명부"]["entries"]] == ["1번단체"]
        assert [r["name"] for r in restored["runs"]] == ["2412 1번단체"]
        # 되살린 산출의 입력이 실제로 열려야 한다.
        again = call("run_restore", name="2412 1번단체", work=str(tmp_path / "w"))
        assert Path(again["roster"]).exists()

    def test_merge_keeps_what_the_archive_does_not_have(self, tmp_path, monkeypatch) -> None:
        call("library_register", kind="금리표", path=str(_kis_book(tmp_path / "a.xlsx")),
             name="예전금리표")
        archive = Path(call("backup_export", work=str(tmp_path))["path"])

        monkeypatch.setenv("PENSION_HOME", str(tmp_path / "다른기기"))
        call("library_register", kind="금리표", path=str(_kis_book(tmp_path / "b.xlsx")),
             name="이기기금리표")
        merged = call("backup_import", path=str(archive))
        assert {e["name"] for e in merged["library"]["금리표"]["entries"]} == {
            "예전금리표", "이기기금리표",
        }

        replaced = call("backup_import", path=str(archive), replace=True)
        assert [e["name"] for e in replaced["library"]["금리표"]["entries"]] == ["예전금리표"]

    def test_rejects_a_random_zip(self, tmp_path) -> None:
        import zipfile

        bogus = tmp_path / "아무거나.zip"
        with zipfile.ZipFile(bogus, "w") as archive:
            archive.writestr("hello.txt", "hi")
        assert "보관함 파일이 아닙니다" in call_error("backup_import", path=str(bogus))


class TestRosterOps:
    def test_scan_and_groups(self, roster_path) -> None:
        found = call("roster_scan", path=str(roster_path), groups=[])["found"]
        assert found, "명부에서 직군 조합을 찾아야 한다"
        pairs = {(f["source"], f["kind"]) for f in found}
        assert ("정규직", "직원") in pairs
        executive = next(f for f in found if f["kind"] == "임원")
        assert executive["suggest"] == "임원"

        result = call("roster_groups", path=str(roster_path))
        # 직군 축 — [기본정보] 의 산출 직군 그대로.
        assert result["groups"] == ["2임원", "1정규직", "3계약직"]
        assert result["scale_rules"] == []

    def _named_rule_column(self, tmp_path, header: str, values: list[str]) -> Path:
        """규정 칸의 머리글을 ``header`` 로 바꾼 재직자명부."""
        from pension.layout import ACTIVE_HEADER_ALIASES, find_header_row
        from pension.rostertemplate import write_roster_template

        path = write_roster_template(tmp_path / "명부.xlsx")
        book = openpyxl.load_workbook(path)
        sheet = book["재직자명부"]
        row = find_header_row(sheet, ACTIVE_HEADER_ALIASES)
        column = next(
            c for c in range(1, sheet.max_column + 1)
            if str(sheet.cell(row, c).value or "").strip() == "규정명"
        )
        sheet.cell(row, column, header)
        for offset, value in enumerate(values):
            sheet.cell(row + 1 + offset, column, value)
        book.save(path)
        return path

    def test_an_unusually_named_rule_column_is_still_found(self, tmp_path) -> None:
        """규정 칸 머리글은 회사마다 다르다 — '지급규정' 이라고 적어 보낸다.

        못 찾으면 조용히 직군으로 물러서고, 사람마다 다른 규정이 통째로
        뭉개진 채 산출이 끝난다. 오류 없이 그럴듯한 숫자가 나오는 쪽이라
        머리글을 넉넉히 알아봐야 한다.

        직군과 지급규정은 별개의 축이다 — 규정명은 ``scale_rules`` 로 오고,
        ``groups`` (직군 축)에는 섞이지 않는다.
        """
        path = self._named_rule_column(tmp_path, "지급규정", ["임원규정", "직원규정"])

        result = call("roster_groups", path=str(path))
        assert result["scale_rules"] == ["임원규정", "직원규정"]
        assert "임원규정" not in result["groups"]

    def test_the_scan_says_which_column_it_read(self, tmp_path) -> None:
        """어느 칸을 읽었는지 말해 주지 않으면 잘못 읽힌 것을 알 길이 없다."""
        path = self._named_rule_column(tmp_path, "퇴직금규정", ["갑규정"])

        result = call("roster_groups", path=str(path))
        assert result["rule_header"] == "퇴직금규정"
        assert result["rule_column"]                 # 엑셀 열 문자 (예: "J")
        assert any("성명" in line for line in result["headers"])

    def test_a_roster_without_a_rule_column_hands_back_its_headers(
        self, roster_path
    ) -> None:
        """못 찾았으면 명부에 적혀 온 머리글을 돌려줘 어느 칸인지 짚게 한다."""
        result = call("roster_groups", path=str(roster_path))
        assert result["rules"] == []
        assert result["rule_column"] == ""

    def test_without_basics_the_groups_come_from_the_roster_column(
        self, tmp_path
    ) -> None:
        """[기본정보] 직군 규칙이 없으면 명부 직군 열에 적혀 온 값이 직군이다."""
        from pension.rostertemplate import write_roster_template

        path = write_roster_template(tmp_path / "명부.xlsx")
        book = openpyxl.load_workbook(path)
        del book["기본정보"]
        book.save(path)

        result = call("roster_groups", path=str(path))
        # 양식 작성 예시 두 줄의 직군: 정규직·임원.
        assert result["groups"] == ["정규직", "임원"]


class TestRunOp:
    def test_full_run(self, tmp_path, roster_path) -> None:
        """편집 state 로 만든 기초율 + 명부 → 산출까지 한 바퀴."""
        state = form.example_state(["1정규직", "2임원", "3계약직"])
        assumptions = str(tmp_path / "기초율.xlsx")
        assert call("state_write", state=state, path=assumptions)["written"] is True

        result = call(
            "run", roster=str(roster_path), assumptions=assumptions,
            work=str(tmp_path), force=True, sensitivity=False, longterm=False,
        )
        assert result["run"] is True
        labels = [row[0] for row in result["summary"]]
        assert "확정급여채무 (DBO)" in labels
        assert (tmp_path / "산출결과.xlsx").exists()
        assert (tmp_path / "개인별결과.xlsx").exists()


class TestRunHistory:
    """산출 내역 — 이름 붙여 보관하고 다음 결산 때 끌어온다."""

    def _saved_run(self, tmp_path, roster_path) -> dict:
        state = form.example_state(["1정규직", "2임원", "3계약직"])
        assumptions = str(tmp_path / "기초율.xlsx")
        call("state_write", state=state, path=assumptions)
        report = call(
            "run", roster=str(roster_path), assumptions=assumptions,
            work=str(tmp_path), force=True, sensitivity=False, longterm=False,
        )
        return call(
            "run_save", name="2412 1번단체", roster=str(roster_path),
            assumptions=assumptions, work=str(tmp_path),
            report=report, options={"force": True},
        )

    def test_save_list_restore(self, tmp_path, roster_path) -> None:
        listing = self._saved_run(tmp_path, roster_path)["runs"]
        assert [r["name"] for r in listing] == ["2412 1번단체"]
        assert listing[0]["dbo"]  # 요약에서 DBO 를 뽑아 보여 준다
        assert listing[0]["has_results"] is True

        restored = call("run_restore", name="2412 1번단체", work=str(tmp_path / "w"))
        assert Path(restored["roster"]).exists()
        assert Path(restored["assumptions"]).exists()
        assert restored["meta"]["options"] == {"force": True}

        results = call("run_results", name="2412 1번단체", work=str(tmp_path / "w"))
        assert "산출결과.xlsx" in results["files"]

    def test_same_name_overwrites(self, tmp_path, roster_path) -> None:
        self._saved_run(tmp_path, roster_path)
        listing = self._saved_run(tmp_path, roster_path)["runs"]
        assert len(listing) == 1

    def test_delete(self, tmp_path, roster_path) -> None:
        self._saved_run(tmp_path, roster_path)
        assert call("run_delete", name="2412 1번단체")["runs"] == []
        assert "없습니다" in call_error("run_restore", name="2412 1번단체")

    def test_prior_link_carries_numbers_and_assumptions(self, tmp_path, roster_path) -> None:
        """저장된 산출을 전기로 끌어오면 DBO·할인율이 서식 없이 그대로 와야 한다."""
        self._saved_run(tmp_path, roster_path)
        prior = call("run_prior", name="2412 1번단체")

        assert prior["values"]["dbo"] > 0
        assert 0 < prior["values"]["discount_rate"] < 1  # 4.5% → 0.045
        assert Path(prior["assumptions"]).exists()

        # 그 값으로 당기를 돌리면 증감분석이 붙는다.
        current = call(
            "run", roster=str(roster_path), assumptions=prior["assumptions"],
            work=str(tmp_path / "당기"), force=True, sensitivity=False, longterm=False,
            prior_dbo=prior["values"]["dbo"],
            prior_rate=prior["values"]["discount_rate"],
            prior_service_cost=prior["values"]["service_cost"],
            prior_assumptions=prior["assumptions"],
        )
        assert current["run"] is True
        assert "보험수리적손익" in [row[0] for row in current["summary"]]

    def test_prior_link_falls_back_to_summary_text(self, tmp_path, roster_path) -> None:
        """숫자를 남기지 않던 예전 저장본도 요약 문자열에서 되짚어야 한다."""
        import json as _json

        self._saved_run(tmp_path, roster_path)
        folder = clients.folder() / "2412 1번단체"
        meta = _json.loads((folder / "meta.json").read_text(encoding="utf-8"))
        meta["report"].pop("values")
        (folder / "meta.json").write_text(
            _json.dumps(meta, ensure_ascii=False), encoding="utf-8"
        )

        prior = call("run_prior", name="2412 1번단체")
        assert prior["values"]["dbo"] > 0
        assert prior["values"]["discount_rate"] == pytest.approx(0.045, abs=1e-4)

    def test_name_is_required_and_sanitized(self, tmp_path, roster_path) -> None:
        assert "산출명" in call_error(
            "run_save", name="  ", roster=str(roster_path),
            assumptions=str(roster_path),
        )
        # 경로 문자가 든 이름도 폴더 이름으로 안전해야 한다.
        state = form.example_state()
        assumptions = str(tmp_path / "a.xlsx")
        call("state_write", state=state, path=assumptions)
        saved = call(
            "run_save", name="24/12: 1번*단체", roster=str(roster_path),
            assumptions=assumptions, work=str(tmp_path), report={},
        )
        assert saved["runs"][0]["name"] == "24 12  1번 단체"


class TestClients:
    """단체 — 앱에서 가장 먼저 고르는 것. 산출 내역이 이 안에 쌓인다."""

    @pytest.fixture(autouse=True)
    def _boot(self):
        """앱을 열면 제일 먼저 목록을 읽는다 — 기본 단체는 그때 생긴다."""
        call("client_list")

    def _saved_run(self, tmp_path, roster_path, name: str, **extra) -> dict:
        state = form.example_state(["1정규직", "2임원", "3계약직"])
        assumptions = str(tmp_path / f"{name}_기초율.xlsx")
        call("state_write", state=state, path=assumptions)
        return call(
            "run_save", name=name, roster=str(roster_path),
            assumptions=assumptions, work=str(tmp_path), report={}, **extra,
        )

    def test_starts_with_one_default_client(self) -> None:
        listing = call("client_list")
        assert listing["client"] == clients.DEFAULT_CLIENT
        assert [c["name"] for c in listing["clients"]] == [clients.DEFAULT_CLIENT]

    def test_add_switches_to_the_new_client(self) -> None:
        assert call("client_add", name="1번단체")["client"] == "1번단체"
        assert call("client_add", name="2번단체")["client"] == "2번단체"
        assert {c["name"] for c in call("client_list")["clients"]} == {
            clients.DEFAULT_CLIENT, "1번단체", "2번단체"}

    def test_runs_do_not_leak_between_clients(self, tmp_path, roster_path) -> None:
        """다른 단체의 산출이 목록에 섞이면 전기 DBO 를 남의 것으로 끌어온다."""
        call("client_add", name="가단체")
        self._saved_run(tmp_path, roster_path, "2312")
        call("client_add", name="나단체")
        self._saved_run(tmp_path, roster_path, "2412")

        assert [r["name"] for r in call("run_list")["runs"]] == ["2412"]
        assert [r["name"] for r in call("run_list", client="가단체")["runs"]] == ["2312"]

        # 이름이 같아도 단체가 다르면 다른 산출이다.
        assert "없습니다" in call_error("run_restore", name="2312")
        assert call("run_restore", name="2312", client="가단체",
                    work=str(tmp_path / "w"))["meta"]["name"] == "2312"

    def test_client_count_follows_its_runs(self, tmp_path, roster_path) -> None:
        call("client_add", name="가단체")
        listing = self._saved_run(tmp_path, roster_path, "2412")
        here = next(c for c in listing["clients"] if c["name"] == "가단체")
        assert here["runs"] == 1

    def test_rename_keeps_the_runs(self, tmp_path, roster_path) -> None:
        call("client_add", name="가단체")
        self._saved_run(tmp_path, roster_path, "2412")
        renamed = call("client_rename", name="가단체", new_name="가나다㈜")

        assert renamed["client"] == "가나다㈜"
        assert [r["name"] for r in renamed["runs"]] == ["2412"]
        assert "가단체" not in {c["name"] for c in renamed["clients"]}

    def test_rename_onto_an_existing_name_is_refused(self) -> None:
        call("client_add", name="가단체")
        call("client_add", name="나단체")
        assert "이미 있습니다" in call_error(
            "client_rename", name="가단체", new_name="나단체")

    def test_remove_needs_force_when_runs_remain(self, tmp_path, roster_path) -> None:
        call("client_add", name="가단체")
        self._saved_run(tmp_path, roster_path, "2412")

        assert "1건" in call_error("client_remove", name="가단체")
        left = call("client_remove", name="가단체", force=True)
        assert "가단체" not in {c["name"] for c in left["clients"]}
        assert left["client"] == clients.DEFAULT_CLIENT

    def test_last_client_cannot_be_removed(self) -> None:
        assert "마지막" in call_error("client_remove", name=clients.DEFAULT_CLIENT)

    def test_name_is_required_and_sanitized(self) -> None:
        assert "단체명" in call_error("client_add", name="   ")
        assert call("client_add", name="가/나:다")["client"] == "가 나 다"

    def test_member_detail_from_saved_run(self, tmp_path, roster_path) -> None:
        """저장된 산출에서도 사번 조회 — 그 산출이 쓴 명부·기초율 그대로."""
        call("client_add", name="가단체")
        self._saved_run(tmp_path, roster_path, "2412")
        detail = call("member_detail", name="2412", employee_id="A001")
        assert detail["rows"][0]["profile"]["성명"] == "김철수"
        assert detail["rows"][0]["result"]["확정급여채무 (DBO)"] > 0
        # 다른 단체에서는 그 산출이 안 보이므로 조회도 안 된다.
        call("client_add", name="나단체")
        assert "없습니다" in call_error("member_detail", name="2412", employee_id="A001")

    def test_legacy_runs_move_into_the_default_client(self, tmp_path, roster_path) -> None:
        """단체를 쓰기 전에 저장한 산출도 그대로 보여야 한다."""
        import shutil

        self._saved_run(tmp_path, roster_path, "2412")
        # 예전 판이 두던 자리(산출내역 바로 아래)로 되돌려 놓는다.
        old = clients.root() / "2412"
        shutil.move(str(clients.folder() / "2412"), str(old))
        shutil.rmtree(clients.root() / clients.DEFAULT_CLIENT)

        listing = call("run_list")
        assert listing["client"] == clients.DEFAULT_CLIENT
        assert [r["name"] for r in listing["runs"]] == ["2412"]
        assert not old.exists()


class TestMemberDetailAndReport:
    """산출 직후의 사번 조회와 계리평가 보고서 — 둘 다 마지막 산출을 쓴다."""

    def _run(self, tmp_path, roster_path, **extra) -> dict:
        state = form.example_state(["1정규직", "2임원", "3계약직"])
        assumptions = str(tmp_path / "기초율.xlsx")
        call("state_write", state=state, path=assumptions)
        return call(
            "run", roster=str(roster_path), assumptions=assumptions,
            work=str(tmp_path), force=True, sensitivity=False, longterm=True,
            **extra,
        )

    def test_member_detail_uses_the_last_run_inputs(self, tmp_path, roster_path) -> None:
        """산출 뒤에는 경로를 다시 주지 않아도 그 명부·기초율로 조회된다."""
        self._run(tmp_path, roster_path)
        detail = call("member_detail", employee_id="A001")
        row = detail["rows"][0]
        assert row["profile"]["성명"] == "김철수"
        assert sum(t["dbo"] for t in row["trace"]) == pytest.approx(
            row["result"]["확정급여채무 (DBO)"])

    def test_member_detail_before_any_run_is_refused(self) -> None:
        assert "먼저 산출을" in call_error("member_detail", employee_id="A001")

    def test_report_html_is_written_and_complete(self, tmp_path, roster_path) -> None:
        self._run(tmp_path, roster_path, prior_dbo=300_000_000, prior_rate=0.04)
        made = call("report_html", kind="severance", client="가나다㈜",
                    work=str(tmp_path))
        page = Path(made["path"]).read_text(encoding="utf-8")
        assert "퇴직급여 확정급여부채 평가보고서" in page
        assert "가나다㈜" in page
        assert made["longterm_available"] is True

        lt = call("report_html", kind="longterm", work=str(tmp_path))
        assert "장기종업원급여부채" in Path(lt["path"]).read_text(encoding="utf-8")

    def test_report_before_any_run_is_refused(self) -> None:
        assert "먼저 산출을" in call_error("report_html", kind="severance")

    def test_dashboard_follows_the_last_run(self, tmp_path, roster_path) -> None:
        """분석 화면은 방금 산출한 회차를 본다 — 다시 산출하면 값도 바뀐다."""
        first = self._run(tmp_path, roster_path)
        board = call("dashboard")
        assert board["totals"]["dbo"] == pytest.approx(
            float(first["values"]["dbo"]))
        assert board["scenarios"][0]["trace"]

        # 기준일을 당겨 다시 산출하면 분석 값도 그 회차로 갈린다.
        again = self._run(tmp_path, roster_path, base_date="2025-06-30")
        moved = call("dashboard")
        assert moved["base_date"] == "2025-06-30"
        assert moved["totals"]["dbo"] == pytest.approx(float(again["values"]["dbo"]))
        assert moved["totals"]["dbo"] != board["totals"]["dbo"]

    def test_dashboard_before_any_run_is_refused(self) -> None:
        assert "먼저 산출을" in call_error("dashboard")


class TestPlanAssetsAndAmendment:
    """사외적립자산과 제도개정 — 재무제표에 바로 들어가는 숫자들."""

    def _pack(self, tmp_path):
        made = call("gen_cases", work=str(tmp_path), seed=4242)
        return next(c for c in made["cases"] if not c["force"])

    def test_plan_assets_roll_and_net_liability(self, tmp_path) -> None:
        case = self._pack(tmp_path)
        report = call(
            "run", roster=case["roster"], assumptions=case["assumptions"],
            work=str(tmp_path), sensitivity=False, longterm=False,
            asset_opening=15_000_000_000, asset_contributions=1_500_000_000,
            asset_paid=900_000_000, asset_closing=17_200_000_000,
            prior_rate=0.045,
        )
        assert report["run"] is True
        rows = dict(report["assets"])
        assert rows["기초 사외적립자산 공정가치"] == 15_000_000_000
        assert rows["기말 사외적립자산 공정가치"] == 17_200_000_000
        # 기초 + 부담금 - 지급 + 이자 + 재측정 = 기말
        assert (
            rows["기초 사외적립자산 공정가치"] + rows["이자수익"]
            + rows["부담금 납입액"] + rows["급여지급액"] + rows["자산 재측정손익"]
        ) == pytest.approx(rows["기말 사외적립자산 공정가치"])
        # 순확정급여부채 = 채무 - 자산
        assert rows["순확정급여부채"] == pytest.approx(
            report["values"]["dbo"] - 17_200_000_000
        )

    def test_general_info_hands_the_screen_its_input_values(self, tmp_path) -> None:
        """산출 화면의 칸에 그대로 넣을 수 있는 형태로 나와야 한다.

        엔진이 알아서 읽는 것만으로는 부족하다. 칸이 비어 있으면 담당자는
        아무것도 읽히지 않은 줄 알고 신탁 명세서를 보고 손으로 다시 적는다.
        """
        case = self._pack(tmp_path)
        info = call("general_info", path=case["roster"])

        fields = info["fields"]
        assert fields["base_date"]
        assert fields["period_start"] < fields["base_date"]
        assert float(fields["asset_opening"]) > 0
        assert float(fields["asset_closing"]) > 0
        assert float(fields["asset_contributions"]) > 0
        assert float(fields["asset_paid"]) >= 0
        assert info["grade"]
        assert any("사외적립자산" in line for line in info["found"])
        assert info["problems"] == []

    def test_it_flags_a_table_that_does_not_balance(self, tmp_path) -> None:
        """회사 표가 스스로 안 맞으면 넣어 주되 반드시 알린다."""
        import openpyxl

        case = self._pack(tmp_path)
        wb = openpyxl.load_workbook(case["roster"])
        ws = wb["사외적립자산"]
        # 기말 잔액을 흔든다. 구성 열(DB퇴직연금)이 합계보다 우선하므로 거기를 고친다.
        closing = next(r for r in range(1, ws.max_row + 1)
                       if str(ws.cell(r, 2).value or "").startswith("기말 잔액"))
        ws.cell(closing, 3, float(ws.cell(closing, 3).value) + 9_103_134)
        broken = str(tmp_path / "안맞는표.xlsx")
        wb.save(broken)

        info = call("general_info", path=broken)
        assert info["fields"]["asset_closing"]
        assert any("맞지 않습니다" in p for p in info["problems"])

    def test_a_roster_without_the_sheet_hands_back_nothing(self, tmp_path) -> None:
        import openpyxl

        case = self._pack(tmp_path)
        wb = openpyxl.load_workbook(case["roster"])
        # 자산·규정·기본정보를 모두 뺀다 — 셋 중 하나만 남아도 읽을 것이 있다.
        for name in ("사외적립자산", "퇴직급여규정", "기본정보"):
            del wb[name]
        stripped = str(tmp_path / "일반사항없음.xlsx")
        wb.save(stripped)

        info = call("general_info", path=stripped)
        assert info["fields"] == {}
        assert info["found"] == []

    def test_the_roster_sheet_fills_the_asset_table_by_itself(self, tmp_path) -> None:
        """명부의 ``사외적립자산`` 에 표가 있으면 손으로 안 넣어도 나와야 한다."""
        case = self._pack(tmp_path)
        report = call("run", roster=case["roster"], assumptions=case["assumptions"],
                      work=str(tmp_path), sensitivity=False, longterm=False)
        rows = dict(report["assets"])
        assert rows["기초 사외적립자산 공정가치"] > 0
        assert rows["기말 사외적립자산 공정가치"] > 0
        assert rows["순확정급여부채"] == pytest.approx(
            report["values"]["dbo"] - rows["기말 사외적립자산 공정가치"]
        )

    def test_a_typed_amount_beats_the_roster_sheet(self, tmp_path) -> None:
        """화면에 넣은 값이 있으면 그쪽이 우선이다 — 명부보다 나중 자료다."""
        case = self._pack(tmp_path)
        report = call(
            "run", roster=case["roster"], assumptions=case["assumptions"],
            work=str(tmp_path), sensitivity=False, longterm=False,
            asset_opening=15_000_000_000, asset_closing=17_200_000_000,
        )
        rows = dict(report["assets"])
        assert rows["기초 사외적립자산 공정가치"] == 15_000_000_000

    def test_no_general_sheet_means_no_asset_table(self, tmp_path) -> None:
        """일반사항이 없는 명부(업로드용 변환본 등)는 표가 나오지 않는다."""
        import openpyxl

        source = openpyxl.load_workbook(case_roster := self._pack(tmp_path)["roster"])
        del source["사외적립자산"]
        stripped = str(tmp_path / "일반사항없음.xlsx")
        source.save(stripped)
        assert case_roster != stripped

        report = call("run", roster=stripped, assumptions=self._pack(tmp_path)[
            "assumptions"], work=str(tmp_path / "b"), sensitivity=False, longterm=False)
        assert report["assets"] == []

    def test_benefit_change_becomes_past_service_cost(self, tmp_path) -> None:
        """지급률을 바꾸면 그 효과가 가정변경이 아니라 과거근무원가로 잡혀야 한다."""
        case = self._pack(tmp_path)
        prior = form.read_state(case["assumptions"])
        current = form.read_state(case["assumptions"])
        for group in current["benefit_rules"]:
            current["benefit_rules"][group] = {"mode": "수식", "formula": "=t*1.5"}
        prior_path = str(tmp_path / "전기.xlsx")
        current_path = str(tmp_path / "당기.xlsx")
        form.write_state(prior, prior_path)
        form.write_state(current, current_path)

        base = call("run", roster=case["roster"], assumptions=prior_path,
                    work=str(tmp_path / "a"), sensitivity=False, longterm=False)
        after = call(
            "run", roster=case["roster"], assumptions=current_path,
            work=str(tmp_path / "b"), sensitivity=False, longterm=False,
            prior_dbo=base["values"]["dbo"], prior_rate=0.045,
            prior_assumptions=prior_path,
        )
        rows = dict(after["rollforward"])
        assert rows["과거근무원가(제도개정)"] > 0
        assert rows["보험수리적손익 - 가정변경"] == pytest.approx(0, abs=1)
        # 표가 실제 기말채무와 맞아야 한다.
        assert rows["기말 확정급여채무"] == pytest.approx(after["values"]["dbo"])

    def test_settlement_gain_is_recognised(self, tmp_path) -> None:
        case = self._pack(tmp_path)
        report = call(
            "run", roster=case["roster"], assumptions=case["assumptions"],
            work=str(tmp_path), sensitivity=False, longterm=False,
            prior_dbo=15_000_000_000, prior_rate=0.045,
            settlement_obligation=500_000_000,
        )
        rows = dict(report["rollforward"])
        paid = -rows["정산지급액(중간정산·전출)"]
        assert rows["정산손익"] == pytest.approx(paid - 500_000_000)


class TestRemainingGapsClosed:
    """감사에서 찾은 나머지 구멍들."""

    def _case(self, tmp_path):
        made = call("gen_cases", work=str(tmp_path), seed=909)
        return next(c for c in made["cases"] if not c["force"])

    def test_transfers_in_and_other_payments_appear(self, tmp_path) -> None:
        """전입 인수액과 퇴직위로금이 증감표에 제 줄로 나와야 한다.

        시험명부2 에는 전입자와 정년퇴직 위로금이 들어 있다.
        """
        made = call("gen_cases", work=str(tmp_path), seed=909)
        case = next(c for c in made["cases"] if c["key"] == "복합제도")
        report = call(
            "run", roster=case["roster"], assumptions=case["assumptions"],
            work=str(tmp_path), sensitivity=False, longterm=False,
            prior_dbo=10_000_000_000, prior_rate=0.045,
        )
        rows = dict(report["rollforward"])
        assert rows["전입 인수액"] > 0
        assert rows["퇴직위로금 등 지급액"] < 0

    def test_unpaid_benefits_raise_the_net_liability(self, tmp_path) -> None:
        case = self._case(tmp_path)
        common = dict(roster=case["roster"], assumptions=case["assumptions"],
                      sensitivity=False, longterm=False,
                      asset_opening=10_000_000_000, asset_closing=11_000_000_000)
        without = call("run", work=str(tmp_path / "a"), **common)
        with_unpaid = call("run", work=str(tmp_path / "b"),
                           unpaid_benefits=300_000_000, **common)

        net_before = dict(without["assets"])["순확정급여부채"]
        rows = dict(with_unpaid["assets"])
        assert rows["미지급 퇴직급여"] == 300_000_000
        assert rows["순확정급여부채"] == pytest.approx(net_before + 300_000_000)

    def test_extra_pay_and_honorary_wage_are_flagged(self, tmp_path) -> None:
        """계산에 쓰이거나 버려지는 임금 칸은 경고로 드러나야 한다."""
        import openpyxl

        case = self._case(tmp_path)
        book = openpyxl.load_workbook(case["roster"])
        ws = book["재직자명부"]
        # 열 번호를 박아 두면 양식이 한 칸 움직일 때 엉뚱한 칸을 채운다.
        # 머리글로 찾는다 — 명부를 읽는 쪽과 같은 방식이다.
        head = {str(c.value).strip(): c.column for c in ws[3] if c.value}
        ws.cell(25, head["추가지급 기본급"], 3_000_000)
        ws.cell(26, head["명예퇴직 기준임금"], 8_000_000)
        marked = str(tmp_path / "표시명부.xlsx")
        book.save(marked)

        report = call("run", roster=marked, assumptions=case["assumptions"],
                      work=str(tmp_path), sensitivity=False, longterm=False)
        assert report["run"] is True
        assert "경고" in report["issues"]


class TestPriorCheck:
    """전기 명부와 맞대어 보기 — 저장해 둔 산출을 상대로."""

    def _saved(self, tmp_path, roster_path, assumptions_path, name="2412 전기"):
        """산출 하나를 저장해 둔 상태를 만든다."""
        call("run", roster=str(roster_path), assumptions=str(assumptions_path),
             work=str(tmp_path), sensitivity=False, longterm=False, force=True)
        call("run_save", name=name, roster=str(roster_path),
             assumptions=str(assumptions_path), work=str(tmp_path))
        return name

    def test_the_same_roster_twice_is_clean(self, tmp_path, roster_path,
                                            assumptions_path) -> None:
        name = self._saved(tmp_path, roster_path, assumptions_path)
        found = call("prior_check", name=name, roster=str(roster_path),
                     work=str(tmp_path))
        assert found["serious"] == []
        assert "이상 없습니다" in found["summary"]
        assert found["matched"] > 0
        assert found["prior_name"] == name

    def test_a_changed_birth_date_is_caught(self, tmp_path, roster_path,
                                            assumptions_path) -> None:
        """당기 명부만 봐서는 알 수 없다. 전기와 맞대야 드러난다."""
        import openpyxl

        from pension.readers import ACTIVE_COLUMNS, ACTIVE_FIRST_ROW, ACTIVE_SHEET

        name = self._saved(tmp_path, roster_path, assumptions_path)

        changed = tmp_path / "당기명부.xlsx"
        wb = openpyxl.load_workbook(roster_path)
        ws = wb[ACTIVE_SHEET]
        ws.cell(ACTIVE_FIRST_ROW, ACTIVE_COLUMNS["birth_date"].index, "1955-01-01")
        wb.save(changed)

        found = call("prior_check", name=name, roster=str(changed), work=str(tmp_path))
        assert "생년월일 변경" in found["counts"]
        assert any("생년월일" in row[2] for row in found["serious"])

    def test_it_needs_a_saved_run(self, tmp_path, roster_path) -> None:
        result = json.loads(api(json.dumps({
            "op": "prior_check", "name": "없는산출", "roster": str(roster_path),
            "work": str(tmp_path)})))
        assert result["ok"] is False


class TestAllocationChoice:
    """지급규정의 '할당 방식' 이 파일을 오가며 살아남는지."""

    def test_round_trip_and_config(self, tmp_path) -> None:
        from pension.config import read_payout_rules
        from pension.workbook import open_workbook

        state = form.example_state(["정규직", "임원"])
        state["payout"]["임원"]["allocation"] = "근속비례"
        path = tmp_path / "기초율.xlsx"
        assert call("state_write", state=state, path=str(path))["written"] is True

        back = call("state_read", path=str(path))["state"]
        assert back["payout"]["임원"]["allocation"] == "근속비례"
        assert back["payout"]["정규직"]["allocation"] == "급여식"

        book = open_workbook(path)
        try:
            rules = {r.mapped_name: r for r in read_payout_rules(book)}
        finally:
            book.close()
        assert rules["임원"].allocation_method == "근속비례"
        assert rules["정규직"].allocation_method in ("", "급여식")

    def test_meta_offers_the_choices(self) -> None:
        assert call("meta")["allocations"] == ["급여식", "근속비례"]
