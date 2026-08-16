"""배포 꾸러미에 같이 넣는 기본 파일들.

받는 사람이 처음 여는 파일이 곧바로 실패하면 안 된다. 양식이 프로그램 자신에게
읽히는지, 기본 기초율만으로 산출이 끝까지 도는지를 여기서 고정한다.
"""

from __future__ import annotations

import datetime as _dt
from pathlib import Path

import openpyxl
import pytest

from pension.samples import (
    ROSTER_TEMPLATE,
    STANDARD_ASSUMPTIONS,
    TEMPLATE_ASSUMPTIONS,
    write_sample_pack,
)
from pension.standard_rates import STANDARD_TABLE, mortality_table


@pytest.fixture
def pack(tmp_path):
    write_sample_pack(tmp_path)
    return tmp_path


class TestRoundTrip:
    """양식이 프로그램 자신에게 읽혀야 한다. 안 그러면 양식이 아니다."""

    def test_roster_template_reads_back(self, pack) -> None:
        from pension.config import read_config
        from pension.errors import IssueLog
        from pension.readers import read_roster
        from pension.workbook import open_workbook

        book = open_workbook(pack / ROSTER_TEMPLATE)
        try:
            config = read_config(book)
            roster = read_roster(book, config, IssueLog())
        finally:
            book.close()

        assert len(roster.active) == 2
        assert len(roster.retired) == 2
        # 머리글로 열을 찾았으므로 값이 제자리에 들어가야 한다.
        assert [m.employee_id for m in roster.active] == ["A0001", "A0002"]
        assert roster.active[0].monthly_wage == 5_000_000
        # 임원인지는 읽는 자리에서 넘겨짚지 않는다. 직군 칸에 '임원' 이라고
        # 적혀 오면 그것을 그대로 들고 있다가, 사람이 정한 직군 매핑을
        # 확인한 뒤 검증 단계에서 임원으로 올린다.
        assert roster.active[1].job_group_raw == "임원"
        assert roster.retired[0].total_payment == 45_000_000

        # 주민등록번호 앞 7자리에서 생년월일과 성별을 읽어 낸다.
        assert roster.active[1].birth_date == _dt.date(1972, 8, 15)
        assert roster.active[1].gender.value == "여자"
        # 넘겨짚은 것이 아니라 **정한** 것이라는 표시가 남아야 한다.
        assert roster.active[1].gender_known

    def test_standard_assumptions_load(self, pack) -> None:
        from pension.assumptions import load_assumptions
        from pension.normalize import Gender

        assumptions = load_assumptions(pack / STANDARD_ASSUMPTIONS)
        assert assumptions.discount.level_rate > 0
        assert assumptions.mortality.qx(Gender.MALE, 60) > 0
        assert assumptions.mortality.qx(Gender.FEMALE, 60) > 0

    def test_blank_template_loads(self, pack) -> None:
        from pension.assumptions import load_assumptions

        assert load_assumptions(pack / TEMPLATE_ASSUMPTIONS).discount.level_rate > 0


class TestOutOfTheBox:
    """기본 파일 두 개만으로 산출이 끝까지 돌아야 한다."""

    def test_calculation_runs_clean(self, pack, tmp_path) -> None:
        from pension.pipeline import RunOptions, run_valuation
        from pension.report import write_report

        run = run_valuation(RunOptions(
            roster_path=pack / ROSTER_TEMPLATE,
            assumptions_path=pack / STANDARD_ASSUMPTIONS,
            output_path=tmp_path / "결과.xlsx",
        ))

        assert not run.issues.has_errors()
        # 작성 예시의 임원은 지급배수 2 를 달고 있지만 그것은 안내(info) 로만
        # 붙는다. 양식 그대로 돌린 산출에 경고가 하나라도 있으면 안 된다.
        assert [i.code for i in run.issues.warnings] == []
        assert run.valuation.headcount == 2
        assert run.valuation.dbo > 0
        write_report(run, tmp_path / "결과.xlsx")
        assert (tmp_path / "결과.xlsx").exists()

    def test_payout_sheet_drives_the_job_groups(self, pack, tmp_path) -> None:
        """기초율의 `지급규정` 시트가 명부 Input 보다 우선한다."""
        from pension.pipeline import load_inputs

        config, roster, _a, _log, _g = load_inputs(
            pack / ROSTER_TEMPLATE, pack / STANDARD_ASSUMPTIONS
        )
        assert config.mapped_names() == ["정규직", "계약직", "임원"]
        # 양식의 작성 예시는 직원 한 줄과 임원 한 줄이다.
        assert {m.job_group for m in roster.active} == {"정규직", "임원"}


class TestMortality:
    """표준사망률은 남녀를 나눠 쓴다. 규모로는 갈리지 않는다."""

    def test_covers_the_supplied_age_range(self) -> None:
        assert [row[0] for row in mortality_table()] == [r[0] for r in STANDARD_TABLE]

    def test_columns_are_male_then_female(self) -> None:
        for (age, *_rates, male, female), row in zip(
            STANDARD_TABLE, mortality_table(), strict=True
        ):
            assert row == [age, male, female]

    def test_women_die_later(self) -> None:
        assert all(row[2] <= row[1] for row in mortality_table())

    def test_rates_increase_with_age(self) -> None:
        for column in (1, 2):
            rates = [row[column] for row in mortality_table()]
            assert rates == sorted(rates)
            assert all(0.0 < r < 1.0 for r in rates)


class TestStandardTable:
    """퇴직률·승급률은 사업장 규모로 갈린 원표를 그대로 쓴다."""

    def test_size_picks_the_column(self) -> None:
        from pension.standard_rates import (
            SIZE_LARGE,
            SIZE_SMALL,
            promotion_table,
            withdrawal_table,
        )

        assert [[r[0], r[1]] for r in STANDARD_TABLE] == withdrawal_table(SIZE_SMALL)
        assert [[r[0], r[2]] for r in STANDARD_TABLE] == withdrawal_table(SIZE_LARGE)
        assert [[r[0], r[3]] for r in STANDARD_TABLE] == promotion_table(SIZE_SMALL)
        assert [[r[0], r[4]] for r in STANDARD_TABLE] == promotion_table(SIZE_LARGE)

    def test_the_two_sizes_really_differ(self) -> None:
        """한 열을 두 번 넣은 것이 아님을 못 박는다."""
        from pension.standard_rates import (
            SIZE_LARGE,
            SIZE_SMALL,
            promotion_table,
            withdrawal_table,
        )

        assert withdrawal_table(SIZE_SMALL) != withdrawal_table(SIZE_LARGE)
        assert promotion_table(SIZE_SMALL) != promotion_table(SIZE_LARGE)

    def test_unknown_size_falls_back_to_the_default(self) -> None:
        from pension.standard_rates import (
            DEFAULT_SIZE,
            normalize_size,
            withdrawal_table,
        )

        assert normalize_size("아무거나") == DEFAULT_SIZE
        assert withdrawal_table("아무거나") == withdrawal_table(DEFAULT_SIZE)
        assert normalize_size("300인↑") == "300인 이상"

    def test_size_is_suggested_from_headcount(self) -> None:
        from pension.standard_rates import SIZE_LARGE, SIZE_SMALL, size_for

        assert size_for(299) == SIZE_SMALL
        assert size_for(300) == SIZE_LARGE

    def test_withdrawal_peaks_at_the_retirement_ramp(self) -> None:
        """55세부터 정년 언저리까지 가파르게 오른다(임금피크·명예퇴직 구간)."""
        for column in (1, 2):
            rates = {r[0]: r[column] for r in STANDARD_TABLE}
            assert rates[55] < rates[58] < rates[61]
            assert rates[61] > rates[50] * 2

    def test_last_row_is_age_70(self) -> None:
        """원표는 110세까지 있지만 70세부터 값이 같다. 표는 계단식으로 읽힌다."""
        assert STANDARD_TABLE[-1][0] == 70

    def test_every_rate_is_a_probability(self) -> None:
        for age, *rates in STANDARD_TABLE:
            assert 15 <= age <= 70
            assert all(0.0 < r < 1.0 for r in rates)


class TestDefaultRoster:
    """기본 명부는 **난수로 만든 가상 명부** 다.

    실제 명부를 넣으면 프로그램을 건네는 순간 그 자료도 같이 건네진다. 성명이
    없어도 생년월일·입사일·30일 평균임금이 사람마다 한 줄씩이면 같은 회사
    안에서는 특정될 수 있다.

    난수로 만들되 **형태는 지킨다.** 작성 예시 두 줄짜리 양식으로는 DC 혼재·
    근속 1년 미만 퇴직자·임원의 직군 표기 같은 것을 볼 수 없다.
    """

    def _roster(self, pack):
        from pension.config import read_config
        from pension.errors import IssueLog
        from pension.readers import read_roster
        from pension.samples import ROSTER_DEFAULT
        from pension.workbook import open_workbook

        book = open_workbook(pack / ROSTER_DEFAULT)
        try:
            return read_roster(book, read_config(book), IssueLog())
        finally:
            book.close()

    def test_reads_back_at_full_size(self, pack) -> None:
        from pension.rostergen import DEFAULT_CASE

        roster = self._roster(pack)
        assert len(roster.active) == DEFAULT_CASE.active
        assert len(roster.retired) == DEFAULT_CASE.retired

    def test_carries_the_shapes_a_two_row_sample_cannot(self, pack) -> None:
        """실제 명부에서 마주치는 형태가 들어 있어야 쓸모가 있다."""
        from pension.normalize import BenefitPlan, EmployeeType

        roster = self._roster(pack)
        # DC 가 섞여 있다 — 산출대상에서 빠지는 사람이 생긴다.
        assert any(m.plan is BenefitPlan.DC for m in roster.active)
        # 임원의 직군이 '정규직' 으로 적혀 온다 — 직군만으로는 갈라낼 수 없다.
        assert any(
            m.employee_type is EmployeeType.EXECUTIVE and m.job_group_raw == "정규직"
            for m in roster.active
        )
        # 근속 1년 미만 퇴직자 — 법정 지급 대상이 아니라 금액이 비어 있다.
        assert any(0 < m.service_years() < 1.0 for m in roster.retired)

    def test_calculates_without_forcing(self, pack, tmp_path) -> None:
        """처음 돌려 보는 명부다. [강행] 을 켜야 도는 자료면 첫인상이 나쁘다."""
        from pension.pipeline import RunOptions, run_valuation
        from pension.samples import ROSTER_DEFAULT

        run = run_valuation(RunOptions(
            roster_path=pack / ROSTER_DEFAULT,
            assumptions_path=pack / STANDARD_ASSUMPTIONS,
            output_path=tmp_path / "결과.xlsx",
        ))
        assert not run.issues.errors
        assert run.valuation.dbo > 0
        # DC 가입자가 빠지므로 산출대상이 재직자보다 적어야 한다.
        assert 0 < run.valuation.headcount < len(run.roster.active)

    def test_the_same_seed_gives_the_same_roster(self, tmp_path) -> None:
        """받는 사람마다 다른 명부가 나오면 '같은 값이 나오나' 를 못 맞춰 본다.

        원 바이트를 비교하지 않는다 — openpyxl 이 저장 시각을 파일 메타데이터
        (``dcterms:created``)에 찍어 넣어서, 두 호출이 초 경계를 걸치면 내용이
        같아도 바이트가 달라진다. 값을 다시 읽어 비교한다.
        """
        from pension.samples import write_default_roster

        def values(path):
            wb = openpyxl.load_workbook(path, data_only=True)
            return {name: [[c.value for c in row] for row in wb[name].iter_rows()]
                    for name in wb.sheetnames}

        one = write_default_roster(tmp_path / "가.xlsx")
        two = write_default_roster(tmp_path / "나.xlsx")
        assert values(one) == values(two)


class TestNoPersonalDataShips:
    """배포물 어디에도 실제 개인정보가 없어야 한다.

    프로그램을 동료에게 건네는 순간 같이 들어 있는 자료도 건네진다. 웹앱 쪽은
    빌드에서 걸러 내고 있었지만, 파이썬 꾸러미와 EXE 에는 그 방어가 없어
    **실제 명부가 그대로 실려 나갔다.**
    """

    def test_the_package_carries_no_data_files(self) -> None:
        import pension

        data = Path(pension.__file__).parent / "data"
        # 자료 폴더 자체는 있어도 된다 — 금리표처럼 **시장 자료** 는 사람을
        # 가리키지 않는다. 막아야 하는 것은 명부다.
        if data.exists():
            allowed = {"금리표_20251231.xlsx"}
            stray = [p for p in data.iterdir()
                     if p.name not in allowed and not p.name.startswith("__")]
            assert stray == [], f"꾸러미에 개인 자료가 남아 있다: {stray}"
            assert not list(data.glob("*.csv"))

    def test_the_exe_bundles_only_the_webapp(self) -> None:
        """실행 파일에 실리는 것은 전체 기능 화면뿐이어야 한다.

        웹앱은 들어가야 한다(옆 폴더에 두면 실행 파일만 복사했을 때 안 열린다).
        **명부 원자료는 들어가면 안 된다** — 그것이 이 시험의 요점이다.
        """
        spec = (Path(__file__).resolve().parent.parent / "pension.spec").read_text(
            encoding="utf-8")
        datas = spec.split("datas=[", 1)[1].split("]", 1)[0]
        assert '"webapp/dist", "webapp"' in datas
        assert "pension/data" not in datas
        assert ".csv" not in datas

    def test_the_default_roster_is_generated_not_stored(self, pack) -> None:
        """씨앗에서 만들어 낸 것이라 파일로 들고 다닐 원자료가 없다."""
        from pension.rostergen import DEFAULT_CASE
        from pension.samples import ROSTER_DEFAULT

        assert (pack / ROSTER_DEFAULT).is_file()
        # 인원수는 지어낸 값이다. 실제 명부의 인원과 같으면 그 자체가 흔적이 된다.
        assert DEFAULT_CASE.active > 100 and DEFAULT_CASE.retired > 0


class TestSingleDiscountRate:
    """수익률곡선기법 — 곡선으로 할인한 채무와 같은 현가를 내는 단일 이자율.

    K-IFRS 1019 문단 85 의 '단일 가중평균 할인율' 이다.
    """

    def test_reproduces_the_present_value(self) -> None:
        from pension.valuation import single_equivalent_rate

        flows = {1.0: 100.0, 5.0: 300.0, 10.0: 600.0}
        target = 100 / 1.03 + 300 / 1.035**5 + 600 / 1.042**10

        rate = single_equivalent_rate(flows, target)
        assert sum(a / (1 + rate) ** t for t, a in flows.items()) == pytest.approx(target)

    def test_lands_between_the_curve_rates(self) -> None:
        """단일 이자율은 곡선의 최저·최고 사이에 있어야 한다."""
        from pension.valuation import single_equivalent_rate

        flows = {1.0: 100.0, 5.0: 300.0, 10.0: 600.0}
        target = 100 / 1.03 + 300 / 1.035**5 + 600 / 1.042**10
        assert 0.03 < single_equivalent_rate(flows, target) < 0.042

    def test_flat_curve_gives_that_rate_back(self) -> None:
        from pension.valuation import single_equivalent_rate

        flows = {1.0: 100.0, 5.0: 300.0, 10.0: 600.0}
        target = sum(a / 1.04**t for t, a in flows.items())
        assert single_equivalent_rate(flows, target) == pytest.approx(0.04)

    def test_no_liability_is_not_an_error(self) -> None:
        """전원 DC 라 채무가 0 이면 할인율이 정의되지 않는다."""
        from pension.valuation import single_equivalent_rate

        assert single_equivalent_rate({}, 0.0) == 0.0
        assert single_equivalent_rate({5.0: 100.0}, 0.0) == 0.0

    def test_matches_the_valuation(self, pack, tmp_path) -> None:
        from pension.pipeline import RunOptions, run_valuation
        from pension.samples import ROSTER_DEFAULT

        run = run_valuation(RunOptions(
            roster_path=pack / ROSTER_DEFAULT,
            assumptions_path=pack / STANDARD_ASSUMPTIONS,
            output_path=tmp_path / "결과.xlsx",
            include_sensitivity=False, include_longterm=False, allow_errors=True,
        ))
        v = run.valuation
        rate = v.single_discount_rate()
        restated = sum(a / (1 + rate) ** t for t, a in v.cash_flows().items() if t > 0)
        assert restated == pytest.approx(v.dbo, rel=1e-9)


class TestYieldCurve:
    """금리표 파일을 할인율 시트로 옮긴다."""

    def _book(self, tmp_path):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "금리표"
        ws.append(["No", "기준일자", "구분", "등급", "1년", "1년6월", "5년", "20년"])
        ws.append([1, None, None, "국고채", 2.55, 2.747, 3.235, 3.347])
        ws.append([2, None, "공모 무보증회사채", "AA0", 3.122, 3.133, 3.62, 5.26])
        ws.append([3, None, "기타이율", "기타1", 0, 0, 0, 0])
        path = tmp_path / "금리표.xlsx"
        wb.save(path)
        return path

    def test_reads_tenors_and_converts_percent(self, tmp_path) -> None:
        from pension.yieldcurve import pick_curve, read_yield_curves

        curves = read_yield_curves(self._book(tmp_path))
        assert [c.grade for c in curves] == ["국고채", "AA0"]   # 0 만 있는 행은 제외

        aa0 = pick_curve(curves, "AA0")
        assert aa0.points[0] == (1.0, pytest.approx(0.03122))
        assert aa0.points[1] == (1.5, pytest.approx(0.03133))
        assert aa0.points[-1] == (20.0, pytest.approx(0.0526))

    def test_default_grade_prefers_investment_grade(self, tmp_path) -> None:
        """등급을 지정하지 않으면 우량회사채를 고른다. 국고채가 첫 행이어도."""
        from pension.yieldcurve import pick_curve, read_yield_curves

        assert pick_curve(read_yield_curves(self._book(tmp_path))).grade == "AA0"

    def test_curve_lands_in_the_assumptions(self, tmp_path) -> None:
        from pension.assumptions import load_assumptions
        from pension.samples import write_standard_assumptions

        path = write_standard_assumptions(
            tmp_path / "기초율.xlsx", yield_curve_path=self._book(tmp_path), grade="AA0"
        )
        discount = load_assumptions(path).discount
        assert discount.flat is None            # 곡선이므로 단일 할인율이 아니다
        assert discount.rate(1) == pytest.approx(0.03122)
        assert discount.rate(20) == pytest.approx(0.0526)


class TestStandardTableWorkbook:
    """고시된 표준률 원표를 그 서식 그대로 내고, 그 서식 그대로 읽는다.

    표준률은 몇 해에 한 번 새로 고시되고, 그때 오는 파일이 원표 서식이다 —
    중도퇴직률·승급률·사망률이 각각 한 시트, 열은 ``No · 연령 · 300인↓ ·
    300인↑``. 사람이 288개 숫자를 이 프로그램 서식으로 옮겨 적으면 어딘가
    한 자리는 틀리고, 틀린 자리는 채무 숫자만 보고는 찾을 수 없다.
    """

    def test_it_writes_the_original_column_shape(self, tmp_path) -> None:
        import openpyxl

        from pension.samples import write_standard_table
        from pension.standard_rates import STANDARD_YEAR

        path = write_standard_table(tmp_path / "원표.xlsx")
        wb = openpyxl.load_workbook(path)
        assert [name for name in wb.sheetnames] == [
            f"Sheet_{STANDARD_YEAR}_표준중도퇴직률",
            f"Sheet_{STANDARD_YEAR}_bu제외_표준승급률",
            f"Sheet_{STANDARD_YEAR}_표준사망률",
        ]
        ws = wb[wb.sheetnames[0]]
        assert [ws.cell(1, c).value for c in range(1, 5)] == ["No", "연령", "300인↓", "300인↑"]
        assert ws.cell(2, 1).value == "001"
        assert ws.cell(2, 2).value == "15"

        mortality = wb[wb.sheetnames[2]]
        assert [mortality.cell(1, c).value for c in range(1, 5)] == [
            "No", "연령", "남자", "여자"]

    def test_the_written_values_are_the_built_in_table(self, tmp_path) -> None:
        from pension.samples import write_standard_table
        from pension.standard_rates import (
            STANDARD_TABLE,
            mortality_table,
            read_raw_workbook,
            withdrawal_table,
        )

        found = read_raw_workbook(write_standard_table(tmp_path / "원표.xlsx"))
        assert found["퇴직률"]["300인 미만"] == withdrawal_table("300인 미만")
        assert found["퇴직률"]["300인 이상"] == withdrawal_table("300인 이상")
        assert len(found["사망률"]["남자"]) == len(STANDARD_TABLE)
        assert [[age, male] for age, male, _f in mortality_table()] == \
            found["사망률"]["남자"]

    def test_a_raw_workbook_loads_as_a_screen_state(self, tmp_path) -> None:
        """등록해서 바로 쓸 수 있어야 한다 — 옮겨 적게 하지 않는 것이 요점이다."""
        from pension import assumption_form as form
        from pension.samples import write_standard_table

        path = write_standard_table(tmp_path / "원표.xlsx")
        assert form.looks_like_raw_table(path)

        small = form.read_state(path, size="300인 미만")
        large = form.read_state(path, size="300인 이상")
        assert small["grids"]["퇴직률"]["rows"][0][1] == "0.394440"
        assert large["grids"]["퇴직률"]["rows"][0][1] == "0.172360"
        # 사망률은 규모로 갈리지 않는다.
        assert small["grids"]["사망률"]["rows"] == large["grids"]["사망률"]["rows"]

    def test_one_table_on_its_own_still_reads(self, tmp_path) -> None:
        """원표는 표마다 파일이 따로 오는 일이 흔하다."""
        import openpyxl

        from pension import assumption_form as form

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet_202312_표준사망률"
        ws.append(["No", "연령", "남자", "여자"])
        ws.append(["001", "20", "0.000050", "0.000030"])
        ws.append(["002", "21", "0.000050", "0.000030"])
        path = tmp_path / "사망률만.xlsx"
        wb.save(path)

        state = form.read_state(path)
        assert state["grids"]["사망률"]["rows"] == [
            ["20", "0.000050", "0.000030"], ["21", "0.000050", "0.000030"]]
        assert state["grids"]["퇴직률"]["rows"] == []

    def test_the_ordinary_workbook_is_not_mistaken_for_a_raw_table(self, tmp_path) -> None:
        """기초율 파일에도 '퇴직률' 시트가 있다. 그쪽을 원표로 읽으면 안 된다."""
        from pension import assumption_form as form
        from pension.samples import write_standard_assumptions

        path = write_standard_assumptions(tmp_path / "기초율.xlsx")
        assert not form.looks_like_raw_table(path)
        # 원표로 잘못 읽으면 지급률 규정이 통째로 사라진다 — 원표에는 없는 것이다.
        assert form.read_state(path)["benefit_rules"]

    def test_a_workbook_with_no_table_says_so(self, tmp_path) -> None:
        import openpyxl
        import pytest

        from pension.standard_rates import read_raw_workbook

        wb = openpyxl.Workbook()
        wb.active.title = "아무것도아님"
        path = tmp_path / "빈것.xlsx"
        wb.save(path)
        with pytest.raises(ValueError, match="표준률 원표"):
            read_raw_workbook(path)


class TestCurveTemplate:
    """금리표(채권) 등록 양식. **이율 칸은 비워 둔다.**"""

    def test_the_rate_cells_are_empty(self, tmp_path) -> None:
        """할인율은 결산일의 시장 자료다. 지어낸 값을 넣어 두면 안 된다."""
        import openpyxl

        from pension.samples import write_curve_template

        path = write_curve_template(tmp_path / "금리표.xlsx")
        ws = openpyxl.load_workbook(path)["금리표"]
        for row in range(3, 7):                    # 3행부터가 자료 (2행이 머리글)
            for column in range(5, ws.max_column + 1):
                assert ws.cell(row, column).value is None

    def test_it_has_the_grades_and_tenors_ready(self, tmp_path) -> None:
        import openpyxl

        from pension.samples import CURVE_TENORS, write_curve_template
        from pension.yieldcurve import INVESTMENT_GRADES

        path = write_curve_template(tmp_path / "금리표.xlsx")
        ws = openpyxl.load_workbook(path)["금리표"]
        headers = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        assert headers[:4] == ["순번", "기준일자", "구분", "등급"]
        assert tuple(headers[4:]) == CURVE_TENORS

        grades = [ws.cell(r, 4).value for r in range(3, 3 + len(INVESTMENT_GRADES))]
        assert set(grades) == set(INVESTMENT_GRADES)

    def test_a_filled_template_is_read_back(self, tmp_path) -> None:
        """채워 넣으면 실제로 읽혀야 양식이라 할 수 있다."""
        import openpyxl

        from pension.samples import write_curve_template
        from pension.yieldcurve import pick_curve, read_yield_curves

        path = write_curve_template(tmp_path / "금리표.xlsx")
        wb = openpyxl.load_workbook(path)
        ws = wb["금리표"]
        row = next(r for r in range(2, 8) if ws.cell(r, 4).value == "AA0")
        ws.cell(row, 5, 3.404)        # 3월
        ws.cell(row, 8, 3.512)        # 1년
        ws.cell(row, 18, 5.260)       # 20년
        wb.save(path)

        curve = pick_curve(read_yield_curves(path), "AA0")
        assert curve is not None
        assert curve.rate_at(0.25) == pytest.approx(0.03404)
        assert curve.rate_at(20) == pytest.approx(0.0526)

    def test_it_ships_with_the_sample_pack(self, tmp_path) -> None:
        from pension.samples import CURVE_TEMPLATE, write_sample_pack

        made = {path.name for path in write_sample_pack(tmp_path)}
        assert CURVE_TEMPLATE in made


class TestByCause:
    """퇴직사유(급부)별로 갈라 놓은 금액.

    사유마다 지급률이 다른 규정에서는 합계만으로 검산이 안 된다. 지급률 한 칸을
    잘못 넣어도 총액은 조금 움직일 뿐이라, 사유별로 갈라 놓아야 눈에 띈다.
    DBO 산출표가 정년·중도·사망 열을 따로 세우는 이유와 같다.
    """

    def _run(self, pack, tmp_path):
        from pension.pipeline import RunOptions, run_valuation

        return run_valuation(RunOptions(
            roster_path=pack / ROSTER_TEMPLATE,
            assumptions_path=pack / STANDARD_ASSUMPTIONS,
            output_path=tmp_path / "결과.xlsx",
        ))

    def test_the_parts_add_up_to_the_whole(self, pack, tmp_path) -> None:
        """나눈 것이지 다시 계산한 것이 아니다 — 합이 총액과 원 단위까지 같아야 한다."""
        valuation = self._run(pack, tmp_path).valuation
        causes = valuation.by_cause()

        assert causes, "사유별 몫이 비어 있다"
        assert sum(s["dbo"] for s in causes.values()) == pytest.approx(valuation.dbo)
        assert sum(s["service_cost"] for s in causes.values()) == pytest.approx(
            valuation.service_cost)

    def test_it_covers_the_three_ways_out(self, pack, tmp_path) -> None:
        from pension.assumptions import CAUSE_DEATH, CAUSE_NORMAL, CAUSE_VOLUNTARY

        causes = self._run(pack, tmp_path).valuation.by_cause()
        assert set(causes) == {CAUSE_NORMAL, CAUSE_VOLUNTARY, CAUSE_DEATH}
        # 정년이 맨 앞 — 산출표를 볼 때의 열 차례다.
        assert list(causes)[0] == CAUSE_NORMAL

    def test_one_person_splits_the_same_way(self, pack, tmp_path) -> None:
        """사번 조회의 사유별 몫도 그 사람 채무와 맞아야 한다."""
        from pension.memberdetail import lookup

        found = lookup(str(pack / ROSTER_TEMPLATE), str(pack / STANDARD_ASSUMPTIONS),
                       "A0001")
        shares = found["by_cause"]
        assert shares
        assert sum(s["dbo"] for s in shares) == pytest.approx(
            found["total"]["확정급여채무 (DBO)"])

    def test_the_result_workbook_carries_it(self, pack, tmp_path) -> None:
        from pension.report import write_report

        run = self._run(pack, tmp_path)
        target = tmp_path / "결과.xlsx"
        write_report(run, target)
        ws = openpyxl.load_workbook(target)["요약"]
        labels = [ws.cell(r, 2).value for r in range(1, ws.max_row + 1)]
        assert any("퇴직사유별" in str(v) for v in labels)
        assert any(v == "합계 (= 확정급여채무)" for v in labels)


class TestBaseDateFromTheScreen:
    """명부에 기준일이 없어도 화면에 넣은 날짜로 산출돼야 한다.

    실제로 그런 명부가 왔다. 화면에 2025-12-31 을 넣어 두었는데도 '산출기준일을
    찾지 못했습니다' 로 멈췄다 — 명부를 읽는 함수가 화면 값을 받기 **전에**
    터졌기 때문이다. 담당자로서는 넣은 값이 왜 무시되는지 알 길이 없다.
    """

    def _roster_without_a_date(self, pack, tmp_path) -> Path:
        book = openpyxl.load_workbook(pack / ROSTER_TEMPLATE)
        del book["기초자료"]         # 기준일이 적힌 유일한 자리
        target = tmp_path / "기준일없는명부.xlsx"
        book.save(target)
        return target

    def test_without_a_date_anywhere_it_says_so(self, pack, tmp_path) -> None:
        from pension.config import read_config
        from pension.workbook import open_workbook

        book = open_workbook(self._roster_without_a_date(pack, tmp_path))
        try:
            with pytest.raises(ValueError, match="산출기준일"):
                read_config(book)
        finally:
            book.close()

    def test_the_screen_date_is_enough(self, pack, tmp_path) -> None:
        import datetime as dt

        from pension.config import read_config
        from pension.workbook import open_workbook

        book = open_workbook(self._roster_without_a_date(pack, tmp_path))
        try:
            config = read_config(book, base_date=dt.date(2025, 12, 31))
        finally:
            book.close()
        assert config.base_date == dt.date(2025, 12, 31)

    def test_the_whole_run_goes_through(self, pack, tmp_path) -> None:
        import datetime as dt

        from pension.pipeline import RunOptions, run_valuation

        run = run_valuation(RunOptions(
            roster_path=self._roster_without_a_date(pack, tmp_path),
            assumptions_path=pack / STANDARD_ASSUMPTIONS,
            output_path=tmp_path / "결과.xlsx",
            base_date=dt.date(2025, 12, 31),
        ))
        assert run.config.base_date == dt.date(2025, 12, 31)
        assert run.valuation.dbo > 0
