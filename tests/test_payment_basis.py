"""지급액을 **채무 전표처리 기준** 과 **실지급 기준** 으로 나눠 받는다.

한 칸으로 받아 왔다. 그러면 회사가 어느 기준으로 적었는지 알 길이 없다.

12월에 퇴직한 사람의 돈이 1월에 나가면 전표는 당기, 현금은 차기다. 원천징수를
뗀 뒤 금액으로 적어 오는 회사도 있다. **어느 쪽이 틀린 것이 아니라 기준이
다른 두 숫자** 이고, 증감표가 안 맞을 때 그 사실을 모르면 원인을 짚을 수가
없다 — 명부를 아무리 들여다봐도 두 기준이 한 칸에 뭉쳐 있으니 보이지 않는다.

산출과 증감표는 **전표 기준** 을 쓴다. 채무가 줄어든 금액이기 때문이다.
"""

from __future__ import annotations

import openpyxl
import pytest

from pension.readers import RETIRED_SHEET


@pytest.fixture
def pack(tmp_path):
    from pension.samples import (
        ROSTER_TEMPLATE, STANDARD_ASSUMPTIONS, write_sample_pack)

    write_sample_pack(tmp_path)
    return tmp_path / ROSTER_TEMPLATE, tmp_path / STANDARD_ASSUMPTIONS


def run(roster, assumptions, tmp_path):
    from pension.pipeline import RunOptions, run_valuation

    return run_valuation(RunOptions(
        roster_path=roster, assumptions_path=assumptions,
        output_path=tmp_path / "결과.xlsx",
    ))


def set_cash(path, amounts):
    """퇴직자명부의 [실지급 기준 지급액] 칸을 채운다."""
    wb = openpyxl.load_workbook(path)
    ws = wb[RETIRED_SHEET]
    column = next(
        col for col in range(1, ws.max_column + 1)
        if str(ws.cell(4, col).value or "").strip() == "실지급 기준 지급액")
    for offset, amount in enumerate(amounts):
        ws.cell(5 + offset, column, amount)
    wb.save(path)
    return path


def test_the_template_asks_for_both(tmp_path) -> None:
    """양식이 두 칸을 다 묻는지. 한 칸만 있으면 기준을 가릴 수 없다."""
    from pension.rostertemplate import write_roster_template

    path = write_roster_template(tmp_path / "양식.xlsx")
    wb = openpyxl.load_workbook(path)
    heads = [str(c.value or "") for c in wb[RETIRED_SHEET][4]]
    assert "퇴직급여 총지급액" in heads
    assert "실지급 기준 지급액" in heads


def test_a_blank_cash_column_changes_nothing(pack, tmp_path) -> None:
    """실지급 칸이 비면 종전과 같다 — 옛 명부가 그대로 읽힌다."""
    roster, assumptions = pack
    result = run(roster, assumptions, tmp_path)
    assert all(m.cash_payment == 0 for m in result.roster.retired)
    assert [i.code for i in result.issues.notices if i.code == "TOI_PAYMENT_BASIS"] == []


def test_the_two_bases_are_read_apart(pack, tmp_path) -> None:
    """두 칸이 서로 다른 값으로 읽히는지."""
    roster, assumptions = pack
    before = run(roster, assumptions, tmp_path)
    booked = [m.total_payment for m in before.roster.retired]
    assert booked and booked[0] > 0

    set_cash(roster, [booked[0] - 3_000_000])
    result = run(roster, assumptions, tmp_path)
    first = result.roster.retired[0]
    assert first.total_payment == booked[0]
    assert first.cash_payment == booked[0] - 3_000_000


def test_a_gap_is_said_once_not_per_person(pack, tmp_path) -> None:
    """갈리면 알리되 **한 줄로** 접는다.

    사람마다 말하면 정작 봐야 할 항목이 묻힌다. 시점이 갈리는 것은 정상이라
    오류도 경고도 아니고, 어느 기준을 썼는지 알려 주는 안내다.
    """
    roster, assumptions = pack
    before = run(roster, assumptions, tmp_path)
    booked = [m.total_payment for m in before.roster.retired]
    set_cash(roster, [amount - 1_000_000 for amount in booked])

    result = run(roster, assumptions, tmp_path)
    said = [i for i in result.issues.notices if i.code == "TOI_PAYMENT_BASIS"]
    assert len(said) == 1, f"한 줄이 아니다: {len(said)}"
    assert "전표" in said[0].message and "실지급" in said[0].message
    assert not any(i.code == "TOI_PAYMENT_BASIS" for i in result.issues.warnings)


def test_the_engine_still_uses_the_booked_basis(pack, tmp_path) -> None:
    """산출은 전표 기준을 쓴다 — 채무가 줄어든 금액이기 때문이다.

    실지급 칸을 아무리 바꿔도 채무·증감표가 흔들리면 안 된다.
    """
    roster, assumptions = pack
    plain = run(roster, assumptions, tmp_path)
    set_cash(roster, [1])
    moved = run(roster, assumptions, tmp_path)
    assert moved.valuation.dbo == pytest.approx(plain.valuation.dbo)
    assert sum(m.total_payment for m in moved.roster.retired) == pytest.approx(
        sum(m.total_payment for m in plain.roster.retired))
