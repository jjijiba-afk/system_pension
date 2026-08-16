"""개인별 산출 결과 내보내기."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from pension.cli import main
from pension.members import (
    COST_SHEET,
    GROUP_SHEET,
    MEMBER_SHEET,
    build_member_rows,
    write_member_export,
)
from pension.pipeline import RunOptions, run_valuation


@pytest.fixture
def run(roster_path: Path, assumptions_path: Path, tmp_path: Path):
    return run_valuation(
        RunOptions(
            roster_path=roster_path,
            assumptions_path=assumptions_path,
            output_path=tmp_path / "결과.xlsx",
            include_sensitivity=False,
            include_longterm=True,
        ),
        lambda *_a: None,
    )


class TestBuildRows:
    def test_one_row_per_person(self, run) -> None:
        rows = build_member_rows(run)
        assert len(rows) == len(run.valuation.members)

    def test_carries_roster_fields_for_cost_allocation(self, run) -> None:
        """원가배분을 하려면 명부를 다시 열지 않고도 원가코드가 있어야 한다."""
        rows = build_member_rows(run)
        assert any(r.cost_code for r in rows)
        assert all(r.plan for r in rows)
        assert all(r.birth_date is not None for r in rows)
        assert all(r.hire_date is not None for r in rows)

    def test_longterm_is_merged_onto_the_same_row(self, run) -> None:
        rows = build_member_rows(run)
        by_id = {r.employee_id: r for r in rows}
        for member in run.longterm.members:
            assert by_id[member.employee_id].longterm_dbo == member.dbo

    def test_cause_split_adds_up_to_the_dbo(self, run) -> None:
        """사유별 몫 셋의 합이 그 사람의 확정급여채무와 같아야 한다."""
        rows = build_member_rows(run)
        assert any(r.dbo_normal for r in rows), "정년 몫이 있어야 한다"
        for row in rows:
            assert (row.dbo_normal + row.dbo_voluntary + row.dbo_death
                    == pytest.approx(row.dbo))

    def test_carries_the_payout_multiple(self, run) -> None:
        """추계액이 큰 것이 임금 탓인지 배수 탓인지 결과만 보고 가려야 한다."""
        rows = build_member_rows(run)
        paid = [r for r in rows if r.accrued_benefit > 0]
        assert paid, "추계액이 있는 사람이 있어야 한다"
        assert all(r.accrued_multiple > 0 for r in paid)

    def test_total_dbo_adds_both_liabilities(self, run) -> None:
        rows = build_member_rows(run)
        for row in rows:
            assert row.total_dbo == pytest.approx(row.dbo + row.longterm_dbo)

    def test_row_totals_match_the_valuation(self, run) -> None:
        rows = build_member_rows(run)
        assert sum(r.dbo for r in rows) == pytest.approx(run.valuation.dbo)
        assert sum(r.service_cost for r in rows) == pytest.approx(run.valuation.service_cost)
        assert sum(r.longterm_dbo for r in rows) == pytest.approx(run.longterm.dbo)

    def test_dc_member_is_kept_with_a_reason(self, run) -> None:
        """산출에서 빠진 사람도 행은 남아야 인원 대사가 된다."""
        rows = build_member_rows(run)
        excluded = [r for r in rows if r.excluded_reason]
        assert excluded
        assert all(r.dbo == 0.0 for r in excluded)


class TestWorkbook:
    def test_writes_three_sheets(self, run, tmp_path: Path) -> None:
        path = write_member_export(run, tmp_path / "개인별.xlsx")
        wb = openpyxl.load_workbook(path)
        assert wb.sheetnames == [MEMBER_SHEET, COST_SHEET, GROUP_SHEET]

    def test_member_sheet_has_a_total_row(self, run, tmp_path: Path) -> None:
        path = write_member_export(run, tmp_path / "개인별.xlsx")
        ws = openpyxl.load_workbook(path)[MEMBER_SHEET]

        rows = build_member_rows(run)
        total_row = 3 + len(rows) + 1
        assert ws.cell(total_row, 1).value == "합계"

        headers = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
        dbo_col = headers.index("확정급여채무") + 1
        assert ws.cell(total_row, dbo_col).value == pytest.approx(run.valuation.dbo)

    def test_cost_summary_totals_match(self, run, tmp_path: Path) -> None:
        path = write_member_export(run, tmp_path / "개인별.xlsx")
        ws = openpyxl.load_workbook(path)[COST_SHEET]

        last = ws.max_row
        assert ws.cell(last, 1).value == "합계"
        # 열: 원가코드 | 인원 | 임금합계 | 추계액 | DBO | ...
        assert ws.cell(last, 5).value == pytest.approx(run.valuation.dbo)

    def test_cost_summary_splits_by_code(self, run, tmp_path: Path) -> None:
        path = write_member_export(run, tmp_path / "개인별.xlsx")
        ws = openpyxl.load_workbook(path)[COST_SHEET]
        codes = [ws.cell(r, 1).value for r in range(2, ws.max_row)]
        assert len(codes) >= 1
        assert "합계" not in codes

    def test_blank_cost_code_is_labelled(self, run, tmp_path: Path) -> None:
        """원가코드가 비어 있어도 행이 사라지면 합계가 맞지 않는다."""
        for member in run.valuation.members:
            member.cost_code = ""
        path = write_member_export(run, tmp_path / "개인별.xlsx")
        ws = openpyxl.load_workbook(path)[COST_SHEET]
        assert ws.cell(2, 1).value == "(미지정)"

    def test_has_autofilter_for_sorting(self, run, tmp_path: Path) -> None:
        path = write_member_export(run, tmp_path / "개인별.xlsx")
        ws = openpyxl.load_workbook(path)[MEMBER_SHEET]
        assert ws.auto_filter.ref is not None


class TestCli:
    def test_members_command_writes_the_file(
        self, roster_path: Path, assumptions_path: Path, tmp_path: Path, capsys
    ) -> None:
        out = tmp_path / "개인별.xlsx"
        code = main(["members", str(roster_path), str(assumptions_path), "-o", str(out)])
        assert code == 0
        assert out.exists()
        assert "개인별 결과" in capsys.readouterr().out

    def test_members_command_stops_on_validation_errors(
        self, roster_path: Path, assumptions_path: Path, tmp_path: Path
    ) -> None:
        wb = openpyxl.load_workbook(roster_path)
        wb["재직자명부"].cell(26, 17, "")  # 제도구분 누락
        wb.save(roster_path)

        out = tmp_path / "개인별.xlsx"
        assert main(["members", str(roster_path), str(assumptions_path), "-o", str(out)]) == 2
        assert not out.exists()

    def test_calc_can_also_write_the_member_file(
        self, roster_path: Path, assumptions_path: Path, tmp_path: Path
    ) -> None:
        out = tmp_path / "결과.xlsx"
        members = tmp_path / "개인별.xlsx"
        code = main([
            "calc", str(roster_path), str(assumptions_path),
            "-o", str(out), "--members", str(members),
        ])
        assert code == 0
        assert out.exists()
        assert members.exists()
