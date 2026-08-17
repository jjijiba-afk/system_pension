"""아이패드 웹앱(브라우저 안 엔진) 끝까지 검증.

Playwright 와 Chromium, 그리고 빌드된 dist 가 있을 때 돈다. 크로미움은 이
상자에 미리 깔린 것을 먼저 쓰고, 없으면 Playwright 가 자기 것을 찾게 둔다 —
경로를 하나로 못박아 두었더니 CI 에서는 늘 조용히 건너뛰어, 브라우저 화면이
한 번도 시험되지 않은 채로 지나갔다.
"""

from __future__ import annotations

import functools
import http.server
import threading
from pathlib import Path

import pytest

from pension.rostertemplate import FIRST_DATA_ROW, HEADER_ROW

playwright_api = pytest.importorskip("playwright.sync_api")

DIST = Path(__file__).resolve().parent.parent / "webapp" / "dist"


def _chromium() -> str | None:
    """쓸 수 있는 크로미움 실행파일. 없으면 ``None``."""
    pinned = Path("/opt/pw-browsers/chromium")
    if pinned.exists():
        return str(pinned)
    try:
        with playwright_api.sync_playwright() as play:
            found = Path(play.chromium.executable_path)
    except Exception:      # noqa: BLE001 — 안 깔렸으면 건너뛴다
        return None
    return str(found) if found.exists() else None


CHROMIUM = _chromium()

pytestmark = [
    pytest.mark.skipif(not DIST.exists(), reason="webapp/build.py 를 먼저 실행"),
    pytest.mark.skipif(CHROMIUM is None, reason="Chromium 없음"),
]


@pytest.fixture(scope="module")
def app_url():
    handler = functools.partial(
        http.server.SimpleHTTPRequestHandler, directory=str(DIST)
    )
    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{server.server_address[1]}/index.html"
    server.shutdown()
    server.server_close()


@pytest.fixture(scope="module")
def shared_dir(tmp_path_factory):
    """테스트끼리 주고받는 파일(보관함 zip)을 두는 곳."""
    return tmp_path_factory.mktemp("웹앱공유")


@pytest.fixture(scope="module")
def browser():
    with playwright_api.sync_playwright() as p:
        browser = p.chromium.launch(executable_path=CHROMIUM)
        yield browser
        browser.close()


@pytest.fixture(scope="module")
def page(browser, app_url):
    """엔진 부팅이 오래 걸리므로 한 번 띄운 페이지를 모듈 전체가 나눠 쓴다."""
    page = browser.new_page()
    page.goto(app_url)
    # 첫 인사가 모달로 떠 있으면 뒤 화면이 통째로 눌리지 않는다. 여기서 한 번
    # 치우고 시작한다 — 이 창의 동작 자체는 test_intro_dialog 가 따로 본다.
    dismiss_intro(page)
    page.wait_for_selector("#run:not([disabled])", timeout=120_000)
    # 재측정요소 분해는 화면 기본값이 **켬** 이지만, 전기를 연결한 회차에서는
    # 명부 전체 산출이 네 번 더 돈다. 이 페이지는 시험 열댓 개가 나눠 쓰는
    # 작업대라, 그 값을 켠 채로 두면 산출을 부르는 시험 전부가 다섯 배로
    # 늘어진다. 기본값이 제대로 걸려 있는지는
    # test_the_calc_screen_opens_with_working_defaults 가 새 페이지에서 따로 본다.
    page.uncheck("#split_remeasurement")
    yield page
    page.close()


#: 산출가정 화면의 구획이 어느 묶음 탭 아래에 있는지. 파이썬 쪽
#: ``EDITOR_GROUPS`` 와 짝이 맞아야 한다(test_webui 가 배치를 따로 검사한다).
_SECTION_GROUP = {
    "할인율": "기초율", "승급률": "기초율", "퇴직률": "기초율", "사망률": "기초율",
    "임금상승률": "직군", "지급률": "퇴직급여", "장기급여지급률": "장기급여",
}


def section(page, sheet: str):
    """묶음 탭을 열고 그 안의 구획을 펼친다.

    탭이 열세 개일 때는 시트 이름이 곧 탭 이름이었다. 네 묶음으로 접은 뒤로는
    묶음을 먼저 열고 구획을 골라야 한다 — 사람이 하는 것과 같은 순서다.
    """
    page.click(f"#ed-subtabs >> text={_SECTION_GROUP[sheet]}")
    box = page.locator(f'#ed-subpages .subpage.on details[data-section="{sheet}"]')
    if not box.evaluate("node => node.open"):
        box.locator("summary").click()
    return box


def grid(page, sheet: str):
    """그 구획 안의 표."""
    return section(page, sheet).locator("tbody")


def grid_row(page, sheet: str, row: int = 1):
    """그 표의 ``row`` 번째 **값** 줄. 머리글과 열 머리 패널은 건너뛴다."""
    return grid(page, sheet).locator("tr[data-row]").nth(row - 1).locator("input")


def dismiss_intro(page) -> None:
    """첫 인사 창을 치운다.

    모달이라 떠 있는 동안에는 뒤 화면이 통째로 눌리지 않는다. 저장소가 빈
    컨텍스트(=새 기기)를 여는 시험은 어디서든 이것을 먼저 거쳐야 한다.
    """
    page.wait_for_selector("#intro-dialog[open]", timeout=30_000)
    page.check("#intro-hide")
    page.click("#intro-close")
    page.wait_for_selector("#intro-dialog", state="hidden", timeout=10_000)


def open_section(page, selector: str):
    """접어 둔 구획을 펼친다. 이미 펼쳐져 있으면 그대로 둔다."""
    box = page.locator(selector)
    if not box.evaluate("node => node.open"):
        box.locator("summary").click()


def unlock_teaching_material(page):
    """교육용(잠긴) 자료를 시험에서 열어 본다.

    실제 비밀번호는 저장소에 없다(해시만 있다). 그래서 브라우저 안에서 도는
    엔진의 해시를 **시험용 비밀번호의 해시로** 갈아 끼운 뒤, 사람이 하듯 안내
    창에 그 비밀번호를 넣는다 — 잠금 화면 자체도 함께 시험하게 된다.
    """
    from pension.rostergen import LOCK_ROUNDS, LOCK_SALT, _derive

    word = "시험용 자물쇠"
    digest = _derive(word)
    assert LOCK_ROUNDS and LOCK_SALT
    page.evaluate(
        """(args) => pyodide.runPython(
            `from pension import rostergen\nrostergen.LOCK_HASH = ${JSON.stringify(args.digest)}`
        )""",
        {"digest": digest},
    )
    page.click("#feat-ask")
    page.wait_for_selector("#lock-dialog[open]", timeout=10_000)
    page.fill("#lock-pass", word)
    page.click("#lock-ok")
    page.wait_for_selector("#feat-open:not([hidden])", timeout=10_000)


def open_calc(page, section_id: str):
    """산출 탭에서 접어 둔 구획을 펼친다. 사람이 하는 것과 같은 순서다."""
    page.click("#tab-calc")
    open_section(page, f"#{section_id}")


def test_upload_run_download(page, tmp_path) -> None:
    """업로드 → 브라우저 안 산출 → 진짜 엑셀 내려받기까지."""
    from pension.samples import write_sample_pack

    files = write_sample_pack(tmp_path)
    roster = next(p for p in files if p.name == "명부_양식.xlsx")
    assumptions = next(p for p in files if p.name == "기초율_기본값.xlsx")

    page.set_input_files("#roster", str(roster))
    page.check("#asrc-file")
    page.set_input_files("#assumptions", str(assumptions))
    page.click("#run")
    page.wait_for_selector("#result", state="visible", timeout=180_000)

    summary = page.inner_text("#summary")
    assert "확정급여채무" in summary

    with page.expect_download() as captured:
        page.click("#dl-result")
    payload = Path(captured.value.path()).read_bytes()
    assert payload[:2] == b"PK"


def test_an_uploaded_roster_comes_back_in_our_template(page, tmp_path) -> None:
    """올린 명부를 우리 양식 그대로 되받는다.

    회사마다 열 순서와 머리글이 달라 매 결산 눈으로 맞춰야 했다. 한 번 올린
    것을 이 모양으로 되받아 다음 해에 그것을 채워 달라고 하면 어긋날 자리가
    없다 — 그 되받기가 브라우저에서 실제로 되는지 여기서 본다.
    """
    import openpyxl

    from pension.samples import write_sample_pack

    files = write_sample_pack(tmp_path)
    roster = next(p for p in files if p.name == "명부_양식.xlsx")

    page.set_input_files("#roster", str(roster))
    with page.expect_download(timeout=120_000) as captured:
        page.click("#roster-export")
    download = captured.value
    assert download.suggested_filename.endswith("_표준양식.xlsx")

    got = tmp_path / "되받은.xlsx"
    download.save_as(got)
    book = openpyxl.load_workbook(got)
    assert "재직자명부" in book.sheetnames
    assert book["재직자명부"].cell(HEADER_ROW, 2).value == "사번"


def test_font_scale_is_small_on_a_pc_and_safe_on_touch(browser, app_url) -> None:
    """글자 크기는 **배율** 로만 건다.

    PC 는 설치판과 같은 9pt(=12px). 손가락 기기는 조금만 조이되 입력칸은
    16px 아래로 내리지 않는다 — iOS 사파리가 그 미만이면 칸을 누를 때마다
    화면을 확대해 버려, 표를 채우는 내내 화면이 튄다.

    px 로 못박으면 브라우저에서 글자를 키워 둔 사람의 설정이 통째로 무시되므로
    ``%`` 인지도 함께 본다.
    """
    def sizes(viewport, **opts):
        view = browser.new_page(viewport=viewport, **opts)
        view.goto(app_url)
        dismiss_intro(view)
        view.wait_for_selector("#run:not([disabled])", timeout=180_000)
        view.click("#tab-edit")
        found = view.evaluate("""() => {
          const px = (s) => parseFloat(
            getComputedStyle(document.querySelector(s)).fontSize);
          return { root: px(":root"), cell: px(".grid input[type=text]") };
        }""")
        view.close()
        return found

    pc = sizes({"width": 1440, "height": 900})
    assert pc["root"] == 12.0, f"PC 는 9pt(12px) 여야 한다: {pc}"

    pad = sizes({"width": 1024, "height": 1366}, has_touch=True, is_mobile=True)
    assert pad["root"] > pc["root"], "손가락 기기가 PC 보다 작아지면 안 된다"
    assert pad["cell"] >= 16.0, f"입력칸이 16px 미만이면 iOS 가 확대한다: {pad}"

    css = (DIST / "app.css").read_text(encoding="utf-8")
    assert ":root { font-size: 87.5%; }" in css, "배율(%)로 걸어야 한다"


def test_the_calc_screen_opens_with_working_defaults(browser, app_url) -> None:
    """[산출] 화면은 매 회차 같은 것을 고르게 된다. 그것을 기본값으로 둔다.

    * 산출기준일 — 지금까지 지나간 가장 가까운 연말. 결산은 끝난 사업연도를
      잰다. 다만 **우리가 채운 칸** 으로 표시해 두어, 명부 [기본정보] 에
      기준일이 있으면 그쪽이 이긴다.
    * 재측정요소 가정별 분해 — 켬.
    * 기초율은 [산출가정 입력 화면] — 다만 그 화면이 비어 있으면 쓸 수 없어
      파일로 물러난다. 채워지면 도로 돌아오고, 사람이 다른 출처를 직접
      고른 뒤에는 건드리지 않는다.
    """
    import datetime as dt

    view = browser.new_page()
    view.goto(app_url)
    dismiss_intro(view)
    view.wait_for_selector("#run:not([disabled])", timeout=180_000)

    today = dt.date.today()
    year = today.year if (today.month, today.day) == (12, 31) else today.year - 1
    start = view.evaluate("""() => ({
      base: document.querySelector('#base_date').value,
      split: document.querySelector('#split_remeasurement').checked,
      file: document.querySelector('#asrc-file').checked,
    })""")
    assert start["base"] == f"{year}-12-31", start
    assert start["split"] is True
    assert start["file"] is True, "편집 화면이 비었으면 파일로 물러나야 한다"

    # 산출가정을 채우면 기본값(입력 화면)으로 돌아온다.
    view.click("#tab-edit")
    view.evaluate("""() => { const r = py('standard_state',
        {groups: ['정규직', '임원'], size: '300인 미만'}); renderState(r.state); }""")
    assert view.evaluate("() => document.querySelector('#asrc-editor').checked") is True

    # 사람이 직접 고른 뒤에는 그 선택을 지킨다.
    view.click("#tab-calc")
    view.check("#asrc-file")
    view.click("#tab-edit")
    view.click("#tab-calc")
    assert view.evaluate("() => document.querySelector('#asrc-file').checked") is True
    view.close()


def test_payout_rules_show_how_many_people_they_reach(page, tmp_path) -> None:
    """[지급규정]에 건 직군이 **실제 명부의 몇 사람에게 닿는지** 보여야 한다.

    직군 이름 한 글자가 명부와 다르면 그 규정은 아무에게도 걸리지 않는데,
    숫자가 안 보이면 산출이 끝날 때까지 알 길이 없다. 0 명이면 눈에 띄어야
    한다 — 오류 없이 그럴듯한 숫자가 나오는 쪽이라 더 그렇다.
    """
    from pension.samples import write_sample_pack

    files = write_sample_pack(tmp_path)
    roster = next(p for p in files if p.name == "명부_양식.xlsx")

    page.set_input_files("#roster", str(roster))
    page.click("#tab-edit")
    page.click("#ed-groups-roster")
    page.wait_for_function(
        "() => document.querySelector('#ed-status')?.textContent.includes('직군')",
        timeout=120_000)

    # 양식 작성 예시는 정규직 1명 · 임원 1명.
    counts = page.evaluate("""() => {
      const out = {};
      for (const tr of [...payoutBody.children].slice(1)) {
        out[tr.dataset.group] = tr.children[1].textContent;
      }
      return out;
    }""")
    assert counts["정규직"] == "1명", counts
    assert counts["임원"] == "1명", counts

    # 명부에 없는 이름을 걸면 0 명으로 뜨고, 짝이 안 맞는다고 알려 준다.
    page.evaluate("() => applyGroups(['정규직', '임원', '없는직군'])")
    zero = page.evaluate("""() => {
      const tr = [...payoutBody.children].find((t) => t.dataset.group === '없는직군');
      return tr.children[1].textContent;
    }""")
    assert zero == "0명"
    # 구획이 접혀 있어도 내용은 채워져 있어야 한다(inner_text 는 숨으면 빈다).
    assert "없는직군" in page.text_content("#payout-coverage")


def test_roster_fills_the_asset_boxes(page, tmp_path) -> None:
    """명부를 고르면 [사외적립자산] 시트의 값이 입력칸에 들어가야 한다.

    엔진이 알아서 읽는 것만으로는 부족하다. 칸이 비어 있으면 담당자는
    아무것도 읽히지 않은 줄 알고 신탁 명세서를 보고 손으로 다시 적는다.
    """
    from pension.rostergen import CASES, write_case_roster

    roster = write_case_roster(CASES[0], tmp_path / "일반사항명부.xlsx")

    page.click("#tab-calc")
    page.set_input_files("#roster", str(roster))
    page.wait_for_selector("#general-filled", state="visible", timeout=60_000)

    # 채워 넣은 구획은 저절로 펼쳐져야 한다 — 접힌 채로 값만 들어가면
    # 담당자가 확인할 기회 없이 그대로 산출된다.
    assert page.locator("#sec-assets").evaluate("node => node.open")
    assert page.input_value("#base_date")
    assert page.input_value("#period_start") < page.input_value("#base_date")
    assert float(page.input_value("#asset_opening")) > 0
    assert float(page.input_value("#asset_closing")) > 0
    assert "예치금" in page.inner_text("#general-filled")

    # 사람이 손댄 칸은 덮어쓰지 않는다 — 명부보다 나중 자료일 수 있다.
    page.fill("#asset_opening", "1")
    page.set_input_files("#roster", str(roster))
    page.wait_for_selector("#general-filled", state="visible", timeout=60_000)
    assert page.input_value("#asset_opening") == "1"

    # 다른 단체로 바꾸면 **우리가 채운 칸은** 그 명부 값으로 다시 채운다.
    # 앞 단체 숫자를 남겨 두면 남의 회사 자산으로 산출된다.
    other = write_case_roster(CASES[1], tmp_path / "다른단체.xlsx")
    before = page.input_value("#asset_closing")
    page.set_input_files("#roster", str(other))
    page.wait_for_selector("#general-filled", state="visible", timeout=60_000)
    assert page.input_value("#asset_closing") != before
    assert page.input_value("#asset_opening") == "1"   # 손댄 칸은 그대로

    for box in ("base_date", "period_start", "asset_opening", "asset_contributions",
                "asset_paid", "asset_closing"):
        page.fill(f"#{box}", "")
    page.set_input_files("#roster", [])


def test_editor_tab_builds_assumptions_and_runs(page, tmp_path) -> None:
    """산출가정 입력 탭에서 만든 가정만으로 산출까지 이어져야 한다."""
    from pension.samples import write_sample_pack

    files = write_sample_pack(tmp_path)
    roster = next(p for p in files if p.name == "명부_양식.xlsx")
    page.set_input_files("#roster", str(roster))

    page.click("#tab-edit")
    page.click("#ed-example")  # 예시 값 채우기 → 할인율이 생겨 편집기 소스가 열린다

    # 저장하면 산출가정.xlsx 가 내려오고, 산출 탭의 소스가 편집기로 바뀐다.
    with page.expect_download() as captured:
        page.click("#ed-save")
    assert Path(captured.value.path()).read_bytes()[:2] == b"PK"

    page.click("#tab-calc")
    assert page.is_checked("#asrc-editor")
    page.click("#run")
    page.wait_for_selector("#result", state="visible", timeout=180_000)
    assert "확정급여채무" in page.inner_text("#summary")


def test_library_registers_curve_into_editor(page, tmp_path) -> None:
    """기본가정 관리에 금리표를 등록하면 편집기 할인율에 적용할 수 있어야 한다."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "KIS_NET금리"
    ws.append(["No", "기준일자", "구분", "등급", "3월", "1년", "5년", "20년"])
    ws.append([1, None, "공모 무보증회사채", "AA0", 3.01, 3.122, 3.62, 5.26])
    source = tmp_path / "금리.xlsx"
    wb.save(source)

    page.click("#tab-lib")
    open_section(page, "#sec-lib-curve")
    page.set_input_files("#lib-curve-file", str(source))
    page.fill("#lib-curve-name", "KIS_E2E")
    page.click("#lib-curve-add")
    # 목록에 표가 있는지로 기다리면 안 된다 — 내장 금리표가 이미 한 줄 있어
    # 곧바로 통과해 버리고, 등록이 끝나기 전에 확인하게 된다.
    page.wait_for_function(
        "() => document.getElementById('lib-curve-list')"
        ".textContent.includes('KIS_E2E')",
        timeout=30_000)
    assert "KIS_E2E" in page.inner_text("#lib-curve-list")

    page.click("#tab-edit")
    page.select_option("#ed-curve", "KIS_E2E")
    page.select_option("#ed-grade", "AA0")
    page.click("#ed-curve-apply")
    assert "할인율에 넣었습니다" in page.inner_text("#ed-status")

    # 3월 만기(0.25년)가 살아 있어야 한다 — 내림하면 안 되는 값.
    assert grid_row(page, "할인율").first.input_value() == "0.25"


def test_run_history_save_and_restore(page, tmp_path) -> None:
    """산출 → 이름 붙여 저장 → 목록 → 입력 불러와 재산출까지."""
    from pension.samples import write_sample_pack

    files = write_sample_pack(tmp_path)
    roster = next(p for p in files if p.name == "명부_양식.xlsx")
    assumptions = next(p for p in files if p.name == "기초율_기본값.xlsx")

    page.click("#tab-calc")
    page.set_input_files("#roster", str(roster))
    page.check("#asrc-file")
    page.set_input_files("#assumptions", str(assumptions))
    page.click("#run")
    page.wait_for_selector("#result", state="visible", timeout=180_000)

    page.fill("#run-name", "2412 1번단체")
    page.click("#run-save")
    page.wait_for_selector("text=저장했습니다", timeout=30_000)

    page.click("#tab-runs")
    listing = page.inner_text("#runs-list")
    assert "2412 1번단체" in listing and "확정급여채무" not in listing

    # 입력을 불러오면 산출 탭으로 돌아오고, 파일을 고르지 않아도 재산출된다.
    page.click("#runs-list >> text=입력 불러오기")
    page.wait_for_selector("#loaded-run-banner", state="visible")
    assert page.is_checked("#asrc-saved")
    page.click("#run")
    page.wait_for_selector("#result", state="visible", timeout=180_000)
    assert "확정급여채무" in page.inner_text("#summary")

    # 저장된 결과 보기 대화상자.
    page.click("#tab-runs")
    page.click("#runs-list >> text=결과 보기")
    page.wait_for_selector("#run-dialog[open]")
    assert "확정급여채무" in page.inner_text("#run-dialog-summary")
    page.click("#run-dialog >> text=닫기")


def test_saved_roster_and_prior_link_and_backup(page, tmp_path, shared_dir) -> None:
    """올렸던 명부 재사용 · 전기 DBO 연결 · 기기 밖 보관함까지."""
    from pension.samples import write_sample_pack

    files = write_sample_pack(tmp_path)
    roster = next(p for p in files if p.name == "명부_양식.xlsx")
    assumptions = next(p for p in files if p.name == "기초율_기본값.xlsx")

    # 전기 산출을 하나 만들어 저장한다.
    page.click("#tab-calc")
    page.set_input_files("#roster", str(roster))
    page.check("#asrc-file")
    page.set_input_files("#assumptions", str(assumptions))

    # 명부를 목록에 저장 — prompt 로 이름을 묻는다.
    page.once("dialog", lambda dialog: dialog.accept("1번단체 명부"))
    page.click("#roster-save")
    page.wait_for_selector("text=목록에 저장했습니다", timeout=30_000)

    page.click("#run")
    page.wait_for_selector("#result", state="visible", timeout=180_000)
    page.fill("#run-name", "2312 1번단체")
    page.click("#run-save")
    page.wait_for_selector("text=저장했습니다", timeout=30_000)

    # 저장된 명부만으로 (파일 재선택 없이) 당기를 돌린다.
    page.reload()
    page.wait_for_selector("#run:not([disabled])", timeout=120_000)
    page.select_option("#roster-saved", "1번단체 명부")
    page.set_input_files("#assumptions", str(assumptions))

    # 전기 산출을 고르면 DBO·할인율이 자동으로 채워진다.
    open_calc(page, "sec-prior")
    page.select_option("#prior-run", "2312 1번단체")
    assert page.input_value("#prior_dbo").replace(",", "").isdigit()
    assert page.input_value("#prior_rate").endswith("%")

    page.click("#run")
    page.wait_for_selector("#result", state="visible", timeout=180_000)
    assert "보험수리적손익" in page.inner_text("#summary")

    # 보관함 내보내기 — 기기 밖에 둘 zip 이 실제로 떨어져야 한다.
    page.click("#tab-lib")
    open_section(page, "#lib-backup")
    with page.expect_download() as captured:
        page.click("#backup-export")
    archive = shared_dir / "보관함.zip"
    captured.value.save_as(str(archive))
    assert archive.read_bytes()[:2] == b"PK"


def test_backup_restores_on_a_clean_device(browser, app_url, shared_dir) -> None:
    """저장소가 빈 기기(=새 브라우저 컨텍스트)에서 보관함으로 되살린다."""
    archive = shared_dir / "보관함.zip"
    if not archive.exists():
        pytest.skip("앞 테스트에서 보관함을 만들지 못했다")

    context = browser.new_context()   # 저장소가 비어 있는 '다른 기기'
    fresh = context.new_page()
    fresh.goto(app_url)
    dismiss_intro(fresh)
    fresh.wait_for_selector("#run:not([disabled])", timeout=120_000)

    fresh.click("#tab-runs")
    assert "저장된 산출이 없습니다" in fresh.inner_text("#runs-list")

    fresh.on("dialog", lambda dialog: dialog.accept())
    fresh.click("#tab-lib")
    open_section(fresh, "#lib-backup")
    fresh.set_input_files("#backup-file", str(archive))
    fresh.click("#backup-replace")
    fresh.click("#tab-runs")
    fresh.wait_for_selector("#runs-list >> text=2312 1번단체", timeout=60_000)

    # 등록 자료(명부)도 함께 돌아와야 한다.
    fresh.click("#tab-calc")
    assert "1번단체 명부" in fresh.inner_text("#roster-saved")
    context.close()


def test_editor_survives_tab_switching(page) -> None:
    """가정을 입력하다 다른 탭에 다녀와도 값이 남아 있어야 한다."""
    page.click("#tab-edit")
    first = grid_row(page, "할인율")
    first.nth(0).fill("3")
    first.nth(1).fill("4.44%")

    # 곧바로 다른 탭으로 — 디바운스가 끝나기 전에 옮긴다.
    page.click("#tab-lib")
    page.click("#tab-runs")
    page.click("#tab-edit")

    kept = grid_row(page, "할인율")
    assert kept.nth(0).input_value() == "3"
    assert kept.nth(1).input_value() == "4.44%"

    # 새로고침해도 남아야 한다(브라우저에 임시 저장).
    page.reload()
    page.wait_for_selector("#run:not([disabled])", timeout=120_000)
    page.click("#tab-edit")
    assert grid_row(page, "할인율").nth(1).input_value() == "4.44%"

    # 접어 둔 구획도 채워졌는지는 제목 옆 숫자로 보여야 한다.
    page.click("#ed-subtabs >> text=기초율")
    summary = page.locator(
        '#ed-subpages .subpage.on details[data-section="사망률"] summary .count'
    )
    assert summary.inner_text().strip() in {"비어 있음", "3줄"}


def test_generator_tab_makes_and_runs_a_case(page) -> None:
    """자료실의 시험 명부 구획에서 만든 명부로 곧바로 산출까지."""
    page.click("#tab-lib")
    open_section(page, "#lib-gen")
    page.fill("#gen-seed", "777")
    page.click("#gen-run")
    page.wait_for_selector("#gen-cases fieldset", timeout=120_000)
    # 잠기지 않은 사례 둘 + 잠긴 사례 안내 하나.
    assert page.locator("#gen-cases button", has_text="이 명부로 산출 준비").count() == 2
    assert "개발중인 메뉴" in page.inner_text("#gen-cases")

    # 특이사항 안내문이 실제 내용을 담고 있어야 한다.
    page.locator("#gen-cases button", has_text="특이사항 보기").first.click()
    page.wait_for_selector("#report-dialog[open]")
    assert "산출 특이사항" in page.inner_text("#report-body")
    page.click("#report-dialog >> text=닫기")

    # 사례마다 zip 이 따로 나온다 — 실습자에게 한 벌씩 나눠 줄 수 있어야 한다.
    with page.expect_download() as captured:
        page.locator("#gen-cases button", has_text="zip 내려받기").first.click()
    assert Path(captured.value.path()).read_bytes()[:2] == b"PK"

    # '이 명부로 산출 준비' → 파일 선택 없이 그대로 산출된다.
    page.locator("#gen-cases button", has_text="이 명부로 산출 준비").first.click()
    page.wait_for_selector("#loaded-run-banner", state="visible")
    page.click("#run")
    page.wait_for_selector("#result", state="visible", timeout=180_000)
    assert "확정급여채무" in page.inner_text("#summary")


def test_preset_round_trip_in_the_browser(page) -> None:
    """가정세트를 저장하고 다시 불러온다."""
    page.click("#tab-edit")
    cells = grid_row(page, "지급률")
    cells.nth(0).fill("10")
    cells.nth(1).fill("13")

    page.once("dialog", lambda dialog: dialog.accept("E2E 규정"))
    page.click("#ed-preset-save")
    page.wait_for_selector("#ed-status:has-text('저장했습니다')", timeout=60_000)

    # 값을 지운 뒤 되불러오면 돌아와야 한다.
    page.click("#ed-example")
    page.select_option("#ed-preset", "E2E 규정")
    page.click("#ed-preset-load")
    page.wait_for_selector("#ed-status:has-text('불러왔습니다')", timeout=60_000)
    back = grid_row(page, "지급률")
    assert back.nth(0).input_value() == "10"
    assert back.nth(1).input_value() == "13"


def test_standard_rates_fill_the_grids(page) -> None:
    """내장 표준률 불러오기 — 15~70세 표가 채워져야 한다."""
    page.click("#tab-edit")
    page.select_option("#ed-rates", "__builtin__")
    page.click("#ed-rates-load")
    assert "불러왔습니다" in page.inner_text("#ed-status")

    assert grid_row(page, "사망률").first.input_value() == "15"


def test_workplace_size_switches_the_standard_table(page) -> None:
    """300인 미만/이상은 다른 표다 — 고른 쪽이 실제로 표에 들어와야 한다."""
    page.click("#tab-edit")
    page.select_option("#ed-rates", "__builtin__")

    def first_withdrawal(size: str) -> str:
        page.select_option("#ed-size", size)
        page.click("#ed-rates-load")
        assert size in page.inner_text("#ed-status")
        return grid_row(page, "퇴직률").nth(1).input_value()

    small = first_withdrawal("300인 미만")
    large = first_withdrawal("300인 이상")
    assert small != large

    # 사망률은 규모로 갈리지 않으므로 그대로여야 한다.
    assert grid_row(page, "사망률").first.input_value() == "15"


def test_benefit_column_splits_into_three_causes(page) -> None:
    """[사유별 차등] 한 번으로 정년·중도·사망 열이 생기고 연결까지 끝나야 한다."""
    def split_box():
        return section(page, "지급률").locator(
            'tr.panel:has(td.name:text-is("사유별 차등")) input[type=checkbox]'
        ).first

    def headers():
        return section(page, "지급률").locator("tbody tr").first.inner_text()

    split_box().check()
    assert "정규직·정년" in headers()

    # 퇴직사유 표까지 자동으로 채워져야 손댈 곳이 없다.
    causes = page.locator('#ed-subpages .subpage.on details[data-section="cause"]')
    if not causes.evaluate("node => node.open"):
        causes.locator("summary").click()
    assert "정규직·정년" in causes.inner_text()

    # 도로 접으면 열도 연결도 사라진다.
    page.once("dialog", lambda dialog: dialog.accept())
    split_box().uncheck()
    assert "정규직·정년" not in headers()
    assert "정규직·정년" not in causes.inner_text()


def test_client_bar_keeps_run_history_apart(page, tmp_path) -> None:
    """단체를 갈아 끼우면 산출 내역과 전기 산출 목록이 그 단체 것만 남는다."""
    from pension.samples import write_sample_pack

    files = write_sample_pack(tmp_path)
    roster = next(p for p in files if p.name == "명부_양식.xlsx")
    assumptions = next(p for p in files if p.name == "기초율_기본값.xlsx")

    # 앞 테스트가 기본 단체에 "2412 1번단체" 를 저장해 두었다.
    page.click("#tab-runs")
    assert "2412 1번단체" in page.inner_text("#runs-list")

    page.once("dialog", lambda dialog: dialog.accept("나단체"))
    page.click("#client-add")
    assert page.input_value("#client-pick") == "나단체"
    assert "아직 저장된 산출이 없습니다" in page.inner_text("#runs-list")

    # 전기 산출 목록에도 앞 단체 것이 남으면 안 된다.
    open_calc(page, "sec-prior")
    assert "2412 1번단체" not in page.inner_text("#prior-run")

    # 이 단체에 한 건 저장하면 여기에만 쌓인다.
    page.set_input_files("#roster", str(roster))
    page.check("#asrc-file")
    page.set_input_files("#assumptions", str(assumptions))
    page.click("#run")
    page.wait_for_selector("#result", state="visible", timeout=180_000)
    page.fill("#run-name", "2412")
    page.click("#run-save")
    page.wait_for_selector("text=저장했습니다", timeout=30_000)

    page.click("#tab-runs")
    listing = page.inner_text("#runs-list")
    assert "2412" in listing and "1번단체" not in listing
    assert page.inner_text("#client-runs") == "1건"

    # 이름을 바꿔도 산출은 따라간다.
    page.once("dialog", lambda dialog: dialog.accept("나단체㈜"))
    page.click("#client-rename")
    assert page.input_value("#client-pick") == "나단체㈜"
    assert "2412" in page.inner_text("#runs-list")

    # 되돌아가면 앞 단체 것이 그대로 있다.
    page.select_option("#client-pick", "기본 단체")
    assert "2412 1번단체" in page.inner_text("#runs-list")

    # 산출이 남은 단체는 한 번 더 물어본 뒤에야 지워진다.
    page.select_option("#client-pick", "나단체㈜")
    page.once("dialog", lambda dialog: dialog.dismiss())
    page.click("#client-remove")
    assert page.input_value("#client-pick") == "나단체㈜"

    page.once("dialog", lambda dialog: dialog.accept())
    page.click("#client-remove")
    assert page.input_value("#client-pick") == "기본 단체"
    assert "2412 1번단체" in page.inner_text("#runs-list")


def test_clients_and_runs_survive_a_reload(page, tmp_path) -> None:
    """단체와 그 안의 산출은 새로고침해도 남아 있어야 한다.

    앱 자료는 메모리 파일시스템에 있다가 ``syncfs`` 로 브라우저 저장소에
    밀어 넣어야 남는다. 그 호출을 한 군데라도 빠뜨리면 화면은 '저장했습니다'
    라고 말하고 새로고침하면 사라진다 — 결산을 마친 뒤에 알게 된다.
    """
    from pension.samples import write_sample_pack

    files = write_sample_pack(tmp_path)
    roster = next(p for p in files if p.name == "명부_양식.xlsx")
    assumptions = next(p for p in files if p.name == "기초율_기본값.xlsx")

    page.once("dialog", lambda dialog: dialog.accept("살아남을단체"))
    page.click("#client-add")
    assert page.input_value("#client-pick") == "살아남을단체"

    page.click("#tab-calc")
    page.set_input_files("#roster", str(roster))
    page.check("#asrc-file")
    page.set_input_files("#assumptions", str(assumptions))
    page.click("#run")
    page.wait_for_selector("#result", state="visible", timeout=180_000)
    page.fill("#run-name", "새로고침시험")
    page.click("#run-save")
    page.wait_for_selector("text=저장했습니다", timeout=30_000)

    # ── 여기서 새로고침 ──────────────────────────────────────────
    page.reload()
    page.wait_for_selector("#run:not([disabled])", timeout=120_000)

    # 고르고 있던 단체가 그대로 열려야 한다. 목록에만 남고 기본 단체로
    # 돌아가 버리면, 그 상태로 산출해 남의 회사 전기값을 끌어온다.
    assert page.input_value("#client-pick") == "살아남을단체"
    assert "살아남을단체" in page.inner_text("#client-pick")

    page.click("#tab-runs")
    assert "새로고침시험" in page.inner_text("#runs-list")
    assert page.inner_text("#client-runs") == "1건"

    # 앞 시험이 만든 단체들도 그대로 있어야 한다.
    names = page.eval_on_selector_all(
        "#client-pick option", "els => els.map(e => e.value)")
    assert "기본 단체" in names

    # 뒷정리 — 다음 시험이 기본 단체에서 시작하도록.
    page.once("dialog", lambda dialog: dialog.accept())
    page.click("#client-remove")
    assert page.input_value("#client-pick") == "기본 단체"
    page.click("#tab-calc")


def test_member_lookup_and_reports(page, tmp_path) -> None:
    """산출 → 사번 조회(연차별 근거) → 계리평가 보고서 미리보기까지.

    사번 조회와 보고서는 **제 탭** 에 있다. PC 본 화면과 탭 구성이 같아야
    두 화면을 오가는 사람이 무엇이 어디 있는지 다시 외우지 않는다.
    """
    from pension.samples import write_sample_pack

    files = write_sample_pack(tmp_path)
    roster = next(p for p in files if p.name == "명부_양식.xlsx")
    assumptions = next(p for p in files if p.name == "기초율_기본값.xlsx")

    page.click("#tab-calc")
    page.set_input_files("#roster", str(roster))
    page.check("#asrc-file")
    page.set_input_files("#assumptions", str(assumptions))
    page.click("#run")
    page.wait_for_selector("#result", state="visible", timeout=180_000)

    # 사번 조회 — 그 사람만 재산출한 연차별 근거가 떠야 한다.
    page.click("#tab-member")
    assert not page.is_visible("#member-empty"), "산출을 마쳤으면 안내문이 빠져야 한다"
    page.fill("#lookup-id", "A0001")
    page.click("#lookup-run")
    # 보고서와 같은 문서 모양으로 곧바로 떠야 한다.
    page.wait_for_selector("#print-dialog[open]", timeout=120_000)
    assert "개인별_산출근거" in page.inner_text("#print-title")
    doc = page.frame_locator("#print-frame").locator("body")
    doc.wait_for(timeout=30_000)
    body = doc.inner_text()
    assert "연차별 계산 근거" in body
    assert "확정급여채무 (DBO)" in body
    page.click("#print-dialog .toolbar button:last-child")
    page.wait_for_selector("#print-dialog[open]", state="detached", timeout=10_000)

    # 퇴직급여 보고서 — 미리보기 iframe 안에 표지 제목이 있어야 한다.
    page.click("#tab-report")
    page.click("#report-sev")
    page.wait_for_selector("#print-dialog[open]", timeout=120_000)
    frame = page.frame_locator("#print-frame")
    assert "확정급여부채 평가보고서" in frame.locator("body").inner_text()
    assert "민감도 분석" in frame.locator("body").inner_text()
    page.click("#print-dialog >> text=닫기")


def test_build_stamp_is_visible_and_matches_the_cache(page) -> None:
    """화면의 빌드 값 = 서비스워커 캐시 이름.

    새 빌드를 올렸는데 옛 캐시가 도는 것을 눈으로 가려낼 수단이다. 치환이
    빠지면 화면에 ``__BUILD__`` 가 그대로 뜨므로 그것도 함께 막는다.
    """
    import re

    stamp = page.inner_text("#build-stamp").strip()
    assert re.fullmatch(r"[0-9a-f]{12}", stamp), f"빌드 값이 이상하다: {stamp}"
    assert f'"pension-{stamp}"' in (DIST / "sw.js").read_text(encoding="utf-8")


def test_the_version_name_is_visible_and_matches_the_release(page) -> None:
    """화면의 버전명 = ``release.json`` 의 버전.

    빌드 값은 열두 자리 해시라 "내 것이 최신인가" 를 눈으로 가릴 수가 없다.
    사람이 읽고 서로 맞대어 볼 수 있는 이름이 따로 있어야 한다.
    """
    import json
    import re

    release = json.loads((DIST / "release.json").read_text(encoding="utf-8"))
    shown = page.inner_text("#app-version").strip()
    assert shown == str(release["version"]), f"버전명이 어긋난다: {shown}"
    assert release["notes"], "업데이트 내역이 비어 있다"
    # 소수점 아래 **두 자리.** 한 자리면 1.9 다음이 1.10 이 되는데, 글자로는
    # 1.10 이 1.9 보다 앞서 보여 어느 쪽이 최신인지 말로 확인할 때 어긋난다.
    assert re.fullmatch(r"\d+\.\d{2}", shown), f"버전은 1.06 꼴이어야 한다: {shown}"


def test_the_update_notice_shows_what_changed_before_it_reloads(page) -> None:
    """[업데이트] — 무엇이 바뀌는지 먼저 보여 주고, 산출이 지워진다고 알린다.

    새 판은 사람이 눌러야 들어온다. 누르기 전에 (1) 어느 버전으로 가는지,
    (2) 무엇이 바뀌는지, (3) 지금 산출이 초기화된다는 사실을 봐야 한다 —
    산출이 몇 분씩 걸리므로 모르고 눌러 날리면 그 시간을 다시 쓴다.
    """
    # 새 판이 올라와 있는 셈 친다 — 서버를 건드리지 않고 화면만 그 상태로 만든다.
    page.evaluate("""() => {
      pendingRelease = {version: "9.9", released: "2099-01-01",
                        notes: ["임원 정년연령 수정 반영"]};
      document.getElementById("update-open").hidden = false;
    }""")
    page.click("#update-open")
    page.wait_for_selector("#update-dialog[open]", timeout=10_000)

    body = page.inner_text("#update-dialog")
    assert "9.9" in body, body
    assert "임원 정년연령 수정 반영" in body, body
    # 대외비·내부 코드가 새어 나가면 안 된다 — 받는 사람은 회사 담당자다.
    assert "sw.js" not in body and "commit" not in body.lower(), body
    assert "초기화" in body and "저장" in body, body

    page.click("#update-later")
    page.wait_for_selector("#update-dialog[open]", state="detached", timeout=10_000)


def test_another_windows_update_does_not_wipe_this_ones_work(page) -> None:
    """다른 창에서 업데이트해도 이 창은 새로고침되지 않아야 한다.

    창을 두 개 열어 두고 한쪽에서 [업데이트] 를 누르면, 서비스워커가 자리를
    넘겨받으면서 **모든 창** 에 그 사실이 통보된다. 예전에는 그 통보를 받는
    즉시 새로고침했는데, 그러면 옆 창에서 몇 분째 돌던 산출이 말없이 사라진다.
    [나중에] 를 눌러 둔 창도 똑같이 당했다.

    지금은 상단 버튼만 뜨고, 그 창에서 직접 누를 때까지 화면은 그대로다.
    """
    # 화면 하나를 여러 시험이 나눠 쓰므로, 앞 시험이 띄워 둔 알림을 먼저 내린다.
    page.evaluate("""() => {
      document.getElementById("update-open").hidden = true;
      window.__stillHere = true;
    }""")
    assert page.locator("#update-open").is_hidden()

    # 다른 창이 업데이트를 마친 셈 친다 — 일꾼이 자리를 넘겨받았다는 통보다.
    page.evaluate(
        "() => navigator.serviceWorker.dispatchEvent(new Event('controllerchange'))")
    page.wait_for_timeout(1_000)

    # 새로고침됐다면 이 표시가 사라진다.
    assert page.evaluate("window.__stillHere === true"), "옆 창이 멋대로 새로고침됐다"
    # 알림을 띄우는 쪽은 `announceUpdate` 가 맡는다. 그 자리가 실제로 화면을
    # 건드릴 수 있는지는 여기서 같이 본다 — 없는 함수를 부르면 통보를 받고도
    # 아무 일이 없다.
    page.evaluate("() => announceUpdate()")
    assert page.locator("#update-open").is_visible()


def test_dashboard_tab_follows_each_run(page, tmp_path) -> None:
    """분석 탭 — 산출할 때마다 그 회차로 다시 그려져야 한다."""
    from pension.samples import write_sample_pack

    files = write_sample_pack(tmp_path)
    roster = next(p for p in files if p.name == "명부_양식.xlsx")
    assumptions = next(p for p in files if p.name == "기초율_기본값.xlsx")

    page.click("#tab-calc")
    page.set_input_files("#roster", str(roster))
    page.check("#asrc-file")
    page.set_input_files("#assumptions", str(assumptions))
    page.click("#run")
    page.wait_for_selector("#result", state="visible", timeout=180_000)

    page.click("#tab-dash")
    page.wait_for_selector("#dash-body", state="visible", timeout=60_000)
    assert page.locator("#dash-strip .stat").count() >= 6
    assert page.locator("#dash-chart rect").count() > 0
    assert page.locator("#dash-trace tr").count() > 2
    first = page.inner_text("#dash-strip")

    # 막대를 누르면 그 해의 계산식이 펼쳐진다.
    page.locator("#dash-chart rect.hit").nth(1).click()
    assert "귀속액" in page.inner_text("#dash-readout")

    # 급부별(정년·중도·사망) 채무. 사유별 지급률이 다른 규정에서는 총액만으로
    # 검산이 안 되므로 갈라 보여야 한다.
    open_section(page, "#sec-dash-group")
    causes = page.inner_text("#dash-causes")
    for word in ("정년", "중도", "사망", "합계"):
        assert word in causes, f"급부별 표에 '{word}' 이(가) 없다"

    # 가정을 흔들면 채무가 움직인다.
    base = page.inner_text("#dash-readout .big")
    page.locator("#dash-chips .chip").nth(1).click()
    assert page.inner_text("#dash-readout .big") != base

    # 사번을 지정해 다른 사람을 해부한다.
    page.fill("#dash-emp", "A0002")
    page.click("#dash-emp-go")
    assert "A0002" in page.inner_text("#dash-who")

    # 기준일을 옮겨 다시 산출하면 분석 값도 그 회차로 갈린다.
    # (앞으로 당기면 퇴직자 퇴사일이 기준일보다 늦어 검증에 걸린다 — 뒤로 민다.)
    open_calc(page, "sec-dates")
    page.fill("#base_date", "2026-06-30")
    page.click("#run")
    page.wait_for_selector("#result", state="visible", timeout=180_000)
    page.click("#tab-dash")
    assert page.inner_text("#dash-when").strip() == "기준일 2026-06-30"
    assert page.inner_text("#dash-strip") != first


def test_new_disclosures_flow_into_the_dashboard(page, tmp_path) -> None:
    """재측정 분해 · 자산인식상한 · 차년도 예측 · 장기급여 증감표까지 한 바퀴."""
    from pension.samples import write_sample_pack

    files = write_sample_pack(tmp_path)
    roster = next(p for p in files if p.name == "명부_양식.xlsx")
    assumptions = next(p for p in files if p.name == "기초율_기본값.xlsx")

    page.click("#tab-calc")
    page.set_input_files("#roster", str(roster))
    page.check("#asrc-file")
    page.set_input_files("#assumptions", str(assumptions))

    # 가정별 분해는 전기 **기초율** 이 있어야 돈다. 한 번 산출해 저장한 뒤
    # 그것을 전기로 연결한다 — 실제 사용 순서와 같다.
    page.click("#run")
    page.wait_for_selector("#result", state="visible", timeout=180_000)
    page.fill("#run-name", "전기연결시험")
    page.click("#run-save")
    page.wait_for_selector("text=저장했습니다", timeout=30_000)

    open_calc(page, "sec-options")
    page.check("#split_remeasurement")
    open_calc(page, "sec-prior")
    page.select_option("#prior-run", "전기연결시험")
    page.fill("#prior_longterm_dbo", "40000000")
    open_calc(page, "sec-assets")
    page.fill("#asset_opening", "200000000")
    page.fill("#asset_contributions", "50000000")
    page.fill("#asset_closing", "900000000")     # 초과적립으로 상한을 물린다
    page.fill("#asset_ceiling", "50000000")
    page.fill("#expected_contributions", "60000000")

    page.click("#run")
    page.wait_for_selector("#result", state="visible", timeout=180_000)
    summary = page.inner_text("#summary")
    assert "자산인식상한 차감액" in summary
    assert "장기급여 재측정(당기손익)" in summary

    page.click("#tab-dash")
    page.wait_for_selector("#dash-body", state="visible", timeout=60_000)

    open_section(page, "#sec-dash-roll")
    roll = page.inner_text("#dash-roll")
    assert "가정변경효과 분해" in roll
    for name in ("사망률", "퇴직률", "임금상승률", "할인율"):
        assert name in roll
    assert "자산인식상한" in roll

    open_section(page, "#sec-dash-next")
    nxt = page.inner_text("#dash-next")
    assert "차년도 예상 퇴직급여 비용" in nxt
    assert "차년도 확정급여채무 예측" in nxt
    assert "예상 부담금 납입액" in nxt

    open_section(page, "#sec-dash-lt")
    assert "재측정요소 (당기손익)" in page.inner_text("#dash-lt")

    # 인쇄를 누르면 접힌 구획이 모두 펼쳐진다(인쇄 대화상자는 막아 둔다).
    page.evaluate("window.print = () => { window.__printed = true; }")
    page.click("#dash-print")
    page.wait_for_timeout(400)
    assert page.evaluate("window.__printed") is True
    assert page.evaluate(
        "[...document.querySelectorAll('#page-dash details.section')]"
        ".every((d) => d.open)")
    assert "기준일" in page.inner_text("#dash-print-title")


def test_the_pc_local_server_boots_the_engine(browser) -> None:
    """PC 의 [전체 기능 화면] — 로컬 서버로 낸 웹앱이 실제로 뜨는지.

    ``.wasm`` 의 MIME 이나 폴더 배치가 틀리면 화면만 뜨고 엔진이 죽는다.
    파일을 그대로 복사한 것이라 조용히 어긋나기 쉬운 대목이다.
    """
    import threading

    from pension import localapp

    server = localapp.serve(DIST)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    page = browser.new_page()
    try:
        page.goto(f"http://127.0.0.1:{server.server_address[1]}/index.html")
        # 엔진이 다 뜨면 [산출 실행] 의 잠금이 풀린다.
        page.wait_for_selector("#run:not([disabled])", timeout=120_000)
        assert page.locator("#tab-dash").count() == 1
    finally:
        page.close()
        server.shutdown()
        server.server_close()


def test_storage_persistence_is_requested_and_reported(page) -> None:
    """자료가 지워질 수 있는지를 화면이 말해 줘야 한다.

    이 앱의 산출 내역은 브라우저 저장소에 있고, 브라우저는 공간이 모자라면
    **말없이 지운다.** 영구 저장을 요청해 두면 그 대상에서 빠지는데, 받아들여
    졌는지는 기기가 정한다. 보장되지 않는데 보장된 줄 알고 내보내기를 건너
    뛰는 것이 제일 나쁘므로, 결과를 그대로 적어 둔다.
    """
    page.click("#tab-lib")
    open_section(page, "#lib-backup")
    note = page.locator("#storage-note")
    page.wait_for_function(
        "() => document.querySelector('#storage-note').textContent.trim().length > 0",
        timeout=30_000)

    # 허용됐든 아니든 **내보내기를 권해야** 한다. 허용은 이 기기에서만이고,
    # 기기를 바꾸면 어차피 파일이 있어야 되살릴 수 있다.
    text = note.inner_text()
    assert "내보내" in text, text
    assert note.get_attribute("class") in ("hint ok-text", "hint bad-text", "hint")

    # 허용되지 않았으면 눈에 띄어야 한다. 회색 안내문에 섞이면 아무도 안 읽는다.
    if "허용되지 않" in text:
        assert note.get_attribute("class") == "hint bad-text"


def test_prior_roster_comparison_runs_before_the_valuation(page, tmp_path) -> None:
    """전기 명부와 맞대어 보기 — 산출 전에 잡아야 할 것을 잡는지.

    당기 명부만 보면 멀쩡한데 전기와 나란히 놓아야 드러나는 것이 있다. 여기서는
    생년월일이 바뀐 사람을 심어 두고, 그것이 **빨간 쪽** 으로 나오는지 본다.
    """
    import openpyxl

    from pension.samples import write_sample_pack

    files = write_sample_pack(tmp_path)
    roster = next(p for p in files if p.name == "명부_양식.xlsx")
    assumptions = next(p for p in files if p.name == "기초율_기본값.xlsx")

    # 전기로 쓸 산출을 하나 만들어 저장한다.
    page.click("#tab-calc")
    page.set_input_files("#roster", str(roster))
    page.check("#asrc-file")
    page.set_input_files("#assumptions", str(assumptions))
    page.click("#run")
    page.wait_for_selector("#result", state="visible", timeout=180_000)
    page.fill("#run-name", "2312 맞대기시험")
    page.click("#run-save")
    page.wait_for_selector("text=저장했습니다", timeout=30_000)

    # 당기 명부 — 첫 사람의 생년월일만 바꾼다.
    changed = tmp_path / "당기명부.xlsx"
    book = openpyxl.load_workbook(roster)
    sheet = book["재직자명부"]
    # 열은 머리글로 찾는다. 번호를 못박으면 양식이 한 칸만 움직여도 빈 칸을
    # 고치게 되고, 그러면 '바뀐 것이 없다' 로 조용히 통과한다.
    birth = next(c.column for c in sheet[HEADER_ROW] if c.value == "생년월일")
    sheet.cell(FIRST_DATA_ROW, birth, "1955-01-01")
    book.save(changed)

    page.click("#tab-calc")
    page.set_input_files("#roster", str(changed))
    open_calc(page, "sec-prior")
    page.select_option("#prior-run", "2312 맞대기시험")
    page.click("#prior-check")

    page.wait_for_function(
        "() => document.querySelector('#prior-check-status').textContent.includes('건')",
        timeout=60_000)
    assert "확인이 필요한" in page.inner_text("#prior-check-status")
    result = page.inner_text("#prior-check-result")
    assert "생년월일" in result
    assert "1955-01-01" in result


def test_intro_dialog_points_at_the_library(browser, app_url) -> None:
    """앱을 처음 열면 [자료실] 안내가 뜨고, 끄면 다시 뜨지 않아야 한다.

    받은 명부가 없는 사람은 첫 화면에서 더 갈 곳이 없다. 양식·시험 명부·문서가
    이미 들어 있다는 것을 눌러 보기 전에는 알 수 없기 때문이다.
    """
    fresh = browser.new_context()
    page = fresh.new_page()
    page.goto(app_url)

    page.wait_for_selector("#intro-dialog[open]", timeout=30_000)
    text = page.inner_text("#intro-dialog")
    for word in ("자료실", "명부 양식", "시험 명부 만들기", "문서", "금리표"):
        assert word in text, f"안내에 '{word}' 가 없다"

    # [자료실 열기] 는 그 탭으로 데려가야 한다.
    page.click("#intro-go")
    page.wait_for_selector("#intro-dialog", state="hidden", timeout=10_000)
    assert page.locator("#page-lib").is_visible()

    # 체크하지 않고 껐으니 다시 열면 또 떠야 한다.
    page.reload()
    page.wait_for_selector("#intro-dialog[open]", timeout=30_000)
    page.check("#intro-hide")
    page.click("#intro-close")
    page.reload()
    page.wait_for_selector("#run:not([disabled])", timeout=120_000)
    assert not page.locator("#intro-dialog[open]").count()

    # 같은 내용이 물음표(사용설명서) 안에도 있어야 한다.
    page.click("#help-open")
    page.wait_for_function(
        "() => document.getElementById('help-body').textContent.includes('자료실')",
        timeout=30_000)
    assert "수령한 명부가 없어도" in page.inner_text("#help-body")

    page.close()
    fresh.close()


def test_feature_pack_shows_a_table_not_a_wall_of_text(page) -> None:
    """특이사항 한 벌의 결과는 **표** 로 나와야 한다.

    자릿수를 맞춘 고정폭 글자표를 그대로 띄우면 좁은 화면에서 한 줄이 세 줄로
    끊겨, 어느 숫자가 어느 명부 것인지 알 수 없다.
    """
    page.click("#tab-lib")
    open_section(page, "#lib-gen")
    unlock_teaching_material(page)
    # 채무까지 재면 명부 열한 벌을 산출해야 해서 브라우저에서 몇 분 걸린다.
    # 화면이 표로 그려지는지만 보면 되므로 여기서는 끄고 만든다.
    page.uncheck("#feat-measure")
    page.click("#feat-run")
    page.wait_for_selector("#feat-cases fieldset", timeout=180_000)

    page.click("#feat-report")
    page.wait_for_selector("#feat-dialog[open]", timeout=10_000)
    body = page.inner_text("#feat-body")
    assert "특이사항 하나" in body
    for word in ("휴직차감", "중간정산", "DC전환"):
        assert word in body, word
    # 명부마다 카드가 하나씩. 기준 + 특이사항들.
    assert page.locator("#feat-dialog .feat-card").count() >= 11

    page.click("#feat-dialog button")
    page.wait_for_selector("#feat-dialog", state="hidden", timeout=10_000)
    page.click("#tab-calc")     # 다음 시험이 산출 화면에서 시작하도록 돌려 놓는다


def test_it_runs_on_a_galaxy_phone(browser, app_url, tmp_path) -> None:
    """갤럭시(안드로이드 크로미움, 360px 세로 화면)에서 끝까지 돌아야 한다.

    같은 프로그램을 동료가 갤럭시로 연다. 아이패드에 맞춰 만든 화면이 좁은
    안드로이드에서 옆으로 밀리거나 엔진이 안 뜨면, 그 사람은 쓸 수가 없다.
    """
    from pension.samples import write_sample_pack

    galaxy = browser.new_context(
        viewport={"width": 360, "height": 800}, device_scale_factor=3,
        is_mobile=True, has_touch=True,
        user_agent="Mozilla/5.0 (Linux; Android 14; SM-S918N) AppleWebKit/537.36 "
                   "(KHTML, like Gecko) Chrome/126.0.0.0 Mobile Safari/537.36",
    )
    page = galaxy.new_page()
    broken: list[str] = []
    page.on("pageerror", lambda e: broken.append(str(e)))
    page.goto(app_url)
    dismiss_intro(page)
    page.wait_for_selector("#run:not([disabled])", timeout=180_000)

    def overflow() -> int:
        return page.evaluate(
            "document.documentElement.scrollWidth"
            " - document.documentElement.clientWidth")

    # 화면 전체가 좌우로 밀리면 안 된다. 넓은 표는 표 안에서만 밀린다.
    for tab in ("tab-calc", "tab-edit", "tab-dash", "tab-report",
                "tab-member", "tab-runs", "tab-lib"):
        page.click(f"#{tab}")
        assert overflow() <= 1, f"{tab} 에서 화면이 {overflow()}px 옆으로 밀린다"

    # 입력칸 글씨는 16px 이어야 한다. 그보다 작으면 아이폰이 포커스 순간
    # 화면을 확대하고 스스로 돌아오지 않는다 — 안드로이드는 확대하지 않지만
    # 같은 화면을 두 기기가 나눠 쓰므로 여기서 함께 못박는다.
    page.click("#tab-calc")
    for box in ("#base_date", "#period_start", "#run-name"):
        size = page.evaluate(
            f"getComputedStyle(document.querySelector('{box}')).fontSize")
        assert size == "16px", f"{box} 가 {size} 라 아이폰에서 확대된다"

    # 산출이 실제로 끝까지 돈다.
    files = write_sample_pack(tmp_path)
    roster = next(p for p in files if p.name == "명부_양식.xlsx")
    assumptions = next(p for p in files if p.name == "기초율_기본값.xlsx")
    page.set_input_files("#roster", str(roster))
    page.check("#asrc-file")
    page.set_input_files("#assumptions", str(assumptions))
    page.click("#run")
    page.wait_for_selector("#result", state="visible", timeout=300_000)
    assert "확정급여채무" in page.inner_text("#summary")
    assert overflow() <= 1, "결과가 나온 뒤 화면이 옆으로 밀린다"

    assert not broken, f"자바스크립트 오류: {broken}"
    page.close()
    galaxy.close()


def test_no_phone_width_pushes_the_screen_sideways(browser, app_url) -> None:
    """갤럭시 기종별 폭에서 화면이 좌우로 밀리지 않아야 한다.

    레이아웃을 정하는 것은 물리 화소가 아니라 **CSS 폭** 이다. S23 은
    1080×2340 이지만 배율이 3 이라 브라우저가 보는 폭은 360 이고, Ultra 는
    412 다. 삼성 기기는 [디스플레이 크기] 를 키우면 이 폭이 더 줄어 320 까지
    내려간다 — 눈이 어두워 글씨를 키운 사람이 곧 가장 좁은 화면을 쓴다.
    """
    for width, label in (
        (320, "화면을 크게 설정한 S23"),
        (360, "S23 · S24 · S25 기본"),
        (384, "S23+ · S24+"),
        (412, "S23 Ultra · S24 Ultra"),
    ):
        context = browser.new_context(
            viewport={"width": width, "height": 780},
            is_mobile=True, has_touch=True)
        page = context.new_page()
        page.goto(app_url)
        dismiss_intro(page)
        page.wait_for_selector("#run:not([disabled])", timeout=180_000)

        for tab in ("tab-calc", "tab-edit", "tab-dash", "tab-report",
                    "tab-member", "tab-runs", "tab-lib"):
            page.click(f"#{tab}")
            over = page.evaluate(
                "document.documentElement.scrollWidth"
                " - document.documentElement.clientWidth")
            assert over <= 1, f"{label}({width}px) 의 {tab} 이 {over}px 밀린다"

        page.close()
        context.close()


def test_sideways_scrolling_stops_inside_the_box_it_started_in(
    browser, app_url
) -> None:
    """옆으로 구르는 상자는 그 힘을 바깥으로 넘기지 않아야 한다.

    문서가 화면보다 넓지 않아도 화면이 밀릴 수 있다. 탭 줄·넓은 표처럼 스스로
    옆으로 구르는 상자 안에서 끝까지 민 뒤 계속 밀면, 브라우저가 남은 힘을
    부모에게 넘긴다(스크롤 연쇄). 넘겨받을 곳이 문서뿐이라 화면이 통째로
    밀리고 오른쪽에 검은 여백이 드러난다 — 아이폰에서 그렇게 보였다.

    문서 폭만 재는 시험으로는 이걸 못 잡는다. 폭은 멀쩡했다. 그래서 **구르는
    상자마다** 연쇄를 끊어 두었는지를 직접 본다.
    """
    context = browser.new_context(
        viewport={"width": 390, "height": 844}, is_mobile=True, has_touch=True)
    page = context.new_page()
    page.goto(app_url)
    dismiss_intro(page)
    page.wait_for_selector("#run:not([disabled])", timeout=180_000)

    leaky = []
    for tab in ("tab-calc", "tab-edit", "tab-dash", "tab-report",
                "tab-member", "tab-runs", "tab-lib"):
        page.click(f"#{tab}")
        # 접힌 부분 안에도 표가 있다. 펴 놓아야 폭과 성질이 잡힌다.
        page.evaluate("() => document.querySelectorAll('details').forEach(d => d.open = true)")
        page.wait_for_timeout(300)
        leaky += page.evaluate("""(tab) => {
          const out = [];
          for (const el of document.querySelectorAll("*")) {
            const style = getComputedStyle(el);
            const scrolls = style.overflowX === "auto" || style.overflowX === "scroll";
            if (!scrolls) continue;
            if (el.scrollWidth <= el.clientWidth) continue;   // 실제로 구르는 것만
            const chain = style.overscrollBehaviorX;
            if (chain !== "contain" && chain !== "none") {
              out.push(`${tab}: ${el.tagName.toLowerCase()}#${el.id}.${el.className} → ${chain}`);
            }
          }
          return out;
        }""", tab)

    page.close()
    context.close()
    assert leaky == [], "옆으로 구르는 힘이 화면 전체로 새어 나간다: " + "; ".join(leaky)


def test_the_home_screen_icon_fits_android(page) -> None:
    """안드로이드 홈 화면 아이콘이 흰 판에 얹히거나 잘리지 않아야 한다.

    안드로이드는 아이콘을 기기 모양대로 **잘라 낸다.** maskable 아이콘이 없으면
    원본을 흰 배경에 축소해 얹어, 남의 앱들과 나란히 두면 그것만 튄다.
    """
    import json

    manifest = json.loads(
        (DIST / "manifest.webmanifest").read_text(encoding="utf-8"))
    sizes = {icon["sizes"] for icon in manifest["icons"]}
    assert "192x192" in sizes, "안드로이드가 먼저 찾는 192 가 없다"
    purposes = {icon.get("purpose") for icon in manifest["icons"]}
    assert "maskable" in purposes

    for icon in manifest["icons"]:
        assert (DIST / icon["src"]).is_file(), icon["src"]

    # maskable 아이콘의 그림은 **가운데 80% 안** 에 들어와야 한다. 안드로이드가
    # 기기 모양대로 잘라 내므로, 가장자리까지 그리면 막대 끝이 잘린다.
    maskable = next(i for i in manifest["icons"] if i.get("purpose") == "maskable")
    size, pixels = _read_png(DIST / maskable["src"])
    stride = size * 3 + 1
    navy = (31, 56, 100)
    marks = [
        (x, y)
        for y in range(size)
        for x in range(size)
        if tuple(pixels[y * stride + 1 + x * 3: y * stride + 4 + x * 3]) != navy
    ]
    assert marks, "아이콘이 바탕색 한 가지뿐이다"
    low, high = size * 0.1, size * 0.9
    assert low <= min(x for x, _ in marks) and max(x for x, _ in marks) <= high
    assert low <= min(y for _, y in marks) and max(y for _, y in marks) <= high


def _read_png(path) -> tuple[int, bytes]:
    """가로세로 같은 24비트 PNG 를 (한 변, 원본 픽셀) 로. 외부 라이브러리 없이."""
    import struct
    import zlib

    raw = path.read_bytes()
    pos, data, size = 8, b"", 0
    while pos < len(raw):
        length = struct.unpack(">I", raw[pos:pos + 4])[0]
        kind = raw[pos + 4:pos + 8]
        body = raw[pos + 8:pos + 8 + length]
        if kind == b"IHDR":
            size = struct.unpack(">II", body[:8])[0]
        elif kind == b"IDAT":
            data += body
        pos += 12 + length
    return size, zlib.decompress(data)


def test_both_screens_have_the_same_tabs(page) -> None:
    """아이패드 화면과 PC 본 화면의 탭이 같아야 한다.

    한 사람이 두 화면을 오간다. 탭 구성이 갈라지면 무엇이 어디 있는지 두 번
    외워야 하고, 한쪽에만 있는 기능을 '없는 것' 으로 여기게 된다. 그래서 이름과
    순서를 여기서 못 박는다 — 한쪽을 고치면 이 시험이 걸린다.
    """
    import re

    web = page.eval_on_selector_all(
        "nav.tabs button", "els => els.map(e => e.textContent.trim())")
    assert web == ["산출", "산출가정 입력", "분석", "계리평가 보고서",
                   "사번 조회", "산출 내역", "자료실"]

    # PC 쪽은 소스에서 읽는다. 임포트하면 tkinter 가 필요한데, 이 시험을 돌리는
    # 곳에 화면이 없을 수 있다 — 탭 이름을 맞대어 보는 데 창까지 띄울 일은 아니다.
    source = (Path(__file__).resolve().parent.parent
              / "src" / "pension" / "desk" / "app.py").read_text(encoding="utf-8")
    block = source.split("self._pages = {", 1)[1].split("}", 1)[0]
    desk = re.findall(r'"([^"]+)":', block)
    assert desk == web


def test_a_new_build_reaches_a_device_that_already_installed_the_app(
    browser, app_url, tmp_path
) -> None:
    """앱을 이미 깔아 둔 기기에 **새 판이 실제로 닿는지.**

    이것이 안 되면 화면을 아무리 고쳐도 쓰는 사람에게는 아무 일도 일어나지
    않는다. 실제로 휴대폰이 옛 화면을 계속 띄웠고, 원인은 화면 파일
    (index.html·app.js·app.css)에 빌드 값이 붙어 있지 않은 채 캐시 우선으로
    나가고 있었던 것이다 — 한 번 캐시에 들어가면 그 뒤로 네트워크를 보지 않는다.

    여기서는 서비스워커를 등록한 기기(=브라우저 컨텍스트)를 만들어 두고,
    서버가 내주는 app.css 를 바꾼 뒤 다시 열어 **바뀐 것이 보이는지** 본다.
    """
    import functools
    import http.server
    import shutil
    import threading

    # dist 를 복사해 서버로 띄운다. 원본을 건드리지 않고 '새 배포' 를 흉내낸다.
    served = tmp_path / "배포본"
    shutil.copytree(DIST, served)
    handler = functools.partial(http.server.SimpleHTTPRequestHandler,
                                directory=str(served))
    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    url = f"http://127.0.0.1:{server.server_address[1]}/index.html"

    context = browser.new_context()
    try:
        first = context.new_page()
        first.goto(url)
        first.wait_for_selector("#run:not([disabled])", timeout=120_000)
        # 일꾼은 자리를 **곧바로 넘겨받지 않는다**(산출 도중에 판이 바뀌지 않게).
        # 그래서 첫 방문은 통제되지 않고, 한 번 더 열어야 통제가 시작된다.
        first.reload()
        first.wait_for_selector("#run:not([disabled])", timeout=120_000)
        first.wait_for_function(
            "() => navigator.serviceWorker.controller !== null", timeout=60_000)
        before = first.evaluate(
            "() => getComputedStyle(document.querySelector('nav.tabs button')).paddingTop")
        first.close()

        # 새 판을 올린다 — 탭 위 여백만 눈에 띄게 바꾼다.
        css = served / "app.css"
        css.write_text(css.read_text(encoding="utf-8")
                       + "\nnav.tabs button { padding-top: 41px; }\n",
                       encoding="utf-8")

        second = context.new_page()
        second.goto(url)
        second.wait_for_selector("#run:not([disabled])", timeout=120_000)
        after = second.evaluate(
            "() => getComputedStyle(document.querySelector('nav.tabs button')).paddingTop")
        second.close()
    finally:
        context.close()
        server.shutdown()
        server.server_close()

    assert before != "41px", "시작부터 41px 이면 이 시험이 아무것도 못 본다"
    assert after == "41px", (
        f"새로 올린 app.css 가 기기에 닿지 않았습니다 (그대로 {after}). "
        "화면 파일은 네트워크를 먼저 봐야 합니다."
    )


def test_help_opens_over_the_screen(page) -> None:
    """물음표를 누르면 화면을 떠나지 않고 사용설명서를 읽는다.

    산출 도중에 물어볼 것이 생기는데 다른 창으로 나가면 입력하던 것을 잃는다.
    겹쳐 뜨는지, 닫으면 하던 화면으로 그대로 돌아오는지 본다.
    """
    page.click("#tab-calc")     # 앞 시험이 어느 탭에 있었든 여기서 시작한다
    page.fill("#base_date", "2025-12-31")          # 하던 입력
    page.click("#help-open")
    page.wait_for_selector("#help-body h1", timeout=15_000)
    assert "사용설명서" in page.inner_text("#help-body h1")
    assert page.locator("#help-body table").count() >= 3   # 표가 살아 있다
    assert "홈 화면에 추가" in page.inner_text("#help-body")   # 웹앱용 문서다

    # 열자마자 입력칸을 잡으면 아이폰이 화면을 확대해 버린다. 읽으러 연
    # 사람에게는 그것이 방해라 자동 포커스를 두지 않는다.
    assert page.evaluate("document.activeElement.id") != "help-find"

    page.fill("#help-find", "확정급여채무")
    page.wait_for_selector("#help-body mark", timeout=5_000)
    assert page.locator("#help-body mark").count() > 0

    page.click("#help-close")
    assert page.locator("#help").is_hidden()
    assert page.input_value("#base_date") == "2025-12-31"


def test_it_runs_with_the_network_cut(browser, tmp_path) -> None:
    """한 번 열어 둔 기기는 **인터넷 없이도** 돌아야 한다.

    이 앱을 쓰는 자리는 사내망이거나 회선이 나쁜 곳이다. 화면만 뜨고 엔진이
    없으면 아무 소용이 없는데, 실제로 그랬다 — 서비스워커가 화면 파일만
    캐시하고 런타임(14MB)·휠은 한 번도 담지 않아, 회선을 끊으면
    `pyodide.asm.wasm` 부터 실패했다. 화면이 다 뜬 뒤 무거운 것을 하나씩
    채우게 고쳤고, 여기서 **서버를 완전히 내린 채** 다시 열어 확인한다.
    """
    import http.server
    import shutil
    import threading

    site = tmp_path / "site"
    shutil.copytree(DIST, site)
    handler = functools.partial(
        http.server.SimpleHTTPRequestHandler, directory=str(site))
    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    url = f"http://127.0.0.1:{server.server_address[1]}/index.html"
    stopped = False
    context = browser.new_context()
    try:
        page = context.new_page()
        page.goto(url)
        page.wait_for_selector("#run:not([disabled])", timeout=180_000)
        # 오프라인 채비가 끝날 때까지 — 런타임과 휠이 모두 캐시에 담겨야 한다.
        page.wait_for_function("""async () => {
            const names = await caches.keys();
            if (!names.length) return false;
            const box = await caches.open(names[0]);
            const urls = (await box.keys()).map((r) => r.url);
            return urls.some((u) => u.includes('pyodide.asm.wasm'))
                && urls.some((u) => u.includes('python_stdlib'))
                && urls.filter((u) => u.includes('/wheels/')).length >= 4;
        }""", timeout=180_000)
        page.close()

        # 서버를 내린다 — 캐시에 없는 것은 이제 어디서도 받을 수 없다.
        server.shutdown()
        server.server_close()
        stopped = True
        context.set_offline(True)

        offline = context.new_page()
        offline.goto(url)
        offline.wait_for_selector("#run:not([disabled])", timeout=240_000)
        assert "준비 완료" in offline.inner_text("#status")
        # 엔진이 실제로 도는지 — 화면만 뜬 것과 구별한다.
        assert offline.evaluate("() => Boolean(META && META.sheets.length)")
    finally:
        context.close()
        if not stopped:
            server.shutdown()
            server.server_close()


def test_a_new_build_replaces_the_old_one(browser, tmp_path) -> None:
    """새 판을 올렸을 때 실제로 그 판이 뜨는지.

    실제로 겪은 일이다 — 두 판을 올렸는데 휴대폰은 그 전 판을 계속 띄웠다.
    원인은 서비스워커가 설치 때 런타임까지(14MB) 한꺼번에 받게 되어 있어서,
    그중 하나만 실패하면 새 일꾼이 통째로 설치되지 않고 옛 일꾼이 그대로 남는
    것이었다. 여기서는 **새 판을 올린 상황을 그대로 만들어** 확인한다.
    """
    import http.server
    import re
    import shutil
    import threading

    site = tmp_path / "site"
    shutil.copytree(DIST, site)

    handler = functools.partial(
        http.server.SimpleHTTPRequestHandler, directory=str(site))
    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    url = f"http://127.0.0.1:{server.server_address[1]}/index.html"
    try:
        context = browser.new_context()
        page = context.new_page()
        page.goto(url)
        page.wait_for_selector("#build-stamp:not(:empty)", timeout=30_000)
        # 첫 방문은 일꾼이 자리를 넘겨받기 전이다 — 한 번 더 열어야 통제된다.
        page.reload()
        page.wait_for_selector("#build-stamp:not(:empty)", timeout=30_000)
        page.wait_for_function(
            "() => navigator.serviceWorker.controller !== null", timeout=30_000)
        before = page.inner_text("#build-stamp").strip()

        # 새 판을 올린 셈 친다 — 화면 파일이 바뀌고 캐시 이름이 갈린다.
        after = "f" * 12
        for name in ("index.html", "sw.js"):
            target = site / name
            target.write_text(
                target.read_text(encoding="utf-8").replace(before, after),
                encoding="utf-8")
        app = site / "app.js"
        app.write_text(app.read_text(encoding="utf-8") + "\n// 새 판\n",
                       encoding="utf-8")

        # 다시 열면 새 일꾼이 들어와 자리를 넘겨받아야 한다.
        for _ in range(3):
            page.goto(url)
            page.wait_for_selector("#build-stamp:not(:empty)", timeout=30_000)
            if page.inner_text("#build-stamp").strip() == after:
                break
            page.wait_for_timeout(1_000)
        assert page.inner_text("#build-stamp").strip() == after, (
            f"옛 판이 그대로 뜬다: {page.inner_text('#build-stamp')}")

        # 새 일꾼이 무거운 것을 설치 때 받지 않는지 — 여기가 막히면 되풀이된다.
        worker = (site / "sw.js").read_text(encoding="utf-8")
        precache = re.search(r"const PRECACHE = (\[.*?\]);", worker, re.S).group(1)
        assert "pyodide/" not in precache
        assert "wheels/" not in precache
        # 새 일꾼이 곧바로 자리를 빼앗으면 산출 도중에 화면과 엔진이 엇갈린다.
        # 주석에는 그 낱말이 나오므로 주석을 걷어내고 **실제 호출** 만 본다.
        code = re.sub(r"/\*.*?\*/", "", worker, flags=re.S)
        code = re.sub(r"//.*", "", code)
        # 넘겨받기는 **사람이 [업데이트] 를 눌렀을 때만** 일어나야 한다. 그래서
        # 호출은 한 군데뿐이고, 그 자리는 'take-over' 신호를 받은 안쪽이다.
        assert code.count("skipWaiting(") == 1, "넘겨받기를 여러 곳에서 부른다"
        where = code.index("skipWaiting(")
        assert "take-over" in code[max(0, where - 300):where], (
            "사람이 누르지 않았는데 자리를 넘겨받는다")
        # 설치·활성화 쪽에서 부르면 산출 중에 판이 갈린다. 호출은 화면이 보낸
        # 신호를 받는 자리(message) 안쪽에만 있어야 한다.
        assert where > code.index('addEventListener("message"')
        # claim 은 반대로 **있어야** 한다 — 첫 방문이 통제되지 않으면 그 방문에서
        # 받은 런타임·휠이 캐시에 담기지 않아 오프라인에서 엔진이 없다.
        assert "clients.claim(" in code
        context.close()
    finally:
        server.shutdown()
        server.server_close()


def test_every_template_downloads_from_the_library(page, tmp_path) -> None:
    """양식을 화면에서 받을 수 있어야 한다.

    회사에 명부를 요청할 때마다 양식 파일을 어디선가 찾아 붙여 보내면, 프로그램이
    바뀐 뒤에도 옛 양식이 돌아다닌다. 눌러서 받는 것이 늘 지금 것이다.
    """
    page.click("#tab-lib")
    page.wait_for_selector("#template-list .lib-line", timeout=30_000)
    buttons = page.locator("#template-list button")
    # 개수를 박아 두면 양식을 하나 뺄 때마다 여기가 깨진다. 목록이 내주는
    # 것을 **전부** 받아 보는 것이 이 시험의 뜻이다.
    assert buttons.count() >= 1

    for index in range(buttons.count()):
        with page.expect_download(timeout=120_000) as got:
            buttons.nth(index).click()
        made = got.value
        assert made.suggested_filename.endswith(".xlsx")
        saved = tmp_path / made.suggested_filename
        made.save_as(str(saved))
        assert saved.stat().st_size > 4_000, made.suggested_filename


def test_the_manuals_download_from_the_library(page, tmp_path) -> None:
    """자료실의 문서를 화면에서 받을 수 있어야 한다.

    받는 사람에게 그대로 보낼 파일이다. 화면에서 읽는 것(물음표)과 **같은
    원본** 이라, 한쪽만 고쳐져 갈라지는 일이 없다.

    이름을 여기 다시 적지 않고 화면에 놓인 것을 **전부** 받아 본다 — 문서를
    하나 더할 때 시험을 고치는 것을 잊으면, 그 문서만 조용히 빠진다.
    """
    page.click("#tab-lib")
    open_section(page, "#lib-docs")

    names = page.eval_on_selector_all(
        "#lib-docs button[data-doc]", "list => list.map((b) => b.dataset.doc)")
    assert len(names) >= 2, names

    for name in names:
        # 화면에서 먼저 읽고(보고서와 같은 창), 그 창에서 원본을 받는다.
        page.click(f'#lib-docs button[data-doc="{name}"]')
        page.wait_for_selector("#print-dialog[open]", timeout=60_000)
        assert page.locator("#print-title").inner_text().strip() == name
        with page.expect_download(timeout=60_000) as got:
            page.click("#print-source")
        made = got.value
        saved = tmp_path / f"{name}.md"
        made.save_as(str(saved))
        text = saved.read_text(encoding="utf-8")
        assert len(text.splitlines()) > 100, name
        # 주소가 실려 나가면 문서만 돌아다녀도 주소까지 같이 돈다.
        assert "http" not in text, name
        page.click("#print-dialog .toolbar button:last-child")
        page.wait_for_selector("#print-dialog[open]", state="detached", timeout=10_000)

    # 물음표 안의 설명과 같은 문서여야 한다.
    manual = (tmp_path / "사용설명서.md").read_text(encoding="utf-8")
    page.click("#help-open")
    page.wait_for_selector("#help:not([hidden])", timeout=10_000)
    shown = page.inner_text("#help-body")
    for heading in ("명부 만들기", "검증 읽는 법", "기억할 것"):
        assert heading in manual and heading in shown, heading
