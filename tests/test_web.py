"""아이패드·다른 PC 에서 쓰는 웹 화면.

브라우저 없이 HTTP 요청으로 검증한다 — 화면이 아니라 규약(업로드 → 산출 →
내려받기)이 맞는지를 본다.
"""

from __future__ import annotations

import urllib.request
import uuid

import pytest

from pension.web import PensionWebServer, _parse_multipart, _run_in_memory


@pytest.fixture
def server_url():
    server = PensionWebServer(("127.0.0.1", 0))
    _run_in_memory(server)
    yield f"http://127.0.0.1:{server.server_address[1]}"
    server.shutdown()
    server.server_close()


def _multipart(fields: dict[str, bytes], files: dict[str, tuple[str, bytes]]):
    boundary = uuid.uuid4().hex
    chunks = []
    for name, value in fields.items():
        chunks.append(
            f'--{boundary}\r\nContent-Disposition: form-data; name="{name}"\r\n\r\n'.encode()
            + value + b"\r\n"
        )
    for name, (filename, payload) in files.items():
        chunks.append(
            f'--{boundary}\r\nContent-Disposition: form-data; name="{name}"; '
            f'filename="{filename}"\r\n'
            f"Content-Type: application/octet-stream\r\n\r\n".encode()
            + payload + b"\r\n"
        )
    chunks.append(f"--{boundary}--\r\n".encode())
    return b"".join(chunks), f"multipart/form-data; boundary={boundary}"


def _post(url: str, body: bytes, content_type: str) -> tuple[int, bytes]:
    request = urllib.request.Request(
        url, data=body, headers={"Content-Type": content_type}, method="POST"
    )
    with urllib.request.urlopen(request) as response:
        return response.status, response.read()


class TestMultipartParser:
    def test_round_trip(self) -> None:
        body, content_type = _multipart(
            {"prior_dbo": b"1,000"}, {"roster": ("a.xlsx", b"PK\x03\x04data")}
        )
        parts = _parse_multipart(body, content_type)
        assert parts["prior_dbo"].body == b"1,000"
        assert parts["roster"].filename == "a.xlsx"
        assert parts["roster"].body == b"PK\x03\x04data"

    def test_binary_payload_with_crlf_survives(self) -> None:
        """엑셀 파일에는 \\r\\n 이 아무 데나 들어 있다. 잘라먹으면 안 된다."""
        payload = b"PK\r\n\r\n" + bytes(range(256)) * 4
        body, content_type = _multipart({}, {"roster": ("b.xlsx", payload)})
        assert _parse_multipart(body, content_type)["roster"].body == payload

    def test_garbage_is_ignored(self) -> None:
        assert _parse_multipart(b"gibberish", "text/plain") == {}


class TestServer:
    def test_form_page(self, server_url) -> None:
        with urllib.request.urlopen(server_url + "/") as response:
            page = response.read().decode()
        assert "산출 실행" in page
        assert "명부 파일" in page

    def test_full_calculation_and_download(
        self, server_url, roster_path, assumptions_path
    ) -> None:
        body, content_type = _multipart(
            {"sensitivity": b"on", "longterm": b"on"},
            {
                "roster": ("명부.xlsx", roster_path.read_bytes()),
                "assumptions": ("기초율.xlsx", assumptions_path.read_bytes()),
            },
        )
        status, page = _post(server_url + "/calc", body, content_type)
        text = page.decode()
        assert status == 200
        assert "확정급여채무" in text

        # 결과·개인별 내려받기 링크가 모두 실제 엑셀이어야 한다.
        import re

        tokens = re.findall(r"/download/([\w-]+)", text)
        assert len(tokens) == 2
        for token in tokens:
            with urllib.request.urlopen(f"{server_url}/download/{token}") as response:
                payload = response.read()
            assert payload[:2] == b"PK", "xlsx(zip) 서명이 아니다"

    def test_validation_errors_stop_the_run(
        self, server_url, roster_path, assumptions_path, tmp_path
    ) -> None:
        import openpyxl

        wb = openpyxl.load_workbook(roster_path)
        wb["재직자명부"].cell(26, 17, "")     # 제도구분 훼손
        broken = tmp_path / "깨진명부.xlsx"
        wb.save(broken)

        body, content_type = _multipart(
            {},
            {
                "roster": ("명부.xlsx", broken.read_bytes()),
                "assumptions": ("기초율.xlsx", assumptions_path.read_bytes()),
            },
        )
        _status, page = _post(server_url + "/calc", body, content_type)
        assert "산출을 중단했습니다" in page.decode()

    def test_force_runs_anyway(self, server_url, roster_path, assumptions_path, tmp_path) -> None:
        import openpyxl

        wb = openpyxl.load_workbook(roster_path)
        wb["재직자명부"].cell(26, 17, "")
        broken = tmp_path / "깨진명부.xlsx"
        wb.save(broken)

        body, content_type = _multipart(
            {"force": b"on"},
            {
                "roster": ("명부.xlsx", broken.read_bytes()),
                "assumptions": ("기초율.xlsx", assumptions_path.read_bytes()),
            },
        )
        _status, page = _post(server_url + "/calc", body, content_type)
        assert "확정급여채무" in page.decode()

    def test_missing_files_are_reported(self, server_url) -> None:
        body, content_type = _multipart({"prior_dbo": b"0"}, {})
        _status, page = _post(server_url + "/calc", body, content_type)
        assert "파일이 없습니다" in page.decode()

    def test_expired_download(self, server_url) -> None:
        request = urllib.request.Request(server_url + "/download/no-such-token")
        with pytest.raises(urllib.error.HTTPError) as caught:
            urllib.request.urlopen(request)
        assert caught.value.code == 404
