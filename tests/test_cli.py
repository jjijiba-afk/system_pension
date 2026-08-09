"""명령행 인터페이스."""

from __future__ import annotations

import io
import sys
from pathlib import Path

import openpyxl
import pytest

from pension.cli import force_utf8_output, main


class TestForceUtf8Output:
    """한글 출력이 콘솔 코드페이지 때문에 죽지 않게 한다.

    한국어 윈도우 콘솔은 CP949, 영문 환경은 CP1252 로 stdout 을 연다. 그대로 두면
    산출을 다 끝내 놓고 결과를 찍다가 ``UnicodeEncodeError`` 로 죽는다.
    """

    def test_sets_utf8_on_reconfigurable_streams(self, capsys) -> None:
        force_utf8_output()
        # 한글이 예외 없이 나가야 한다.
        print("확정급여채무 산출")
        assert "확정급여채무" in capsys.readouterr().out

    def test_survives_streams_without_reconfigure(self, monkeypatch) -> None:
        """PyInstaller 창 모드에서는 stdout 이 None 이거나 최소 객체일 수 있다."""
        monkeypatch.setattr(sys, "stdout", None)
        monkeypatch.setattr(sys, "stderr", io.StringIO())
        force_utf8_output()  # 예외가 나면 안 된다

    def test_survives_reconfigure_failure(self, monkeypatch) -> None:
        class Stubborn:
            def reconfigure(self, **_kw):
                raise OSError("이미 닫힌 스트림")

        monkeypatch.setattr(sys, "stdout", Stubborn())
        monkeypatch.setattr(sys, "stderr", Stubborn())
        force_utf8_output()


class TestCommands:
    def test_version(self, capsys) -> None:
        with pytest.raises(SystemExit) as info:
            main(["--version"])
        assert info.value.code == 0
        assert "1.0.0" in capsys.readouterr().out

    def test_template_creates_every_sheet(self, tmp_path: Path, capsys) -> None:
        out = tmp_path / "기초율.xlsx"
        assert main(["template", str(out)]) == 0
        assert out.exists()

        wb = openpyxl.load_workbook(out)
        assert set(wb.sheetnames) == {
            "할인율", "임금상승률", "승급률", "퇴직률",
            "사망률", "지급률", "지급률규정", "장기급여지급률", "장기급여규정",
            "퇴직사유",
        }

    def test_template_uses_rule_names_from_the_roster(
        self, roster_path: Path, tmp_path: Path
    ) -> None:
        out = tmp_path / "기초율.xlsx"
        assert main(["template", str(out), "--roster", str(roster_path)]) == 0

        wb = openpyxl.load_workbook(out)
        ws = wb["퇴직률"]
        headers = [ws.cell(1, c).value for c in range(2, ws.max_column + 1)]
        assert headers == ["정규임원", "정규직", "계약직"]

    def test_check_reports_clean_roster(self, roster_path: Path, capsys) -> None:
        assert main(["check", str(roster_path)]) == 0
        assert "오류 0건" in capsys.readouterr().out

    def test_check_writes_csv(self, roster_path: Path, tmp_path: Path) -> None:
        csv_path = tmp_path / "검증.csv"
        main(["check", str(roster_path), "-o", str(csv_path)])
        assert csv_path.exists()
        assert "심각도" in csv_path.read_text(encoding="utf-8-sig")

    def test_check_returns_nonzero_on_errors(self, roster_path: Path) -> None:
        wb = openpyxl.load_workbook(roster_path)
        wb["재직자명부"].cell(26, 17, "")  # 제도구분 누락
        wb.save(roster_path)
        assert main(["check", str(roster_path)]) == 1

    def test_calc_writes_a_report(
        self, roster_path: Path, assumptions_path: Path, tmp_path: Path, capsys
    ) -> None:
        out = tmp_path / "결과.xlsx"
        code = main(["calc", str(roster_path), str(assumptions_path), "-o", str(out)])
        assert code == 0
        assert out.exists()
        assert "확정급여채무" in capsys.readouterr().out

    def test_calc_stops_on_validation_errors(
        self, roster_path: Path, assumptions_path: Path, tmp_path: Path
    ) -> None:
        wb = openpyxl.load_workbook(roster_path)
        wb["재직자명부"].cell(26, 17, "")
        wb.save(roster_path)

        out = tmp_path / "결과.xlsx"
        assert main(["calc", str(roster_path), str(assumptions_path), "-o", str(out)]) == 2
        assert not out.exists()

    def test_calc_force_flag_overrides(
        self, roster_path: Path, assumptions_path: Path, tmp_path: Path
    ) -> None:
        wb = openpyxl.load_workbook(roster_path)
        wb["재직자명부"].cell(26, 17, "")
        wb.save(roster_path)

        out = tmp_path / "결과.xlsx"
        code = main([
            "calc", str(roster_path), str(assumptions_path), "-o", str(out), "--force",
        ])
        assert code == 0
        assert out.exists()

    def test_upload_writes_both_sheets(self, roster_path: Path, tmp_path: Path) -> None:
        out = tmp_path / "업로드.xlsx"
        assert main(["upload", str(roster_path), "-o", str(out)]) == 0

        wb = openpyxl.load_workbook(out)
        assert wb.sheetnames == ["UpLoad_Jae", "UpLoad_Toi"]
        assert wb["UpLoad_Jae"].max_row == 6  # 머리글 + 재직자 5명

    def test_missing_file_returns_nonzero(self, tmp_path: Path, capsys) -> None:
        assert main(["check", str(tmp_path / "없는명부.xlsx")]) == 2


def test_percent_argument_accepts_both_forms(
    roster_path: Path, assumptions_path: Path, tmp_path: Path
) -> None:
    """``--prior-rate 4.8%`` 와 ``--prior-rate 0.048`` 이 같아야 한다."""
    from pension.cli import _percent

    assert _percent("4.8%") == pytest.approx(0.048)
    assert _percent("0.048") == pytest.approx(0.048)
    assert _percent("4.8") == pytest.approx(0.048)


class TestCheckWithAssumptions:
    """check 에 기초율을 주면 calc 와 같은 조건으로 검증한다.

    직군 매핑을 기초율의 `지급규정` 시트에 둔 경우, check 가 명부의 Input 만
    보면 직군 미매칭 오류를 내고 calc 는 통과하는 — 서로 다른 판정이 나오는
    함정이 있었다.
    """

    def _roster_without_input(self, tmp_path):
        import openpyxl

        wb = openpyxl.Workbook()
        for sheet in ("재직자명부", "퇴직자명부"):
            ws = wb.create_sheet(sheet)
            ws.cell(1, 1, "작성기준일")
            ws.cell(1, 2, "2025-12-31")
            for col, title in (
                (2, "순번"), (3, "사번"), (4, "임직원구분"), (5, "직군"), (6, "성명"),
                (7, "성별"), (8, "생년월일"), (9, "입사일자"), (10, "퇴사일"),
                (13, "제도구분"), (14, "30일 평균임금"),
            ):
                ws.cell(3, col, title)
        ws = wb["재직자명부"]
        # 직군이 '과장' — Input 시트가 없으니 지급규정 없이는 잠정 규칙뿐이다.
        for col, value in ((2, 1), (3, "A1"), (4, "정규사원"), (5, "과장"), (7, "남"),
                           (8, "19850101"), (9, "20100101"), (13, "DB"), (14, 3_000_000)):
            ws.cell(4, col, value)
        del wb["Sheet"]
        path = tmp_path / "명부.xlsx"
        wb.save(path)
        return path

    def _assumptions_with_mapping(self, tmp_path):
        from pension.assumptions import write_assumptions
        from pension.config import PAYOUT_SHEET

        return write_assumptions(
            tmp_path / "기초율.xlsx",
            {
                "할인율": (["연차", "할인율"], [[1, 0.045]]),
                "지급률": (["근속연수", "정규직"], [[0, 1.0]]),
                PAYOUT_SHEET: (
                    [f"c{i}" for i in range(1, 23)],
                    [["과장", "정규직", 60, 60, 2, *[""] * 8, 0, 0, 2, "",
                      "일할", "그대로", 0, "반올림", "정규사원"]],
                ),
            },
            None, None,
        )

    def test_matches_calc_when_assumptions_are_given(self, tmp_path, capsys) -> None:
        roster = self._roster_without_input(tmp_path)
        assumptions = self._assumptions_with_mapping(tmp_path)

        assert main(["check", str(roster), str(assumptions)]) == 0
        assert "오류 0건" in capsys.readouterr().out

    def test_grouped_warning_summary_kicks_in_when_noisy(self, tmp_path, capsys) -> None:
        """경고가 15건을 넘으면 코드별 요약으로 접는다."""
        import openpyxl

        roster = self._roster_without_input(tmp_path)
        wb = openpyxl.load_workbook(roster)
        ws = wb["재직자명부"]
        for i in range(2, 22):   # DC 20명, 임금 공란 → 경고 20건
            for col, value in ((2, i), (3, f"A{i}"), (4, "정규사원"), (5, "과장"),
                               (7, "남"), (8, "19900101"), (9, "20150101"), (13, "DC")):
                ws.cell(3 + i, col, value)
        wb.save(roster)

        main(["check", str(roster), str(self._assumptions_with_mapping(tmp_path))])
        out = capsys.readouterr().out
        assert "경고 요약" in out
        assert "JAE_WAGE_MISSING_DC" in out
