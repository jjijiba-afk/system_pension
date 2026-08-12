"""퇴직사유(중도·사망·정년)별 지급 차등.

자료요청서 6번에는 사유가 세 줄로 갈려 있고, 실제로 다르게 적어 보내는 회사가
많다. 스터디 자료에서 확인한 것들::

    사망 정액 가산   사망시   '동일, 정액 가산 있음'
    15번 케이스  사망시   '근속10년 미만 기본급 3개월분, 10년 이상 5개월분 가산'
                          '재직 중 사망으로 퇴직 시 1년 미만도 1년으로 계산'
    임원 배수 별도   정년시   '임원은 배수 별도'
    21번 케이스  정년시   기본급 추가지급

사유를 뭉뚱그리면 어느 규정을 적용할지 정할 수 없다. 여기서는 **갈라 놓아도
합계는 그대로인지**, 그리고 **갈라 놓아야만 되는 것들이 되는지** 를 본다.
"""

from __future__ import annotations

import pytest
from test_valuation import BASE_DATE, make_assumptions, make_member

from pension.assumptions import (
    ATTRIB_IMMEDIATE,
    ATTRIB_SERVICE,
    CAUSE_DEATH,
    CAUSE_NORMAL,
    CAUSE_VOLUNTARY,
    BenefitScale,
    CauseBenefit,
    CauseBenefits,
    RateCurve,
    load_assumptions,
    write_template,
)
from pension.config import CalculationConfig, JobGroupRule
from pension.valuation import value_member


@pytest.fixture
def config() -> CalculationConfig:
    return CalculationConfig(
        base_date=BASE_DATE,
        job_group_rules=[
            JobGroupRule(
                "정규직", "정규직", severance_nra=60, longterm_nra=60, over_nra_add_age=2
            )
        ],
    )


def with_causes(assumptions, **entries: CauseBenefit):
    """``{사유: 규정}`` 을 기본 지급률 규정('법정')에 붙인다."""
    assumptions.exit_causes = CauseBenefits(
        {("법정", cause): entry for cause, entry in entries.items()}
    )
    return assumptions


def valued(config, *, member=None, **kwargs):
    member = member if member is not None else make_member(
        age=40, past_service=10.0, wage=5_000_000, nra=60
    )
    # 규정명을 못 찾으면 법정 퇴직금을 쓴다. 사유별 규정을 붙일 열쇠가 필요해
    # 이름만 '법정' 으로 지정한다 — 지급률 표는 여전히 비어 있다.
    member.rules.severance_benefit = "법정"
    return value_member(member, config, kwargs["assumptions"])


class TestNoCauseRulesChangesNothing:
    """사유별 규정을 안 쓰면 종전 결과가 그대로 나와야 한다."""

    def test_totals_are_unchanged(self, config: CalculationConfig) -> None:
        member = make_member(age=40, past_service=10.0, wage=5_000_000, nra=60)
        plain = make_assumptions(discount=0.045, salary=0.03, withdrawal=0.08, mortality=0.003)
        before = value_member(member, config, plain)

        member.rules.severance_benefit = "법정"
        after = value_member(member, config, with_causes(plain))
        assert after.dbo == pytest.approx(before.dbo, rel=1e-12)
        assert after.service_cost == pytest.approx(before.service_cost, rel=1e-12)

    def test_split_decrements_still_sum_to_the_combined_probability(
        self, config: CalculationConfig
    ) -> None:
        """중도 + 사망 = ``1-(1-w)(1-q)``.

        갈라 놓느라 확률이 새면 급여가 통째로 사라지거나 두 번 잡힌다.
        지급률을 상수로 두고 할인·임금상승을 끄면 미래급여 현가는 급여액과 같다.
        """
        member = make_member(age=55, past_service=10.0, wage=1_000_000, nra=60)
        assumptions = make_assumptions(discount=0.0, withdrawal=0.10, mortality=0.02)
        assumptions.severance_benefit = BenefitScale(
            curves={"고정": RateCurve({0: 5.0})}, statutory_when_missing=False
        )
        member.rules.severance_benefit = "고정"

        result = value_member(member, config, assumptions)
        assert result.expected_benefit_pv == pytest.approx(5_000_000, rel=1e-9)


class TestDeathGrant:
    """사망 정액 가산 — '동일, 정액 가산 있음'."""

    def test_flat_grant_raises_the_obligation(self, config: CalculationConfig) -> None:
        plain = make_assumptions(discount=0.045, withdrawal=0.05, mortality=0.01)
        base = valued(config, assumptions=plain).dbo

        with_grant = valued(
            config,
            assumptions=with_causes(
                make_assumptions(discount=0.045, withdrawal=0.05, mortality=0.01),
                **{CAUSE_DEATH: CauseBenefit(extra_amount=50_000_000)},
            ),
        ).dbo
        assert with_grant > base

    def test_the_grant_only_costs_what_death_is_worth(
        self, config: CalculationConfig
    ) -> None:
        """5천만원이 통째로 얹히면 안 된다 — 사망확률과 할인이 걸려야 한다."""
        assumptions = with_causes(
            make_assumptions(discount=0.045, withdrawal=0.05, mortality=0.01),
            **{CAUSE_DEATH: CauseBenefit(extra_amount=50_000_000)},
        )
        plain = make_assumptions(discount=0.045, withdrawal=0.05, mortality=0.01)
        added = valued(config, assumptions=assumptions).dbo - valued(
            config, assumptions=plain
        ).dbo
        assert 0 < added < 50_000_000

    def test_no_mortality_means_no_grant(self, config: CalculationConfig) -> None:
        """사망률을 0 으로 두면 정액 가산금은 채무에 들어오지 않는다."""
        plain = make_assumptions(discount=0.045, withdrawal=0.05, mortality=0.0)
        base = valued(config, assumptions=plain).dbo
        granted = valued(
            config,
            assumptions=with_causes(
                make_assumptions(discount=0.045, withdrawal=0.05, mortality=0.0),
                **{CAUSE_DEATH: CauseBenefit(extra_amount=50_000_000)},
            ),
        ).dbo
        assert granted == pytest.approx(base, rel=1e-12)

    def test_a_withdrawal_grant_does_not_leak_into_death(
        self, config: CalculationConfig
    ) -> None:
        """중도퇴직에 붙인 가산이 사망 급여까지 올리면 사유를 가른 뜻이 없다."""
        member = make_member(age=59, past_service=10.0, wage=1_000_000, nra=60)
        member.rules.severance_benefit = "법정"
        only_death = with_causes(
            make_assumptions(discount=0.0, withdrawal=0.0, mortality=0.5),
            **{CAUSE_VOLUNTARY: CauseBenefit(extra_amount=9_000_000)},
        )
        # 중도퇴직률이 0 이므로 중도퇴직 가산은 한 푼도 들어오지 않는다.
        plain = make_assumptions(discount=0.0, withdrawal=0.0, mortality=0.5)
        assert value_member(member, config, only_death).dbo == pytest.approx(
            value_member(member, config, plain).dbo, rel=1e-12
        )


class TestDeathGrantAttribution:
    """근속이 늘어도 안 늘어나는 급여는 더 쌓을 것이 없다(문단 70)."""

    def test_flat_grant_is_fully_attributed_by_default(
        self, config: CalculationConfig
    ) -> None:
        """재직 중 사망하면 근속이 1년이든 20년이든 같은 금액이다."""
        young = make_member(age=30, past_service=1.0, wage=1_000_000, nra=60)
        young.rules.severance_benefit = "법정"
        old = make_member(age=30, past_service=20.0, wage=1_000_000, nra=60)
        old.rules.severance_benefit = "법정"

        def grant_share(member):
            grant = with_causes(
                make_assumptions(discount=0.0, withdrawal=0.0, mortality=0.02),
                **{CAUSE_DEATH: CauseBenefit(extra_amount=10_000_000)},
            )
            plain = make_assumptions(discount=0.0, withdrawal=0.0, mortality=0.02)
            return (
                value_member(member, config, grant).dbo
                - value_member(member, config, plain).dbo
            )

        # 근속이 스무 배 달라도 귀속된 정액 가산금은 같아야 한다.
        assert grant_share(young) == pytest.approx(grant_share(old), rel=1e-9)

    def test_service_basis_can_be_chosen_instead(self, config: CalculationConfig) -> None:
        """근속비례로 두면 근속이 짧은 사람에게 덜 귀속된다."""
        def grant_share(past_service: float) -> float:
            member = make_member(
                age=30, past_service=past_service, wage=1_000_000, nra=60
            )
            member.rules.severance_benefit = "법정"
            grant = with_causes(
                make_assumptions(discount=0.0, withdrawal=0.0, mortality=0.02),
                **{
                    CAUSE_DEATH: CauseBenefit(
                        extra_amount=10_000_000, attribution=ATTRIB_SERVICE
                    )
                },
            )
            plain = make_assumptions(discount=0.0, withdrawal=0.0, mortality=0.02)
            return (
                value_member(member, config, grant).dbo
                - value_member(member, config, plain).dbo
            )

        assert grant_share(1.0) < grant_share(20.0)

    def test_immediate_attribution_adds_nothing_to_service_cost(
        self, config: CalculationConfig
    ) -> None:
        """이미 전액 귀속했으면 한 해 더 일해도 쌓일 것이 없다."""
        member = make_member(age=40, past_service=10.0, wage=1_000_000, nra=60)
        member.rules.severance_benefit = "법정"
        grant = with_causes(
            make_assumptions(discount=0.0, withdrawal=0.0, mortality=0.02),
            **{CAUSE_DEATH: CauseBenefit(extra_amount=10_000_000)},
        )
        plain = make_assumptions(discount=0.0, withdrawal=0.0, mortality=0.02)
        assert value_member(member, config, grant).service_cost == pytest.approx(
            value_member(member, config, plain).service_cost, rel=1e-12
        )

    def test_default_basis_differs_by_cause(self) -> None:
        entry = CauseBenefit(extra_amount=1.0)
        assert entry.attribution_basis(CAUSE_DEATH) == ATTRIB_IMMEDIATE
        assert entry.attribution_basis(CAUSE_NORMAL) == ATTRIB_SERVICE
        assert entry.attribution_basis(CAUSE_VOLUNTARY) == ATTRIB_SERVICE


class TestServiceFloor:
    """15·12번 케이스 — '사망 시 1년 미만도 1년으로 계산'."""

    def test_short_service_is_lifted_for_death(self, config: CalculationConfig) -> None:
        member = make_member(age=30, past_service=0.5, wage=1_000_000, nra=60)
        member.rules.severance_benefit = "법정"
        floored = with_causes(
            make_assumptions(discount=0.0, withdrawal=0.0, mortality=0.05),
            **{CAUSE_DEATH: CauseBenefit(min_service=1.0)},
        )
        plain = make_assumptions(discount=0.0, withdrawal=0.0, mortality=0.05)
        assert (
            value_member(member, config, floored).dbo
            > value_member(member, config, plain).dbo
        )

    def test_long_service_is_untouched(self, config: CalculationConfig) -> None:
        member = make_member(age=40, past_service=10.0, wage=1_000_000, nra=60)
        member.rules.severance_benefit = "법정"
        floored = with_causes(
            make_assumptions(discount=0.045, withdrawal=0.05, mortality=0.01),
            **{CAUSE_DEATH: CauseBenefit(min_service=1.0)},
        )
        plain = make_assumptions(discount=0.045, withdrawal=0.05, mortality=0.01)
        assert value_member(member, config, floored).dbo == pytest.approx(
            value_member(member, config, plain).dbo, rel=1e-12
        )

    def test_the_floor_clears_the_eligibility_gate(self, config: CalculationConfig) -> None:
        """가입자격 1년을 못 채운 사람도 사망 시에는 1년으로 봐 지급 대상이다."""
        member = make_member(age=30, past_service=0.5, wage=1_000_000, nra=60)
        member.rules.severance_benefit = "법정"
        member.min_service_years = 1.0
        floored = with_causes(
            make_assumptions(discount=0.0, withdrawal=0.0, mortality=0.05),
            **{CAUSE_DEATH: CauseBenefit(min_service=1.0)},
        )
        assert value_member(member, config, floored).dbo > 0.0


class TestAlternateRuleAndExtraScale:
    """12번 케이스 — 정년퇴직만 배수가 다르다. 15번 — 근속에 따라 가산이 달라진다."""

    def test_normal_retirement_can_use_a_different_scale(
        self, config: CalculationConfig
    ) -> None:
        member = make_member(age=59, past_service=10.0, wage=1_000_000, nra=60)
        member.rules.severance_benefit = "법정"
        assumptions = make_assumptions(discount=0.0, withdrawal=0.0, mortality=0.0)
        assumptions.severance_benefit = BenefitScale(
            curves={"정년배수": RateCurve({0: 30.0})}, statutory_when_missing=True
        )
        with_causes(
            assumptions, **{CAUSE_NORMAL: CauseBenefit(benefit_rule="정년배수")}
        )
        # 탈퇴가 정년 하나뿐이므로 급여 = 30배 × 임금, 귀속은 배수가 상수라 1.0.
        assert value_member(member, config, assumptions).dbo == pytest.approx(
            30_000_000, rel=1e-9
        )

    def test_extra_scale_can_depend_on_service(self, config: CalculationConfig) -> None:
        """'근속 10년 미만 3개월분, 10년 이상 5개월분 가산'."""
        def death_extra(past_service: float) -> float:
            member = make_member(
                age=30, past_service=past_service, wage=1_000_000, nra=60
            )
            member.rules.severance_benefit = "법정"
            assumptions = make_assumptions(discount=0.0, withdrawal=0.0, mortality=0.02)
            assumptions.severance_benefit = BenefitScale(
                curves={"사망가산": RateCurve({0: 3.0, 10: 5.0})},
                statutory_when_missing=True,
            )
            with_causes(
                assumptions, **{CAUSE_DEATH: CauseBenefit(extra_rule="사망가산")}
            )
            plain = make_assumptions(discount=0.0, withdrawal=0.0, mortality=0.02)
            return (
                value_member(member, config, assumptions).dbo
                - value_member(member, config, plain).dbo
            )

        # 즉시 귀속은 **오늘 근속** 으로 잰다. 5년차는 3개월분, 15년차는 5개월분.
        assert death_extra(15.0) == pytest.approx(death_extra(5.0) * 5 / 3, rel=1e-9)

    def test_crossing_the_step_shows_up_as_service_cost(
        self, config: CalculationConfig
    ) -> None:
        """올해 10년째가 되는 사람은 가산이 3 → 5 로 오른다. 그 몫이 근무원가다."""
        def death_unit(past_service: float) -> float:
            member = make_member(
                age=30, past_service=past_service, wage=1_000_000, nra=60
            )
            member.rules.severance_benefit = "법정"
            assumptions = make_assumptions(discount=0.0, withdrawal=0.0, mortality=0.02)
            assumptions.severance_benefit = BenefitScale(
                curves={"사망가산": RateCurve({0: 3.0, 10: 5.0})},
                statutory_when_missing=True,
            )
            with_causes(
                assumptions, **{CAUSE_DEATH: CauseBenefit(extra_rule="사망가산")}
            )
            plain = make_assumptions(discount=0.0, withdrawal=0.0, mortality=0.02)
            return (
                value_member(member, config, assumptions).service_cost
                - value_member(member, config, plain).service_cost
            )

        assert death_unit(9.5) > 0.0     # 올해 계단을 넘는다
        assert death_unit(15.0) == pytest.approx(0.0, abs=1e-9)   # 이미 넘었다
        assert death_unit(2.0) == pytest.approx(0.0, abs=1e-9)    # 아직 멀었다



class TestSheetRoundTrip:
    def test_template_carries_the_sheet(self, tmp_path) -> None:
        path = write_template(tmp_path / "기초율.xlsx", job_groups=["정규직"])
        assert load_assumptions(path).exit_causes.is_empty()

    def test_rows_are_read_back(self, tmp_path) -> None:
        import openpyxl

        path = write_template(tmp_path / "기초율.xlsx", job_groups=["정규직"])
        wb = openpyxl.load_workbook(path)
        ws = wb["퇴직사유"]
        ws.cell(2, 1, "정규직")
        ws.cell(2, 2, CAUSE_DEATH)
        ws.cell(2, 4, "사망가산")
        ws.cell(2, 5, 50_000_000)
        ws.cell(2, 6, 1)
        ws.cell(3, 1, "정규직")
        ws.cell(3, 2, CAUSE_NORMAL)
        ws.cell(3, 3, "정년배수")
        ws.cell(3, 7, ATTRIB_SERVICE)
        wb.save(path)

        causes = load_assumptions(path).exit_causes
        death = causes.get("정규직", CAUSE_DEATH)
        assert death.extra_rule == "사망가산"
        assert death.extra_amount == 50_000_000
        assert death.min_service == 1.0
        assert death.attribution == ""          # 비우면 사유 기본값

        normal = causes.get("정규직", CAUSE_NORMAL)
        assert normal.benefit_rule == "정년배수"
        assert normal.attribution == ATTRIB_SERVICE
        assert causes.rule_names() == ["사망가산", "정년배수"]

    def test_unknown_cause_is_refused(self, tmp_path) -> None:
        """'명예퇴직' 을 사망으로 갈음하면 안 된다 — 모르면 멈춘다."""
        import openpyxl

        path = write_template(tmp_path / "기초율.xlsx", job_groups=["정규직"])
        wb = openpyxl.load_workbook(path)
        ws = wb["퇴직사유"]
        ws.cell(2, 1, "정규직")
        ws.cell(2, 2, "명예퇴직")
        ws.cell(2, 5, 1_000_000)
        wb.save(path)

        with pytest.raises(ValueError, match="명예퇴직"):
            load_assumptions(path)

    def test_unknown_attribution_is_refused(self, tmp_path) -> None:
        import openpyxl

        path = write_template(tmp_path / "기초율.xlsx", job_groups=["정규직"])
        wb = openpyxl.load_workbook(path)
        ws = wb["퇴직사유"]
        ws.cell(2, 1, "정규직")
        ws.cell(2, 2, CAUSE_DEATH)
        ws.cell(2, 5, 1_000_000)
        ws.cell(2, 7, "적당히")
        wb.save(path)

        with pytest.raises(ValueError, match="적당히"):
            load_assumptions(path)

    def test_a_row_with_nothing_filled_in_is_not_a_rule(self, tmp_path) -> None:
        """사유만 적고 값을 안 넣은 줄을 규정으로 세면 '차등이 있다' 고 오해한다."""
        import openpyxl

        path = write_template(tmp_path / "기초율.xlsx", job_groups=["정규직"])
        wb = openpyxl.load_workbook(path)
        ws = wb["퇴직사유"]
        ws.cell(2, 1, "정규직")
        ws.cell(2, 2, CAUSE_DEATH)
        wb.save(path)

        assert load_assumptions(path).exit_causes.is_empty()


class TestSplitBenefitColumn:
    """지급률 열 하나를 정년·중도·사망 세 열로 가르기.

    엔진은 처음부터 사유별 차등을 받았지만 넣는 길이 2단 우회였다 — 열을
    손으로 만들고, [퇴직사유] 탭에서 그 이름을 지목해야 했다. 여기서는
    **한 번에 갈라 주는 조작이 종전과 똑같은 파일을 낳는지** 를 본다.
    """

    def base(self):
        from pension import assumption_form as form

        state = form.example_state(["정규직", "임원"])
        state["grids"]["지급률"]["rows"] = [["0", "1.0", "1.5"], ["10", "2.0", "3.0"]]
        return state

    def test_makes_three_columns_and_wires_them(self) -> None:
        from pension import assumption_form as form

        state = form.split_benefit_by_cause(self.base(), "정규직")
        columns = form.grid_columns(state, "지급률")
        for cause in ("중도", "사망", "정년"):
            assert f"정규직·{cause}" in columns
        assert form.cause_split_rules(state) == ["정규직"]
        # 퇴직사유 줄이 그 열을 가리켜야 엔진이 집는다.
        linked = {(r[0], r[1]): r[2] for r in state["exit_causes"]}
        assert linked[("정규직", "정년")] == "정규직·정년"

    def test_new_columns_inherit_the_original_rule(self) -> None:
        """가르는 목적은 '셋이 다르다' 를 적는 것이지 처음부터 짜는 게 아니다."""
        from pension import assumption_form as form

        state = self.base()
        state["benefit_rules"]["정규직"] = {"mode": "누진", "formula": ""}
        state = form.split_benefit_by_cause(state, "정규직")
        for cause in ("중도", "사망", "정년"):
            assert state["benefit_rules"][f"정규직·{cause}"]["mode"] == "누진"

    def test_untouched_groups_stay_put(self) -> None:
        from pension import assumption_form as form

        state = form.split_benefit_by_cause(self.base(), "정규직")
        assert form.cause_split_rules(state) == ["정규직"]
        assert "임원·정년" not in form.grid_columns(state, "지급률")

    def test_existing_cause_row_keeps_its_other_values(self) -> None:
        """가산액을 먼저 적어 둔 회사가 있다. 줄을 새로 만들면 그게 사라진다."""
        from pension import assumption_form as form

        state = self.base()
        state["exit_causes"] = [["정규직", "사망", "", "", "50000000", "1", ""]]
        state = form.split_benefit_by_cause(state, "정규직")
        death = next(r for r in state["exit_causes"] if r[1] == "사망")
        assert death[2] == "정규직·사망"
        assert death[4] == "50000000" and death[5] == "1"

    def test_merge_removes_the_columns_and_their_values(self) -> None:
        from pension import assumption_form as form

        state = form.split_benefit_by_cause(self.base(), "정규직")
        state["grids"]["지급률"]["rows"] = [
            ["0", "1.0", "1.5", "1.0", "1.0", "2.0"],
            ["10", "2.0", "3.0", "2.0", "2.0", "4.0"],
        ]
        merged = form.merge_benefit_causes(state, "정규직")
        assert form.grid_columns(merged, "지급률") == ["정규직", "임원"]
        # 값 칸까지 같이 들어내야 남은 값이 옆 열로 밀리지 않는다.
        assert merged["grids"]["지급률"]["rows"] == [["0", "1.0", "1.5"],
                                                  ["10", "2.0", "3.0"]]
        assert merged["exit_causes"] == []
        assert "정규직·정년" not in merged["benefit_rules"]

    def test_merge_keeps_a_cause_row_that_still_has_content(self) -> None:
        from pension import assumption_form as form

        state = self.base()
        state["exit_causes"] = [["정규직", "사망", "", "", "50000000", "", ""]]
        state = form.split_benefit_by_cause(state, "정규직")
        merged = form.merge_benefit_causes(state, "정규직")
        assert merged["exit_causes"] == [["정규직", "사망", "", "", "50000000", "", ""]]

    def test_splitting_twice_is_refused(self) -> None:
        from pension import assumption_form as form

        state = form.split_benefit_by_cause(self.base(), "정규직")
        with pytest.raises(ValueError, match="이미"):
            form.split_benefit_by_cause(state, "정규직·정년")

    def test_unknown_column_is_a_clear_error(self) -> None:
        from pension import assumption_form as form

        with pytest.raises(ValueError, match="없습니다"):
            form.split_benefit_by_cause(self.base(), "없는직군")

    def test_the_split_survives_a_workbook_round_trip(self, tmp_path) -> None:
        """파일에서 되읽었을 때도 갈라 놓은 상태로 떠야 한다."""
        from pension import assumption_form as form

        state = form.split_benefit_by_cause(self.base(), "정규직")
        assert form.state_problems(state) == []
        form.write_state(state, tmp_path / "기초율.xlsx")
        back = form.read_state(tmp_path / "기초율.xlsx")
        assert form.cause_split_rules(back) == ["정규직"]

    def test_the_engine_reads_the_split_column(self, tmp_path) -> None:
        """화면이 만든 파일을 엔진이 사유별 규정으로 집어야 의미가 있다."""
        from pension import assumption_form as form
        from pension.assumptions import load_assumptions

        state = form.split_benefit_by_cause(self.base(), "정규직")
        columns = form.grid_columns(state, "지급률")
        row = ["0"] + [""] * len(columns)
        row[columns.index("정규직·정년") + 1] = "2.0"
        state["grids"]["지급률"]["rows"] = [row]

        form.write_state(state, tmp_path / "기초율.xlsx")
        causes = load_assumptions(tmp_path / "기초율.xlsx").exit_causes
        assert causes.get("정규직", CAUSE_NORMAL).benefit_rule == "정규직·정년"
        assert causes.get("정규직", CAUSE_DEATH).benefit_rule == "정규직·사망"


class TestTheFinalYearStillHasMidYearExits:
    """마지막 해라고 중도퇴직·사망이 멈추지는 않는다.

    정년까지 3년 남은 사람은 0.5·1.5·2.5년에 중도로 나갈 수 있고, 그 해를
    넘긴 사람만 3.0년에 정년을 맞는다. 마지막 해를 통째로 정년으로 두면
    2.5년의 급부가 통째로 사라지고, 정년 지급률이 더 높은 회사에서는 그만큼
    채무가 부푼다.
    """

    def _trace(self, config, **kwargs):
        member = make_member(age=54, past_service=24.0, wage=7_856_560, nra=57)
        assumptions = make_assumptions(
            discount=0.03, salary=0.02, withdrawal=0.08, mortality=0.0004)
        trace: list[dict] = []
        result = value_member(member, config, assumptions, trace=trace, **kwargs)
        return result, trace

    def test_every_year_has_a_mid_year_exit(self, config) -> None:
        _result, trace = self._trace(config)
        assert sorted({row["timing"] for row in trace}) == [0.5, 1.5, 2.5, 3.0]

    def test_the_last_year_pays_all_three_causes(self, config) -> None:
        _result, trace = self._trace(config)
        final = {row["cause"] for row in trace if row["t"] == 3}
        assert final == {CAUSE_VOLUNTARY, CAUSE_DEATH, CAUSE_NORMAL}

    def test_only_the_survivors_reach_the_retirement_age(self, config) -> None:
        """정년 몫은 마지막 해의 탈퇴를 뺀 나머지다 — 이중으로 세면 안 된다."""
        _result, trace = self._trace(config)
        assert sum(row["exit_probability"] for row in trace) == pytest.approx(1.0)

    def test_the_normal_retirement_share_shrank_by_the_last_year_exits(
        self, config
    ) -> None:
        _result, trace = self._trace(config)
        normal = next(r for r in trace if r["cause"] == CAUSE_NORMAL)
        leavers = sum(r["exit_probability"] for r in trace if r["t"] == 3
                      and r["cause"] != CAUSE_NORMAL)
        assert leavers > 0
        # 마지막 해 시작 시점의 재직확률에서 그 해 탈퇴자를 뺀 것이 정년 몫이다.
        assert normal["survival"] == pytest.approx(
            normal["exit_probability"] + leavers)
