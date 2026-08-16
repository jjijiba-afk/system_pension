"""회사에 보내는 명부 양식.

이 파일은 우리가 아니라 **인사·회계 담당자** 가 읽는다. 여기서 어긋나면 산출이
틀리는 게 아니라 자료가 안 오거나, 엉뚱하게 채워져 온다. 그래서 시험도
"프로그램이 읽히는가" 가 아니라 "사람이 무엇을 어디에 적어야 하는지 알 수
있는가" 를 본다.
"""

from __future__ import annotations

import openpyxl
import pytest

from pension import rostertemplate as tpl
from pension.rostertemplate import (
    ACTIVE,
    BLOCK_ROW,
    BLOCKS,
    FIRST_DATA_ROW,
    HEADER_ROW,
    PART_ROW,
    PARTS,
    RETIRED,
    build_workbook,
    write_roster_template,
)


@pytest.fixture(scope="module")
def blank(tmp_path_factory):
    path = tmp_path_factory.mktemp("양식") / "명부양식.xlsx"
    write_roster_template(path)
    book = openpyxl.load_workbook(path)
    yield book
    book.close()


class TestTheTwoParts:
    """퇴직급여와 기타장기는 열 자리로 갈려 있어야 한다.

    근속포상이 없는 회사가 대부분이다. 장기급여 열이 퇴직급여 열 사이에 끼어
    있으면 "우리도 이 칸을 채워야 하나" 를 열마다 되묻게 되고, 되물어 볼 곳이
    없으면 짐작해서 채운다. 짐작이 들어온 명부는 우리가 알아볼 방법이 없다.
    """

    @pytest.mark.parametrize("columns", [ACTIVE, RETIRED], ids=["재직자", "퇴직자"])
    def test_the_long_term_part_is_one_unbroken_run_at_the_right_end(
        self, columns
    ) -> None:
        parts = [BLOCKS[block][0] for block, *_rest in columns]
        first = parts.index("기타장기")
        assert set(parts[first:]) == {"기타장기"}, (
            "기타장기 파트가 끊겨 있다. 통째로 지우면 퇴직급여 열까지 사라진다")
        assert set(parts[:first]) == {"퇴직급여"}

    @pytest.mark.parametrize("columns", [ACTIVE, RETIRED], ids=["재직자", "퇴직자"])
    def test_every_block_belongs_to_a_declared_part(self, columns) -> None:
        for block, *_rest in columns:
            assert BLOCKS[block][0] in PARTS

    @pytest.mark.parametrize("sheet", ["재직자명부", "퇴직자명부"])
    def test_the_sheet_shows_the_part_above_the_block(self, blank, sheet) -> None:
        ws = blank[sheet]
        bands = [c.value for c in ws[PART_ROW] if c.value]
        assert bands == ["퇴직급여", "기타장기"], (
            "맨 윗줄이 파트 두 개로 갈려 보여야 한다")
        assert "필수" in [c.value for c in ws[BLOCK_ROW] if c.value]

    def test_dropping_the_long_term_part_still_reads(self, blank, tmp_path) -> None:
        """기타장기를 통째로 지운 명부도 그대로 읽혀야 한다.

        "없으면 지우고 보내세요" 라고 적어 놓고 지운 것이 안 읽히면, 그 말을
        믿은 회사만 손해를 본다.
        """
        from pension.layout import (
            ACTIVE_HEADER_ALIASES,
            find_data_start,
            find_header_row,
            normalize_header,
        )

        path = tmp_path / "기타장기없음.xlsx"
        book = build_workbook()
        ws = book["재직자명부"]
        parts = [BLOCKS[block][0] for block, *_rest in ACTIVE]
        first = parts.index("기타장기") + 1          # 1부터 세는 열 번호
        ws.delete_cols(first, len(ACTIVE) - first + 1)
        book.save(path)

        reopened = openpyxl.load_workbook(path)
        sheet = reopened["재직자명부"]
        assert find_header_row(sheet, ACTIVE_HEADER_ALIASES) == HEADER_ROW
        assert find_data_start(sheet, HEADER_ROW) == FIRST_DATA_ROW

        known = {normalize_header(name): key
                 for key, names in ACTIVE_HEADER_ALIASES.items() for name in names}
        found = {known.get(normalize_header(c.value))
                 for c in sheet[HEADER_ROW] if c.value}
        assert "hire_date" in found, "퇴직급여 열까지 함께 사라졌다"
        assert "longterm_target" not in found
        reopened.close()


class TestTheInputSheet:
    """담당자가 채우는 칸은 [기초자료] 한 장에 모은다.

    여기저기 흩어 두면 시트를 오가며 채우다가 한 장을 통째로 빼먹는다.
    [예치금] 만 따로 두는 것은 그쪽이 신탁·보험 명세서를 보고 옮기는 일이라
    자료의 출처가 다르기 때문이다.
    """

    def test_the_workbook_has_exactly_these_sheets(self, blank) -> None:
        assert blank.sheetnames == [
            "작성요령", "기초자료", "예치금", "재직자명부", "퇴직자명부", "추가명부",
        ]

    def test_the_five_blocks_are_all_on_one_sheet(self, blank) -> None:
        ws = blank["기초자료"]
        bands = [str(c.value) for row in ws.iter_rows(max_col=1)
                 for c in row if c.value and str(c.value).startswith(("①", "②", "③",
                                                                     "④", "⑤"))]
        assert [b[0] for b in bands] == ["①", "②", "③", "④", "⑤"], bands
        joined = " ".join(bands)
        for want in ("회사", "직군 규칙", "퇴직급여 지급규정", "특이사항", "장기급여"):
            assert want in joined, f"[{want}] 블록이 없다"

    def test_the_deposit_sheet_covers_every_item(self, blank) -> None:
        """낱말은 우리식이되 항목은 빠짐없어야 한다."""
        ws = blank["예치금"]
        text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row
                          if c.value)
        for want in ("부담금 납입액", "이자수익", "운용관리수수료",
                     "자산관리수수료", "기말 잔액", "검증", "자산 분류",
                     "활성시장 공시가격", "자산인식상한", "기중 장기근속 지급액"):
            assert want in text, f"[예치금] 에 '{want}' 이(가) 없다"


class TestTheLongTermRuleBlock:
    """근속포상 규정을 적을 자리.

    퇴직급여규정 안에 한 칸으로 끼워 두었더니, 없는 회사는 "뭘 적으라는 건지"
    로 읽고 있는 회사는 한 칸에 못 적어 규정집을 통째로 첨부해 보냈다. 어느
    쪽이든 우리가 되물어야 했다.
    """

    def test_it_comes_last_so_it_can_be_left_blank(self, blank) -> None:
        """기타장기는 맨 아래에 둔다 — 없는 회사가 거기서 그냥 멈추면 된다."""
        ws = blank["기초자료"]
        marks = {str(c.value)[0]: c.row for row in ws.iter_rows(max_col=1)
                 for c in row if c.value and str(c.value)[:1] in "①②③④⑤"}
        assert marks["⑤"] > max(marks[m] for m in "①②③④")

    def test_it_asks_for_the_schedule_row_by_row(self, blank) -> None:
        ws = blank["기초자료"]
        heads = {c.value for row in ws.iter_rows() for c in row if c.value}
        for want in ("규정명", "근속년수", "지급 내용", "지급기준 (금액 환산)"):
            assert want in heads, f"'{want}' 을 묻지 않는다"

    def test_the_payout_words_match_what_the_engine_takes(self) -> None:
        """지급방법 낱말이 산출가정과 달라지면 우리가 손으로 옮기며 뜻을 바꾼다."""
        from pension.assumptions import LONGTERM_TYPES

        assert tpl.LONGTERM_KINDS == LONGTERM_TYPES

    def test_the_timings_offered_are_the_ones_the_engine_knows(self) -> None:
        from pension.assumptions import LONGTERM_TIMINGS

        note = dict((label, note) for label, _s, note in tpl.LONGTERM_RULE_ROWS)
        for timing in LONGTERM_TIMINGS:
            assert timing in note["지급시점"]

    def test_a_blank_form_still_shows_what_an_answer_looks_like(self, blank) -> None:
        """예시가 없으면 '대상' 칸은 대개 비어서 돌아온다."""
        ws = blank["기초자료"]
        written = [c.value for row in ws.iter_rows(min_col=2, max_col=6)
                   for c in row if isinstance(c.value, str)]
        assert any("휴가" in v for v in written)
        assert any("소멸" in v or "이월" in v for v in written)

    def test_it_points_back_at_the_roster_columns(self, blank) -> None:
        """규정 블록과 명부가 서로를 가리켜야 한 바퀴가 닫힌다."""
        ws = blank["기초자료"]
        text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row
                         if c.value)
        assert "장기급여 대상" in text
        assert "장기급여 기산일" in text


class TestTheGuideStaysHonest:
    def test_every_sheet_is_named_in_the_guide(self, blank) -> None:
        ws = blank["작성요령"]
        named = {c.value for row in ws.iter_rows(max_col=1) for c in row if c.value}
        for sheet in blank.sheetnames:
            if sheet != "작성요령":
                assert sheet in named, f"[{sheet}] 을 작성요령이 소개하지 않는다"

    def test_the_column_table_lists_every_column(self, blank) -> None:
        ws = blank["작성요령"]
        listed = [c.value for row in ws.iter_rows(min_col=3, max_col=3)
                  for c in row if c.value]
        for _block, label, *_rest in ACTIVE + RETIRED:
            assert label in listed, f"'{label}' 열이 작성요령에 없다"

    def test_the_header_row_is_frozen(self, blank) -> None:
        """수백 줄을 채우는 동안 열 이름이 화면에서 사라지면 칸을 밀려 적는다."""
        for sheet in ("재직자명부", "퇴직자명부"):
            assert blank[sheet].freeze_panes == f"A{FIRST_DATA_ROW}"


class TestColoursAreNotTransparent:
    """색이 투명하게 저장되면 표는 그려지는데 글자만 사라진다.

    openpyxl 은 여섯 자리로 준 색을 ARGB 여덟 자리로 늘리며 앞에 ``00`` 을
    붙인다. 그 자리는 알파 채널이고 ``00`` 은 완전 투명이다. 엑셀 대부분은
    무시하고 그리지만 그렇지 않은 버전에서는 남색 머리글 위의 흰 글자가
    통째로 안 보인다 — 받는 사람 눈에는 파일이 깨진 것으로 읽힌다.
    """

    def _transparent(self, path) -> list[str]:
        import re
        import zipfile

        styles = zipfile.ZipFile(path).read("xl/styles.xml").decode("utf-8")
        return re.findall(r'rgb="00[0-9A-Fa-f]{6}"', styles)

    def test_the_blank_form_has_no_transparent_colour(self, tmp_path) -> None:
        from pension.rostertemplate import write_roster_template

        path = write_roster_template(tmp_path / "양식.xlsx")
        assert self._transparent(path) == []

    def test_every_file_we_hand_over_is_opaque(self, tmp_path) -> None:
        """명부만 고치면 기초율 양식에서 같은 일이 난다."""
        from pension.samples import write_sample_pack

        for path in write_sample_pack(tmp_path):
            assert self._transparent(path) == [], f"{path.name} 에 투명한 색이 있다"

    def test_the_header_ink_survives_the_round_trip(self, tmp_path) -> None:
        import openpyxl

        from pension.rostertemplate import write_roster_template

        path = write_roster_template(tmp_path / "양식.xlsx")
        cell = openpyxl.load_workbook(path)["재직자명부"].cell(HEADER_ROW, 1)
        assert cell.value == "순번"
        assert cell.font.color.rgb == "FFFFFFFF", "머리글 글자가 흰색 불투명이어야 한다"
        assert cell.fill.fgColor.rgb.startswith("FF"), "채움색도 불투명이어야 한다"


class TestTheSheetReadsBackWhole:
    """우리가 만든 양식을 우리가 다시 읽을 수 있어야 한다.

    표 제목·항목 이름을 낱말로 찾는 구조라, 안내 문구에 그 낱말을 쓰면 안내
    줄이 표로 잡힌다. 실제로 ② 블록 안내에 '검증' 이라고 적었더니 그 줄이
    검증 줄로 읽혀 기말 잔액이 통째로 0 이 됐다 — 오류 없이 그럴듯한 숫자가
    나오는 쪽이다.
    """

    def _info(self, tmp_path):
        from pension.general_info import read_general_info
        from pension.rostertemplate import write_roster_template

        path = write_roster_template(tmp_path / "양식.xlsx")
        return read_general_info(openpyxl.load_workbook(path))

    def test_the_deposit_table_balances(self, tmp_path) -> None:
        assets = self._info(tmp_path).assets
        assert assets.opening > 0 and assets.closing > 0
        assert round(assets.difference) == 0, "검산줄이 0 이 아니다"

    def test_every_block_of_the_deposit_sheet_is_read(self, tmp_path) -> None:
        info = self._info(tmp_path)
        assert info.obligation.benefits_paid > 0, "① 추계액 증감을 못 읽었다"
        assert info.assets.contributions > 0, "② 예치금 증감을 못 읽었다"
        assert info.assets.breakdown, "③ 예치금 구성을 못 읽었다"
        assert info.assets.quoted, "③ 활성시장 공시가격 칸을 못 읽었다"

    def test_the_breakdown_adds_up_to_the_closing_balance(self, tmp_path) -> None:
        assets = self._info(tmp_path).assets
        assert sum(assets.breakdown.values()) == pytest.approx(assets.closing)

    def test_the_basics_block_is_read(self, tmp_path) -> None:
        from pension.config import read_config
        from pension.rostertemplate import write_roster_template

        path = write_roster_template(tmp_path / "양식.xlsx")
        config = read_config(openpyxl.load_workbook(path))
        assert config.base_date.isoformat() == "2025-12-31"
        assert [r.source_name for r in config.job_group_rules] == [
            "정규직", "계약직", "임원"]

    def test_the_rule_table_does_not_swallow_the_blocks_below_it(
        self, tmp_path
    ) -> None:
        """세 블록이 한 시트에 있다. 직군 표가 아래 표까지 삼키면 있지도 않은
        직군이 스무 개쯤 생긴다."""
        from pension.config import read_config
        from pension.rostertemplate import write_roster_template

        path = write_roster_template(tmp_path / "양식.xlsx")
        config = read_config(openpyxl.load_workbook(path))
        assert len(config.job_group_rules) == 3


class TestTheLookIsConsistent:
    """한 줄만 선이 끊겨 있으면 그 줄이 표 안인지 밖인지 알 수 없다.

    받는 사람은 이 파일 하나로 무엇을 어디에 적는지 알아내야 한다. 선이
    들쭉날쭉하면 "여기는 안 채워도 되나" 로 읽혀, 비워서 돌아온다.
    """

    def test_no_row_is_half_boxed(self, blank) -> None:
        for name in blank.sheetnames:
            ws = blank[name]
            for row in ws.iter_rows():
                written = [c for c in row if c.value is not None]
                if len(written) < 2:
                    continue
                bare = [c.coordinate for c in written
                        if not (c.border and c.border.left.style)]
                assert not bare or len(bare) == len(written), (
                    f"[{name}] {row[0].row}행의 {bare} 만 선이 없다")

    def test_the_face_is_the_same_everywhere(self, blank) -> None:
        """글꼴이 섞이면 한 사람이 만든 파일로 보이지 않는다."""
        for name in blank.sheetnames:
            ws = blank[name]
            faces = {c.font.name for row in ws.iter_rows() for c in row
                     if c.value is not None and c.font and c.font.name}
            assert faces <= {tpl.FACE}, f"[{name}] 에 다른 글꼴 {faces}"

    def test_no_markdown_leaks_into_the_cells(self, blank) -> None:
        """엑셀은 별표를 굵게 바꿔 주지 않는다. 그대로 별표로 보인다."""
        for name in blank.sheetnames:
            ws = blank[name]
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        assert "**" not in cell.value, (
                            f"[{name}] {cell.coordinate} 에 별표가 그대로 있다")

    def test_the_block_colours_match_the_roster(self, blank) -> None:
        """입력 시트의 띠와 명부의 블록이 같은 색이어야 한 벌로 보인다."""
        assert tuple(tpl.BAND_COLOURS) == tuple(
            colour for _part, colour, _ink, _note in BLOCKS.values()
            if colour in tpl.BAND_COLOURS
        ) or set(tpl.BAND_COLOURS) <= {v[1] for v in BLOCKS.values()}


class TestFrozenRowsCarryNoMerge:
    """고정(freeze)된 머리글 행에는 병합 칸을 두지 않는다.

    그 조합이 엑셀 제한된 보기에서 그리기를 깨뜨려, 멀쩡한 파일인데 머리글
    글자가 군데군데 사라져 보였다 — 어떤 칸은 윗줄만, 어떤 칸은 아랫줄만
    그려졌다. 파트·블록 띠는 '선택 영역의 가운데로' 정렬로 같은 모양을 낸다.
    """

    def test_roster_sheets_have_zero_merged_cells(self, blank) -> None:
        for name in ("재직자명부", "퇴직자명부", "추가명부"):
            ws = blank[name]
            assert ws.freeze_panes == f"A{FIRST_DATA_ROW}"
            assert not ws.merged_cells.ranges, (
                f"[{name}] 에 병합 칸 {len(ws.merged_cells.ranges)}개 — "
                "고정 행과 만나면 제한된 보기에서 글자가 사라진다")

    def test_the_bands_still_read_as_one_ribbon(self, blank) -> None:
        """병합을 뺐어도 띠 전체가 칠해지고 가운데 정렬로 이어져야 한다."""
        ws = blank["재직자명부"]
        for column in range(1, len(ACTIVE) + 1):
            cell = ws.cell(PART_ROW, column)
            assert cell.alignment.horizontal == "centerContinuous", cell.coordinate
            assert cell.fill.fgColor.rgb.startswith("FF")
