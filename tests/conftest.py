"""테스트용 워크북 생성기.

원본 통합문서와 같은 시트 구조·행 위치를 가진 명부를 만들어, 실제 파일 없이도
전 과정을 돌려볼 수 있게 한다.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl
import pytest

from pension.assumptions import write_template
from pension.readers import ACTIVE_FIRST_ROW, RETIRED_FIRST_ROW

BASE_DATE = dt.date(2025, 12, 31)

ACTIVE_ROWS = [
    # 사번, 임직원, 직군, 성명, 성별, 생년월일, 입사일, 중간정산일, 평균임금, 제도, 장기대상
    ("A001", "직원", "정규직", "김철수", "남자", "1980-03-15", "2005-04-01", "", 5_000_000, "DB", "Y"),
    ("A002", "직원", "정규직", "이영희", "여자", "1990-11-02", "2015-01-05", "", 4_200_000, "DB", "Y"),
    ("A003", "임원", "임원", "박대표", "남자", "1968-07-20", "1995-03-02", "", 12_000_000, "DB", "N"),
    ("A004", "직원", "계약직", "최기술", "남자", "1995-01-09", "2021-06-14", "", 3_600_000, "퇴직금제도", "N"),
    ("A005", "직원", "정규직", "정미나", "여자", "1985-05-30", "2010-09-01", "2018-01-01", 4_800_000, "DC", "N"),
]

RETIRED_ROWS = [
    # 사번, 임직원, 직군, 성명, 성별, 생년월일, 입사일, 퇴사일, 사유, 제도, 총지급, 사외자산
    ("R001", "직원", "정규직", "한퇴직", "남자", "1975-02-11", "2001-03-01", "2025-06-30", 1, "DB", 92_000_000, 80_000_000),
    ("R002", "직원", "계약직", "서만료", "여자", "1992-12-01", "2022-01-03", "2025-03-31", "계약만료", "퇴직금제도", 11_500_000, 0),
]


def _write_input(ws) -> None:
    ws["A1"] = "■ 명부 작성 Input 사항"
    ws["B3"], ws["C3"] = "산출기준일", BASE_DATE
    ws["B5"], ws["C5"] = "평균임금 첵크금액", 1_000_000

    headers = ["명부직군", "변환직군명", "퇴직급여 정년연령", "장기급여 정년연령", "정년연령 초과자 plus 연령"]
    for col, title in enumerate(headers, start=2):
        ws.cell(11, col, title)

    # 규정명은 기초율 워크북의 열 머리글과 일치해야 한다.
    rules = [
        ("임원", "2임원", 65, 65, 2, "정규임원", "정규임원"),
        ("정규직", "1정규직", 60, 60, 2, "정규직", "정규직"),
        ("계약직", "3계약직", 60, 60, 1, "계약직", "계약직"),
    ]
    for offset, (src, mapped, nra, jnra, add, withdrawal, salary) in enumerate(rules):
        row = 12 + offset
        ws.cell(row, 2, src)
        ws.cell(row, 3, mapped)
        ws.cell(row, 4, nra)
        ws.cell(row, 5, jnra)
        ws.cell(row, 6, add)
        # 퇴직급여 지급률(G열)은 비워 두어 법정 퇴직금(배수 = 근속연수)을 쓰게 하고,
        # 장기급여 지급률(H열)만 규정을 지정한다.
        ws.cell(row, 8, withdrawal)   # 장기급여 지급률 규정
        ws.cell(row, 9, withdrawal)   # 퇴직급여 중도퇴직률 규정
        ws.cell(row, 10, salary)      # 퇴직급여 승급률 규정
        ws.cell(row, 11, withdrawal)  # 장기급여 중도퇴직률 규정
        ws.cell(row, 12, salary)      # 장기급여 승급률 규정
        ws.cell(row, 13, withdrawal)  # 퇴직자 퇴직률 규정
        ws.cell(row, 14, withdrawal)


def _write_active(ws) -> None:
    ws.cell(22, 2, "작성기준일")
    ws.cell(22, 3, BASE_DATE)
    ws.cell(23, 3, "사번")
    for offset, values in enumerate(ACTIVE_ROWS):
        row = ACTIVE_FIRST_ROW + offset
        (emp_id, kind, group, name, sex, birth, hire, settle, wage, plan, longterm) = values
        ws.cell(row, 3, emp_id)
        ws.cell(row, 4, kind)
        ws.cell(row, 5, group)
        ws.cell(row, 6, name)
        ws.cell(row, 7, sex)
        ws.cell(row, 8, birth)
        ws.cell(row, 9, hire)
        ws.cell(row, 10, settle)
        ws.cell(row, 11, wage)
        ws.cell(row, 17, plan)
        ws.cell(row, 19, longterm)
        ws.cell(row, 36, "제조원가")


def _write_retired(ws) -> None:
    ws.cell(18, 2, "작성기간")
    ws.cell(19, 3, "사번")
    for offset, values in enumerate(RETIRED_ROWS):
        row = RETIRED_FIRST_ROW + offset
        (emp_id, kind, group, name, sex, birth, hire, exit_date,
         reason, plan, total, fund) = values
        ws.cell(row, 3, emp_id)
        ws.cell(row, 4, kind)
        ws.cell(row, 5, group)
        ws.cell(row, 6, name)
        ws.cell(row, 7, sex)
        ws.cell(row, 8, birth)
        ws.cell(row, 9, hire)
        ws.cell(row, 10, exit_date)
        ws.cell(row, 12, reason)
        ws.cell(row, 13, plan)
        ws.cell(row, 14, total)
        ws.cell(row, 15, fund)
        ws.cell(row, 20, "N")
        ws.cell(row, 24, "판관비")


# ── 일반사항 ──────────────────────────────────────────────────
# 실제 자료요청서의 행 배치를 그대로 옮겼다. 금액도 실제 파일의 것을 쓴다 —
# 검산줄이 맞아떨어지는 조합이라야 시험이 의미가 있다.

OBLIGATION_ROWS = (
    ("(+)증가", "계열사 전입", 1_742_800_303.0),
    ("", "합병", 0.0),
    ("(-)감소", "퇴직금 지급액", 129_623_571.0),
    ("", "중간정산금", 0.0),
    ("", "DC전환", 1_016_614_322.0),
    ("", "퇴직위로금", 0.0),
    ("", "계열사 전출", 0.0),
    ("", "사업처분/분할", 0.0),
)

ASSET_ROWS = (
    ("(+)증가", "부담금납입", 1_710_353_902.0),
    ("", "이자수익", 415_974_575.0),
    ("", "계열사 전입", 0.0),
    ("", "합병", 0.0),
    ("(-)감소", "퇴직금", 129_623_571.0),
    ("", "중간정산금", 0.0),
    ("", "DC전환", 1_016_614_322.0),
    ("", "계열사 전출", 0.0),
    ("", "사업처분/분할", 0.0),
    ("", "운용관리수수료", 12_704_896.0),
    ("", "자산관리수수료", 17_786_853.0),
)

ASSET_OPENING = 12_044_373_875.0
ASSET_CLOSING = 12_993_972_710.0


def write_general_sheet(
    ws,
    *,
    period=(dt.date(2025, 1, 1), dt.date(2025, 12, 31)),
    grade: str = "AA+",
    payout: dict[str, str] | None = None,
    obligation=OBLIGATION_ROWS,
    assets=ASSET_ROWS,
    opening: float = ASSET_OPENING,
    closing: float = ASSET_CLOSING,
    national_pension: float = 0.0,
    breakdown=(("⑴ 현금 및 현금등가물", ASSET_CLOSING),),
    longterm: tuple[float, float] | None = None,
) -> None:
    """``일반사항`` 시트를 자료요청서 서식대로 채운다."""
    if period is not None:
        ws.cell(22, 2, "2.")
        ws.cell(22, 3, "대상 회계기간")
        ws.cell(23, 3, "기시")
        ws.cell(23, 4, "기말")
        ws.cell(24, 3, period[0])
        ws.cell(24, 4, period[1])

    ws.cell(49, 3, "할인율 회사채 신용등급")
    ws.cell(53, 3, grade)

    if obligation is not None:
        ws.cell(65, 3, "1) 퇴직급여추계액 변동내역 (발생기준 작성)")
        ws.cell(66, 3, "구분")
        ws.cell(66, 5, "추계액")
        for offset, (group, label, amount) in enumerate(obligation):
            row = 67 + offset
            if group:
                ws.cell(row, 3, group)
            ws.cell(row, 4, label)
            ws.cell(row, 5, amount)

    head = 67 + (len(obligation) if obligation else 0) + 1
    if assets is not None:
        ws.cell(head, 3, "2) 사외적립자산 변동내역 (현금기준 작성)")
        ws.cell(head + 1, 3, "구분")
        ws.cell(head + 1, 5, "DB퇴직연금,\n퇴직보험")
        ws.cell(head + 1, 6, "국민연금전환금")
        ws.cell(head + 1, 7, "합계")
        ws.cell(head + 2, 3, dt.datetime(2025, 1, 1))
        ws.cell(head + 2, 5, opening)
        ws.cell(head + 2, 7, opening)
        for offset, (group, label, amount) in enumerate(assets):
            row = head + 3 + offset
            if group:
                ws.cell(row, 3, group)
            ws.cell(row, 4, label)
            ws.cell(row, 5, amount)
            ws.cell(row, 7, amount)
        last = head + 3 + len(assets)
        ws.cell(last, 3, dt.datetime(2025, 12, 31))
        ws.cell(last, 5, closing - national_pension)
        if national_pension:
            ws.cell(last, 6, national_pension)
        ws.cell(last, 7, closing)
        ws.cell(last + 1, 3, "검증")

        detail = last + 3
        ws.cell(detail, 3, "3) 사외적립자산 세부내역")
        ws.cell(detail + 2, 3, "구분")
        ws.cell(detail + 2, 5, "금액")
        for offset, (label, amount) in enumerate(breakdown):
            ws.cell(detail + 3 + offset, 3, label)
            ws.cell(detail + 3 + offset, 5, amount)
        ws.cell(detail + 3 + len(breakdown), 3, "합계")
        ws.cell(detail + 3 + len(breakdown), 5, sum(a for _, a in breakdown))

    for key, row in PAYOUT_ROWS.items():
        ws.cell(row, 5, (payout or {}).get(key, ""))

    if longterm is not None:
        # 표 제목이 항목 이름을 그대로 품는다 — 실제 서식 그대로 둔다. 제목을
        # 먼저 집으면 금액 칸이 비어 0 원으로 읽힌다.
        ws.cell(133, 3, "2) 기중 장기근속 지급액")
        ws.cell(134, 3, "구분")
        ws.cell(134, 5, "금액")
        ws.cell(135, 3, "(-)감소")
        ws.cell(135, 4, "장기근속 지급액")
        ws.cell(135, 5, longterm[0])
        ws.cell(136, 3, "(+)증가")
        ws.cell(136, 4, "장기근속 받은금액")
        ws.cell(136, 5, longterm[1])


#: 6번 '회사의 퇴직금 지급규정' 항목의 행 배치.
PAYOUT_ROWS = {
    "eligibility": 110, "service_period": 111, "formula": 112, "base_wage": 113,
    "staff_nra": 114, "executive_nra": 115,
}


@pytest.fixture
def roster_path(tmp_path: Path) -> Path:
    """원본과 같은 시트 구조를 가진 명부 워크북."""
    wb = openpyxl.Workbook()
    del wb["Sheet"]
    _write_input(wb.create_sheet("Input"))
    _write_active(wb.create_sheet("재직자명부"))
    _write_retired(wb.create_sheet("퇴직자명부"))

    path = tmp_path / "명부.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def assumptions_path(tmp_path: Path) -> Path:
    """테스트용 기초율 워크북(양식 생성기가 만든 예시 값 사용)."""
    path = tmp_path / "기초율.xlsx"
    write_template(path, job_groups=["정규직", "정규임원", "계약직"])
    return path
