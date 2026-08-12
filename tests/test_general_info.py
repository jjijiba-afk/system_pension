"""``일반사항`` 6번 항목에서 지급규정 초안 읽기.

실제 케이스 6건의 문구를 그대로 놓고 시험한다. 자유서술이라 기계가 확정할 수
없으므로, **읽은 것과 못 읽은 것이 정확히 갈리는지** 가 핵심이다. 못 읽은 항목을
기본값으로 조용히 메우면 담당자가 확인 없이 넘어간다.
"""

from __future__ import annotations

import datetime as _dt

import openpyxl
import pytest
from conftest import (
    ASSET_CLOSING,
    ASSET_OPENING,
    ASSET_ROWS,
    write_general_sheet,
)

from pension.actuarial import (
    FRACTION_DOWN,
    SERVICE_ANNUAL,
    SERVICE_DAILY,
    SERVICE_MONTHLY,
)
from pension.general_info import read_general_info
from pension.workbook import open_workbook

# 실제 케이스에서 그대로 가져온 문구
CASES = {
    "1": {
        "eligibility": "전 임직원",
        "service_period": "근속기간 1년 이상 (월할 계산)",
        "formula": "근속일수*(월평균임금/365)*지급률",
        "base_wage": "퇴직전 3개월 임금의 30일평균임금",
        "staff_nra": "만 60세",
        "executive_nra": "없음",
    },
    "2": {
        "eligibility": "입사후 1년 이상",
        "service_period": "입사후 1년 이상",
        "formula": "직전 3개월 평균임금",
        "base_wage": "ROUND(평균임금*근속년월*근속개수1,-1)",
        "staff_nra": "만 60세",
        "executive_nra": "정년 없음",
    },
    "4": {
        "eligibility": "1년 이상 근속한 전직원",
        "service_period": "근로기준법에 따른 산정법 - 일수",
        "formula": "(근속년수*평균임금)+위로금[대상자에한함]",
        "base_wage": "퇴직전 3개월 임금의 30일 평균임금",
        "staff_nra": "만 60세",
        "executive_nra": "없음",
    },
    "5": {
        "eligibility": "임원·정규직사원은 전 직원 대상/계약사원은 대상에서 제외",
        "service_period": "대상기간은 입사일부터 퇴직일까지로 하며 근속 3년 이상이 대상이 된다"
                          "/1년이 되지 않는 단수개월은 절사 예) 3년 10개월인 경우에는 3년",
        "formula": "기초금액 × 근속연수별 지급률 / 근속연수별 지급률은 근속연수의 1/2",
        "base_wage": "기초금액 = 퇴직년도 1개월 임금",
        "staff_nra": "정규직 60세, 계약직 60세",
        "executive_nra": "제약 無",
    },
    "10": {
        "eligibility": "",
        "service_period": "",
        "formula": "",
        "base_wage": "",
        "staff_nra": "정규직 60세",
        "executive_nra": "39세",
    },
}

_ROWS = {
    "eligibility": 110, "service_period": 111, "formula": 112, "base_wage": 113,
    "staff_nra": 114, "executive_nra": 115,
}


def _workbook(values: dict[str, str], tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "일반사항"
    for key, row in _ROWS.items():
        ws.cell(row, 5, values.get(key, ""))
    # 명부 시트도 있어야 실제 파일과 비슷하다.
    roster = wb.create_sheet("2)재직자명부")
    roster.cell(20, 2, "작성기준일")
    roster.cell(20, 3, _dt.datetime(2025, 12, 31))
    path = tmp_path / "명부.xlsx"
    wb.save(path)
    return open_workbook(path)


class TestCaseByCase:
    def test_case1_monthly(self, tmp_path) -> None:
        book = _workbook(CASES["1"], tmp_path)
        try:
            draft = read_general_info(book).draft
        finally:
            book.close()
        assert draft.min_service_years == 1.0
        assert draft.staff_nra == 60
        assert draft.service_basis == SERVICE_MONTHLY
        assert draft.executive_unlimited is True
        assert draft.executive_nra == 60   # '없음' → 직원 정년과 동일
        assert not draft.unread

    def test_case2_rounding_to_ten_won(self, tmp_path) -> None:
        book = _workbook(CASES["2"], tmp_path)
        try:
            draft = read_general_info(book).draft
        finally:
            book.close()
        assert draft.rounding_unit == 10
        assert draft.min_service_years == 1.0
        # 기간 산정방법은 어디에도 안 적혀 있다 — 지어내면 안 된다.
        assert draft.service_basis is None
        assert "근속기간 산정방법" in draft.unread

    def test_case4_statutory_days(self, tmp_path) -> None:
        book = _workbook(CASES["4"], tmp_path)
        try:
            draft = read_general_info(book).draft
        finally:
            book.close()
        assert draft.service_basis == SERVICE_DAILY
        assert draft.min_service_years == 1.0

    def test_case5_three_years_and_truncation(self, tmp_path) -> None:
        book = _workbook(CASES["5"], tmp_path)
        try:
            draft = read_general_info(book).draft
        finally:
            book.close()
        assert draft.min_service_years == 3.0
        assert draft.service_fraction == FRACTION_DOWN
        # '1년이 되지 않는 단수개월은 절사' 는 연 단위로 센다는 뜻이다.
        assert draft.service_basis == SERVICE_ANNUAL

    def test_case10_executive_age_is_taken_literally(self, tmp_path) -> None:
        book = _workbook(CASES["10"], tmp_path)
        try:
            draft = read_general_info(book).draft
        finally:
            book.close()
        assert draft.staff_nra == 60
        assert draft.executive_nra == 39
        assert draft.executive_unlimited is False
        assert "가입자격(최소 근속연수)" in draft.unread

    def test_case3_blank_section_is_reported(self, tmp_path) -> None:
        """규정이 통째로 비어 있으면 기본값으로 메우지 않고 알린다."""
        book = _workbook({}, tmp_path)
        try:
            info = read_general_info(book)
        finally:
            book.close()
        assert info.has_payout_section is False
        assert info.draft.min_service_years is None
        assert info.draft.staff_nra is None
        assert any("비어 있" in item for item in info.draft.unread)


class TestParsing:
    @pytest.mark.parametrize(
        ("wording", "expected"),
        [
            ("없음", True), ("정년 없음", True), ("제약 無", True),
            ("해당없음", True), ("만 60세", False), ("39세", False),
        ],
    )
    def test_unlimited_executive_wordings(self, wording, expected, tmp_path) -> None:
        values = dict(CASES["1"], executive_nra=wording)
        book = _workbook(values, tmp_path)
        try:
            draft = read_general_info(book).draft
        finally:
            book.close()
        assert draft.executive_unlimited is expected

    def test_evidence_is_kept_for_review(self, tmp_path) -> None:
        book = _workbook(CASES["5"], tmp_path)
        try:
            draft = read_general_info(book).draft
        finally:
            book.close()
        assert "가입자격" in draft.evidence
        assert "3년" in draft.evidence["가입자격"]

    def test_missing_sheet_is_not_an_error(self, tmp_path) -> None:
        wb = openpyxl.Workbook()
        wb.active.title = "2)재직자명부"
        path = tmp_path / "명부만.xlsx"
        wb.save(path)

        book = open_workbook(path)
        try:
            info = read_general_info(book)
        finally:
            book.close()
        assert info.draft.min_service_years is None
        assert info.draft.unread


# ── 2·4·5·7번 항목: 담당자가 이미 채워 보낸 표 ──────────────────

_ASSET_ROWS = ASSET_ROWS
_ASSET_OPENING = ASSET_OPENING
_ASSET_CLOSING = ASSET_CLOSING


def _full_workbook(tmp_path, **kwargs):
    """자료요청서 전 항목이 채워진 워크북."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "일반사항"
    write_general_sheet(ws, payout=CASES["1"], **kwargs)

    roster = wb.create_sheet("2)재직자명부")
    roster.cell(20, 2, "작성기준일")
    roster.cell(20, 3, _dt.datetime(2025, 12, 31))
    path = tmp_path / "자료요청서.xlsx"
    wb.save(path)
    return open_workbook(path)


def _read(book):
    try:
        return read_general_info(book)
    finally:
        book.close()


class TestAccountingPeriod:
    def test_period_is_read_from_section_two(self, tmp_path) -> None:
        info = _read(_full_workbook(tmp_path))
        assert info.period_start == _dt.date(2025, 1, 1)
        assert info.period_end == _dt.date(2025, 12, 31)

    def test_absent_period_stays_none(self, tmp_path) -> None:
        info = _read(_workbook(CASES["1"], tmp_path))
        assert info.period_start is None
        assert info.period_end is None


class TestCreditGrade:
    @pytest.mark.parametrize("grade", ["AAA", "AA+", "AA0", "AA-", "A0", "국고채"])
    def test_grade_is_picked_from_the_allowed_list(self, grade, tmp_path) -> None:
        info = _read(_full_workbook(tmp_path, grade=grade))
        assert info.credit_grade == grade

    def test_free_text_is_not_taken_as_a_grade(self, tmp_path) -> None:
        """'미정' 을 등급으로 삼으면 할인율이 조용히 틀어진다."""
        info = _read(_full_workbook(tmp_path, grade="미정"))
        assert info.credit_grade == ""


class TestObligationMovement:
    def test_every_line_lands_in_its_own_field(self, tmp_path) -> None:
        book = _read(_full_workbook(tmp_path))
        movement = book.obligation
        assert movement.transfer_in == 1_742_800_303.0
        assert movement.benefits_paid == 129_623_571.0
        assert movement.dc_converted == 1_016_614_322.0
        assert movement.merger_in == 0.0
        assert movement.settlement_paid == 0.0
        assert movement.other_paid == 0.0
        assert movement.transfer_out == 0.0
        assert movement.disposal == 0.0
        assert movement.is_empty() is False

    def test_dc_transfer_is_not_counted_as_an_ordinary_payment(self, tmp_path) -> None:
        """DC전환은 정산이지 지급이 아니다 — 섞이면 증감표가 어긋난다."""
        movement = _read(_full_workbook(tmp_path)).obligation
        assert movement.benefits_paid != movement.dc_converted

    def test_blank_table_is_empty(self, tmp_path) -> None:
        movement = _read(_workbook(CASES["1"], tmp_path)).obligation
        assert movement.is_empty() is True

    def test_row_order_does_not_matter(self, tmp_path) -> None:
        """회사가 줄을 넣고 빼도 라벨로 찾으므로 값이 흔들리면 안 된다."""
        shuffled = (
            ("(-)감소", "퇴직위로금", 55_000_000.0),
            ("", "계열사 전입", 1_742_800_303.0),
            ("", "안내문구", 0.0),
            ("", "퇴직금 지급액", 129_623_571.0),
        )
        movement = _read(_full_workbook(tmp_path, obligation=shuffled)).obligation
        assert movement.other_paid == 55_000_000.0
        assert movement.transfer_in == 1_742_800_303.0
        assert movement.benefits_paid == 129_623_571.0


class TestAssetMovement:
    def test_balances_reconcile(self, tmp_path) -> None:
        """기초 + 납입 + 이자 + 전입 − 유출 = 기말. 검산줄이 맞아야 한다."""
        assets = _read(_full_workbook(tmp_path)).assets
        computed = (
            assets.opening
            + assets.contributions
            + assets.actual_return
            + assets.total_received
            - assets.total_paid
        )
        assert round(computed - assets.closing) == 0

    def test_fees_are_kept_apart(self, tmp_path) -> None:
        assets = _read(_full_workbook(tmp_path)).assets
        assert assets.management_fee == 12_704_896.0
        assert assets.custody_fee == 17_786_853.0
        # 수수료도 자산을 줄이므로 유출 합계에 들어간다.
        assert assets.total_paid == pytest.approx(
            129_623_571.0 + 1_016_614_322.0 + 12_704_896.0 + 17_786_853.0
        )

    def test_opening_and_closing_come_from_the_date_rows(self, tmp_path) -> None:
        """기초·기말 줄에는 항목 이름 대신 날짜가 적혀 라벨로는 못 찾는다."""
        assets = _read(_full_workbook(tmp_path)).assets
        assert assets.opening == _ASSET_OPENING
        assert assets.closing == _ASSET_CLOSING

    def test_national_pension_column_is_separate(self, tmp_path) -> None:
        book = _full_workbook(tmp_path, national_pension=300_000_000.0)
        assets = _read(book).assets
        assert assets.national_pension == 300_000_000.0
        # 합계 열이 기말 잔액이다 — DB 열만 읽으면 전환금이 사라진다.
        assert assets.closing == _ASSET_CLOSING

    def test_breakdown_labels_lose_their_numbering(self, tmp_path) -> None:
        book = _full_workbook(
            tmp_path,
            breakdown=(
                ("⑴ 현금 및 현금등가물", 10_000_000_000.0),
                ("⑵ 지분상품", 2_993_972_710.0),
            ),
        )
        assets = _read(book).assets
        assert assets.breakdown == {
            "현금 및 현금등가물": 10_000_000_000.0,
            "지분상품": 2_993_972_710.0,
        }

    def test_total_row_is_not_a_category(self, tmp_path) -> None:
        assets = _read(_full_workbook(tmp_path)).assets
        assert "합계" not in assets.breakdown

    def test_blank_table_is_empty(self, tmp_path) -> None:
        assets = _read(_workbook(CASES["1"], tmp_path)).assets
        assert assets.is_empty() is True

    def test_outflows_written_as_negative_still_reduce(self, tmp_path) -> None:
        """'(-)감소' 칸을 음수로 적어 오는 회사가 있다.

        부호를 그대로 빼면 유출이 유입으로 뒤집혀, 자산이 유출액의 두 배만큼
        부풀려진다. 방향은 항목 이름이 이미 정해 놓았으므로 크기만 쓴다.
        """
        signed = tuple(
            (group, label, -amount if group == "(-)감소" or label in {
                "중간정산금", "DC전환", "계열사 전출", "사업처분/분할",
                "운용관리수수료", "자산관리수수료",
            } else amount)
            for group, label, amount in _ASSET_ROWS
        )
        assets = _read(_full_workbook(tmp_path, assets=signed)).assets
        assert assets.benefits_paid == 129_623_571.0
        assert assets.management_fee == 12_704_896.0
        assert round(assets.difference) == 0

    def test_investment_loss_keeps_its_sign(self, tmp_path) -> None:
        """운용손실은 실제로 음수다 — 여기까지 절댓값을 씌우면 안 된다."""
        losing = tuple(
            (group, label, -415_974_575.0 if label == "이자수익" else amount)
            for group, label, amount in _ASSET_ROWS
        )
        assets = _read(
            _full_workbook(tmp_path, assets=losing, closing=_ASSET_CLOSING)
        ).assets
        assert assets.actual_return == -415_974_575.0

    def test_component_columns_beat_a_shifted_total(self, tmp_path) -> None:
        """합계 열이 한 칸 밀려 적힌 파일이 실제로 있었다.

        합계를 먼저 믿으면 수수료가 두 줄에 걸쳐 두 번 잡힌다. 구성 열이
        적혀 있으면 그것을 더한다.
        """
        book = _full_workbook(tmp_path)
        book.close()
        wb = openpyxl.load_workbook(tmp_path / "자료요청서.xlsx")
        ws = wb["일반사항"]
        fee_row = next(
            r for r in range(60, 100) if ws.cell(r, 4).value == "운용관리수수료"
        )
        ws.cell(fee_row, 7, 0)                      # 합계가 한 줄 밀렸다
        ws.cell(fee_row + 1, 5, 0)
        ws.cell(fee_row + 1, 7, 12_704_896.0)
        shifted = tmp_path / "밀림.xlsx"
        wb.save(shifted)

        assets = _read(open_workbook(shifted)).assets
        assert assets.management_fee == 12_704_896.0
        assert assets.custody_fee == 0.0

    def test_reported_difference_is_zero_when_the_table_balances(self, tmp_path) -> None:
        assert round(_read(_full_workbook(tmp_path)).assets.difference) == 0

    def test_broken_table_reports_its_gap(self, tmp_path) -> None:
        """회사 표가 스스로 안 맞으면 조용히 쓰지 말고 차이를 알려야 한다."""
        assets = _read(
            _full_workbook(tmp_path, closing=_ASSET_CLOSING + 9_103_134.0)
        ).assets
        assert round(assets.difference) == -9_103_134


class TestLongtermAmounts:
    def test_paid_and_received_are_read(self, tmp_path) -> None:
        info = _read(_full_workbook(tmp_path, longterm=(12_000_000.0, 3_000_000.0)))
        assert info.longterm_paid == 12_000_000.0
        assert info.longterm_received == 3_000_000.0

    def test_absent_rows_stay_zero(self, tmp_path) -> None:
        info = _read(_full_workbook(tmp_path))
        assert info.longterm_paid == 0.0
        assert info.longterm_received == 0.0


class TestPayoutSectionStillWorks:
    def test_tables_do_not_disturb_the_rule_draft(self, tmp_path) -> None:
        """5번 표를 읽느라 6번 지급규정이 밀리면 안 된다."""
        draft = _read(_full_workbook(tmp_path)).draft
        assert draft.min_service_years == 1.0
        assert draft.staff_nra == 60


class TestTheStandardTemplateReadsBackWhole:
    """회사에 보내는 표준 양식은 **적은 것이 전부 읽혀야** 한다.

    물어봐 놓고 읽지 않으면 담당자는 엑셀에 채운 것을 화면에 손으로 한 번 더
    옮겨야 한다. 그 자리에서 자릿수를 틀린다.
    """

    def _book(self, path, **numbers):
        import openpyxl

        from pension.rostertemplate import build_workbook

        build_workbook(**numbers).save(path)
        return openpyxl.load_workbook(path, data_only=True)

    def test_the_blank_template_gives_period_grade_and_a_balanced_table(
        self, tmp_path
    ) -> None:
        book = self._book(tmp_path / "양식.xlsx")
        info = read_general_info(book)

        assert info.period_start is not None and info.period_end is not None
        assert info.period_start < info.period_end
        assert info.credit_grade, "신용등급이 [기본정보] 에 있는데 읽히지 않았다"
        assert round(info.assets.difference) == 0
        # 세부내역이 아래 표까지 훑어 들어가면 합이 기말과 어긋난다.
        assert sum(info.assets.breakdown.values()) == pytest.approx(
            info.assets.closing, rel=1e-9)

    def test_the_ceiling_and_unpaid_boxes_are_read(self, tmp_path) -> None:
        book = self._book(
            tmp_path / "채운양식.xlsx",
            numbers={
                "opening": (1_000_000_000, 0), "closing": (1_000_000_000, 0),
                "asset": {}, "obligation": {},
                "breakdown": [("현금 및 현금등가물", 1_000_000_000)],
                "extras": {"자산인식상한 (문단 64)": 250_000_000,
                           "기준일 현재 미지급 퇴직급여": 3_000_000},
            },
        )
        assets = read_general_info(book).assets
        assert assets.asset_ceiling == 250_000_000
        assert assets.unpaid_benefits == 3_000_000

    def test_an_unwritten_ceiling_is_not_zero(self, tmp_path) -> None:
        """0 은 '상한이 0 원', 빈칸은 '미적용' 이다. 섞으면 자산을 통째로 깎는다."""
        book = self._book(tmp_path / "빈양식.xlsx")
        assert read_general_info(book).assets.asset_ceiling is None
