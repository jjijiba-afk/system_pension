"""[전년명부] 가 있으면 당기 명부와 사람 단위로 맞대어 본다.

전기 산출 결과가 있으면 화면에서 그것과 맞대면 된다. 그런데 **첫 해에 맡은
회사나 다른 곳에서 넘겨받은 회사는 맞댈 상대가 없다.** 그러면 당기 명부가
스스로 맞다고 말하는 것 외에 확인할 길이 없고, 사람이 통째로 빠져도 알 수
없다 — 명부에 없으니 빠졌다는 사실 자체가 명부에 안 적힌다.

회사가 작년 파일을 그대로 한 장 더 붙여 주면 그 자리를 메운다. 우리 산출
결과가 아니라 **회사 자료만으로** 신규·퇴사 인원을 검증하는 셈이다.
"""

from __future__ import annotations

import openpyxl
import pytest

from pension.readers import PRIOR_SHEET


@pytest.fixture
def pack(tmp_path):
    from pension.samples import (
        ROSTER_TEMPLATE, STANDARD_ASSUMPTIONS, write_sample_pack)

    write_sample_pack(tmp_path)
    return tmp_path / ROSTER_TEMPLATE, tmp_path / STANDARD_ASSUMPTIONS


def run(roster, assumptions, tmp_path):
    from pension.pipeline import RunOptions, run_valuation

    return run_valuation(RunOptions(
        roster_path=roster, assumptions_path=assumptions,
        output_path=tmp_path / "결과.xlsx",
    ))


def copy_actives_into_prior(path, *, drop: int = 0):
    """재직자명부를 전년명부로 옮겨 붙인다. ``drop`` 만큼 뒤에서 뺀다.

    회사가 작년 파일을 그대로 붙이는 상황을 흉내낸다.
    """
    wb = openpyxl.load_workbook(path)
    source, target = wb["재직자명부"], wb[PRIOR_SHEET]
    width = source.max_column
    for offset, row in enumerate(range(1, source.max_row + 1 - drop)):
        for col in range(1, width + 1):
            target.cell(offset + 1, col, source.cell(row, col).value)
    wb.save(path)
    return path


class TestTheSheetIsOptional:
    def test_no_sheet_means_no_extra_words(self, pack, tmp_path) -> None:
        """양식을 그대로 돌린 산출에는 아무 말도 붙지 않는다.

        전년명부는 **있으면 좋은** 자료다. 없다고 경고하면 대부분의 산출에
        지워지지 않는 잔소리가 하나 붙는다.
        """
        roster, assumptions = pack
        run_result = run(roster, assumptions, tmp_path)
        codes = [i.code for i in run_result.issues.warnings]
        assert "JAE_PRIOR_SHEET" not in codes, codes
        assert run_result.roster.prior == []

    def test_an_empty_sheet_is_the_same_as_none(self, pack, tmp_path) -> None:
        """빈 채로 돌아와도 마찬가지다 — 이 시트는 비어 있는 것이 정상이다."""
        roster, assumptions = pack
        run_result = run(roster, assumptions, tmp_path)
        assert not run_result.issues.has_errors()


class TestItCatchesAMissingPerson:
    def test_the_same_people_raise_nothing(self, pack, tmp_path) -> None:
        """작년과 올해가 같은 사람들이면 짚을 것이 없다."""
        roster, assumptions = pack
        run_result = run(copy_actives_into_prior(roster), assumptions, tmp_path)
        assert run_result.roster.prior, "전년명부를 못 읽었다"
        codes = [i.code for i in run_result.issues.warnings]
        assert "JAE_PRIOR_SHEET" not in codes, codes

    def test_someone_who_vanished_is_named(self, pack, tmp_path) -> None:
        """작년에 있던 사람이 올해 어디에도 없으면 사번까지 짚는다.

        재직자에서 빠뜨렸는데 퇴직자명부에도 안 옮긴 경우다. 당기 명부만
        보아서는 영영 드러나지 않는다.
        """
        roster, assumptions = pack
        copy_actives_into_prior(roster)

        # 당기 재직자에서 한 사람을 지운다. 전년명부에는 그대로 남아 있다.
        # 지울 줄은 **엔진이 읽은 대로** 고른다 — 자리를 손으로 세면 서식이
        # 바뀔 때 빈 줄을 지우고도 시험이 통과한다.
        before = run(roster, assumptions, tmp_path)
        victim = before.roster.active[0]
        gone = victim.employee_id
        wb = openpyxl.load_workbook(roster)
        wb["재직자명부"].delete_rows(victim.row)
        wb.save(roster)

        run_result = run(roster, assumptions, tmp_path)
        said = [i for i in run_result.issues.warnings if i.code == "JAE_PRIOR_SHEET"]
        assert said, "작년 사람이 사라졌는데 아무 말이 없다"
        assert any(str(gone) in (i.employee_id or "") for i in said), \
            f"사번을 짚지 않았다: {[i.employee_id for i in said]}"


def test_the_template_carries_the_sheet(tmp_path) -> None:
    """양식에 시트가 있는지. 없으면 아무도 안 채운다."""
    from pension.rostertemplate import write_roster_template

    path = write_roster_template(tmp_path / "양식.xlsx")
    wb = openpyxl.load_workbook(path)
    assert PRIOR_SHEET in wb.sheetnames
    # 예시 줄을 넣으면 그 사람이 전기 재직자로 섞여 든다. 머리글만 있어야 한다.
    ws = wb[PRIOR_SHEET]
    filled = [ws.cell(r, 2).value for r in range(5, ws.max_row + 1)]
    assert not any(filled), f"예시 줄이 남아 있다: {filled}"
    # 열은 재직자명부와 같은 것이어야 한다 — 회사가 작년 파일을 그대로 붙인다.
    assert [c.value for c in wb[PRIOR_SHEET][4]] == [
        c.value for c in wb["재직자명부"][4]]


def test_the_guide_mentions_it(tmp_path) -> None:
    """작성요령이 이 시트를 설명하는지. 설명이 없으면 빈 장으로만 보인다."""
    from pension.rostertemplate import write_roster_template

    path = write_roster_template(tmp_path / "양식.xlsx")
    wb = openpyxl.load_workbook(path)
    text = "\n".join(
        str(cell.value) for row in wb["작성요령"].iter_rows()
        for cell in row if cell.value is not None)
    assert PRIOR_SHEET in text


def test_relaying_a_roster_keeps_the_prior_sheet(tmp_path) -> None:
    """되받은 파일이 곧 다음 결산에 보낼 양식이다 — 전년명부를 흘리면 안 된다.

    회사가 채워 보낸 자료가 사라진 파일을 돌려주면, 다음 해에는 그 시트가 빈
    채로 돌아온다. 검산 한 겹이 조용히 없어진다.
    """
    from pension.rosterexport import relayout_roster
    from pension.samples import ROSTER_TEMPLATE, write_sample_pack

    write_sample_pack(tmp_path)
    source = copy_actives_into_prior(tmp_path / ROSTER_TEMPLATE)
    target = tmp_path / "표준양식.xlsx"
    relayout_roster(source, target)

    wb = openpyxl.load_workbook(target)
    assert PRIOR_SHEET in wb.sheetnames
    ids = [wb[PRIOR_SHEET].cell(r, 2).value
           for r in range(5, wb[PRIOR_SHEET].max_row + 1)]
    assert any(ids), "전년명부가 빈 채로 돌아갔다"
