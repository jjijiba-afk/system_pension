"""명부 직급·직군을 산출 직군으로 묶기.

명부의 `직군` 열에 무엇이 오는지는 회사가 정한다. 실제 파일에서 본 것만 해도
고용형태(`정규직`/`계약직`), 직급(`사원`/`과장`/`대표이사`), 그리고 직군은 비고
`임직원구분` 에만 `정사원`/`촉탁사원`/`임원（주재원）` 이 적힌 경우가 있었다.
"""

from __future__ import annotations

import datetime as _dt

import pytest

from pension.config import CalculationConfig, JobGroupRule
from pension.jobgroup import (
    DEFAULT_GROUPS,
    RosterGroup,
    suggest_group,
    suggest_mapping,
)
from pension.normalize import EmployeeType, normalize_employee_type

BASE = _dt.date(2025, 12, 31)


class TestSuggestGroup:
    """제안은 제안일 뿐이지만, 흔한 것은 맞혀야 손이 덜 간다."""

    @pytest.mark.parametrize(
        ("job", "kind", "expected"),
        [
            ("정규직", "직원", "정규직"),
            ("계약직", "직원", "계약직"),
            ("정규직", "임원", "임원"),
            # 직급이 직군 열에 들어온 명부
            ("사원", "정사원", "정규직"),
            ("과장", "촉탁사원", "계약직"),
            ("대표이사", "임원（주재원）", "임원"),
            ("부장", "임원", "임원"),
            # 직군만 있고 임직원구분이 빈 명부
            ("기간제 근로자", "", "계약직"),
            ("전무", "", "임원"),
            ("생산직", "", "정규직"),
        ],
    )
    def test_common_shapes(self, job: str, kind: str, expected: str) -> None:
        assert suggest_group(job, kind) == expected

    def test_unlisted_group_falls_back_to_the_first(self) -> None:
        """묶음을 `생산직/관리직` 으로만 두면 '임원' 을 제안할 수 없다."""
        assert suggest_group("대표이사", "임원", ["생산직", "관리직"]) == "생산직"

    def test_custom_group_names_are_honoured(self) -> None:
        groups = ["생산직", "관리직", "일반직", "계약직", "임원"]
        assert suggest_group("과장", "촉탁사원", groups) == "계약직"
        assert suggest_group("대표이사", "임원", groups) == "임원"
        # 생산직/관리직은 규정을 봐야 갈리므로 프로그램이 정하지 않는다.
        assert suggest_group("사원", "정사원", groups) == "생산직"

    def test_empty_group_list_still_answers(self) -> None:
        assert suggest_group("사원", "정사원", []) == "정규직"

    def test_suggestion_is_keyed_by_pair(self) -> None:
        found = [
            RosterGroup("과장", "정사원", "직원", active=6),
            RosterGroup("과장", "촉탁사원", "직원", active=2),
        ]
        mapping = suggest_mapping(found, DEFAULT_GROUPS)
        assert mapping[("과장", "정사원")] == "정규직"
        assert mapping[("과장", "촉탁사원")] == "계약직"


class TestExecutiveNormalization:
    """VBA 는 완전일치만 보아 `임원（주재원）` 을 직원으로 분류했다.

    임원은 정년·지급배수가 달라 그대로 두면 채무가 어긋난다.
    """

    @pytest.mark.parametrize(
        "token", ["임원", "임원（주재원）", "임원(주재원)", "임원A", "이사", "등기임원", "Y", "2"]
    )
    def test_executive(self, token: str) -> None:
        assert normalize_employee_type(token) is EmployeeType.EXECUTIVE

    @pytest.mark.parametrize("token", ["직원", "정사원", "촉탁사원", "비임원", "Senior Manager", ""])
    def test_staff(self, token: str) -> None:
        assert normalize_employee_type(token) is EmployeeType.STAFF


class TestLookup:
    """`지급규정` 한 행이 곧 하나의 조회 키다."""

    def _config(self) -> CalculationConfig:
        return CalculationConfig(
            base_date=BASE,
            job_group_rules=[
                JobGroupRule("과장", "계약직", 60, 60, 2, employee_type_filter="촉탁사원"),
                JobGroupRule("과장", "정규직", 60, 60, 2, employee_type_filter="정사원"),
                JobGroupRule("과장", "임원", 65, 65, 2, employee_type_filter="임원（주재원）"),
                JobGroupRule("사원", "정규직", 60, 60, 2, employee_type_filter="정사원"),
            ]
        )

    def test_raw_employee_type_splits_the_same_job(self) -> None:
        """정규화하면 `정사원`·`촉탁사원` 이 똑같이 '직원' 이라 구분이 사라진다."""
        config = self._config()
        assert config.find_job_group("과장", "직원", "정사원")[1].mapped_name == "정규직"
        assert config.find_job_group("과장", "직원", "촉탁사원")[1].mapped_name == "계약직"
        assert config.find_job_group("과장", "임원", "임원（주재원）")[1].mapped_name == "임원"

    def test_normalized_type_is_the_fallback(self) -> None:
        """명부 원문이 규정에 없으면 임원/직원으로라도 찾아본다."""
        config = CalculationConfig(
            base_date=BASE,
            job_group_rules=[
                JobGroupRule("정규직", "정규직", 60, 60, 2, employee_type_filter="직원"),
                JobGroupRule("정규직", "임원", 65, 65, 2, employee_type_filter="임원"),
            ]
        )
        # 명부에는 'Senior Manager' 라 적혀 있지만 규정에는 그런 칸이 없다.
        assert config.find_job_group("정규직", "직원", "Senior Manager")[1].mapped_name == "정규직"
        assert config.find_job_group("정규직", "임원", "임원")[1].mapped_name == "임원"

    def test_job_only_rule_still_matches(self) -> None:
        """임직원구분을 안 적은 기존 파일도 그대로 읽힌다."""
        config = CalculationConfig(
            base_date=BASE, job_group_rules=[JobGroupRule("정규직", "정규직", 60, 60, 2)]
        )
        assert config.find_job_group("정규직", "직원", "Senior")[1].mapped_name == "정규직"

    def test_unknown_job_is_none(self) -> None:
        assert self._config().find_job_group("파견직", "직원", "파견") is None

    def test_mapped_names_dedupe_in_order(self) -> None:
        assert self._config().mapped_names() == ["계약직", "정규직", "임원"]


class TestScanRoster:
    """실제 명부 두 장을 훑어 조합과 인원수를 센다."""

    def _book(self, tmp_path, rows):
        import openpyxl

        wb = openpyxl.Workbook()
        for sheet, kind in (("재직자명부", "active"), ("퇴직자명부", "retired")):
            ws = wb.create_sheet(sheet)
            ws.cell(1, 2, "순번")
            ws.cell(1, 3, "사번")
            ws.cell(1, 4, "임직원구분")
            ws.cell(1, 5, "직군")
            ws.cell(1, 6, "성명")
            ws.cell(1, 8, "생년월일")
            for index, (job, raw, where) in enumerate(rows, start=1):
                if where != kind:
                    continue
                ws.cell(index + 1, 2, index)
                ws.cell(index + 1, 3, f"A{index}")
                ws.cell(index + 1, 4, raw)
                ws.cell(index + 1, 5, job)
                ws.cell(index + 1, 8, "19900101")
        del wb["Sheet"]
        path = tmp_path / "명부.xlsx"
        wb.save(path)
        return path

    def test_counts_and_ordering(self, tmp_path) -> None:
        from pension.jobgroup import scan_roster
        from pension.workbook import open_workbook

        path = self._book(tmp_path, [
            ("과장", "정사원", "active"),
            ("과장", "정사원", "active"),
            ("과장", "촉탁사원", "active"),
            ("대표이사", "임원（주재원）", "active"),
            ("과장", "정사원", "retired"),
        ])
        book = open_workbook(path)
        try:
            found = scan_roster(book)
        finally:
            book.close()

        by_key = {g.key: g for g in found}
        assert by_key[("과장", "정사원")].active == 2
        assert by_key[("과장", "정사원")].retired == 1
        assert by_key[("과장", "촉탁사원")].headcount == 1
        assert by_key[("대표이사", "임원（주재원）")].normalized_type == "임원"
        # 인원 많은 순
        assert found[0].key == ("과장", "정사원")


class TestWorkbookRoundTrip:
    """가정 입력 화면이 저장한 `지급규정` 이 산출에서 그대로 조회 키가 된다."""

    def test_mapping_survives_the_workbook(self, tmp_path) -> None:
        import openpyxl

        from pension.config import PAYOUT_SHEET, read_payout_rules

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = PAYOUT_SHEET
        # 머리글 한 줄은 건너뛰므로 내용은 상관없다. 22열이 임직원구분.
        ws.append([f"열{c}" for c in range(1, 23)])
        # (명부직군, 변환직군명, …, 임직원구분)
        ws.append(["과장", "정규직", 60, 60, 2, *[""] * 8, 1, 0, 2, "", "일할", "유지", 0,
                   "반올림", "정사원"])
        ws.append(["과장", "계약직", 60, 60, 2, *[""] * 8, 1, 0, 2, "", "일할", "유지", 0,
                   "반올림", "촉탁사원"])
        path = tmp_path / "기초율.xlsx"
        wb.save(path)

        rules = read_payout_rules(openpyxl.load_workbook(path, data_only=True))
        config = CalculationConfig(base_date=BASE, job_group_rules=rules)

        assert config.find_job_group("과장", "직원", "정사원")[1].mapped_name == "정규직"
        assert config.find_job_group("과장", "직원", "촉탁사원")[1].mapped_name == "계약직"
        assert config.mapped_names() == ["정규직", "계약직"]


class TestPipelineUsesPayoutSheet:
    """가정 입력 화면이 저장한 `지급규정` 이 실제 산출에 반영되는지.

    직군 배정은 명부를 읽는 도중에 일어난다. 시트를 늦게 읽으면 이미 배정이
    끝난 뒤라 아무 효과가 없다.
    """

    def _roster(self, tmp_path):
        import openpyxl

        wb = openpyxl.Workbook()
        for sheet in ("재직자명부", "퇴직자명부"):
            ws = wb.create_sheet(sheet)
            # Input 시트가 없는 자료요청서 원본을 흉내낸다. 기준일은 명부 머리에서 찾는다.
            ws.cell(1, 1, "작성기준일")
            ws.cell(1, 2, "2025-12-31")
            for col, title in (
                (2, "순번"), (3, "사번"), (4, "임직원구분"), (5, "직군"), (6, "성명"),
                (7, "성별"), (8, "생년월일"), (9, "입사일자"), (10, "정산일자"),
                (13, "제도구분"), (14, "30일 평균임금"),
            ):
                ws.cell(3, col, title)
        ws = wb["재직자명부"]
        rows = [("과장", "정사원"), ("과장", "촉탁사원")]
        for index, (job, kind) in enumerate(rows, start=1):
            ws.cell(index + 3, 2, index)
            ws.cell(index + 3, 3, f"A{index}")
            ws.cell(index + 3, 4, kind)
            ws.cell(index + 3, 5, job)
            ws.cell(index + 3, 6, f"홍길{index}")
            ws.cell(index + 3, 7, "남")
            ws.cell(index + 3, 8, "19850101")
            ws.cell(index + 3, 9, "20100101")
            ws.cell(index + 3, 13, "DB")
            ws.cell(index + 3, 14, 3_000_000)
        del wb["Sheet"]
        path = tmp_path / "명부.xlsx"
        wb.save(path)
        return path

    def _assumptions(self, tmp_path):
        from pension.assumptions import write_assumptions
        from pension.config import PAYOUT_SHEET

        header = [f"열{c}" for c in range(1, 23)]
        row = lambda job, target, kind: [  # noqa: E731
            job, target, 60, 60, 2, *[""] * 8, 0, 0, 2, "", "일할", "그대로", 0, "반올림", kind
        ]
        return write_assumptions(
            tmp_path / "기초율.xlsx",
            {
                "할인율": (["연차", "할인율"], [[1, 0.045]]),
                "지급률": (["근속연수", "정규직", "계약직"], [[0, 1.0, 1.0]]),
                PAYOUT_SHEET: (header, [
                    row("과장", "정규직", "정사원"),
                    row("과장", "계약직", "촉탁사원"),
                ]),
            },
            None, None,
        )

    def test_payout_sheet_overrides_the_roster(self, tmp_path) -> None:
        from pension.pipeline import load_inputs

        config, roster, _assumptions, _log, _g = load_inputs(
            self._roster(tmp_path), self._assumptions(tmp_path)
        )
        assert not config.inferred
        by_kind = {m.employee_type_raw: m.job_group for m in roster.active}
        assert by_kind == {"정사원": "정규직", "촉탁사원": "계약직"}

    def test_without_the_sheet_the_roster_wins(self, tmp_path) -> None:
        """`지급규정` 이 없던 기존 기초율 파일도 그대로 돌아야 한다."""
        from pension.assumptions import write_assumptions
        from pension.pipeline import load_inputs

        path = write_assumptions(
            tmp_path / "기초율.xlsx",
            {"할인율": (["연차", "할인율"], [[1, 0.045]]),
             "지급률": (["근속연수", "과장"], [[0, 1.0]])},
            None, None,
        )
        config, roster, _a, _log, _g = load_inputs(self._roster(tmp_path), path)
        # Input 시트가 없으니 명부에서 직군을 끌어낸 잠정 설정이다.
        assert config.inferred
        assert {m.job_group for m in roster.active} == {"과장"}
