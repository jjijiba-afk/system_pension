"""받아 온 명부를 지금 양식으로 되돌려 주기.

회사가 보내오는 명부는 열 순서도 머리글도 제각각이다. 매 결산 눈으로 맞추는
대신 한 번 올린 것을 우리 양식으로 되받아, 다음 해에 그것을 채워 달라고 한다.

여기서 지키는 것은 하나다 — **값은 손대지 않는다.** 날짜를 고쳐 쓰거나 빈
성별을 채우거나 규정명을 메우면, 산출이 넘겨짚은 값이 회사가 적어 보낸 것처럼
되돌아가 다음 해에 원본 행세를 한다.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl
import pytest

from pension.layout import ACTIVE_HEADER_ALIASES, find_header_row
from pension.rosterexport import relayout_roster


def _messy(path: Path) -> Path:
    """실제로 받아 보는 모양의 명부.

    머리글이 우리 표준 표기가 아니고, 열 순서가 뒤죽박죽이고, 회사가 쓰는
    열('본부', '호봉')이 가운데 끼어 있고, 머리글 위에 안내 줄이 두 개 있다.
    """
    book = openpyxl.Workbook()
    del book["Sheet"]

    active = book.create_sheet("2)재직자명부")
    active["A1"] = "2025년 결산용 재직자 명부"
    active["A2"] = "인사팀 작성"
    headers = ["사원번호", "본부", "성명", "생년월일", "성별", "입사일",
               "호봉", "직군", "지급규정", "월평균임금", "제도구분"]
    for index, head in enumerate(headers, start=1):
        active.cell(3, index, head)
    rows = [
        ["A001", "영업본부", "김하나", dt.datetime(1980, 5, 3), "여",
         dt.datetime(2005, 3, 2), 12, "정규직", "직원규정", 4_100_000, "DB"],
        ["A002", "생산본부", "이두리", dt.datetime(1975, 11, 20), "남",
         dt.datetime(1999, 1, 4), 21, "임원", "임원규정", 9_000_000, "DB"],
    ]
    for offset, line in enumerate(rows):
        for index, value in enumerate(line, start=1):
            active.cell(4 + offset, index, value)

    retired = book.create_sheet("3)퇴직자명부")
    for index, head in enumerate(
        ["사원번호", "성명", "생년월일", "입사일", "퇴사일", "퇴직사유", "직군"],
        start=1,
    ):
        retired.cell(3, index, head)
    retired.append([])
    for index, value in enumerate(
        ["T001", "박세찌", dt.datetime(1970, 2, 2), dt.datetime(2000, 4, 1),
         dt.datetime(2025, 6, 30), "정년", "정규직"],
        start=1,
    ):
        retired.cell(4, index, value)

    book.save(path)
    return path


def _sheet_map(path: Path, sheet: str) -> tuple[list[str], list[list]]:
    """되받은 명부의 ``(머리글, 자료 줄들)``."""
    book = openpyxl.load_workbook(path)
    ws = book[sheet]
    row = find_header_row(ws, ACTIVE_HEADER_ALIASES)
    heads = [str(ws.cell(row, c).value or "") for c in range(1, ws.max_column + 1)]
    body = [
        [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        for r in range(row + 1, ws.max_row + 1)
        if any(ws.cell(r, c).value is not None for c in range(1, ws.max_column + 1))
    ]
    return heads, body


@pytest.fixture
def relaid(tmp_path) -> tuple[Path, object]:
    source = _messy(tmp_path / "받은명부.xlsx")
    target = tmp_path / "표준양식.xlsx"
    return target, relayout_roster(source, target)


class TestRelayout:
    def test_the_columns_land_in_our_order(self, relaid) -> None:
        """'사원번호'·'월평균임금' 처럼 다르게 적어 와도 우리 자리로 간다."""
        target, _report = relaid
        heads, _body = _sheet_map(target, "재직자명부")

        assert heads[:12] == [
            "순번", "사번", "성명", "주민등록번호 앞7자리", "생년월일", "성별",
            "입사일자", "중간정산일", "직군", "규정명", "30일 평균임금", "추계액",
        ]

    def test_the_values_are_carried_over_untouched(self, relaid) -> None:
        """값은 옮기기만 한다. 이것이 깨지면 되받은 명부를 믿을 수 없다."""
        target, _report = relaid
        heads, body = _sheet_map(target, "재직자명부")
        at = {head: index for index, head in enumerate(heads)}
        first = body[0]

        assert first[at["사번"]] == "A001"
        assert first[at["성명"]] == "김하나"
        assert first[at["성별"]] == "여"
        assert first[at["직군"]] == "정규직"
        assert first[at["규정명"]] == "직원규정"
        assert first[at["30일 평균임금"]] == 4_100_000
        assert first[at["생년월일"]] == dt.datetime(1980, 5, 3)

    def test_an_unknown_column_is_kept_on_the_right(self, relaid) -> None:
        """회사가 쓰는 열을 지워 돌려주면 그쪽이 자기 자료를 잃는다."""
        target, report = relaid
        heads, body = _sheet_map(target, "재직자명부")

        assert "본부" in heads and "호봉" in heads
        assert heads.index("본부") > heads.index("비고"), "표준 열 오른쪽에 붙어야 한다"
        assert body[0][heads.index("본부")] == "영업본부"
        assert report.active.carried == ["본부", "호봉"]

    def test_the_rows_are_renumbered(self, relaid) -> None:
        """순번은 우리가 다시 매긴다 — 원본에 없을 때가 더 많다."""
        target, _report = relaid
        heads, body = _sheet_map(target, "재직자명부")
        assert [line[heads.index("순번")] for line in body] == [1, 2]

    def test_the_retired_sheet_comes_too(self, relaid) -> None:
        target, report = relaid
        heads, body = _sheet_map(target, "퇴직자명부")
        assert body[0][heads.index("사번")] == "T001"
        assert body[0][heads.index("퇴직사유")] == "정년"
        assert report.retired.rows == 1

    def test_the_workbook_has_every_sheet_of_the_template(self, relaid) -> None:
        """되받은 것이 곧 다음 결산에 보낼 양식이다 — 안내 시트까지 있어야 한다."""
        target, _report = relaid
        assert openpyxl.load_workbook(target).sheetnames == [
            "작성요령", "기본정보", "퇴직급여규정", "장기급여규정", "사외적립자산",
            "재직자명부", "퇴직자명부", "추가명부",
        ]

    def test_it_reports_what_it_did(self, relaid) -> None:
        _target, report = relaid
        assert report.active.rows == 2
        assert ("사원번호", "사번") in report.active.matched
        assert ("월평균임금", "30일 평균임금") in report.active.matched
        assert "재직자명부 2명" in report.summary()


class TestRoundTrip:
    def test_our_own_template_survives_a_round_trip(self, tmp_path) -> None:
        """우리 양식으로 받은 것을 되돌리면 열이 그대로여야 한다.

        여기서 열이 하나라도 밀리면, 매년 되받을수록 명부가 조금씩 어긋난다.
        """
        from pension.rostergen import CASES, write_case_roster

        source = write_case_roster(CASES[0], tmp_path / "원본.xlsx")
        target = tmp_path / "되받은.xlsx"
        report = relayout_roster(source, target)

        before, _ = _sheet_map(source, "재직자명부")
        after, _ = _sheet_map(target, "재직자명부")
        assert after == before
        assert report.active.carried == []

    def test_the_result_still_reads_and_calculates(self, tmp_path) -> None:
        """되받은 명부로 산출이 돌아야 한다 — 안 돌면 돌려줄 이유가 없다."""
        from pension.pipeline import RunOptions, run_valuation
        from pension.rostergen import CASES, write_case_assumptions, write_case_roster

        spec = CASES[0]
        source = write_case_roster(spec, tmp_path / "원본.xlsx")
        target = tmp_path / "되받은.xlsx"
        relayout_roster(source, target)
        assumptions = write_case_assumptions(spec, tmp_path / "기초율.xlsx")

        before = run_valuation(RunOptions(
            roster_path=source, assumptions_path=assumptions,
            output_path=tmp_path / "전.xlsx"))
        after = run_valuation(RunOptions(
            roster_path=target, assumptions_path=assumptions,
            output_path=tmp_path / "후.xlsx"))

        assert not after.issues.has_errors()
        assert after.valuation.headcount == before.valuation.headcount
        assert after.valuation.dbo == pytest.approx(before.valuation.dbo)


class TestBadInput:
    def test_a_roster_without_headers_is_refused(self, tmp_path) -> None:
        """머리글이 없으면 어느 칸이 무엇인지 짚을 근거가 없다.

        짐작으로 옮기면 통째로 어긋난 명부를 돌려주게 된다 — 오류 없이
        그럴듯해 보이는 쪽이라 여기서 멈춰야 한다.
        """
        book = openpyxl.Workbook()
        del book["Sheet"]
        sheet = book.create_sheet("재직자명부")
        sheet.append(["A001", "김하나", 4_100_000])
        path = tmp_path / "머리글없음.xlsx"
        book.save(path)

        with pytest.raises(ValueError, match="머리글"):
            relayout_roster(path, tmp_path / "나온것.xlsx")
