"""명부 열 배치 인식.

같은 계열 서식인데도 통합문서마다 열 구성이 다르다. 재직자명부가 30열인 것이
보통인데, ``연봉제 전환 추계일``·``누진적용 근속연수``·``누진적용 율`` 처럼
회사가 쓰는 열이 중간에 끼어들면 그 뒤가 통째로 밀린다.

열 번호를 고정해 두면 이 명부를 읽을 때 ``누진적용 율`` 이 제도구분 자리로,
제도구분이 임금피크 자리로 들어간다. **오류 없이 그럴듯한 숫자가 나오는** 가장
위험한 형태다.

그래서 열은 번호가 아니라 **머리글 문자열로 찾는다**. 머리글을 못 찾으면
기본 위치로 물러서되, 무엇을 어떻게 찾았는지 :class:`LayoutReport` 로 남겨
담당자가 확인할 수 있게 한다.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Any

from .errors import IssueLog, Severity

__all__ = [
    "ACTIVE_HEADER_ALIASES",
    "RETIRED_HEADER_ALIASES",
    "LayoutReport",
    "ResolvedLayout",
    "find_data_start",
    "find_header_row",
    "resolve_layout",
]

#: 머리글 비교 전에 걷어낼 것들 — 줄바꿈, 공백, 괄호 주석, 단위 표기.
_BRACKETS = re.compile(r"[（(\[].*?[)）\]]", re.DOTALL)
_NOISE = re.compile(r"[\s　·,./\\'\"-]+")


def normalize_header(value: object) -> str:
    """머리글을 비교용 문자열로 정리한다.

    ``"임직원구분\\n(임원,직원)"`` → ``"임직원구분"``,
    ``"30일 평균임금"`` → ``"30일평균임금"``,
    ``"1日기본급"`` → ``"1일기본급"`` (한자 日 도 통일한다).
    """
    text = str(value or "").strip()
    if not text:
        return ""
    text = _BRACKETS.sub("", text)
    text = _NOISE.sub("", text)
    return text.replace("日", "일").replace("年", "년").replace("月", "월")


#: 필드명 → 머리글 별칭들. 앞의 것이 표준 표기다.
ACTIVE_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "employee_id": ("사번", "사원번호", "사원코드"),
    "employee_type": ("임직원구분", "임직원"),
    "job_group": ("직군", "직군구분"),
    "name": ("성명", "이름"),
    "gender": ("성별",),
    "resident_number": ("주민등록번호앞7자리", "주민등록번호", "주민번호", "주민등록번호앞자리"),
    "birth_date": ("생년월일", "생년월일자"),
    "hire_date": ("입사일자", "입사일"),
    "settlement_date": ("중간정산일", "중간정산일자"),
    "monthly_wage": ("30일평균임금", "월평균임금", "평균임금"),
    "honorary_wage": ("명예퇴직산정용임금", "명예퇴직임금", "명예퇴직기준임금"),
    "accrued_benefit": ("퇴직급여추계액", "추계액", "회사추계액"),
    "next_accrued_benefit": ("차년도추계액", "차년도퇴직급여추계액", "1년후추계액",
                             "익년도추계액", "차기추계액"),
    "daily_base_pay": ("1일기본급", "일기본급", "1일통상임금"),
    "leave_days": ("휴직차감일수", "휴직일수", "차감일수"),
    "plan": ("퇴직급여제도구분", "제도구분"),
    # 머리글 비교는 대소문자를 가리므로 두 표기를 모두 적어 둔다.
    "db_ratio": ("DB비율", "db비율", "DB비중", "db비중", "확정급여비율"),
    "remaining_contract_years": ("잔여계약기간", "잔여계약연수", "계약잔여기간"),
    # 날짜로 받는 쪽이 낫다 — 연수는 기준일이 바뀌는 순간 틀리고, 적는 사람이
    # 반올림해 넣는다. 둘 다 오면 날짜를 쓴다.
    "contract_end_date": ("계약종료일", "계약만료일", "근로계약종료일",
                          "계약종료일자", "계약만료일자", "계약기간종료일"),
    "settlement_amount": ("중간정산지급금액", "중간정산금액"),
    # 지급률 근속만 움직이는 가산·차감. 실제 서식은 '군경력 등 가산 근속연수' /
    # '차감근속연수(+로 입력)' 로 온다 — 괄호는 정규화에서 걷힌다.
    "service_add_years": ("가산근속연수", "군경력등가산근속연수", "가산근속년수",
                          "근속가산연수"),
    "service_deduct_years": ("차감근속연수", "차감근속년수", "근속차감연수"),
    "longterm_target": ("장기급여산출대상여부", "장기급여대상여부", "장기급여대상"),
    "wage_peak_age": ("임금피크연령", "임금피크"),
    "transfer_in_date": ("전입일", "전입일자"),
    # 장기근속포상 근속의 기산일. 실제 자료요청서들이 이 칸을 따로 받는다 —
    # 퇴직급여 기산일(중간정산일)과 다르기 때문이다.
    "longterm_start_date": ("장기급여기산일", "장기근속포상기산일", "장기포상기산일",
                            "장기근속기산일"),
    # 실제 자료요청서 서식에는 '퇴직급여 정년연령' / '장기급여 정년연령' 이
    # 두 줄로 있는데, 여기서 일부러 **안 읽는다**. 그 칸을 채워 보내면서도
    # 정년은 [지급규정] 화면의 직군 규칙으로 잡는 것이 실무 관행이라,
    # 읽어 버리면 담당자가 화면에 넣은 정년이 조용히 덮인다.
    # 사람마다 정년이 다르면 직군을 나누거나 이 별칭을 살려 쓴다.
    "declared_nra": ("정년연령", "정년연령세"),
    "longterm_amount": ("장기종업원급여지급금액", "장기급여지급액", "장기급여기지급액"),
    "transfer_in_amount": ("전입액", "전입받은금액", "전입인수액"),
    "payout_multiple": ("퇴직금지급배수", "임원지급배수", "지급배수", "개인지급배수",
                        "가산지급률", "가산감소지급률"),
    "note": ("비고",),
    # [추가명부] 전용. 재직자명부에는 이 열이 없어 그냥 안 잡힌다.
    "event_kind": ("사건구분", "사건유형", "구분사유", "변동구분"),
    "event_date": ("사건일", "사건일자", "발생일", "발생일자", "사건발생일"),
    "event_payment": ("지급액", "사건지급액", "정산지급액", "사건지급금액"),
    # 이 칸의 머리글은 회사마다 제각각이다 — 사람마다 다른 규정이 걸리는
    # 회사에서 여기가 안 잡히면 지급률을 넣을 열 자체가 안 생기고, 전원이
    # 직군 하나로 뭉뚱그려져 법정 퇴직금으로 떨어진다. 넉넉히 적어 둔다.
    "severance_benefit": (
        "퇴직급여지급률규정", "지급률규정", "규정명", "지급규정", "퇴직금지급규정",
        "퇴직급여규정", "퇴직금규정", "급여규정", "적용규정", "지급률규정명",
        "규정", "규정구분",
    ),
    "longterm_benefit": ("장기급여지급률규정",),
    "severance_withdrawal": ("퇴직급여중도사망퇴직률산출규정", "퇴직급여중도퇴직률산출규정",
                             "퇴직급여중도퇴직률규정", "퇴직률규정"),
    "severance_salary_increase": ("퇴직급여승급률산출규정", "퇴직급여승급률규정", "승급률규정"),
    "longterm_withdrawal": ("장기급여중도사망퇴직률산출규정", "장기급여중도퇴직률산출규정",
                            "장기급여중도퇴직률규정", "장기급여퇴직률규정"),
    "longterm_salary_increase": ("장기급여승급률산출규정", "장기급여승급률규정"),
    "extra_pay_base_date": ("추가지급기준일", "추가지급기산일"),
    "extra_pay_base_wage": ("추가지급기본급",),
    "cost_code": ("원가코드", "원가구분", "원가구분코드"),
    # 일부 서식에만 있는 열. 없어도 되지만 있으면 잡아 둔다.
    "progressive_service": ("누진적용근속연수",),
    "progressive_rate": ("누진적용율", "누진직용율"),
    "annual_salary_date": ("연봉제전환추계일",),
    "group_hire_date": ("그룹입사일자", "그룹입사일"),
    "period_start": ("지급률기산일", "지급구간시작일", "적용기간시작일"),
    "period_end": ("지급률종료일", "지급구간종료일", "적용기간종료일"),
}

RETIRED_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "employee_id": ("사번", "사원번호"),
    "employee_type": ("임직원구분", "임직원"),
    "job_group": ("직군",),
    "name": ("성명", "이름"),
    "gender": ("성별",),
    "resident_number": ("주민등록번호앞7자리", "주민등록번호", "주민번호", "주민등록번호앞자리"),
    "birth_date": ("생년월일",),
    "hire_date": ("입사일", "입사일자"),
    "exit_date": ("퇴사일", "퇴사일자"),
    "fund_payment_date": ("사외적립자산지급일", "사외자산지급일"),
    "reason": ("지급퇴직사유구분", "지급사유구분", "퇴직사유구분", "퇴직사유"),
    "plan": ("퇴직급여제도구분", "제도구분"),
    "total_payment": ("퇴직급여총지급금액", "총지급금액", "퇴직급여총지급액"),
    "fund_payment": ("퇴직급여사외자산지급금액", "사외자산지급금액", "사외자산지급액"),
    "national_pension_payment": ("퇴직급여국민연금전환금지급금액", "국민연금전환금지급금액", "국민연금전환금"),
    "longterm_payment": ("장기종업원급여지급금액", "장기급여지급금액", "장기급여지급액"),
    "other_payment": ("퇴직금이외의지급금액", "퇴직위로금등지급금액", "퇴직금이외지급금액", "퇴직위로금등"),
    "transfer_out_payment": ("전출분할사업처분지급금액", "전출지급금액", "전출지급액"),
    "longterm_target": ("장기급여산출대상여부", "장기급여대상여부", "장기급여대상"),
    "note": ("비고", "비고당기전입자가당기퇴직시전입액"),
    "severance_withdrawal": ("퇴직급여중도사망퇴직률산출규정", "퇴직급여중도퇴직률규정", "퇴직률규정"),
    "longterm_withdrawal": ("장기급여중도사망퇴직률산출규정", "장기급여중도퇴직률규정", "장기급여퇴직률규정"),
    "cost_code": ("원가코드", "원가구분", "원가구분코드"),
}

#: 이 열들을 못 찾으면 산출할 수 없다.
REQUIRED_ACTIVE = ("employee_id", "job_group", "birth_date", "hire_date", "monthly_wage", "plan")
REQUIRED_RETIRED = ("employee_id", "job_group", "birth_date", "hire_date", "exit_date")


@dataclass(slots=True)
class LayoutReport:
    """열을 어떻게 찾았는지에 대한 기록. 담당자 확인용."""

    sheet: str = ""
    header_row: int = 0
    data_start_row: int = 0
    matched: dict[str, tuple[int, str]] = field(default_factory=dict)
    """필드명 → (열 번호, 실제 머리글)."""
    fallback: dict[str, int] = field(default_factory=dict)
    """머리글을 못 찾아 기본 위치를 쓴 필드."""
    missing: list[str] = field(default_factory=list)
    """머리글도 없고 기본 위치도 없는 필드."""

    def lines(self) -> list[str]:
        out = [f"{self.sheet}: 머리글 {self.header_row}행, 데이터 {self.data_start_row}행부터"]
        if self.fallback:
            out.append(f"  머리글을 못 찾아 기본 위치를 쓴 열: {', '.join(sorted(self.fallback))}")
        if self.missing:
            out.append(f"  찾지 못한 열: {', '.join(self.missing)}")
        return out


@dataclass(slots=True)
class ResolvedLayout:
    """확정된 열 배치."""

    columns: dict[str, int]
    header_row: int
    data_start_row: int
    report: LayoutReport

    def index(self, field_name: str) -> int | None:
        return self.columns.get(field_name)


def find_header_row(sheet, aliases: dict[str, tuple[str, ...]], *, limit: int = 40) -> int:
    """머리글 행을 찾는다.

    ``사번`` 과 ``생년월일`` 이 같은 행에 있으면 머리글 행으로 본다. 안내문에
    같은 낱말이 섞여 있어도 두 개가 한 행에 나란히 오는 곳은 머리글뿐이다.
    """
    want_id = {normalize_header(a) for a in aliases.get("employee_id", ())}
    want_birth = {normalize_header(a) for a in aliases.get("birth_date", ())}

    best = 0
    for row in range(1, min(sheet.max_row, limit) + 1):
        seen = {
            normalize_header(sheet.cell(row, col).value)
            for col in range(1, min(sheet.max_column, 60) + 1)
        }
        if seen & want_id and seen & want_birth:
            best = row
            break
    return best


def find_data_start(sheet, header_row: int, *, limit: int = 12) -> int:
    """데이터가 시작하는 행.

    머리글 아래에는 ``TYPE`` 행과 ``작성 샘플`` 행이 한두 줄 끼어 있고, 그 수가
    통합문서마다 다르다. **순번이 숫자인 첫 행** 을 데이터 시작으로 본다.

    순번이 어느 열인지는 서식마다 다르다 — 첫 열을 비워 두고 둘째 열에 순번을
    적는 것이 있고, 첫 열이 곧 순번인 것이 있다. 앞 세 열을 함께 보아 어느 쪽이든
    잡는다. 여기서 한 줄을 잘못 잡으면 그만큼의 사람이 조용히 빠진다.

    순번 열이 아예 없는 명부도 온다(사번부터 시작하는 인사시스템 출력이 그렇다).
    그때 예전에는 ``머리글 + 3`` 으로 물러섰는데, 그러면 맨 앞 두 사람이 아무
    말 없이 빠졌다 — 사람이 빠져도 총액만 조금 작아져 눈에 띄지 않는 쪽이다.
    지금은 **글자라도 적힌 첫 줄** 로 물러선다. 안내 줄을 사람으로 잘못 세면
    검증에서 시끄럽게 걸리므로, 조용히 빠지는 것보다 낫다.
    """
    first_filled = 0
    for row in range(header_row + 1, min(sheet.max_row, header_row + limit) + 1):
        for column in (1, 2, 3):
            value = sheet.cell(row, column).value
            if isinstance(value, bool):
                continue
            if isinstance(value, (int, float)) and value >= 1:
                return row
        if not first_filled and any(
            str(sheet.cell(row, column).value or "").strip()
            for column in range(1, min(sheet.max_column, 12) + 1)
        ):
            first_filled = row
    return first_filled or header_row + 3


def _match(seen: dict[str, int], aliases: tuple[str, ...]) -> int | None:
    """머리글 별칭으로 열을 찾는다.

    실제 명부의 머리글에는 설명이 덧붙는다. 예를 들어 지급사유 열은
    ``"지급(퇴직)사유 구분\n1:중도퇴직\n2:사망퇴직 …"`` 처럼 적혀 있어 완전일치로는
    찾지 못한다. 그래서 완전일치를 먼저 보고, 없으면 접두 일치로 찾는다.
    """
    keys = [normalize_header(a) for a in aliases]
    for key in keys:
        if key in seen:
            return seen[key]
    for key in keys:
        for header, col in seen.items():
            if header.startswith(key):
                return col
    return None


def resolve_layout(
    sheet,
    aliases: dict[str, tuple[str, ...]],
    defaults: dict[str, int],
    required: tuple[str, ...],
    log: IssueLog,
) -> ResolvedLayout:
    """머리글을 읽어 열 배치를 확정한다.

    :param aliases: 필드명 → 머리글 별칭.
    :param defaults: 머리글을 못 찾았을 때 쓸 기본 열 번호.
    :param required: 없으면 산출할 수 없는 필드.
    """
    report = LayoutReport(sheet=sheet.title)
    header_row = find_header_row(sheet, aliases)

    if header_row == 0:
        # 머리글을 아예 못 찾았다. 기본 배치로 진행하되 반드시 알린다.
        report.header_row = 0
        report.data_start_row = defaults.get("_data_start", 26)
        report.fallback = dict(defaults)
        log.add(
            Severity.WARNING,
            "LAYOUT_HEADER_NOT_FOUND",
            f"'{sheet.title}' 에서 머리글 행을 찾지 못해 기본 열 배치를 씁니다. "
            "열이 밀려 있으면 결과가 틀릴 수 있으니 확인하세요",
            sheet=sheet.title,
        )
        columns = {k: v for k, v in defaults.items() if not k.startswith("_")}
        return ResolvedLayout(columns, 0, report.data_start_row, report)

    report.header_row = header_row
    report.data_start_row = find_data_start(sheet, header_row)

    # 머리글 → 열 번호. 같은 이름이 여러 번 나오면 처음 것을 쓴다.
    seen: dict[str, int] = {}
    for col in range(1, sheet.max_column + 1):
        key = normalize_header(sheet.cell(header_row, col).value)
        if key and key not in seen:
            seen[key] = col

    columns: dict[str, int] = {}
    for field_name, names in aliases.items():
        col = _match(seen, names)
        if col is not None:
            columns[field_name] = col
            report.matched[field_name] = (col, str(sheet.cell(header_row, col).value or "").strip())
        else:
            # 머리글 행을 제대로 찾았다면, 이름이 없는 열은 그 서식에 **없는**
            # 것이다. 기본 위치로 물러서면 엉뚱한 열을 그 값으로 읽게 된다.
            # (구 서식은 28열이 '추가지급 기본급' 인데 신 서식은 '지급률 규정'
            #  이라, 폴백하면 기본급이 규정명으로 들어간다.)
            report.missing.append(field_name)

    for field_name in required:
        if field_name not in columns:
            log.error(
                "LAYOUT_REQUIRED_COLUMN_MISSING",
                f"'{sheet.title}' 에서 '{aliases[field_name][0]}' 열을 찾지 못했습니다",
                sheet=sheet.title,
                row=header_row,
            )

    return ResolvedLayout(columns, header_row, report.data_start_row, report)


def describe(layout: ResolvedLayout, aliases: dict[str, tuple[str, ...]]) -> list[list[Any]]:
    """열 인식 결과를 리포트 표로. (필드, 열, 머리글, 판정)"""
    from openpyxl.utils import get_column_letter

    rows: list[list[Any]] = []
    for field_name, names in aliases.items():
        label = names[0]
        if field_name in layout.report.matched:
            col, header = layout.report.matched[field_name]
            rows.append([label, get_column_letter(col), header, "머리글 일치"])
        elif field_name in layout.report.fallback:
            col = layout.report.fallback[field_name]
            rows.append([label, get_column_letter(col), "", "기본 위치"])
        else:
            rows.append([label, "", "", "없음"])
    return rows
