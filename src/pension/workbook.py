"""통합문서 열기 — ``.xlsx`` 와 ``.xls`` 를 같은 방식으로 다룬다.

실제로 받는 명부는 상당수가 예전 ``.xls`` (BIFF) 서식이다. openpyxl 은 이 형식을
읽지 못하므로 ``xlrd`` 로 읽어 openpyxl 과 같은 모양으로 감싼다. 덕분에
:mod:`pension.readers` 등 나머지 코드는 두 형식을 구분할 필요가 없다.

시트 이름도 통합문서마다 다르다. 같은 명부인데 ``재직자명부`` 와
``2)재직자명부`` 가 섞여 오므로 별칭으로 찾아 준다.
"""

from __future__ import annotations

import datetime as _dt
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Protocol

__all__ = ["Cell", "Sheet", "Workbook", "find_sheet", "open_workbook", "save_workbook"]


class Cell(Protocol):
    """openpyxl 셀 중 우리가 쓰는 부분."""

    value: Any


class Sheet(Protocol):
    """openpyxl 워크시트 중 우리가 쓰는 부분."""

    title: str
    max_row: int
    max_column: int

    def cell(self, row: int, column: int) -> Cell: ...


class Workbook(Protocol):
    sheetnames: list[str]

    def __getitem__(self, name: str) -> Sheet: ...
    def close(self) -> None: ...


# ────────────────────────────────────────────────────────────────
# .xls (BIFF) 어댑터
# ────────────────────────────────────────────────────────────────


@dataclass(frozen=True, slots=True)
class _XlsCell:
    value: Any


class _XlsSheet:
    """xlrd 시트를 openpyxl 워크시트처럼 보이게 감싼다.

    행·열 번호는 openpyxl 과 같이 1부터 센다(xlrd 는 0부터).
    """

    def __init__(self, sheet, datemode: int) -> None:
        self._sheet = sheet
        self._datemode = datemode

    @property
    def title(self) -> str:
        return self._sheet.name

    @property
    def max_row(self) -> int:
        return self._sheet.nrows

    @property
    def max_column(self) -> int:
        return self._sheet.ncols

    def cell(self, row: int, column: int) -> _XlsCell:
        if not (1 <= row <= self._sheet.nrows and 1 <= column <= self._sheet.ncols):
            return _XlsCell(None)
        return _XlsCell(self._convert(row - 1, column - 1))

    def _convert(self, row: int, col: int) -> Any:
        import xlrd

        kind = self._sheet.cell_type(row, col)
        value = self._sheet.cell_value(row, col)

        if kind in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
            return None
        if kind == xlrd.XL_CELL_DATE:
            parts = xlrd.xldate_as_tuple(value, self._datemode)
            if parts[:3] == (0, 0, 0):  # 시각만 있는 셀
                return _dt.time(*parts[3:])
            return _dt.datetime(*parts)
        if kind == xlrd.XL_CELL_BOOLEAN:
            return bool(value)
        if kind == xlrd.XL_CELL_ERROR:
            return None
        if kind == xlrd.XL_CELL_TEXT:
            return value.strip()
        return value


class _XlsWorkbook:
    def __init__(self, path: Path) -> None:
        import xlrd

        self._book = xlrd.open_workbook(path)
        self._sheets = {
            sheet.name: _XlsSheet(sheet, self._book.datemode)
            for sheet in self._book.sheets()
        }

    @property
    def sheetnames(self) -> list[str]:
        return list(self._sheets)

    def __getitem__(self, name: str) -> _XlsSheet:
        try:
            return self._sheets[name]
        except KeyError:
            raise KeyError(f"시트를 찾을 수 없습니다: {name}") from None

    def __contains__(self, name: str) -> bool:
        return name in self._sheets

    def close(self) -> None:
        release = getattr(self._book, "release_resources", None)
        if release is not None:
            release()


# ────────────────────────────────────────────────────────────────
# 열기
# ────────────────────────────────────────────────────────────────


def open_workbook(path: str | Path) -> Workbook:
    """명부·기초율 통합문서를 연다.

    :param path: ``.xlsx`` / ``.xlsm`` / ``.xls`` 경로.
    :raises FileNotFoundError: 파일이 없을 때.
    :raises ValueError: 지원하지 않는 확장자거나 파일이 손상되었을 때.
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"파일을 찾을 수 없습니다: {path}")

    suffix = path.suffix.lower()
    if suffix == ".xls":
        try:
            import xlrd  # noqa: F401
        except ImportError as exc:  # pragma: no cover - 설치 구성 문제
            raise ValueError(
                ".xls 형식을 읽으려면 xlrd 가 필요합니다. "
                "엑셀에서 .xlsx 로 저장한 뒤 다시 시도하셔도 됩니다"
            ) from exc
        try:
            return _XlsWorkbook(path)
        except Exception as exc:
            raise ValueError(f"파일을 읽지 못했습니다: {path.name} ({exc})") from exc

    if suffix in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        import openpyxl

        try:
            return openpyxl.load_workbook(path, data_only=True)
        except Exception as exc:
            raise ValueError(f"파일을 읽지 못했습니다: {path.name} ({exc})") from exc

    raise ValueError(
        f"지원하지 않는 파일 형식입니다: {path.suffix} "
        "(.xlsx, .xlsm, .xls 만 읽을 수 있습니다)"
    )


def uncalculated_formulas(
    path: str | Path, *, limit: int = 30
) -> tuple[list[str], int]:
    """수식이 적혀 있는데 **계산된 값이 없는** 칸.

    통합문서는 ``data_only=True`` 로 연다 — 수식이 아니라 값을 쓰기 위해서다.
    그런데 엑셀은 계산 결과를 파일에 **함께 저장해 둘 때만** 그 값을 돌려준다.
    파이썬·다른 도구로 만들었거나 엑셀에서 한 번도 열어 저장하지 않은 파일은
    그 자리가 비어 있고, 우리는 ``None`` 을 받는다.

    ``None`` 은 빈 칸과 구별되지 않으므로 **숫자 칸이면 조용히 0 이 된다.**
    임금·날짜는 뒤에서 검증이 잡지만 추계액·중간정산 지급금액처럼 0 이어도
    말이 되는 칸은 그냥 지나간다. 특히 추계액이 0 이 되면 `추계액대비` 검산이
    **소리 없이 꺼진다** — 지급률을 잘못 걸어도 알아챌 마지막 수단이 사라진다.

    :returns: ``(칸 목록, 전체 건수)``. 목록은 ``limit`` 개까지만 담는다.
        ``.xls`` 는 수식을 꺼내 볼 수 없어 항상 빈 결과다.
    """
    path = Path(path)
    if path.suffix.lower() not in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        return [], 0

    import openpyxl

    try:
        formulas = openpyxl.load_workbook(path, data_only=False, read_only=True)
        values = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception:
        # 여기서 못 읽어도 산출을 막지 않는다. 본 읽기가 따로 오류를 낸다.
        return [], 0

    found: list[str] = []
    total = 0
    try:
        for name in formulas.sheetnames:
            if name not in values.sheetnames:
                continue
            fsheet, vsheet = formulas[name], values[name]
            for frow, vrow in zip(fsheet.iter_rows(), vsheet.iter_rows()):
                for fcell, vcell in zip(frow, vrow):
                    text_value = fcell.value
                    if not isinstance(text_value, str) or not text_value.startswith("="):
                        continue
                    if vcell.value is not None:
                        continue
                    total += 1
                    if len(found) < limit:
                        found.append(f"{name}!{fcell.coordinate}")
    finally:
        formulas.close()
        values.close()
    return found, total


def find_sheet(workbook: Workbook, *aliases: str) -> Sheet | None:
    """별칭 중 하나에 해당하는 시트를 찾는다.

    같은 명부인데도 통합문서마다 시트 이름이 ``재직자명부`` / ``2)재직자명부``
    처럼 다르다. 번호 접두사와 공백을 무시하고 맞춰 본다.
    """
    wanted = {_simplify(a) for a in aliases}
    for name in workbook.sheetnames:
        if _simplify(name) in wanted:
            return workbook[name]
    return None


def _simplify(name: str) -> str:
    """``2)재직자명부`` → ``재직자명부``.

    번호 접두사와 공백·괄호를 걷어내고, 뒤에 붙은 꼬리표도 뗀다. 실제 명부에서
    ``2)재직자명부-2025`` 처럼 처럼 연도나 메모를 붙여 오는
    일이 잦은데, 그것 때문에 시트를 못 찾으면 파일이 통째로 안 열린다.
    """
    text = str(name).strip()
    # 앞의 "2)" "3." "1_" 같은 번호 표기를 떼어 낸다.
    index = 0
    while index < len(text) and (text[index].isdigit() or text[index] in ")].-_ "):
        index += 1
    text = text[index:].replace(" ", "") or text.replace(" ", "")
    # 뒤의 "-2025" "-해당X" 같은 꼬리표를 뗀다.
    for mark in ("-", "_", "("):
        head = text.split(mark)[0]
        if head:
            text = head
    return text


def save_workbook(book: Any, path: str | Path) -> Path:
    """워크북을 저장하고 **색이 투명해지는 것을 막는다.**

    openpyxl 은 ``Font(color="FFFFFF")`` 처럼 여섯 자리로 준 색을 ARGB 여덟
    자리로 늘리면서 앞에 ``00`` 을 붙인다. 그 자리는 알파 채널이고 ``00`` 은
    **완전 투명** 이다. 엑셀 대부분은 이 자리를 무시하고 그리지만, 그렇지
    않은 버전에서는 남색 머리글 위의 흰 글자가 통째로 사라진다 — 표는
    그려지는데 열 이름만 안 보이는, 파일이 깨진 것처럼 보이는 모양이다.

    저장한 뒤 서식표의 알파를 모두 ``FF``(불투명)로 올린다. 색값 자체는
    건드리지 않으므로 보이던 것이 달라지지는 않는다.
    """
    import re
    import shutil
    import zipfile

    path = Path(path)
    book.save(path)

    backup = path.with_suffix(path.suffix + ".alpha")
    shutil.copy2(path, backup)
    try:
        with zipfile.ZipFile(backup) as source, \
                zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "xl/styles.xml":
                    data = re.sub(
                        rb'rgb="00([0-9A-Fa-f]{6})"', rb'rgb="FF\1"', data)
                target.writestr(item, data)
    finally:
        backup.unlink(missing_ok=True)
    return path
