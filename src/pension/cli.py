"""명령행 인터페이스.

인자 없이 실행하면 GUI 가 뜬다(EXE 를 더블클릭한 경우). 하위 명령을 주면
배치 실행이 된다::

    pension calc 명부.xlsm 기초율.xlsx -o 산출결과.xlsx
    pension check 명부.xlsm                 # 검증만
    pension template 기초율.xlsx            # 기초율 양식 생성
    pension samples 기본자료                 # 명부 양식·기초율 기본값 한 벌 생성
    pension library add 금리표 KIS.xlsx      # 한 번 등록해 두고 모든 단체에 재사용
    pension upload 명부.xlsm -o 업로드.xlsx  # 업로드 명부만 생성
    pension app                              # 이 PC 브라우저로 전체 기능 화면
    pension web --host 0.0.0.0               # 아이패드 등에서 접속하는 웹 화면

``app`` 과 ``web`` 은 다른 화면이다. ``app`` 은 분석 그래프·보고서·산출 내역이
다 있는 전체 화면을 **이 PC 에만**(127.0.0.1) 열고, 계산은 브라우저 안에서
돈다. ``web`` 은 다른 기기에서 붙어 쓰는 간단한 산출 폼이고 계산은 서버 PC 가
한다 — 인증이 없으므로 사내망 밖으로 열면 안 된다.
"""

from __future__ import annotations

import argparse
import contextlib
import sys
from pathlib import Path

from . import __version__
from .errors import PensionDataError, Severity
from .standard_rates import DEFAULT_SIZE, SIZES, STANDARD_YEAR

__all__ = ["main"]


def force_utf8_output() -> None:
    """표준 출력을 UTF-8 로 맞춘다.

    한국어 윈도우의 명령프롬프트는 기본 코드페이지가 949(CP949)이고, GitHub
    Actions 같은 영문 환경은 1252 다. 어느 쪽이든 파이썬은 그 코드페이지로
    stdout 을 열기 때문에, 한글이 섞인 메시지를 ``print`` 하면
    ``UnicodeEncodeError: 'charmap' codec can't encode characters`` 로 죽는다.
    산출은 다 끝내 놓고 결과를 찍다가 실패하는 셈이라 특히 나쁘다.

    ``errors="replace"`` 를 함께 주어, 콘솔이 정말 표현하지 못하는 글자가 있어도
    프로그램이 멈추지는 않게 한다.

    창 모드로 빌드된 실행 파일(``console=False``)에서는 ``sys.stdout`` 이
    ``None`` 일 수 있으므로 존재 여부를 확인한다.
    """
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if reconfigure is None:
            continue
        # 이미 닫혔거나 재설정할 수 없는 스트림이면 그냥 둔다. 출력 인코딩
        # 때문에 산출 자체를 막을 이유는 없다.
        with contextlib.suppress(ValueError, OSError):
            reconfigure(encoding="utf-8", errors="replace")


def _percent(value: str) -> float:
    token = value.strip().rstrip("%")
    number = float(token)
    return number / 100.0 if value.strip().endswith("%") or number > 1.0 else number


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="pension",
        description="연금계리 산출 시스템 — K-IFRS 1019 확정급여채무(PUC)",
    )
    parser.add_argument("--version", action="version", version=f"%(prog)s {__version__}")
    sub = parser.add_subparsers(dest="command")

    calc = sub.add_parser("calc", help="확정급여채무 산출")
    calc.add_argument("roster", type=Path, help="명부 워크북 (.xlsm/.xlsx)")
    calc.add_argument("assumptions", type=Path, help="기초율 워크북")
    calc.add_argument("-o", "--output", type=Path, required=True, help="결과 엑셀 경로")
    calc.add_argument("--no-sensitivity", action="store_true", help="민감도분석 생략")
    calc.add_argument("--no-longterm", action="store_true", help="장기종업원급여 생략")
    calc.add_argument("--force", action="store_true", help="검증 오류가 있어도 산출")
    calc.add_argument(
        "--members", type=Path, metavar="경로",
        help="개인별 결과를 별도 엑셀로도 저장 (생략하면 결과 파일 안 시트로만)",
    )
    calc.add_argument("--prior-dbo", type=float, default=0.0, help="전기말 확정급여채무")
    calc.add_argument("--prior-rate", type=_percent, default=0.0, help="전기말 할인율 (예: 4.5%%)")
    calc.add_argument("--prior-assumptions", type=Path, help="전기 기초율 워크북")

    check = sub.add_parser("check", help="명부 검증만 실행")
    check.add_argument("roster", type=Path)
    check.add_argument(
        "assumptions", type=Path, nargs="?",
        help="기초율 워크북. 주면 그 안의 '지급규정' 시트(직군 매핑)를 적용해 "
             "검증합니다 — calc 와 같은 조건이 됩니다",
    )
    check.add_argument("-o", "--output", type=Path, help="검증 결과 CSV 경로")

    template = sub.add_parser("template", help="기초율 양식 생성")
    template.add_argument("output", type=Path)
    template.add_argument("--roster", type=Path, help="직군명을 가져올 명부 워크북")

    samples = sub.add_parser(
        "samples", help="명부 양식·기초율 기본값 등 기본 파일 한 벌 생성"
    )
    samples.add_argument(
        "directory", type=Path, nargs="?", default=Path("기본자료"),
        help="만들 폴더 (생략하면 ./기본자료)",
    )
    samples.add_argument(
        "--job-groups", default="", help="직군을 쉼표로. 생략하면 정규직,계약직,임원"
    )
    samples.add_argument(
        "--yield-curve", type=Path, metavar="금리표",
        help="할인율로 쓸 채권 금리표(KIS-Net 등). 만기별 현물이자율을 그대로 옮깁니다",
    )
    samples.add_argument(
        "--grade", default="",
        help="금리표에서 고를 등급 (기본: AA0 → AA+ → AA- 순으로 찾음)",
    )
    samples.add_argument(
        "--size", default=DEFAULT_SIZE, choices=list(SIZES),
        help="표준률 원표의 사업장 규모. 승급률·중도퇴직률이 이 축으로 갈립니다",
    )
    samples.add_argument(
        "--cases", action="store_true",
        help="난수로 만든 시험용 명부 3종(표준·복합제도·자료불량)과 짝 기초율도 함께",
    )
    samples.add_argument(
        "--seed", type=int, default=20251231,
        help="시험용 명부 난수 씨앗. 같은 값이면 같은 명부가 나옵니다",
    )
    samples.add_argument(
        "--case-base-date", metavar="YYYY-MM-DD", default="",
        help="시험용 명부의 산출기준일 (기본 2025-12-31)",
    )

    lib = sub.add_parser("library", help="금리표·표준률을 시스템에 등록/조회")
    lib.add_argument("action", choices=["list", "add", "remove"], help="할 일")
    lib.add_argument("kind", nargs="?", choices=["금리표", "표준률"], help="등록 종류")
    lib.add_argument("file", nargs="?", type=Path, help="add 할 파일 / remove 할 이름")
    lib.add_argument("--name", default="", help="목록에 표시할 이름")

    upload = sub.add_parser("upload", help="업로드 명부만 생성")
    upload.add_argument("roster", type=Path)
    upload.add_argument("-o", "--output", type=Path, required=True)
    upload.add_argument("--force", action="store_true", help="검증 오류가 있어도 생성")

    members = sub.add_parser("members", help="개인별 산출 결과만 엑셀로 내보내기")
    members.add_argument("roster", type=Path, help="명부 워크북")
    members.add_argument("assumptions", type=Path, help="기초율 워크북")
    members.add_argument("-o", "--output", type=Path, required=True)
    members.add_argument("--force", action="store_true", help="검증 오류가 있어도 산출")

    web = sub.add_parser("web", help="웹 화면 실행 (아이패드·다른 PC 에서 접속)")
    web.add_argument("--host", default="127.0.0.1",
                     help="접속을 허용할 주소. 아이패드에서 쓰려면 0.0.0.0 (기본: 이 PC 만)")
    web.add_argument("--port", type=int, default=8035, help="포트 (기본 8035)")

    app = sub.add_parser(
        "app",
        help="전체 기능 화면 열기 (분석 그래프·보고서·산출 내역·단체 관리)",
    )
    app.add_argument("--port", type=int, default=None,
                     help="포트 (기본 8036). 저장해 둔 산출 내역은 주소마다 "
                          "따로 보관되므로 늘 같은 포트로 여는 편이 좋습니다")
    app.add_argument("--path", type=Path, help="웹앱 폴더를 직접 지정")
    app.add_argument("--no-browser", action="store_true",
                     help="브라우저를 열지 않고 주소만 알려 줍니다")
    app.add_argument("--check", action="store_true",
                     help="화면을 띄우지 않고 웹앱이 제자리에 있는지만 확인합니다")

    sub.add_parser("gui", help="GUI 실행")

    editor = sub.add_parser("assumptions", help="산출 가정 입력 화면 실행")
    editor.add_argument(
        "--roster", type=Path, help="직군명을 미리 채워 올 명부 워크북"
    )
    editor.add_argument(
        "--open", type=Path, dest="existing", help="열어 둘 기초율 파일"
    )
    return parser


def _print_issue_list(issues: list, limit: int = 30) -> None:
    """이슈 목록 출력. 오류는 낱낱이, 경고가 많으면 코드별로 묶는다.

    실제 명부에서는 같은 성격의 경고가 수십 건씩 나온다(근속 1년 미만 퇴직자
    37건, DC 가입자 32건…). 이것을 한 줄씩 나열하면 화면이 그것으로 차서 정작
    조치가 필요한 오류가 스크롤 위로 밀려난다. 오류는 행 단위 조치가 필요하니
    전부 보여 주되, 경고는 많아지면 "무엇이 몇 건" 으로 접는다.
    """
    errors = [i for i in issues if i.severity is Severity.ERROR]
    warnings = [i for i in issues if i.severity is Severity.WARNING]
    notices = [i for i in issues if i.severity is Severity.INFO]

    for issue in errors[:limit]:
        print(f"  {issue}")
    if len(errors) > limit:
        print(f"  … 외 오류 {len(errors) - limit}건 (검증 리포트 참조)")

    if len(warnings) <= 15:
        for issue in warnings:
            print(f"  {issue}")
    else:
        print(f"\n  경고 요약 ({len(warnings)}건 — 상세는 결과 파일의 검증리포트 시트)")
        grouped: dict[str, list] = {}
        for issue in warnings:
            grouped.setdefault(issue.code, []).append(issue)
        for code, group in sorted(grouped.items(), key=lambda kv: -len(kv[1])):
            print(f"    {len(group):>4}건  {code:<28} 예) {group[0].message}")

    # 안내는 확인할 것이 없으므로 끝에 모아 보여 준다.
    for issue in notices:
        print(f"  {issue}")
    print(f"\n오류 {len(errors)}건 / 경고 {len(warnings)}건 / 안내 {len(notices)}건")


def _print_issues(log, limit: int = 30) -> None:
    _print_issue_list(list(log), limit)


def _cmd_calc(args: argparse.Namespace) -> int:
    from .pipeline import PriorPeriod, RunOptions, run_valuation
    from .report import write_report
    from .sensitivity import DEFAULT_SHOCKS

    options = RunOptions(
        roster_path=args.roster,
        assumptions_path=args.assumptions,
        output_path=args.output,
        include_sensitivity=not args.no_sensitivity,
        include_longterm=not args.no_longterm,
        shocks=DEFAULT_SHOCKS,
        allow_errors=args.force,
        prior=PriorPeriod(
            dbo=args.prior_dbo,
            discount_rate=args.prior_rate,
            assumptions_path=str(args.prior_assumptions) if args.prior_assumptions else "",
        ),
    )

    def progress(message: str, fraction: float) -> None:
        print(f"[{fraction * 100:3.0f}%] {message}")

    try:
        run = run_valuation(options, progress)
    except PensionDataError as exc:
        print(f"\n{exc}\n", file=sys.stderr)
        _print_issue_list(exc.issues)
        print("\n--force 를 주면 오류를 남긴 채 산출할 수 있습니다.", file=sys.stderr)
        return 2

    write_report(run, args.output)

    v = run.valuation
    print("\n── 산출 결과 " + "─" * 44)
    print(f"  산출기준일     {run.config.base_date}")
    if run.assumptions.discount.flat is None:
        print(f"  적용 할인율    {v.single_discount_rate():.3%}  (수익률곡선기법 단일할인율)")
    else:
        print(f"  적용 할인율    {run.assumptions.discount.level_rate:.3%}")
    print(f"  산출대상 인원  {v.headcount:,}명")
    print(f"  확정급여채무   {v.dbo:>18,.0f} 원")
    print(f"  당기근무원가   {v.service_cost:>18,.0f} 원")
    print(f"  이자원가(차기) {v.interest_cost:>18,.0f} 원")
    print(f"  듀레이션       {v.duration:>18,.1f} 년")
    if run.longterm is not None:
        print(f"  장기급여채무   {run.longterm.dbo:>18,.0f} 원")
    if run.rollforward is not None:
        print(f"  보험수리적손익 {run.rollforward.actuarial_gain_loss:>18,.0f} 원")
    print(f"  명부 검증      오류 {len(run.issues.errors)}건 / 경고 {len(run.issues.warnings)}건")

    excluded = v.exclusion_summary()
    if excluded:
        total = sum(excluded.values())
        print(f"\n  산출 제외 {total:,}명")
        for reason, count in excluded.items():
            print(f"    {count:>5,}명  {reason}")
        if v.headcount == 0:
            print(
                "\n  ※ 산출대상이 한 명도 없어 확정급여채무가 0 입니다. "
                "위 사유를 확인하세요."
            )
    print(f"\n결과 파일: {args.output}")
    if args.members:
        from .members import write_member_export

        write_member_export(run, args.members)
        print(f"개인별 결과: {args.members}")
    return 0


def _cmd_check(args: argparse.Namespace) -> int:
    """명부 검증.

    기초율 워크북을 함께 주면 그 ``지급규정`` 시트를 적용한다. 안 주면 명부의
    ``Input`` 만 보게 되는데, 직군 매핑을 지급규정에 둔 경우 check 는 직군
    미매칭 오류를 내고 calc 는 통과하는 — 서로 다른 판정이 나오는 함정이 있었다.
    """
    import csv
    from dataclasses import replace as _replace

    from .config import read_config
    from .errors import IssueLog
    from .readers import read_roster
    from .validation import validate_roster
    from .workbook import open_workbook

    payout_rules = []
    if args.assumptions:
        from .config import read_payout_rules

        assumptions_wb = open_workbook(args.assumptions)
        try:
            payout_rules = read_payout_rules(assumptions_wb)
        finally:
            assumptions_wb.close()

    wb = open_workbook(args.roster)
    try:
        config = read_config(wb)
        if payout_rules:
            config = _replace(config, job_group_rules=payout_rules, inferred=False)
        log = IssueLog()
        roster = read_roster(wb, config, log)
    finally:
        wb.close()
    validate_roster(roster, config, log)

    print(f"산출기준일 {config.base_date} · 재직 {len(roster.active):,}명 / 퇴직 {len(roster.retired):,}명\n")
    _print_issues(log)

    if args.output:
        with open(args.output, "w", newline="", encoding="utf-8-sig") as fh:
            writer = csv.writer(fh)
            writer.writerow(["심각도", "시트", "행", "열", "순번", "사번", "코드", "내용", "값"])
            for issue in log:
                writer.writerow([
                    issue.severity.value, issue.sheet, issue.row, issue.column,
                    issue.seq, issue.employee_id, issue.code, issue.message, issue.value,
                ])
        print(f"\n검증 결과: {args.output}")

    return 1 if log.has_errors() else 0


def _job_groups_from_roster(roster: Path | None) -> list[str]:
    """명부 ``Input`` 시트가 참조하는 규정명. 읽지 못하면 빈 목록."""
    if not roster or not roster.exists():
        return []


    from .config import read_config
    from .workbook import open_workbook

    try:
        wb = open_workbook(roster)
    except (OSError, ValueError, KeyError):
        return []
    try:
        return read_config(wb).referenced_rule_names()
    except (ValueError, KeyError):
        return []
    finally:
        wb.close()


def _cmd_template(args: argparse.Namespace) -> int:
    from .assumptions import write_template

    groups = _job_groups_from_roster(args.roster)
    path = write_template(args.output, job_groups=groups)
    print(f"기초율 양식을 만들었습니다: {path}")
    if groups:
        print(f"  규정명 열: {', '.join(groups)}")
    return 0


def _cmd_samples(args: argparse.Namespace) -> int:
    from .jobgroup import DEFAULT_GROUPS
    from .samples import write_sample_pack

    groups = tuple(g.strip() for g in args.job_groups.split(",") if g.strip())
    paths = write_sample_pack(
        args.directory,
        job_groups=groups or DEFAULT_GROUPS,
        yield_curve_path=args.yield_curve,
        grade=args.grade,
        size=args.size,
    )

    print(f"기본 파일을 만들었습니다: {args.directory}")
    for path in paths:
        print(f"  {path.name}")

    if args.cases:
        from .rostergen import CASES, write_case_pack

        base_date = None
        if args.case_base_date:
            import datetime as _dt

            base_date = _dt.date.fromisoformat(args.case_base_date)
        made = write_case_pack(
            args.directory / "시험명부", seed=args.seed, base_date=base_date,
        )
        print(f"\n시험용 명부 {len(CASES)}종 (재직 290명 안팎, 난수 씨앗 {args.seed}):")
        for spec in CASES:
            print(f"  {spec.title}.xlsx + {spec.title}_기초율.xlsx — {spec.summary}")
        print(f"  → {args.directory / '시험명부'} ({len(made)}개 파일)")

    print(
        "\n바로 돌려 보려면:\n"
        f"  pension calc {paths[0]} {paths[1]} -o 산출결과.xlsx\n"
    )
    if args.yield_curve:
        print(f"할인율은 {Path(args.yield_curve).name} 의 현물이자율 곡선을 옮겼습니다.")
    else:
        print("할인율은 자리값(4.5%)입니다. --yield-curve 로 결산일 금리표를 주세요.")
    print(
        f"퇴직률·승급률은 {STANDARD_YEAR} 의 '{args.size}' 열입니다. "
        "회사 경험률이 있으면 그쪽이 우선입니다.\n"
        "사망률은 재직자 기준 표준사망률(남녀 구분)이라 그대로 쓸 수 있습니다."
    )
    return 0


def _cmd_web(args: argparse.Namespace) -> int:
    from .web import serve

    return serve(args.host, args.port)


def _cmd_app(args: argparse.Namespace) -> int:
    """전체 기능 화면(웹앱)을 이 PC 브라우저로 연다."""
    from .localapp import MissingAppError, app_root, serve

    # 빌드가 웹앱을 실행 파일 안에 제대로 묶었는지 확인하는 길. 화면을 띄우면
    # 사람이 창을 닫아 줘야 끝나므로 자동 확인에 쓸 수 없다.
    if args.check:
        try:
            print(f"확인: 전체 기능 화면이 제자리에 있습니다 — {app_root()}")
            return 0
        except MissingAppError as exc:
            print(str(exc))
            return 1

    try:
        server = serve(args.path, args.port)
    except MissingAppError as exc:
        print(str(exc))
        return 1

    url = f"http://127.0.0.1:{server.server_address[1]}/index.html"
    print("전체 기능 화면")
    print(f"  주소:  {url}")
    print("  분석 그래프 · 계리평가 보고서 · 산출 내역 · 단체 관리가 모두 있습니다.")
    print("  계산은 브라우저 안에서 돕니다 — 명부가 이 PC 밖으로 나가지 않습니다.")
    print("  멈추려면 Ctrl+C\n")

    if not args.no_browser:
        import webbrowser

        webbrowser.open(url)

    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\n화면을 닫았습니다.")
    finally:
        server.server_close()
    return 0


def _cmd_library(args: argparse.Namespace) -> int:
    from .library import CURVE_KIND, RATES_KIND, entries, library_dir, register, remove

    kinds = [args.kind] if args.kind else [CURVE_KIND, RATES_KIND]

    if args.action == "list":
        print(f"등록 폴더: {library_dir()}")
        for kind in kinds:
            found = entries(kind)
            print(f"\n[{kind}] {len(found)}건")
            for entry in found:
                print(f"  {entry.label}")
            if not found:
                print("  (없음)")
        return 0

    if not args.kind or args.file is None:
        print("종류와 파일(또는 이름)을 지정하세요.", file=sys.stderr)
        return 2

    if args.action == "add":
        entry = register(args.kind, args.file, name=args.name)
        print(f"등록했습니다: [{entry.kind}] {entry.name}\n  {entry.path}")
        return 0

    if remove(args.kind, str(args.file)):
        print(f"지웠습니다: [{args.kind}] {args.file}")
        return 0
    print(f"'{args.file}' 을(를) 찾지 못했습니다.", file=sys.stderr)
    return 1


def _cmd_upload(args: argparse.Namespace) -> int:
    import openpyxl

    from .config import read_config
    from .errors import IssueLog
    from .readers import read_roster
    from .upload import ACTIVE_UPLOAD_HEADERS, RETIRED_UPLOAD_HEADERS, build_upload
    from .validation import validate_roster
    from .workbook import open_workbook

    wb = open_workbook(args.roster)
    try:
        config = read_config(wb)
        log = IssueLog()
        roster = read_roster(wb, config, log)
    finally:
        wb.close()
    validate_roster(roster, config, log)

    if log.has_errors() and not args.force:
        print(f"검증 오류 {len(log.errors)}건으로 중단했습니다.\n", file=sys.stderr)
        _print_issues(log)
        return 2

    active, retired = build_upload(roster, config)
    out = openpyxl.Workbook()
    del out["Sheet"]
    for title, headers, rows in (
        ("UpLoad_Jae", ACTIVE_UPLOAD_HEADERS, active),
        ("UpLoad_Toi", RETIRED_UPLOAD_HEADERS, retired),
    ):
        ws = out.create_sheet(title)
        ws.append(list(headers))
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
    out.save(args.output)

    print(f"재직자 {len(active):,}명 / 퇴직자 {len(retired):,}명 → {args.output}")
    if log.warnings:
        print(f"경고 {len(log.warnings)}건이 있습니다. 'pension check' 로 확인하세요.")
    return 0


def _cmd_members(args: argparse.Namespace) -> int:
    """산출한 뒤 개인별 결과만 별도 엑셀로 내보낸다."""
    from .members import build_member_rows, write_member_export
    from .pipeline import RunOptions, run_valuation

    options = RunOptions(
        roster_path=args.roster,
        assumptions_path=args.assumptions,
        output_path=args.output,
        include_sensitivity=False,
        include_longterm=True,
        allow_errors=args.force,
    )

    def progress(message: str, fraction: float) -> None:
        print(f"[{fraction * 100:3.0f}%] {message}")

    try:
        run = run_valuation(options, progress)
    except PensionDataError as exc:
        print(f"\n{exc}\n", file=sys.stderr)
        _print_issue_list(exc.issues)
        print("\n--force 를 주면 오류를 남긴 채 산출할 수 있습니다.", file=sys.stderr)
        return 2

    write_member_export(run, args.output)
    rows = build_member_rows(run)
    included = [r for r in rows if not r.excluded_reason]
    print(f"\n개인별 결과 {len(rows):,}명 (산출대상 {len(included):,}명)")
    print(f"  확정급여채무 합계 {sum(r.dbo for r in rows):>18,.0f} 원")
    print(f"  채무 합계       {sum(r.total_dbo for r in rows):>18,.0f} 원")
    print(f"\n파일: {args.output}")
    return 0


def _cmd_assumptions(args: argparse.Namespace) -> int:
    """산출 가정 입력 화면을 띄운다."""
    import tkinter as tk

    from .editor import AssumptionsEditor

    groups = None
    if args.roster:
        groups = _job_groups_from_roster(args.roster) or None

    root = tk.Tk()
    root.withdraw()
    window = AssumptionsEditor(root, job_groups=groups, on_close=lambda _e: root.destroy())
    if args.existing:
        try:
            window.load_workbook(args.existing)
            window.path = args.existing
            window.status.configure(text=f"불러왔습니다: {args.existing.name}")
        except (OSError, ValueError) as exc:
            print(f"기초율 파일을 읽지 못했습니다: {exc}")
    root.mainloop()
    return 0


def _cmd_gui(_args: argparse.Namespace) -> int:
    from .gui import main as gui_main

    return gui_main()


def main(argv: list[str] | None = None) -> int:
    """진입점. 인자가 없으면 본 화면을 띄운다.

    본 화면은 **전체 기능 화면 하나** 다. 산출·산출가정 입력·분석·보고서·
    산출 내역이 한 창 안에 다 있다. 예전의 작은 입력 창은 `gui` 하위 명령으로
    남겨 두었다 — 그 창에만 익숙한 사람이 있을 수 있어서다.
    """
    force_utf8_output()
    argv = list(sys.argv[1:] if argv is None else argv)
    if not argv:
        from .localapp import MissingAppError, run_app

        try:
            return run_app()
        except MissingAppError:
            # 웹앱이 없는 설치본이면 옛 입력 창이라도 띄운다.
            return _cmd_gui(argparse.Namespace())

    parser = _build_parser()
    args = parser.parse_args(argv)

    handlers = {
        "calc": _cmd_calc,
        "check": _cmd_check,
        "template": _cmd_template,
        "samples": _cmd_samples,
        "library": _cmd_library,
        "web": _cmd_web,
        "app": _cmd_app,
        "assumptions": _cmd_assumptions,
        "members": _cmd_members,
        "upload": _cmd_upload,
        "gui": _cmd_gui,
    }
    handler = handlers.get(args.command)
    if handler is None:
        parser.print_help()
        return 1

    try:
        return handler(args)
    except FileNotFoundError as exc:
        print(f"오류: {exc}", file=sys.stderr)
        return 2
    except ValueError as exc:
        print(f"오류: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main())
