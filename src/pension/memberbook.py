"""사번 하나의 산출 근거를 엑셀 한 부로.

화면(사번 조회)에서 보는 것과 **같은 자료**(:func:`pension.memberdetail.lookup`)
를 그대로 옮긴다. 별도로 다시 계산하지 않으므로 화면과 파일의 숫자가 어긋날
길이 없다.

인쇄(PDF)만으로는 부족해서 만들었다. 감사인·회사 담당자는 받은 숫자를 **자기
파일에 붙여 다시 계산해 본다.** 표가 그림으로 굳어 있으면 한 줄씩 손으로 옮겨
적어야 하고, 옮기다 틀리면 그 틀린 값으로 우리 산출을 의심하게 된다.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from .workbook import save_workbook

__all__ = ["write_member_book"]

_MONEY = "#,##0"
_RATE = "0.000000"
_YEARS = "#,##0.00"

#: 연차별 근거 표의 열. (자료 열쇠, 머리글, 서식)
_TRACE: tuple[tuple[str, str, str], ...] = (
    ("t", "연차", ""),
    ("timing", "시점(년)", _YEARS),
    ("age", "연령", _YEARS),
    ("service", "근속", _YEARS),
    ("wage", "월평균임금", _MONEY),
    ("multiple", "지급률", _YEARS),
    ("cause", "퇴직사유", ""),
    ("withdrawal", "중도퇴직률", _RATE),
    ("mortality", "사망률", _RATE),
    ("survival", "연초 재직확률", _RATE),
    ("exit_probability", "탈퇴확률", _RATE),
    ("benefit", "퇴직 시 급여", _MONEY),
    ("attributed", "귀속 급여", _MONEY),
    ("unit", "당기 1년치", _MONEY),
    ("discount", "할인계수", _RATE),
    ("dbo", "확정급여채무", _MONEY),
    ("service_cost", "당기근무원가", _MONEY),
)


def _style():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color="BFBFBF")
    return {
        "title": Font(bold=True, size=13, color="1F3864"),
        "head": Font(bold=True, color="FFFFFF"),
        "head_fill": PatternFill("solid", fgColor="44546A"),
        "section": Font(bold=True, size=11, color="1F3864"),
        "section_fill": PatternFill("solid", fgColor="D9E2F3"),
        "total": Font(bold=True),
        "total_fill": PatternFill("solid", fgColor="FFF2CC"),
        "note": Font(color="5B6478", size=9),
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "wrap": Alignment(vertical="top", wrap_text=True),
        "center": Alignment(horizontal="center", vertical="center", wrap_text=True),
    }


def _money_like(label: str) -> bool:
    """금액으로 보이는 항목 이름인지. 서식을 이름으로 고른다.

    자료가 ``{이름: 값}`` 뭉치라 열 번호로 서식을 걸 수가 없다. 이름으로 고르면
    항목이 늘어도 따라온다.
    """
    if any(word in label for word in ("배수", "지급률", "듀레이션", "연수", "근속", "비중")):
        return False
    return any(word in label for word in
               ("임금", "채무", "원가", "추계액", "금액", "급여", "가산", "DBO"))


def _pairs(ws, row: int, title: str, items: list[tuple[str, Any]], st) -> int:
    """``이름 | 값`` 두 칸짜리 구획 하나."""
    cell = ws.cell(row, 2, title)
    cell.font = st["section"]
    cell.fill = st["section_fill"]
    ws.cell(row, 3).fill = st["section_fill"]
    row += 1
    for label, value in items:
        ws.cell(row, 2, label).font = st["total"]
        target = ws.cell(row, 3, value)
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            target.number_format = _MONEY if _money_like(label) else _YEARS
        else:
            target.alignment = st["wrap"]
        for column in (2, 3):
            ws.cell(row, column).border = st["border"]
        row += 1
    return row + 1


def _trace_sheet(wb, title: str, trace: list[dict], st) -> None:
    """연차별 근거 한 장. 줄 합계가 곧 그 사람의 채무다."""
    ws = wb.create_sheet(title)
    for index, (_key, head, _fmt) in enumerate(_TRACE, start=1):
        cell = ws.cell(1, index, head)
        cell.font = st["head"]
        cell.fill = st["head_fill"]
        cell.alignment = st["center"]
        cell.border = st["border"]
        ws.column_dimensions[ws.cell(1, index).column_letter].width = max(
            11, min(16, len(head) + 5))

    for offset, line in enumerate(trace, start=2):
        for index, (key, _head, fmt) in enumerate(_TRACE, start=1):
            cell = ws.cell(offset, index, line.get(key))
            cell.border = st["border"]
            if fmt:
                cell.number_format = fmt

    end = len(trace) + 1
    if not trace:
        return
    # 합계는 수식으로 둔다. 박아 넣은 숫자는 눌러 봐도 무엇을 더한 것인지
    # 알 수 없어, 받는 사람이 검산할 수가 없다.
    row = end + 1
    ws.cell(row, 1, "합계").font = st["total"]
    for index, (key, _head, fmt) in enumerate(_TRACE, start=1):
        if key not in ("dbo", "service_cost"):
            continue
        letter = ws.cell(1, index).column_letter
        cell = ws.cell(row, index, f"=SUM({letter}2:{letter}{end})")
        cell.number_format = fmt
        cell.font = st["total"]
        cell.fill = st["total_fill"]
    ws.freeze_panes = ws.cell(2, 1)


def write_member_book(detail: dict[str, Any], path: str | Path) -> Path:
    """:func:`pension.memberdetail.lookup` 이 낸 것을 엑셀로 쓴다.

    :param detail: 조회 결과 그대로.
    :param path: 저장할 ``.xlsx`` 경로.
    """
    from openpyxl import Workbook

    st = _style()
    wb = Workbook()
    ws = wb.active
    ws.title = "산출근거"
    for column, width in (("A", 3), ("B", 30), ("C", 30)):
        ws.column_dimensions[column].width = width

    rows = detail.get("rows") or []
    first = rows[0] if rows else {}
    who = (first.get("profile", {}).get("사번")
           or detail.get("employee_id") or "")
    ws["B2"] = f"개인별 산출 근거 · {who}"
    ws["B2"].font = st["title"]
    ws["B3"] = (f"산출기준일 {detail.get('base_date', '')}   ·   "
                f"할인율 {detail.get('discount_rate', 0):.3%}")
    ws["B3"].font = st["note"]

    row = 5
    for index, block in enumerate(rows, start=1):
        if len(rows) > 1:
            ws.cell(row, 2, f"지급 구간 {index} / {len(rows)}").font = st["section"]
            row += 1
        if block.get("excluded"):
            ws.cell(row, 2, f"산출 제외: {block['excluded']}").font = st["total"]
            row += 2
            continue
        row = _pairs(ws, row, "인적사항", list(block["profile"].items()), st)
        row = _pairs(ws, row, "적용한 규정·가정", list(block["applied"].items()), st)
        row = _pairs(ws, row, "산출 결과", list(block["result"].items()), st)
        shares = [(f"{s['cause']} — 확정급여채무", s["dbo"]) for s in block["by_cause"]]
        shares += [(f"{s['cause']} — 그중 가산", s.get("extra", 0.0))
                   for s in block["by_cause"] if s.get("extra")]
        if shares:
            row = _pairs(ws, row, "퇴직사유별 몫", shares, st)

    for index, block in enumerate(rows, start=1):
        if block.get("trace"):
            name = "연차별 근거" if len(rows) == 1 else f"연차별 근거 {index}"
            _trace_sheet(wb, name, block["trace"], st)

    for index, block in enumerate(detail.get("longterm") or [], start=1):
        if block.get("excluded"):
            continue
        name = "장기급여" if index == 1 else f"장기급여 {index}"
        sheet = wb.create_sheet(name)
        sheet.column_dimensions["B"].width = 30
        sheet.column_dimensions["C"].width = 24
        line = 2
        items = [(k, v) for k, v in block.items()
                 if k not in ("result", "trace", "excluded")]
        items += list(block["result"].items())
        line = _pairs(sheet, line, "장기종업원급여", items, st)

    if detail.get("retired"):
        sheet = wb.create_sheet("퇴직자명부 내역")
        sheet.column_dimensions["B"].width = 30
        sheet.column_dimensions["C"].width = 24
        line = 2
        for entry in detail["retired"]:
            line = _pairs(sheet, line, "퇴직 내역", list(entry.items()), st)

    if detail.get("total") and len(rows) > 1:
        row = _pairs(ws, row, "합계 (지급 구간 전체)",
                     list(detail["total"].items()), st)

    return save_workbook(wb, path)
